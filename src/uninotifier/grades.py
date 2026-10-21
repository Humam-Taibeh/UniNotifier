"""
Automated grade finder.

Scans downloaded spreadsheets / Word documents / text files for the
student's university number and, if present, extracts that row's marks.

Returns {"summary": "18/20", "details": [(header, value), ...]} or None.
"""

import csv
import logging
import re
from pathlib import Path

log = logging.getLogger("uninotifier")

SCANNABLE_SUFFIXES = {".xlsx", ".xlsm", ".xls", ".docx", ".txt", ".csv"}

_TOTAL_HEADER = re.compile(
    r"total|grade|mark|score|final|sum|المجموع|مجموع|العلامة|علامة|الدرجة|درجة|النهائي", re.I)
_IDENTITY_HEADER = re.compile(
    r"\bname|\bid\b|\bno\.?$|number|serial|seq|student|اسم|رقم|الرقم|تسلسل|الطالب|م$|#", re.I)
_MAX_IN_HEADER = re.compile(r"/\s*(\d+(?:\.\d+)?)|\(\s*(\d+(?:\.\d+)?)\s*\)|out of\s*(\d+(?:\.\d+)?)", re.I)


def _norm(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return re.sub(r"\s+", " ", str(value)).strip()


def _is_number(text: str) -> bool:
    return bool(re.fullmatch(r"-?\d+(?:\.\d+)?", text))


def _read_rows(path: Path) -> list[list[str]]:
    suffix = path.suffix.lower()
    rows: list[list[str]] = []
    if suffix in (".xlsx", ".xlsm"):
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        for ws in wb.worksheets:
            rows += [[_norm(c) for c in r] for r in ws.iter_rows(values_only=True)]
            rows.append([])
        wb.close()
    elif suffix == ".xls":
        import xlrd
        wb = xlrd.open_workbook(path)
        for ws in wb.sheets():
            rows += [[_norm(c) for c in ws.row_values(i)] for i in range(ws.nrows)]
            rows.append([])
        wb.release_resources()
    elif suffix == ".docx":
        import docx
        doc = docx.Document(str(path))
        for table in doc.tables:
            rows += [[_norm(c.text) for c in r.cells] for r in table.rows]
            rows.append([])
        rows += [[_norm(p.text)] for p in doc.paragraphs if p.text.strip()]
    elif suffix in (".txt", ".csv"):
        text = path.read_text(encoding="utf-8-sig", errors="replace")
        for line in text.splitlines():
            if "," in line or ";" in line or "\t" in line:
                rows.append([_norm(c) for c in next(csv.reader([line], delimiter=_delim(line)))])
            else:
                rows.append([_norm(line)])
    return rows


def _delim(line: str) -> str:
    return "\t" if "\t" in line else ";" if line.count(";") > line.count(",") else ","


def _find_header(rows: list[list[str]], idx: int) -> list[str]:
    """Nearest row above idx that looks like column titles (3+ text cells)."""
    for j in range(idx - 1, max(idx - 60, -1), -1):
        texts = [c for c in rows[j] if c and not _is_number(c)]
        if len(texts) >= 3:
            return rows[j]
    return []


def _with_max(header: str, value: str) -> str:
    m = _MAX_IN_HEADER.search(header or "")
    if m:
        return f"{value}/{next(g for g in m.groups() if g)}"
    return value


def extract_grade(path: Path, student_id: str) -> dict | None:
    if path.suffix.lower() not in SCANNABLE_SUFFIXES:
        return None
    try:
        rows = _read_rows(path)
    except Exception:
        log.exception("Could not read %s for grades", path.name)
        return None

    for idx, row in enumerate(rows):
        if student_id not in row and not any(student_id == c.replace(".0", "") for c in row):
            continue
        header = _find_header(rows, idx)
        details = []
        for col, cell in enumerate(row):
            if not cell or cell == student_id or not _is_number(cell):
                continue
            title = header[col] if col < len(header) else f"col {col + 1}"
            if _IDENTITY_HEADER.search(title) and not _TOTAL_HEADER.search(title):
                continue
            if col == 0 and not header:
                continue  # a bare leading index
            details.append((title, cell))

        if not details:
            # e.g. a plain text line: report the raw row.
            raw = " | ".join(c for c in row if c)
            return {"summary": raw[:80], "details": []}

        totals = [(h, v) for h, v in details if _TOTAL_HEADER.search(h)]
        h, v = (totals[-1] if totals else details[-1])
        return {"summary": _with_max(h, v),
                "details": [(h2, _with_max(h2, v2)) for h2, v2 in details]}
    return None


def scan_paths(path: Path, student_id: str) -> dict | None:
    """Scan a file, or every scannable file inside an extracted folder."""
    if path.is_file():
        return extract_grade(path, student_id)
    for p in sorted(path.rglob("*")):
        if p.is_file():
            found = extract_grade(p, student_id)
            if found:
                return found
    return None
