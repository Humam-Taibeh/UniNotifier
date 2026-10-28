"""Offline checks of the pure-logic modules (no portal, no Telegram).

Run:  .venv\\Scripts\\python tests\\test_architecture.py
"""

import importlib
import json
import logging
import sys
import os
import re
import shutil
import tempfile
import time
from datetime import date, datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

# The whole run lives in a throw-away home, set up BEFORE anything of the project is imported:
# a made-up .env (your real one is never read, so no real token, password or id is even in
# memory) and an empty data folder (your state is never read or written). Both are passed on
# in the environment, so every child process a test starts is isolated the same way.
_HOME = Path(tempfile.mkdtemp(prefix="uninotifier-test-home-"))
(_HOME / ".env").write_text(
    "TELEGRAM_BOT_TOKEN=123456:test-token\nTELEGRAM_CHAT_ID=42\nADMIN_USER_ID=42\n"
    "ZUJ_USERNAME=tester@example.invalid\nZUJ_PASSWORD=not-a-real-password\n"
    f"ARCHIVE_ROOT={_HOME / 'archive'}\n", encoding="utf-8")       # never the real Documents/University
for _key in ("TELEGRAM_BOT_TOKEN", "TELEGRAM_CHAT_ID", "ADMIN_USER_ID", "ZUJ_USERNAME", "ZUJ_PASSWORD",
             "PORTAL_USERNAME", "PORTAL_PASSWORD", "STUDENT_ID", "ARCHIVE_ROOT", "HEALTHCHECK_URL",
             "BOT_HEALTHCHECK_URL", "BOT_POLL_TIMEOUT", "UNINOTIFIER_MANUAL", "CAMBRIDGE_ASSIGNMENTS_URL",
             "CAMBRIDGE_UNTIL", "CAMBRIDGE_EMAIL", "CAMBRIDGE_PASSWORD"):
    os.environ.pop(_key, None)                      # nothing inherited from a shell that loaded the real one
_FIXTURES = Path(__file__).resolve().parent / "fixtures"
os.environ["UNINOTIFIER_ENV"] = str(_HOME / ".env")
os.environ["UNINOTIFIER_DATA"] = str(_HOME / "data")
# your timetable (schedule.json, assets/) is git-ignored and yours alone: the tests read a fixed copy
# of a real one and stand-in timetable files, so a fresh clone passes exactly like this PC
(_HOME / "schedule.json").write_bytes((_FIXTURES / "schedule.json").read_bytes())
os.environ["UNINOTIFIER_SCHEDULE"] = str(_HOME / "schedule.json")
(_HOME / "assets").mkdir()
(_HOME / "assets" / "Weekly_Academic_Schedule.pdf").write_bytes(b"%PDF-1.4\n% stand-in timetable\n%%EOF\n")
(_HOME / "assets" / "Weekly_Schedule.png").write_bytes(bytes.fromhex(        # a 1x1 PNG
    "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c489"
    "0000000d4944415478da63f8ffff3f0005fe02fea7d6a4ab0000000049454e44ae426082"))
os.environ["UNINOTIFIER_ASSETS"] = str(_HOME / "assets")
# what the screens show comes from a fixed copy of the tracker state (tests/fixtures), not from yours
(_HOME / "data").mkdir()
(_HOME / "data" / "saved_items.json").write_bytes((_FIXTURES / "saved_items.json").read_bytes())

from uninotifier import paths  # noqa: E402

assert paths.ENV_FILE == _HOME / ".env" and paths.DATA == _HOME / "data" and paths.SCHEDULE_FILE == _HOME / "schedule.json" \
    and paths.ASSETS == _HOME / "assets", "isolation must be set before any import"

# No test may reach the outside world or start a browser by accident. Both guards are installed
# before the project is imported (scan.py binds sync_playwright at import). A blocked attempt is
# refused (to the code it looks like no network) AND recorded: the code's own catch-alls may
# swallow the refusal, so the runner checks GUARD_HITS after every test and fails the run.
# Loopback stays open for the tests' own local servers; test_mfa_flow drives a local page with
# _REAL_SYNC_PLAYWRIGHT on purpose.
import socket  # noqa: E402

import playwright.sync_api as _playwright_api  # noqa: E402

GUARD_HITS: list[str] = []
_LOOPBACK = {"127.0.0.1", "::1", "localhost"}
_REAL_CONNECT, _REAL_CONNECT_EX, _REAL_GETADDRINFO = socket.socket.connect, socket.socket.connect_ex, socket.getaddrinfo
_REAL_SYNC_PLAYWRIGHT = _playwright_api.sync_playwright


def _outside(address) -> str | None:
    host = address[0] if isinstance(address, tuple) else None
    return str(host) if host is not None and str(host) not in _LOOPBACK else None


def _guarded_connect(self, address):
    host = _outside(address)
    if host:
        GUARD_HITS.append(f"socket connect to {host}")
        raise ConnectionRefusedError(f"test guard: no network in tests ({host})")
    return _REAL_CONNECT(self, address)


def _guarded_connect_ex(self, address):
    host = _outside(address)
    if host:
        GUARD_HITS.append(f"socket connect to {host}")
        raise ConnectionRefusedError(f"test guard: no network in tests ({host})")
    return _REAL_CONNECT_EX(self, address)


def _guarded_getaddrinfo(host, *args, **kwargs):
    if host is not None and str(host) not in _LOOPBACK:
        GUARD_HITS.append(f"DNS lookup of {host}")
        raise socket.gaierror(socket.EAI_NONAME, f"test guard: no network in tests ({host})")
    return _REAL_GETADDRINFO(host, *args, **kwargs)


def _guarded_sync_playwright(*args, **kwargs):
    GUARD_HITS.append("sync_playwright() — a browser was started")
    raise RuntimeError("test guard: no browser in tests (use _REAL_SYNC_PLAYWRIGHT on purpose)")


socket.socket.connect, socket.socket.connect_ex = _guarded_connect, _guarded_connect_ex
socket.getaddrinfo = _guarded_getaddrinfo
_playwright_api.sync_playwright = _guarded_sync_playwright

from uninotifier import archive  # noqa: E402
from uninotifier import grades
from uninotifier import media
from uninotifier import moodle
from uninotifier import notifier
from uninotifier import schedule

SID = "209900001"
# Nothing a test logs may reach the real monitor.log / bot.log: the root logger is set up
# here first, so monitor.py's and bot.py's own logging setup (basicConfig) does nothing.
logging.basicConfig(handlers=[logging.NullHandler()], force=True)
logging.disable(logging.INFO)

from uninotifier import daily  # noqa: E402
from uninotifier import finals  # noqa: E402
from uninotifier import gpa  # noqa: E402
from uninotifier import lecturelinks  # noqa: E402
from uninotifier import lectures  # noqa: E402
from uninotifier import reminders  # noqa: E402

# Every state file a test could write — reminders, GPA, Finals Watch, lecture links,
# delivery claims — lives in a temporary folder for the whole run, never next to the
# real ones; a test that forgets to restore a path can't reach your data either.
_REMINDER_TMP = Path(tempfile.mkdtemp(prefix="uninotifier-test-"))
reminders.STATE_FILE = _REMINDER_TMP / "reminders.json"
reminders.SOURCES_FILE = _REMINDER_TMP / "reminder_sources.json"
reminders.CLAIM_DIR = _REMINDER_TMP / "reminders_sent"
gpa.STATE_FILE = _REMINDER_TMP / "gpa.json"
finals.STATE_FILE = _REMINDER_TMP / "finals.json"
lecturelinks.LINKS_FILE = _REMINDER_TMP / "lecture_links.json"
lectures.CLAIM_DIR = _REMINDER_TMP / "lecture_sent"
lectures.STATE_FILE = _REMINDER_TMP / "saved_items.json"
daily.CLAIM_DIR = _REMINDER_TMP / "daily_sent"

# The fail-loud layer must never reach Telegram or the real files from a test: its messages
# land in OUTBOX, its hold-back record, the bot's pulse and the schedule's last good copy
# live in the temp folder.
from uninotifier import access  # noqa: E402
from uninotifier import failsafe  # noqa: E402
from uninotifier import telegram  # noqa: E402
from uninotifier import heartbeat  # noqa: E402

failsafe.STATE_FILE = _REMINDER_TMP / "failsafe.json"
failsafe._credentials = lambda: ("123456:test-token", "42")
OUTBOX: list[str] = []
failsafe.send = lambda token, chat, text, *a, **k: OUTBOX.append(text) or True
heartbeat.FILE = _REMINDER_TMP / "bot_heartbeat.json"
telegram.STATS_FILE = _REMINDER_TMP / "telegram_stats.json"
from uninotifier import fileids  # noqa: E402

fileids.STORE = _REMINDER_TMP / "file_ids.json"
telegram.MIN_SEND_INTERVAL = 0


def _no_network(url, **kwargs):
    raise AssertionError(f"a test reached Telegram: {url.split('/bot')[0]} (patch telegram.post)")


_REAL_POST = telegram.post
telegram.post = _no_network          # no test may talk to the real Telegram
schedule.LAST_GOOD_FILE = _REMINDER_TMP / "schedule.last_good.json"


class _Package:
    """One handle on a package that is split into submodules (uninotifier.monitor,
    uninotifier.bot): reading or overriding a name reaches the submodule that defines
    it, so a test can still say `monitor.DATA_FILE = …` or `bot.send_view = …` and the
    code that uses the name sees the change. An imported name (is_authenticated) is
    overridden in every part that imported it; a name nobody has is an error."""

    def __init__(self, package: str):
        import ast
        import importlib
        import pkgutil
        pkg = importlib.import_module(package)
        object.__setattr__(self, "_modules", [importlib.import_module(f"{package}.{m.name}")
                                              for m in pkgutil.iter_modules(pkg.__path__)])
        homes = {}
        for module in self._modules:
            tree = ast.parse(Path(module.__file__).read_text(encoding="utf-8"))
            for node in tree.body:
                names = ([node.name] if isinstance(node, (ast.FunctionDef, ast.ClassDef)) else
                         [t.id for t in getattr(node, "targets", []) if isinstance(t, ast.Name)] +
                         ([node.target.id] if isinstance(node, ast.AnnAssign) else []))
                for name in names:
                    homes.setdefault(name, module)
        object.__setattr__(self, "_homes", homes)

    def _home(self, name):
        if name in self._homes:
            return self._homes[name]
        for module in self._modules:                  # an imported name (a module such as daily)
            if hasattr(module, name):
                return module
        raise AttributeError(name)

    def __getattr__(self, name):
        return getattr(self._home(name), name)

    def __setattr__(self, name, value):
        if name in self._homes:
            setattr(self._homes[name], name, value)
            return
        users = [m for m in self._modules if hasattr(m, name)]   # an imported name: every part that uses it
        if not users:
            raise AttributeError(f"nothing in the package defines {name}")
        for module in users:
            setattr(module, name, value)


_PACKAGES: dict = {}


def _package(name: str) -> _Package:
    if name not in _PACKAGES:
        _PACKAGES[name] = _Package(name)
    return _PACKAGES[name]


def test_grades():
    import docx
    import openpyxl

    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Course marks"])
        ws.append(["No.", "Student ID", "Name", "Quiz (10)", "Mid (30)", "Total /40"])
        ws.append([1, 202300001, "Someone", 5, 20, 25])
        ws.append([2, 209900001.0, "Me", 9, 27, 36])
        wb.save(tmp / "marks.xlsx")
        g = grades.extract_grade(tmp / "marks.xlsx", SID)
        assert g and g["summary"] == "36/40", g
        assert ("Quiz (10)", "9/10") in g["details"] and ("Mid (30)", "27/30") in g["details"], g

        d = docx.Document()
        t = d.add_table(rows=2, cols=3)
        for c, v in zip(t.rows[0].cells, ["ID", "Name", "Mark (20)"]):
            c.text = v
        for c, v in zip(t.rows[1].cells, [SID, "Me", "17"]):
            c.text = v
        d.save(tmp / "mid.docx")
        g = grades.extract_grade(tmp / "mid.docx", SID)
        assert g and g["summary"] == "17/20", g

        (tmp / "list.csv").write_text(f"id,name,grade\n{SID},Me,88\n", encoding="utf-8")
        g = grades.extract_grade(tmp / "list.csv", SID)
        assert g and g["summary"] == "88", g

        (tmp / "other.csv").write_text("id,name,grade\n1,X,5\n", encoding="utf-8")
        assert grades.extract_grade(tmp / "other.csv", SID) is None
        assert grades.extract_grade(tmp / "slides.pdf", SID) is None


def test_schedule():
    sch = schedule.load_schedule()
    # Sunday 2026-10-11
    # in-person days: no morning summary, no lecture reminders
    assert schedule.due_messages(sch, datetime(2026, 10, 11, 7, 30)) == []
    assert schedule.due_messages(sch, datetime(2026, 10, 11, 8, 40)) == []
    assert schedule.due_messages(sch, datetime(2026, 10, 11, 9, 45)) == []
    # Thursday: no classes
    assert schedule.due_messages(sch, datetime(2026, 10, 15, 8, 0)) == []
    # Monday is an activity, not a lecture: no alert
    assert schedule.due_messages(sch, datetime(2026, 10, 12, 18, 15)) == []

    # Wednesday's online lecture (7:00 PM): a text-only heads-up at 6:00, the alert at 7:00
    def wed(h, m):
        return schedule.due_messages(sch, datetime(2026, 10, 14, h, m))
    assert wed(17, 59) == []
    (key, kind, text, course), = wed(18, 0)
    assert kind == "pre" and course == "Arabic Communication Skills" and "pre-19:00" in key
    assert text.startswith("⏰ *Reminder:* Arabic Communication Skills starts in 1 hour"), text
    assert "🕘 7:00–8:30 PM" in text and "Join" not in text
    assert "starts in 45 min" in wed(18, 15)[0][2] and "starts in 1 min" in wed(18, 59)[0][2]
    (key2, kind2, text2, course2), = wed(19, 0)
    assert kind2 == "start" and key2 != key and text2.startswith("🎥 *Lecture starting now:* Arabic Communication Skills")
    assert "Lecture in progress" in wed(19, 12)[0][2] and "started 12 min ago" in wed(19, 12)[0][2]
    assert wed(20, 30) == [] and wed(23, 0) == []
    assert wed(18, 15)[0][0] == wed(18, 44)[0][0], "one key per lecture and day: sent once"


def test_media():
    assert media.classify({"type": "resource"}) == "file"
    assert media.classify({"type": "url"}) == "link"
    assert media.classify({"type": "assign"}) == "activity"
    assert media.classify({"type": "h5pactivity"}) == "unknown"
    assert media.is_video_url("https://www.youtube.com/watch?v=abc")
    assert not media.is_video_url("https://example.com/page")
    item = {"name": "Lecture_1", "course": "Big Data", "url": "https://x/mod/url/view.php?id=1", "type": "url"}
    assert "🎥" in media.build_link_message(item, "https://youtu.be/abc")
    assert "Unclassified Item" in media.build_unclassified_message(item)
    assert "*Type:* Video / Lecture Recording" in media.build_link_message(item, "https://youtu.be/abc")
    assert "*Type:* Assignment" in media.build_item_message({**item, "type": "assign"}, None)
    assert "File Updated" in media.build_update_message({**item, "type": "resource"}, 1.0, 2.0)
    upd = media.build_update_message({**item, "type": "resource"}, 1_700_000_000, 1_700_000_000 + 3 * 86400 + 4 * 3600,
                                     old_name="Chapter 1.pdf", new_name="Chapter 1 (updated 2026-10-12).pdf",
                                     where="`Big Data/01_Lectures` (Lectures & Slides)")
    assert "*Previous file:* `Chapter 1.pdf`" in upd and "*New file:* `Chapter 1 (updated 2026-10-12).pdf`" in upd
    assert "*Saved to:* `Big Data/01_Lectures`" in upd
    assert "*Changed:* 3 days 4 hours after the previous version" in upd, upd
    assert (media.elapsed(30), media.elapsed(60), media.elapsed(7500), media.elapsed(86400)) == (
        "under a minute", "1 minute", "2 hours 5 minutes", "1 day")
    rem = media.build_reminder_message({"type": "quiz", "course": "c", "name": "q", "timestamp": 2e9}, "1h")
    assert "Quiz / Exam" in rem and "Detected" in rem
    msg = media.build_item_message(item, None, grade={"summary": "36/40", "details": []})
    assert "Marks Detected" in msg and "36/40" in msg


def test_ui():
    base = {"name": "HW 1", "course": "20261-Big Data-AllSec-3-", "url": "https://x/mod/assign/view.php?id=1"}
    heads = {"assign": "New Assignment", "resource": "New File", "folder": "New File",
             "quiz": "New Quiz / Exam", "forum": "New Announcement", "page": "New Page",
             "label": "New Course Info", "section": "New Section Description"}
    for kind, head in heads.items():
        msg = media.build_item_message({**base, "type": kind}, None)
        assert head in msg and "New Content" not in msg, msg
        assert "تم رصد عنصر جديد" not in msg, msg          # the duplicated footer is gone
        assert "*Course:* Big Data\n" in msg and "20261" not in msg and "AllSec" not in msg, msg
    assert "New Video" in media.build_link_message(base, "https://youtu.be/abc")
    assert "*Course:* Arabic Communication Skills" in media.course_line("20261-الاتصال والتواصل باللغة العربية-ش31-3-ك")
    assert media.clean_course_name("20261-Linear Algebra (1)-AllSec-3-F") == "Linear Algebra (1)"
    assert notifier.PORTAL_BUTTON_TEXT == "🌐 Open in ZUJ Portal ↗️"

    due = {"timestamp": 2e9}
    assert "Status:* ⏳ Pending" in media.build_item_message({**base, "type": "assign"}, due)
    assert "Status:* ✅ Submitted" in media.build_item_message({**base, "type": "assign"}, {**due, "submitted": True})
    assert "Status:* ✅ Submitted" in media.build_item_message({**base, "type": "assign"}, None, status=media.STATUS_SUBMITTED)


def test_submission_status():
    ok = moodle.parse_submission_status
    assert ok("Submission status Submitted for grading Grading status Not graded")
    assert ok("حالة التسليم تم التسليم للتقييم")
    assert ok("", has_submitted_cell=True)
    assert not ok("Submission status No submissions have been made yet")
    assert not ok("Submission status Draft (not submitted)")
    assert not ok("")


def test_deadline_flag_and_reminders():
    ev = [{"modulename": "assign", "eventtype": "due", "timesort": 2e9, "name": "HW",
           "action": {"url": "https://x/mod/assign/view.php?id=7"}, "course": {"fullname": "c"}}]
    saved = moodle.update_deadlines({}, ev)
    saved["assign-7"]["submitted"] = True
    assert moodle.update_deadlines(saved, ev)["assign-7"]["submitted"] is True   # survives a refresh
    monitor = _package("uninotifier.monitor")
    sent = []
    monitor.send_telegram = lambda text, url=None: sent.append(text) or True
    soon = {"assign-7": {**saved["assign-7"], "timestamp": time.time() + 1800}}
    monitor.send_deadline_reminders(soon)
    assert sent == [], "reminder sent for a submitted assignment"
    soon["assign-7"]["submitted"] = False
    monitor.send_deadline_reminders(soon)
    assert len(sent) == 2   # 24h and 1h windows both due


def test_bot_views():
    bot = _package("uninotifier.bot")
    text, rows = bot.view_menu()
    assert [[b["text"] for b in row] for row in rows] == [
        ["📚 My Courses & Schedule", "⏳ Upcoming Tasks & Reminders"],
        ["📊 GPA & Grade Hub", "📜 Academic Record & Balance"],
        ["📖 System Guide"]], rows
    assert [b.get("callback_data") for row in rows for b in row] == [
        "courses", "tasks", "academic", "record", "help"], "no external site buttons"
    assert len(text.splitlines()) == 2 and text.startswith("🎓 *UniNotifier System* • First Semester 2026/2027")
    assert text.splitlines()[1].endswith("• 6 Courses Tracked"), text
    # My Courses & Schedule is a menu, not a wall of text; Tracked Items and Campus Guide are gone
    text, rows = bot.view_courses()
    assert [[b["text"] for b in row] for row in rows] == [
        ["📋 Course List & Details"], ["📅 Weekly Timetable"], ["🔙 Main Menu"]]
    assert [rows[0][0]["callback_data"], rows[1][0]["callback_data"]] == ["courselist", "schedule"]
    assert "6 courses · 18 credit hours" in text and len(text.splitlines()) <= 5
    everywhere = {b.get("callback_data") for view in bot.VIEWS.values() for row in view()[1] for b in row}
    assert not {"items", "campus"} & everywhere and not {"items", "campus"} & {c for c, _ in bot.COMMANDS}
    assert not hasattr(bot, "view_items") and not hasattr(bot, "view_campus")
    # Lecture Links is gone: the link rides on the lecture reminder itself
    assert not hasattr(bot, "view_links") and not hasattr(bot, "handle_setlink")
    assert "links" not in everywhere and "setlink" not in {c for c, _ in bot.COMMANDS}
    # no scholarship / discount wording anywhere, and no "How It Works" in the GPA hub
    from uninotifier import gpa_ui
    names = bot.course_names()
    screens = [view()[0] for view in bot.VIEWS.values()] + [
        gpa_ui.view_need_picker(names)[0], gpa_ui.view_targets("c")[0], gpa_ui.view_targets("s")[0],
        gpa_ui.view_home(names)[0], gpa_ui.view_plan(names, "s", 60)[0], gpa_ui.view_plan(names, "c", 75)[0]]
    for screen in screens:
        assert not any(w in screen.lower() for w in ("scholarship", "discount", "standing:")), screen
    assert not [b for row in bot.view_academic()[1] for b in row if b.get("callback_data", "").startswith("help")]
    moved = {b.get("callback_data") for view in (bot.view_academic, bot.view_record) for row in view()[1] for b in row}
    assert {"gpa:need", "finals:status", "portal:transcript", "portal:balance"} <= moved
    assert not {"gpa:home", "portal:grades"} & moved, "Course Adjustments and This Term's Grades left the GPA hub"
    # no Quick Check anywhere: not in a view, not in the / menu
    assert "check" not in everywhere | {b.get("callback_data") for row in bot.view_tasks()[1] for b in row}
    assert "check" not in {c for c, _ in bot.COMMANDS}
    courses = bot.view_course_list()[0]
    assert "20261" not in courses and "Section 01" in courses and "Dr. Sam Taylor" in courses
    names = [l for l in courses.splitlines() if l[:1].isdigit()]
    assert len(names) == 6 and not any("؀" <= ch <= "ۿ" for l in names for ch in l), names
    assert bot.SCHEDULE_PDF.is_file() and bot.SCHEDULE_IMAGE.is_file()
    assert "🕘 Sun · Tue · 11:00 AM–12:30 PM" in courses and "🕘 Sun · Tue · 3:30–5:00 PM" in courses, courses
    assert "14:00" not in courses and "Room 449" in courses
    assert "Upcoming Tasks & Reminders" in bot.view_tasks()[0]
    assert notifier.clock12(0, 5) == "12:05 AM" and notifier.clock12(12, 0) == "12:00 PM"
    assert notifier.format_timestamp(1.7e9).endswith(("AM", "PM"))
    assert bot.remaining_text(time.time() + 90000).startswith("1d")


COURSE_PAGE = """<html><body><ul>
<li data-for="section" data-id="10" data-number="0"><div data-for="sectioninfo">
  <div class="section_availability"></div></div>
  <ul data-for="cmlist">
  <li class="activity label modtype_label" id="module-19793" data-for="cmitem" data-id="19793">
    <button data-action="toggle-manual-completion" data-toggletype="manual:mark-done" data-cmid="19793">Mark as done</button>
    <div class="activity-altcontent text-break"><div class="no-overflow">
      <h3><img src="https://x/pluginfile.php/1/mod_label/intro/banner.jpg?time=5"></h3>
      <p>Arabic Communication</p><p>Lecture time :&nbsp;</p><p>Office hours :</p></div></div></li>
  <li class="activity forum modtype_forum" id="module-20770" data-for="cmitem" data-id="20770">
    <button data-action="toggle-manual-completion" data-toggletype="manual:undo" data-cmid="20770">Done</button></li>
  </ul></li>
<li data-for="section" data-id="11" data-number="1"><div data-for="sectioninfo">
  <div class="summarytext"><p>Join the lecture: <a href="https://teams.microsoft.com/l/meetup-join/abc">Teams</a></p></div></div>
  <ul data-for="cmlist">
  <li class="activity page modtype_page" id="module-21747" data-for="cmitem" data-id="21747">
    <button data-action="toggle-manual-completion" data-toggletype="manual:mark-done" data-cmid="21747">Mark as done</button>
    <div class="activity-altcontent text-break activity-description"><div class="no-overflow">
      <h4>Overview</h4><p>Aims to improve skills.</p><h4>Grades</h4>
      <ol><li>30% activities</li><li>30% midterm</li><li>40% final</li></ol></div></div></li>
  </ul></li>
<li data-for="section" data-id="12" data-number="2"><div data-for="sectioninfo"><p>&nbsp;.</p></div></li>
</ul></body></html>"""


def test_richtext_and_course_page():
    from uninotifier import richtext
    page = moodle.parse_course_page(COURSE_PAGE)
    assert page["manual"] == {"19793": False, "20770": True, "21747": False}
    assert set(page["bodies"]) == {"19793", "21747"}
    assert list(page["summaries"]) == ["11"]   # empty / "." summaries are ignored
    desc = richtext.to_text(page["bodies"]["21747"])
    assert desc.startswith("*Overview*") and "*Grades*" in desc, desc
    assert "1. 30% activities\n2. 30% midterm\n3. 40% final" in desc, desc
    label = page["bodies"]["19793"]
    assert richtext.first_line(label) == "Arabic Communication"
    assert richtext.image_paths(label) == ["/pluginfile.php/1/mod_label/intro/banner.jpg"]
    assert "Lecture time :" in richtext.to_text(label)
    summary = richtext.to_text(page["summaries"]["11"])
    assert "[Teams](https://teams.microsoft.com/l/meetup-join/abc)" in summary, summary
    assert richtext.links(page["summaries"]["11"]) == ["https://teams.microsoft.com/l/meetup-join/abc"]
    # fingerprints: markup / cache-buster churn is not an edit, a text change is
    assert richtext.digest(label) == richtext.digest(label.replace("?time=5", "?time=9").replace("<p>", "<p class='x'>"))
    assert richtext.digest(label) != richtext.digest(label.replace("Office hours :", "Office hours : Sun 10-12"))
    assert not richtext.is_meaningful("<p>.</p>") and richtext.is_meaningful(label)
    assert richtext.to_text("<p>a_b *c*</p>") == "a\\_b \\*c\\*"


def test_display_names_and_completion_policy():
    item = {"name": ".", "type": "page", "section": "مقدمة المادة", "cmid": "1"}
    assert moodle._display_name(item, "") == "مقدمة المادة"
    assert moodle._display_name({**item, "type": "label", "name": "x ..."}, "<p>Office hours</p>") == "Office hours"
    assert moodle._display_name({**item, "name": "Week 1"}, "") == "Week 1"
    assert media.completion_rows({"type": "assign", "manual": True, "cmid": "5"}) == []
    assert media.completion_rows({"type": "forum", "manual": False, "cmid": "5"}) == []
    assert media.completion_rows({"type": "forum", "manual": True, "cmid": "5"}) == [
        [{"text": "✅ Mark as Done", "callback_data": "done:5:1"}]]
    assert media.completion_rows({"type": "page", "manual": True, "done": True, "cmid": "5"}) == [
        [{"text": "↩️ Undo Done", "callback_data": "done:5:0"}]]
    assert "quiz" in media.NEVER_DONE_TYPES and "assign" not in media.AUTO_DONE_TYPES
    msg = media.build_item_message({"name": "n", "type": "page", "course": "c", "url": "u",
                                    "manual": True, "done": True}, None)
    assert "Progress:* ✅ Marked as done" in msg


def test_forum_messages_and_state():
    forum = {"name": "Announcements", "course": "20261-Big Data-AllSec-3-", "type": "forum"}
    post = {"id": "9", "subject": "Midterm moved", "author": "Dr. Hani", "time": 2e9,
            "is_reply": False, "attachments": [("a.pdf", "u")], "html": "<p>Sunday</p>"}
    msg = media.build_forum_post_message(forum, post)
    assert "New Announcement" in msg and "Midterm moved" in msg and "Attachments:* 1" in msg
    assert "New Forum Reply" in media.build_forum_post_message(forum, {**post, "is_reply": True})
    assert "New Forum Topic" in media.build_forum_post_message({**forum, "name": "Q&A"}, post)
    monitor = _package("uninotifier.monitor")
    p = lambda i: {"id": str(i)}
    saved = {"forums": {"5-1": ["1", "2"]}, "content": {"5-1": {}, "5-2": {}}}
    fresh = monitor.new_forum_posts(saved, {"5-1": [p(1), p(2), p(3)], "5-2": [p(4)],
                                            "5-3": [p(5)], "5-4": None}, notify_new=True)
    # known forum -> only the unseen post; known-but-never-read forum -> baseline;
    # brand-new forum -> everything; unreadable -> nothing
    assert fresh == {"5-1": [p(3)], "5-3": [p(5)]}, fresh


def test_edit_detection():
    monitor = _package("uninotifier.monitor")
    saved = {"content": {"a": {"hash": "1"}, "b": {"hash": "2"}, "c": {}}}
    content = {"a": {"hash": "1"}, "b": {"hash": "3"}, "c": {"hash": "4"}, "d": {"hash": "5"}}
    assert monitor.edited_items(saved, content) == ["b"]
    assert monitor.public({"name": "x", "_body": "<p>", "hash": "1"}) == {"name": "x", "hash": "1"}


def test_failure_policy():
    monitor = _package("uninotifier.monitor")
    orig = monitor.send_telegram, monitor.save_data      # restored below: other tests rely on the real ones
    try:
        _failure_policy_checks(monitor)
    finally:
        monitor.send_telegram, monitor.save_data = orig


def _failure_policy_checks(monitor):
    sent, saved = [], {"alerts": {}, "failures": {}}
    monitor.send_telegram = lambda text, url=None: sent.append(text) or True
    monitor.save_data = lambda data: None
    for _ in range(2):
        monitor.record_failure(saved, "network", "down")
    assert sent == [], "a transient network error must not alert"
    monitor.record_failure(saved, "network", "down")
    assert len(sent) == 1 and "3 consecutive checks" in sent[0]
    monitor.record_failure(saved, "network", "down")
    assert len(sent) == 1, "cooldown"
    monitor.clear_failures(saved)
    assert "Back to normal" in sent[-1] and saved["failures"] == {}
    monitor.record_failure(saved, "mfa", "renew")      # needs the user: immediate
    assert sent[-1] == "renew"
    assert monitor.failure_kind(Exception("Page.goto: net::ERR_NAME_NOT_RESOLVED")) == "network"
    assert monitor.failure_kind(KeyError("data")) == "unexpected"


def test_notifier_helpers():
    token = "1234567890:FAKEfakeFAKEfakeFAKEfakeFAKEfake123"
    line = f"HTTPSConnectionPool url: /bot{token}/sendMessage"
    assert token not in notifier.redact(line) and "/bot<redacted>/sendMessage" in notifier.redact(line)
    chunks = notifier.split_text("\n".join(f"line {n} " + "x" * 50 for n in range(300)), 1000)
    assert all(len(c) <= 1000 for c in chunks) and len(chunks) > 10
    assert "\n".join(chunks).count("line ") == 300
    assert notifier.split_text("y" * 2500, 1000) == ["y" * 1000, "y" * 1000, "y" * 500]
    kb = notifier._keyboard("https://p", [[{"text": "✅ Mark as Done", "callback_data": "done:1:1"}]])
    assert '"done:1:1"' in kb and kb.index("done:1:1") < kb.index("https://p")


FAKE_MICROSOFT = """<html><body><div id="lightbox"></div><script>
const screens = {
  number: `<div id="idDiv_SAOTCAS_Title">Approve sign in request</div>
           <div id="idRichContext_DisplaySign">47</div>
           <input type="checkbox" id="idChkBx_SAOTCAS_TD">
           <a id="signInAnotherWay" href="#" onclick="show('proofs')">I can't use my Microsoft Authenticator app right now</a>`,
  proofs: `<div id="idDiv_SAOTCS_Proofs">
           <div role="button" data-value="PhoneAppNotification" onclick="show('number')">Approve a request</div>
           <div role="button" data-value="OneWaySMS" onclick="show('code')">Text +XXX XXXXXXX12</div></div>`,
  code: `<div id="idDiv_SAOTCC_Description">We texted your phone +XXX XXXXXXX12.</div>
         <input id="idTxtBx_SAOTCC_OTC" name="otc"><span id="idSpan_SAOTCC_Error_OTC" style="display:none">Wrong</span>
         <input type="button" id="idSubmit_SAOTCC_Continue" value="Verify"
                onclick="if (document.getElementById('idTxtBx_SAOTCC_OTC').value === '123456') show('done');
                         else document.getElementById('idSpan_SAOTCC_Error_OTC').style.display = 'inline'">`,
  timeout: `<div id="idDiv_SAASTO_Title">We didn't hear from you</div>
            <a id="idA_SAASTO_Resend" href="#" onclick="show('number')">Send another request</a>`,
  denied: `<div id="idDiv_SAASDS_Title">Request denied</div>`,
  done: `<div id="finished">ok</div>`,
};
function show(name) { document.getElementById('lightbox').innerHTML = screens[name]; return false; }
</script></body></html>"""


def test_mfa_flow():
    from uninotifier import mfa
    from uninotifier.zuj_auth import MFARequiredError
    sync_playwright = _REAL_SYNC_PLAYWRIGHT            # a local page, on purpose: the guard is for accidents

    sent = []
    with tempfile.TemporaryDirectory() as tmp:
        mfa.MFA_REQUEST, mfa.MFA_RESPONSE = Path(tmp) / "req.json", Path(tmp) / "resp.json"
        notifier_upload, notifier_call, notifier_send = notifier.upload, notifier.call_api, notifier.send_message
        notifier.upload = lambda token, method, data, *a: sent.append(("photo", data["caption"], data["reply_markup"])) or {"message_id": len(sent)}
        notifier.call_api = lambda token, method, data=None, timeout=20: sent.append((method, data)) or {"message_id": len(sent)}
        notifier.send_message = lambda token, chat, text, url=None, rows=None: sent.append(("text", text, rows)) or True
        try:
            with sync_playwright() as pw:
                browser = pw.chromium.launch(headless=True)
                page = browser.new_page()
                page.set_content(FAKE_MICROSOFT)

                # 1. number matching -> the number goes to Telegram, once
                page.evaluate("show('number')")
                handler = mfa.TelegramMFA("t", "c")
                handler.step(page)
                handler.step(page)
                photos = [s for s in sent if s[0] == "photo"]
                assert len(photos) == 1, sent
                assert "4️⃣7️⃣" in photos[0][1] and "(*47*)" in photos[0][1], photos
                assert mfa.SMS_CALLBACK in photos[0][2]
                assert page.is_checked("#idChkBx_SAOTCAS_TD"), "'don't ask again' should be ticked"
                assert mfa.read_request()["expects"] == "choice"

                # 2. user taps "Text me a code instead" -> another way -> SMS -> code screen
                assert mfa.answer_request(choice="sms")
                handler.step(page)                       # clicks "I can't use my Authenticator"
                assert mfa.detect(page) == "proofs"
                handler.step(page)                       # chooses OneWaySMS
                assert mfa.detect(page) == "code"
                handler.step(page)                       # asks for the code on Telegram
                assert any("Verification Code Needed" in s[1] and "XXXXXXX12" in s[1]
                           for s in sent if s[0] == "text"), sent
                assert mfa.read_request()["expects"] == "code"

                # 3. a wrong code is reported, the right one gets through
                mfa.answer_request(code="000000")
                handler.step(page)
                handler.step(page)
                assert any("rejected that code" in s[1] for s in sent if s[0] == "text"), sent
                mfa.answer_request(code="123456")
                handler.step(page)
                assert page.locator("#finished").count() == 1
                handler.succeeded()
                assert "Signed in" in sent[-1][1] and not mfa.MFA_REQUEST.exists()

                # 4. timeout: one fresh request, then give up (already notified)
                handler = mfa.TelegramMFA("t", "c")
                page.evaluate("show('timeout')")
                handler.step(page)
                assert mfa.detect(page) == "number" and "fresh request" in sent[-1][1]
                page.evaluate("show('timeout')")
                try:
                    handler.step(page)
                    raise AssertionError("second timeout should give up")
                except MFARequiredError as e:
                    assert e.notified and "No response" in sent[-1][1]

                # 5. denied, and 6. nobody answers before the deadline
                for screen, expected in (("denied", "denied"), ("number", "expired")):
                    handler = mfa.TelegramMFA("t", "c", wait_seconds=0 if screen == "number" else 300)
                    page.evaluate(f"show('{screen}')")
                    if screen == "number":
                        handler.started, handler.deadline = time.time(), time.monotonic() - 1
                    try:
                        handler.step(page)
                        raise AssertionError(f"{screen} should end the attempt")
                    except MFARequiredError as e:
                        assert e.notified and expected in sent[-1][1].lower(), sent[-1]
                        buttons = [b["callback_data"] for row in sent[-1][2] for b in row]
                        assert buttons == ["check"], sent[-1]   # "🔄 Try again now"
                browser.close()
        finally:
            notifier.upload, notifier.call_api, notifier.send_message = notifier_upload, notifier_call, notifier_send
    assert mfa.keycaps("05") == "0️⃣5️⃣"


def test_mfa_policy_and_bot_relay():
    bot = _package("uninotifier.bot")
    from uninotifier import mfa
    monitor = _package("uninotifier.monitor")
    saved = {"mfa": {}}
    monitor.MANUAL_RUN = False
    assert monitor.mfa_allowed(saved, datetime(2026, 10, 12, 14, 0))
    assert not monitor.mfa_allowed(saved, datetime(2026, 10, 12, 3, 0)), "quiet hours"
    saved["mfa"]["last_attempt"] = datetime(2026, 10, 12, 13, 30).timestamp()
    assert not monitor.mfa_allowed(saved, datetime(2026, 10, 12, 14, 0)), "hourly retry"
    assert monitor.mfa_allowed(saved, datetime(2026, 10, 12, 14, 31))
    monitor.MANUAL_RUN = True
    assert monitor.mfa_allowed(saved, datetime(2026, 10, 12, 3, 0)), "/check always may"
    monitor.MANUAL_RUN = False

    replies = []
    bot.send_view = lambda text, rows: replies.append(text)
    with tempfile.TemporaryDirectory() as tmp:
        mfa.MFA_REQUEST, mfa.MFA_RESPONSE = Path(tmp) / "req.json", Path(tmp) / "resp.json"
        assert not bot.handle_mfa_code("123456"), "no pending request: not a code"
        mfa._write_json(mfa.MFA_REQUEST, {"id": "r1", "created": time.time(), "expects": "code"})
        assert not bot.handle_mfa_code("hello")
        assert bot.handle_mfa_code("123 456")
        assert json.loads(mfa.MFA_RESPONSE.read_text()) == {"id": "r1", "code": "123456"}
        assert "Code received" in replies[-1]
        mfa._write_json(mfa.MFA_REQUEST, {"id": "r2", "created": time.time() - 3600, "expects": "code"})
        assert not bot.handle_mfa_code("654321"), "a stale request is ignored"


def test_lecture_links():
    bot = _package("uninotifier.bot")
    from uninotifier import lecturelinks as ll
    monitor = _package("uninotifier.monitor")

    teams = "https://teams.microsoft.com/l/meetup-join/19%3ameeting_abc%40thread.v2/0?context=%7b%7d"
    zoom = "https://us06web.zoom.us/j/8123456789?pwd=xyz"
    assert ll.meeting_urls([f"Join: {teams}.", zoom, "https://example.com/x"]) == [teams, zoom]
    assert ll.meeting_urls(["https://meet.google.com/abc-defg-hij"]) == ["https://meet.google.com/abc-defg-hij"]
    assert ll.meeting_urls(["https://elearning.zuj.edu.jo/mod/page/view.php?id=1"]) == []
    assert ll.online_courses() == ["Arabic Communication Skills"]
    assert ll.online_slots("Arabic Communication Skills") == "Wed 7:00 PM–8:30 PM"

    arabic = "20261-الاتصال والتواصل باللغة العربية-ش31-3-ك"
    with tempfile.TemporaryDirectory() as tmp:
        ll.LINKS_FILE = Path(tmp) / "links.json"
        label = {"name": "Instructor info", "course": arabic, "url": "u1", "type": "label", "cmid": "1", "course_id": "8805",
                 "_body": f'<p>Lecture link: <a href="{teams}">Teams</a></p>'}
        other = {"name": "Week 1", "course": "20261-Big Data-AllSec-3-", "url": "u2", "type": "label", "cmid": "2", "course_id": "8804",
                 "_body": f"<p>{zoom}</p>"}
        content = {"8805-1": label, "8804-2": other}

        changes = ll.update_from_portal(content, {})
        assert [(c["course"], c["url"]) for c in changes] == [("Arabic Communication Skills", teams)]
        assert "Instructor info" in changes[0]["origin"]
        assert changes[0]["portal_url"] == "u1"
        assert ll.get("Big Data") is None, "in-person courses are ignored"
        assert ll.update_from_portal(content, {}) == [], "an unchanged link is not re-announced"

        # /setlink wins over the old label link; a NEW portal link wins over /setlink
        ll.set_link("Arabic Communication Skills", "https://teams.live.com/meet/manual1", "manual")
        assert ll.update_from_portal(content, {}) == []
        assert ll.get("Arabic Communication Skills").endswith("manual1")
        new_link = "https://teams.microsoft.com/l/meetup-join/19%3ameeting_NEW%40thread.v2/0"
        forum = {"name": "Announcements", "course": arabic, "url": "f", "type": "forum"}
        post = {"id": "88", "subject": "Link for this week", "time": 2e9, "url": "https://x/mod/forum/discuss.php?d=5#p88",
                "html": f'<p><a href="{new_link}">join</a></p>'}
        changes = ll.update_from_portal({**content, "8805-9": forum}, {"8805-9": [post]})
        assert changes[0]["url"] == new_link and "Link for this week" in changes[0]["origin"]
        assert changes[0]["shot"] == ("https://x/mod/forum/discuss.php?d=5", "#p88")
        assert changes[0]["portal_url"].endswith("#p88")
        assert ll.get("Arabic Communication Skills") == new_link

        # a URL activity pointing at a meeting is recognised through its target
        url_item = {"name": "Online lecture", "course": arabic, "url": "u3", "target": zoom, "type": "page", "cmid": "3", "course_id": "8805"}
        assert ll.update_from_portal({"8805-3": url_item}, {})[0]["url"] == zoom

        # a portal link comes with where it was published; the screenshot is kept per link
        ll.save_shot("Arabic Communication Skills", b"PNG-fake")
        assert ll.shot_bytes("Arabic Communication Skills") == b"PNG-fake"
        assert ll.entry("Arabic Communication Skills")["portal_url"] == "u3"
        assert ll.join_rows(ll.entry("Arabic Communication Skills")) == [
            [{"text": "🎥 Join Meeting", "url": zoom}], [{"text": "🌐 Open Course Page", "url": "u3"}]]
        ll.set_link("Arabic Communication Skills", zoom, "manual")
        assert "portal_url" not in ll.entry("Arabic Communication Skills"), "a hand-set link has no source post"
        assert ll.shot_bytes("Arabic Communication Skills") is None
        assert ll.join_rows(ll.entry("Arabic Communication Skills")) == [[{"text": "🎥 Join Meeting", "url": zoom}]]

        # a URL activity's card is shot on the course page (its own page would redirect away)
        act = {"type": "url", "cmid": "55", "course_id": "8805", "url": "https://x/mod/url/view.php?id=55"}
        assert moodle.screenshot_target(act) == (moodle.course_url("8805"), "#module-55")
        assert ll._course_anchor(act) == moodle.course_url("8805") + "#module-55"

        # re-establish a portal link with its source, as the monitor would
        ll.set_link("Arabic Communication Skills", zoom, "portal", '"Online lecture"')
        links = ll.load()
        links["Arabic Communication Skills"]["portal_url"] = "u3"
        ll._save(links)
        ll.save_shot("Arabic Communication Skills", b"PNG-fake")

        # lecture messages: text-only heads-up, then the actionable alert with buttons + screenshot
        from uninotifier import lectures
        orig = (lectures.notifier.send_message, lectures.notifier.send_png_file, lectures.CLAIM_DIR)
        lectures.CLAIM_DIR = Path(tmp) / "claims"
        sent, files = [], []
        lectures.notifier.send_message = lambda token, chat, text, url=None, rows=None: sent.append((text, rows)) or True
        lectures.notifier.send_png_file = lambda token, chat, png, name, caption="", url=None, rows=None: (
            files.append((png, caption, rows)) or True)
        try:
            assert lectures.send_due("t", "c", datetime(2026, 10, 14, 17, 59)) == 0
            assert lectures.send_due("t", "c", datetime(2026, 10, 14, 18, 0)) == 1
            assert lectures.send_due("t", "c", datetime(2026, 10, 14, 18, 0)) == 0, "claimed: not sent twice"
            assert lectures.send_due("t", "c", datetime(2026, 10, 14, 18, 20)) == 0
            assert len(sent) == 1 and sent[0][1] is None and files == []
            assert sent[0][0].startswith("⏰ *Reminder:* Arabic Communication Skills starts in 1 hour")

            assert lectures.send_due("t", "c", datetime(2026, 10, 14, 19, 0)) == 1
            assert lectures.send_due("t", "c", datetime(2026, 10, 14, 19, 5)) == 0
            assert len(sent) == 1 and len(files) == 1, "the alert is the image + caption, not a second text"
            png, caption, rows = files[0]
            assert png == b"PNG-fake" and caption.startswith("🎥 *Lecture starting now:*")
            assert "Tap *Join Meeting*" in caption and "attached" in caption
            assert rows == [[{"text": "🎥 Join Meeting", "url": zoom}],
                            [{"text": "🌐 Open Course Page", "url": "u3"}]], rows

            # no link yet: the start alert says so and has no buttons
            ll.set_link("Arabic Communication Skills", None, "manual")
            lectures.CLAIM_DIR = Path(tmp) / "claims2"
            assert lectures.send_due("t", "c", datetime(2026, 10, 14, 19, 0)) == 1
            assert "No meeting link has been posted yet" in sent[-1][0] and sent[-1][1] is None

            # a failed send releases its claim and is retried on the next tick
            lectures.CLAIM_DIR = Path(tmp) / "claims3"
            ok = {"v": False}
            lectures.notifier.send_message = lambda token, chat, text, url=None, rows=None: ok["v"]
            assert lectures.send_due("t", "c", datetime(2026, 10, 14, 18, 0)) == 0
            ok["v"] = True
            assert lectures.send_due("t", "c", datetime(2026, 10, 14, 18, 0)) == 1
        finally:
            lectures.notifier.send_message, lectures.notifier.send_png_file, lectures.CLAIM_DIR = orig
        # the monitor's fallback goes through the same claims
        calls = []
        orig_due = lectures.send_due
        lectures.send_due = lambda token, chat, now=None: calls.append((token, chat))
        try:
            monitor.send_schedule_messages({})
        finally:
            lectures.send_due = orig_due
        assert len(calls) == 1
        ll.set_link("Arabic Communication Skills", zoom, "manual")

        # /setlink is gone: links come from the portal and ride on the lecture reminder
        replies = []
        bot.send_view = lambda text, rows: replies.append((text, rows))
        bot.handle_command("/setlink https://teams.microsoft.com/l/meetup-join/19%3amine")
        assert replies[-1][0].startswith("❓ Unknown command")
        assert ll.load()["Arabic Communication Skills"]["url"] == zoom, "nothing changed"
        assert "/setlink" not in lectures.NO_LINK_NOTE


def test_portal_parsers():
    from uninotifier import portal

    heads = ["رمز المقرر", "اسم المقرر", "الحالة", "عدد الساعات المعتمدة", "العلامة", "ناجح؟", "معادة"]
    ok, no = "✓", "✗"
    raw = [
        {"legend": "الفصل الأول 2026/2027 20261", "heads": heads,
         "text": "الحالة منتظم المعدل الفصلي المعدل التراكمي 71.3",
         "rows": [["1", "Algebra", "مسجلة", "3", "", "قيد التنفيذ", no]]},
        {"legend": "الفصل الثاني 2025/2026 20252", "heads": heads,
         "text": "المعدل الفصلي 76 المعدل التراكمي 71.3",
         "rows": [["2", "Web", "مسجلة", "3", "88", ok, no], ["2", "Web", "مسجلة", "3", "", "قيد التنفيذ", no],
                  ["3", "Stats", "مسجلة", "3", "74", ok, no]]},
        {"legend": "الفصل الأول 2024/2025 20241", "heads": heads, "text": "المعدل الفصلي 58",
         "rows": [["4", "Old", "مسجلة", "3", "35", no, ok], ["5", "Ethics", "مسجلة", "2", "95", ok, no]]},
        {"legend": "الفصل الثاني 2024/2025 20242", "heads": heads, "text": "", "rows": [["لا يوجد بيانات"]]},
        {"legend": "الفصل الأول 2023/2024 20231", "heads": heads, "text": "المعدل الفصلي 70",
         "rows": [["3", "Stats", "مسجلة", "3", "60", ok, no]]},          # passed twice: counted once
    ]
    semesters = portal.parse_transcript(raw)
    assert [s["term"] for s in semesters] == [
        "First semester 2026/2027", "Second semester 2025/2026", "First semester 2024/2025",
        "Second semester 2024/2025", "First semester 2023/2024"]
    assert semesters[0]["average"] is None and semesters[0]["cumulative"] == 71.3
    assert semesters[1]["average"] == 76 and len(semesters[1]["courses"]) == 2, "the stale duplicate row is hidden"
    assert semesters[0]["courses"][0]["passed"] is None and semesters[2]["courses"][0]["passed"] is False
    assert semesters[2]["courses"][0]["repeated"] is True and semesters[3]["courses"] == []
    assert portal.passed_hours(semesters) == 3 + 3 + 2, "88 + 74 + 95 (hours 3 + 3 + 2); a failed course adds nothing"
    lines = "\n".join(portal.transcript_lines(semesters))
    assert "In Progress" in lines and "✅ *88*" in lines and "🔁 *35*" in lines and "average 76 · cumulative 71.3" in lines
    assert lines.index("In Progress") < lines.index("Second semester 2025/2026"), "the semester in progress comes first"
    assert "⏳ \u2068Algebra\u2069 · 3h" in lines, "an in-progress course shows ⏳, its name isolated from its neighbours"
    assert lines.rstrip().endswith("⏳ in progress_"), "a one-line legend closes the blocks"
    assert "Second semester 2024/2025" not in lines, "an empty semester is skipped"
    assert portal.english_term("الفصل الصيفي 2023/2024 20233") == "Summer semester 2023/2024"

    text = "الرصيد الكلي 1250 مدين"
    tables = [[["", "الفصل الدراسي", "الحالة", "عدد الساعات", "الرصيد"],
               ["", "الفصل الأول 2026/2027", "منتظم", "18", "د.أ1250"],
               ["", "الفصل الثاني 2025/2026", "منتظم", "12", "د.أ0"]]]
    info = portal.parse_balance(text, tables)
    assert info["total"] == 1250 and info["rows"][0] == ("First semester 2026/2027", 18.0, 1250.0)
    out = portal.balance_lines(info)
    assert out[0] == "💰 *Total balance:* 1,250 JD — owed to the university (debit)", out
    assert out[1] == "• First semester 2026/2027 · 18h · 1,250 JD" and "All other semesters: 0 JD" in out[-1]
    assert portal.parse_balance("الرصيد الكلي 40.5 دائن", [])["total"] == -40.5
    assert "in your favour" in portal.balance_lines({"total": -40.5, "rows": []})[0]
    assert "settled" in portal.balance_lines({"total": 0.0, "rows": []})[0]

    empty = [[["رمز المقرر", "اسم المقرر", "العلامة"], ["لا يوجد نتائج للمساقات المسجلة حاليا"]]]
    out = portal.results_lines({"term": "الفصل الأول 2026/2027", "text": "لا يوجد نتائج"}, empty)
    assert out[0] == "ℹ️ No results are published yet for First semester 2026/2027."
    full = [[["رمز المقرر", "اسم المقرر", "العلامة"], ["1", "Algebra", "80"]]]
    out = portal.results_lines({"term": "الفصل الأول 2026/2027", "text": ""}, full)
    assert out[0] == "📅 *First semester 2026/2027*" and out[1] == "*Code | Course | Mark*" and out[2] == "1 | Algebra | 80"
    assert set(portal.SECTIONS) == {"grades", "transcript", "balance"}
    assert all(s["path"].startswith("/faces/ui/pages/student/") for s in portal.SECTIONS.values())


def test_http_layer():
    import http.server
    import threading

    from uninotifier import httpctx
    from uninotifier.httpctx import HttpContext

    class Handler(http.server.BaseHTTPRequestHandler):
        def log_message(self, *args):
            pass

        def _reply(self, status=200, body=b"", headers=None):
            self.send_response(status)
            for k, v in (headers or {}).items():
                self.send_header(k, v)
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            if self.command != "HEAD":
                self.wfile.write(body)

        def do_GET(self):
            if self.path == "/my/":
                self._reply(body=b'<script>M.cfg = {"sesskey":"abc123XY"}</script>')
            elif self.path == "/redir":
                self._reply(302, headers={"Location": "/file.bin"})
            elif self.path == "/file.bin":
                self._reply(body=b"DATA", headers={"Last-Modified": "Wed, 01 Jan 2025 00:00:00 GMT"})
            elif self.path == "/cookie":
                self._reply(body=self.headers.get("Cookie", "").encode())
            elif self.path == "/missing":
                self._reply(404, b"nope")
            else:
                self._reply(404)

        do_HEAD = do_GET

        def do_POST(self):
            size = int(self.headers.get("Content-Length", 0))
            payload = self.rfile.read(size)
            self._reply(body=json.dumps({"got": payload.decode(), "ctype": self.headers.get("Content-Type")}).encode())

    server = http.server.ThreadingHTTPServer(("127.0.0.1", 0), Handler)
    threading.Thread(target=server.serve_forever, daemon=True).start()
    base = f"http://127.0.0.1:{server.server_port}"
    try:
        http_ctx = HttpContext.from_cookie_list([{"name": "MoodleSession", "value": "s1", "domain": "127.0.0.1", "path": "/"}])
        assert http_ctx.request is http_ctx
        assert http_ctx.get(f"{base}/cookie").text() == "MoodleSession=s1", "saved cookies are sent"

        resp = http_ctx.get(f"{base}/redir")                      # followed, like Playwright's default
        assert resp.ok and resp.body() == b"DATA" and resp.url.endswith("/file.bin")
        resp = http_ctx.get(f"{base}/redir", max_redirects=0)     # not followed: the Location is readable
        assert not resp.ok and resp.status == 302 and resp.headers.get("location") == "/file.bin"
        head = http_ctx.head(f"{base}/redir")
        assert head.ok and head.headers.get("content-length") == "4" and head.headers.get("last-modified")
        assert not http_ctx.get(f"{base}/missing").ok

        out = http_ctx.post(f"{base}/ajax", data='[{"a":1}]', headers={"Content-Type": "application/json"}).json()
        assert out == {"got": '[{"a":1}]', "ctype": "application/json"}

        old_base, httpctx.BASE_URL = httpctx.BASE_URL, base
        try:
            assert http_ctx.sesskey() == "abc123XY"
            assert HttpContext.from_cookie_list([]).sesskey() == "abc123XY"   # the fake portal needs no login
        finally:
            httpctx.BASE_URL = old_base

        # a failing request raises like Playwright's: the monitor's network-error rule recognises it
        # (after httpctx's own retries — sped up and kept out of monitor.log here)
        server.shutdown()
        server.server_close()
        backoff, httpctx.BACKOFF_SECONDS = httpctx.BACKOFF_SECONDS, 0
        logging.disable(logging.WARNING)
        try:
            http_ctx.get(f"{base}/my/", timeout=2000)
        except Exception as e:
            assert monitor_failure_kind(e) == "network", repr(e)
        else:
            raise AssertionError("expected a connection error")
        finally:
            httpctx.BACKOFF_SECONDS = backoff
            logging.disable(logging.INFO)
    finally:
        server.shutdown()

    with tempfile.TemporaryDirectory() as tmp:
        state = Path(tmp) / "auth_state.json"
        state.write_text(json.dumps({"cookies": [
            {"name": "MoodleSession", "value": "old", "domain": "elearning.zuj.edu.jo", "path": "/"},
            {"name": "other", "value": "keep", "domain": "login.microsoftonline.com", "path": "/"}],
            "origins": [{"origin": "x"}]}), encoding="utf-8")
        ctx = HttpContext.from_state_file(state)
        assert HttpContext.from_state_file(Path(tmp) / "none.json") is None
        ctx._session().cookies.set("MoodleSession", "new", domain="elearning.zuj.edu.jo", path="/")
        ctx.save_cookies(state)
        saved = json.loads(state.read_text(encoding="utf-8"))
        assert saved["cookies"][0]["value"] == "new" and saved["cookies"][1]["value"] == "keep"
        assert saved["origins"] == [{"origin": "x"}], "the rest of the saved browser state is untouched"


def monitor_failure_kind(exc):
    monitor = _package("uninotifier.monitor")
    return monitor.failure_kind(exc)


def test_healthcheck():
    from uninotifier import healthcheck

    calls = []
    orig_post = healthcheck.requests.post

    class Reply:
        def __init__(self, fail=False):
            self.fail = fail

        def raise_for_status(self):
            if self.fail:
                raise healthcheck.requests.HTTPError("500 for https://hc-ping.com/secret-uuid")

    def fake_post(url, data=None, timeout=None):
        calls.append((url, data))
        return Reply(fail=fail["on"])

    fail = {"on": False}
    healthcheck.requests.post = fake_post
    saved_env = os.environ.pop(healthcheck.ENV_NAME, None)
    try:
        assert healthcheck.url() is None and not healthcheck.start() and calls == [], "unset = disabled"
        for bad in ("http://hc-ping.com/x", "not a url", "https://"):
            os.environ[healthcheck.ENV_NAME] = bad
            assert healthcheck.url() is None, bad

        os.environ[healthcheck.ENV_NAME] = " https://hc-ping.com/uuid-1/ "
        assert healthcheck.url() == "https://hc-ping.com/uuid-1"
        healthcheck.start()
        healthcheck.report(0, "6 courses, 8 items, 9.1 s")
        healthcheck.report(1, "exit code 1 — see monitor.log")
        assert calls == [("https://hc-ping.com/uuid-1/start", b""),
                         ("https://hc-ping.com/uuid-1", b"6 courses, 8 items, 9.1 s"),
                         ("https://hc-ping.com/uuid-1/log", "exit code 1 — see monitor.log".encode())], calls

        # an unreachable service never raises, retries, and never writes the secret URL to the log
        fail["on"] = True
        calls.clear()
        records = []
        handler = logging.Handler()
        handler.emit = lambda record: records.append(record.getMessage())
        healthcheck.log.addHandler(handler)
        healthcheck.log.propagate = False        # keep the simulated failure out of monitor.log
        logging.disable(logging.NOTSET)
        try:
            assert healthcheck.success("x") is False
        finally:
            healthcheck.log.removeHandler(handler)
            healthcheck.log.propagate = True
            logging.disable(logging.INFO)
        assert len(calls) == healthcheck.RETRIES and records
        assert all("secret-uuid" not in r and "hc-ping" not in r for r in records), records
    finally:
        healthcheck.requests.post = orig_post
        os.environ.pop(healthcheck.ENV_NAME, None)
        if saved_env is not None:
            os.environ[healthcheck.ENV_NAME] = saved_env


def test_fast_scan_helpers():
    monitor = _package("uninotifier.monitor")

    done = monitor.pmap(lambda x: 10 // x, [5, 0, 2, 1])
    assert [(i, r) for i, r, e in done] == [(5, 2), (0, None), (2, 5), (1, 10)], "order is kept"
    assert isinstance(done[1][2], ZeroDivisionError) and done[0][2] is None, "an error is returned, not raised"

    class Fake:
        def __init__(self, url, ok=True):
            self.url, self.ok = url, ok

        def text(self):
            return PAGE

    class FakeContext:
        class request:
            @staticmethod
            def get(url, timeout=0):
                return Fake(url)

    PAGE = '<div class="submissionstatustable"><table><tr><td>Submission status</td><td class="submissionstatussubmitted">Submitted for grading</td></tr></table></div>'
    assert moodle.fetch_submission_status(FakeContext, "https://x/mod/assign/view.php?id=1") is True
    PAGE = '<div class="submissionsummarytable"><table><tr><td>Submission status</td><td>No attempt</td></tr></table></div>'
    assert moodle.fetch_submission_status(FakeContext, "https://x/mod/assign/view.php?id=1") is False
    PAGE = "<p>no status block here</p>"
    assert moodle.fetch_submission_status(FakeContext, "https://x/mod/assign/view.php?id=1") is None
    PAGE = "<p>login</p>"
    FakeContext.request.get = staticmethod(lambda url, timeout=0: Fake("https://x/login/index.php"))
    try:
        moodle.fetch_submission_status(FakeContext, "https://x/mod/assign/view.php?id=1")
    except moodle.SessionExpiredError:
        pass
    else:
        raise AssertionError("an expired session must surface")

    # check_submissions: only upcoming or brand-new assignments are read; results flag the deadline
    now = time.time()
    content = {"1-10": {"type": "assign", "url": "https://x/mod/assign/view.php?id=10", "name": "HW1"},
               "1-11": {"type": "assign", "url": "https://x/mod/assign/view.php?id=11", "name": "HW2 (old)"},
               "1-12": {"type": "assign", "url": "https://x/mod/assign/view.php?id=12", "name": "HW3 (new)"},
               "1-13": {"type": "page", "url": "https://x/mod/page/view.php?id=13", "name": "P"}}
    deadlines = {"assign-10": {"timestamp": now + 3600}, "assign-11": {"timestamp": now - 3600}}
    saved = {"content": {"1-10": {}, "1-11": {}}}
    read = []
    orig = moodle.fetch_submission_status
    moodle.fetch_submission_status = lambda ctx, url: read.append(url) or (None if url.endswith("12") else True)
    try:
        result = monitor.check_submissions(object(), content, deadlines, saved)
    finally:
        moodle.fetch_submission_status = orig
    assert sorted(read) == ["https://x/mod/assign/view.php?id=10", "https://x/mod/assign/view.php?id=12"]
    assert result == {"assign-10": True} and deadlines["assign-10"]["submitted"] is True
    assert "submitted" not in deadlines["assign-11"], "a past deadline is not re-read"

    # the browser is started lazily: building the object costs nothing
    browser = monitor.Browser()
    assert not browser.started
    browser.close()                                    # closing a never-started browser is harmless
    assert monitor.HTTP_WORKERS >= 2 and monitor.DEEP_CHECK_SECONDS >= 3600
    assert "deep" in monitor.STATE_KEYS


def _sem(term, courses, cumulative=None):
    """A transcript semester as portal.parse_transcript returns it; courses are
    (code, name, hours, mark) — mark None while in progress."""
    return {"term": term, "average": None, "cumulative": cumulative, "courses": [
        {"code": c, "name": n, "hours": h, "mark": m,
         "passed": None if m is None else m >= 50, "repeated": False} for c, n, h, m in courses]}


def _record():
    return [
        _sem("First semester 2026/2027", [("0101221", "جبر خطي (1)", 3, None), ("0135220", "قواعد البيانات", 3, None),
                                          ("0101272", "تحليل عددي (1)", 3, None)], 50.0),
        _sem("First semester 2025/2026", [("0112110", "رياضيات متقطعة", 3, 80), ("0101221", "جبر خطي (1)", 3, 35),
                                          ("0142251", "قواعد البيانات", 3, 35)], 50.0),   # the database course had another code
    ]


def test_gradebook_record():
    from uninotifier import gradebook

    sems = _record()
    record, hours = gradebook.cumulative(gradebook.attempts(sems))
    assert (record, hours) == (50.0, 9.0), "failed attempts count, at 35"
    assert gradebook.recorded(12) == 35 and gradebook.recorded(35) == 35 and gradebook.recorded(61) == 61

    # a retake replaces the old attempt — by code, or by name when the code changed — and hours aren't repeated
    marks = {"0101221": 70.0, "0135220": 45.0, "0101272": 33.0}
    rep = gradebook.season_report(sems, marks)
    assert rep["complete"] and rep["published"] == 3 and rep["total"] == 3
    assert rep["before"] == (50.0, 9.0)
    after, after_hours = rep["after"]
    assert after_hours == 12.0 and abs(after - (80 + 70 + 45 + 35) / 4) < 1e-9, "33 is recorded as 35"
    assert abs(rep["semester_average"] - 50.0) < 1e-9
    assert rep["replaced"] == [("0101221", "جبر خطي (1)", 35.0, 70.0), ("0135220", "قواعد البيانات", 35.0, 45.0)]
    assert rep["floored"] == [("تحليل عددي (1)", 33)]
    assert [c["passed"] for c in rep["courses"]] == [True, False, False]

    # partly published: the record only changes for what is out
    partial = gradebook.season_report(sems, {"0101221": 70.0})
    assert not partial["complete"] and partial["published"] == 1 and partial["semester_average"] == 70.0
    assert partial["after"] == (((80 + 70 + 35) * 3) / 9, 9.0)
    # a mark already on the transcript needs no overlay
    sems2 = _record()
    sems2[0]["courses"][0]["mark"] = 90
    assert gradebook.season_report(sems2, {})["published"] == 1
    assert gradebook.season_report([], {})["total"] == 0


def test_results_page_parser():
    from uninotifier import portal

    grids = [{"heads": ["x"], "rows": []},
             {"heads": ["رمز المقرر", "اسم المقرر", "النشاط", "الحالة", "العلامة", "ناجح؟"],
              "rows": [["0101221", "جبر خطي (1)", "نظري", "مسجلة", "70", "\u2713"],
                       ["0135220", "قواعد البيانات", "نظري", "مسجلة", "", "\u2717"],
                       ["لا يوجد نتائج للمساقات المسجلة حاليا"]]}]
    rows = portal.parse_results(grids)
    assert rows == [
        {"code": "0101221", "name": "جبر خطي (1)", "activity": "نظري", "status": "مسجلة", "mark": 70.0, "passed": True},
        {"code": "0135220", "name": "قواعد البيانات", "activity": "نظري", "status": "مسجلة", "mark": None, "passed": False}]
    assert portal.parse_results([{"heads": ["رمز المقرر", "العلامة"], "rows": [["لا يوجد نتائج"]]}]) == []
    lines = portal.results_lines({"term": "الفصل الأول 2026/2027", "text": ""}, [], rows)
    assert lines[0] == "📅 *First semester 2026/2027*" and lines[1].endswith("(نظري) · 70 ✅") and lines[2].endswith("· —")


def test_finals_watch():
    from uninotifier import finals
    from uninotifier import portal

    sch = {"finals": {"from": "2027-01-04", "to": "2027-01-31", "interval_minutes": 20}}
    assert finals.in_window(date(2027, 1, 4), sch) and finals.in_window(date(2027, 1, 31), sch)
    assert not finals.in_window(date(2027, 1, 3), sch) and not finals.in_window(date(2027, 2, 1), sch)
    assert not finals.in_window(date(2027, 1, 10), {}), "no window configured: only the switch turns it on"
    assert finals.interval_seconds(sch) == 1200 and finals.interval_seconds({}) == 1800

    with tempfile.TemporaryDirectory() as tmp:
        finals.STATE_FILE = Path(tmp) / "finals.json"
        state = finals.load()
        assert state["enabled"] is False
        noon = datetime(2026, 10, 12, 12, 0)
        assert not finals.active(noon, state, {}) and not finals.due(noon, state, {})
        state["enabled"] = True
        assert finals.active(noon, state, {}) and finals.due(noon, state, {})
        state["last_check"] = noon.timestamp() - 600
        assert not finals.due(noon, state, {}), "checked 10 min ago: not due"
        state["last_check"] = noon.timestamp() - 1801
        assert finals.due(noon, state, {})

        def results(marks, term_marks=None):
            sems = _record()
            for c in sems[0]["courses"]:
                if term_marks and c["code"] in term_marks:
                    c["mark"] = term_marks[c["code"]]
            transcript = portal.Result("transcript", True, figures={"cgpa": 50.0}, data={"semesters": sems})
            grades = portal.Result("grades", True, data={"results": [
                {"code": k, "name": "", "activity": "", "status": "", "mark": v, "passed": None}
                for k, v in marks.items()]})
            return grades, transcript

        state = finals.load()
        # 1) first reading of the term: just the status, nothing announced
        msgs = finals.evaluate(state, *results({}))
        assert len(msgs) == 1 and "Finals Watch" in msgs[0]["text"] and "0 of 3" in msgs[0]["text"]
        assert finals.evaluate(state, *results({})) == [], "nothing new: silence"

        # 2) a grade appears: one alert, in English, noting the retake it replaces
        msgs = finals.evaluate(state, *results({"0101221": 70}))
        assert len(msgs) == 1 and msgs[0]["text"].startswith("🎓 *Grade Released*")
        text = msgs[0]["text"]
        assert "Linear Algebra (1)" in text and "*Mark:* 70  ✅ Passed" in text
        assert "replaces your earlier 35" in text and "1 of 3 grades published" in text
        assert finals.evaluate(state, *results({"0101221": 70})) == [], "the same grade is not announced twice"

        # 3) the last two appear (one below the pass mark, one below the floor): two alerts, then the summary
        msgs = finals.evaluate(state, *results({"0101221": 70, "0135220": 45, "0101272": 33}))
        assert [m["text"].split("\n")[0] for m in msgs] == ["🎓 *Grade Released*", "🎓 *Grade Released*",
                                                            "🏁 *Semester Results — First semester 2026/2027*"]
        assert "❌ Below the 50 pass mark" in msgs[0]["text"] or "❌ Below the 50 pass mark" in msgs[1]["text"]
        assert any("Recorded as 35" in m["text"] for m in msgs[:2])
        summary = msgs[2]["text"]
        assert "*Semester Average:* 50%" in summary
        assert "*New Cumulative GPA:* 57.5%  (+7.5)" in summary, summary
        assert "*Hours in your record:* 12  (was 9" in summary and "Marks under 35 are recorded as 35" in summary
        assert "Linear Algebra (1) 35 → 70" in summary and "Database Systems 35 → 45" in summary
        assert "the portal's own field still shows 50" in summary
        assert msgs[2]["rows"] == [[{"text": "📊 GPA & Grade Hub", "callback_data": "academic"}]]
        assert state["base"] == [50.0, 9.0], "the record before this term: what Grade Needed builds on"
        assert finals.evaluate(state, *results({"0101221": 70, "0135220": 45, "0101272": 33})) == []

        # 4) the portal corrects a mark: announced as an update, and the summary is redone
        msgs = finals.evaluate(state, *results({"0101221": 70, "0135220": 55, "0101272": 33}))
        assert msgs[0]["text"].startswith("🎓 *Grade Updated*") and "55  (was 45)" in msgs[0]["text"]
        assert "*New Cumulative GPA:* 60%" in msgs[-1]["text"]

        # marks the results page lacks are taken from the transcript's current block
        fresh = finals.load()
        fresh["term"] = ""
        msgs = finals.evaluate(fresh, None, results({}, {"0101221": 70, "0135220": 45, "0101272": 33})[1])
        assert [m["text"].split("\n")[0][:12] for m in msgs] == ["🏁 *Finals Wa", "🏁 *Semester "], \
            "everything already out when the watch starts: the status, then the summary at once"

        # a new term starts over
        sems = _record()
        sems[0]["term"] = "Second semester 2026/2027"
        transcript = portal.Result("transcript", True, data={"semesters": sems})
        msgs = finals.evaluate(state, None, transcript)
        assert state["term"] == "Second semester 2026/2027" and "0 of 3" in msgs[0]["text"]

        # an unreadable transcript is an error for the caller, not silence
        try:
            finals.evaluate(state, None, portal.Result("transcript", True, data={}))
        except portal.PortalError:
            pass
        else:
            raise AssertionError("expected PortalError")

        # check(): saves the state on success, leaves it alone when the portal fails
        state = finals.load()
        fetched = []
        msgs, state = finals.check(lambda sections: fetched.append(sections) or dict(
            zip(("grades", "transcript"), results({"0101221": 70}))), state)
        assert fetched == [["grades", "transcript"]] and state["last_check"] > 0
        assert finals.load()["last_check"] == state["last_check"]
        before = finals.load()

        def boom(sections):
            raise portal.PortalError("The portal refused the sign-in: Incorrect username or password.")
        try:
            finals.check(boom, finals.load())
        except portal.PortalError:
            pass
        else:
            raise AssertionError("expected PortalError")
        assert finals.load() == before

        status = finals.status_text(finals.load(), sch)
        assert "Finals Watch" in status and "2027-01-04 to 2027-01-31" in status and "1 of 3" in status


def test_finals_commands_and_file_destination():
    from uninotifier import archive
    bot = _package("uninotifier.bot")
    from uninotifier import finals

    with tempfile.TemporaryDirectory() as tmp:
        finals.STATE_FILE = Path(tmp) / "finals.json"
        replies = []
        orig = (bot.send_view, bot.edit_view, bot.portal.credentials)
        bot.send_view = lambda text, rows: replies.append((text, rows))
        bot.edit_view = lambda mid, text, rows: replies.append((f"edit:{mid}", text, rows))
        bot.portal.credentials = lambda: ("u", "p")
        try:
            bot.handle_command("/finals")
            assert "Finals Watch" in replies[-1][0] and "*Off*" in replies[-1][0]
            assert [b["callback_data"] for row in replies[-1][1] for b in row] == ["finals:check", "finals:on", "academic", "menu"]
            assert replies[-1][1][0][0]["text"] == "🔄 Check Grades Now" and "Background Checks" in replies[-1][1][1][0]["text"]
            bot.handle_command("/finals on")
            assert "*On*" in replies[-1][0] and finals.load()["enabled"] and finals.load()["last_check"] == 0
            assert replies[-1][1][1][0]["callback_data"] == "finals:off"
            bot.handle_finals("off", 5)                                   # a button edits its message
            assert replies[-1][0] == "edit:5" and "*Off*" in replies[-1][1] and not finals.load()["enabled"]
            bot.portal.credentials = lambda: None
            bot.handle_command("/finals on")
            assert "not configured" in replies[-1][0], "turning it on without a portal login explains .env"
            bot.handle_command("/finals check")
            assert "not configured" in replies[-1][0]
        finally:
            bot.send_view, bot.edit_view, bot.portal.credentials = orig
        assert not bot._portal_lock.locked()
        assert "finals" in {c for c, _ in bot.COMMANDS}
        assert "finals:status" not in [b.get("callback_data") for row in bot.view_menu()[1] for b in row]
        assert "finals:status" in [b.get("callback_data") for row in bot.view_academic()[1] for b in row], "nested in the hub"

    # where a new file is filed is part of its Telegram message
    root = archive.ARCHIVE_ROOT
    assert archive.destination(root / "Big Data" / archive.LECTURES / "Chapter 1.pdf") == \
        "`Big Data/01_Lectures` (Lectures & Slides)"
    assert archive.destination(root / "Big Data" / archive.ASSIGNMENTS / "HW1" / "a.zip") == \
        "`Big Data/02_Assignments` (Assignments & Projects)"
    assert archive.destination(root / archive.GENERAL / "My timetable.pdf") is None, "the personal folder is not a filing place"
    assert archive.destination(Path("C:/Windows/x.txt")) is None and archive.destination(root / "Big Data") is None


def test_finals_check_ends_on_the_grades_screen():
    """Finals Watch is the one place for this term's grades: Check Grades Now shows progress in the
    message that was pressed (no extra message), ends on the screen listing every course with its mark
    or ⏳, announcements land first and the screen returns under them, and /grades and the old
    "This Term's Grades" button lead here."""
    from uninotifier import finals, portal
    bot = _package("uninotifier.bot")

    def results(marks):
        sems = _record()
        grades = portal.Result("grades", True, data={"results": [
            {"code": k, "name": "", "activity": "", "status": "", "mark": v, "passed": None} for k, v in marks.items()]})
        return {"grades": grades, "transcript": portal.Result("transcript", True, figures={"cgpa": 50.0}, data={"semesters": sems})}

    shown, announced = [], []
    reading = {"marks": {}, "error": None}

    def fetch_many(sections, ask, **kwargs):
        if reading["error"]:
            raise reading["error"]
        return results(reading["marks"])

    with tempfile.TemporaryDirectory() as tmp:
        keep = (finals.STATE_FILE, bot.send_view, bot.edit_view, bot.portal.fetch_many, bot.portal.credentials,
                notifier.send_message, bot.api)
        finals.STATE_FILE = Path(tmp) / "finals.json"
        bot.api = lambda method, **data: True                            # the chat action of the worker
        bot.send_view = lambda text, rows: shown.append(("send", None, text, rows))
        bot.edit_view = lambda mid, text, rows: shown.append(("edit", mid, text, rows))
        bot.portal.fetch_many, bot.portal.credentials = fetch_many, lambda: ("u", "p")
        notifier.send_message = lambda token, chat, text, url=None, rows=None, **kw: announced.append(text) or True

        def check(message_id=7):
            shown.clear()
            announced.clear()
            assert bot._portal_lock.acquire(blocking=False)
            bot.run_finals_check(True, message_id)                      # the worker, run in line
            assert not bot._portal_lock.locked()

        try:
            before = bot.finals_status_view()[0]
            assert "not read yet" in before and "Check Grades Now" in before

            # first reading: the "started" announcement, then the screen at the bottom — grades listed, none published
            check()
            assert len(announced) == 1 and "Finals Watch" in announced[0]
            kind, mid, text, rows = shown[-1]
            assert kind == "send" and len(shown) == 1, "announcements landed first: the screen returns under them"
            assert "0 of 3 published" in text and text.count("⏳") == 3 and "Linear Algebra (1)" in text
            assert rows[0][0]["callback_data"] == "finals:check"

            # a grade appears: announced, then the screen shows the mark
            reading["marks"] = {"0101221": 70}
            check()
            assert announced and announced[0].startswith("🎓 *Grade Released*")
            text = shown[-1][2]
            assert "1 of 3 published" in text and "✅ *70* · Linear Algebra (1)" in text and text.count("⏳") == 2

            # nothing new: no announcement, the screen is updated WHERE IT IS (the pressed message), nothing else is sent
            check()
            assert announced == [] and len(shown) == 1 and shown[0][:2] == ("edit", 7), shown
            assert "1 of 3 published" in shown[0][2] and "🕒 *Last check:*" in shown[0][2]
            check(None)                                                 # typed (/grades): the screen goes to the bottom
            assert shown[0][0] == "send" and "Finals Watch" in shown[0][2]
            assert [c["code"] for c in finals.load()["courses"]] == ["0101221", "0101272", "0135220"]
            assert _legacy_markdown_ok(shown[0][2])

            # a failed reading: the message that was pressed says why and offers to retry
            reading["error"] = portal.PortalError("The portal refused the sign-in: Incorrect username or password.")
            finals_state = finals.load()
            finals_state["failure_notice"] = 0
            finals.save(finals_state)
            check()
            kind, mid, text, rows = shown[-1]
            assert (kind, mid) == ("edit", 7) and "Incorrect username" in text
            assert "finals:check" in [b["callback_data"] for row in rows for b in row if "callback_data" in b]
            reading["error"] = None

            # the button shows progress in its own message (no new one), and /grades + the old button lead here
            shown.clear()
            keep_run, keep_handle = bot.finals_flow.run_finals_check, bot.finals_flow.handle_finals
            started = []
            bot.finals_flow.run_finals_check = lambda manual, message_id=None: (started.append((manual, message_id)),
                                                                                  bot._portal_lock.release())
            try:
                bot.handle_finals("check", 9)
                assert shown == [("edit", 9, "🏁 *Finals Watch*\n\n⏳ Reading your grades from the portal…", [])], shown
                for _ in range(100):
                    if started:
                        break
                    time.sleep(0.01)
                assert started == [(True, 9)]
                calls = []
                bot.finals_flow.handle_finals = lambda arg, message_id=None: calls.append((arg, message_id))
                bot.handle_command("/grades")
                bot.handle_callback({"id": "g1", "data": "portal:grades", "message": {"message_id": 11}})
                assert calls == [("check", None), ("check", 11)], calls
                assert "grades" not in bot.PORTAL_COMMANDS and "Check Grades Now" in dict(bot.COMMANDS)["grades"]
            finally:
                bot.finals_flow.run_finals_check, bot.finals_flow.handle_finals = keep_run, keep_handle
        finally:
            (finals.STATE_FILE, bot.send_view, bot.edit_view, bot.portal.fetch_many, bot.portal.credentials,
             notifier.send_message, bot.api) = keep
            bot._portal_lock.acquire(blocking=False)
            bot._portal_lock.release()


def _sample_transcript():
    """A parsed transcript: a term in progress, two finished terms (one with a failed attempt a later
    retake replaced), Arabic and Latin names, a name with Markdown characters."""
    from uninotifier import portal
    heads = ["", "رمز المقرر", "اسم المقرر", "الحالة", "عدد الساعات", "العلامة", "ناجح؟", "معادة"]
    ok, no = "✓", "✗"

    def row(code, name, hours, mark, passed, repeated=False):
        return ["1", code, name, "مسجلة", str(hours), str(mark), passed, ok if repeated else ""]

    raw = [
        {"legend": "الفصل الأول 2026/2027 20261", "heads": heads, "text": "المعدل التراكمي 71.3",
         "rows": [row("0135220", "قواعد البيانات", 3, "", ""), row("0101221", "جبر خطي (1)", 3, "", ""),
                  row("0135321", "Big_Data *2*", 3, "", "")]},
        {"legend": "الفصل الثاني 2025/2026 20252", "heads": heads, "text": "المعدل الفصلي 76 المعدل التراكمي 71.3",
         "rows": [row("0125220", "برمجة تطبيقات الانترنت (HTML)", 3, 88, ok), row("0135214", "الاحصاء", 3, 74, ok)]},
        {"legend": "الفصل الأول 2024/2025 20241", "heads": heads, "text": "المعدل الفصلي 58",
         "rows": [row("0101221", "جبر خطي (1)", 3, 35, no, True), row("0142251", "قواعد البيانات", 3, 35, no),
                  row("0420155", "القانون في حياتنا", 2, 95, ok), row("0112212", "تراكيب البيانات", 3, 30, no)]},
    ]
    return portal.parse_transcript(raw)


def test_transcript_layout():
    """The transcript reads like a report: an executive summary (GPA, passed hours with what is in
    progress or still to repeat, attempted hours, status — and for a guest the major and the hours
    left), then the semester in progress first and every other semester with ✅ ❌ 🔁 marks."""
    from uninotifier import transcript

    sems = _sample_transcript()
    state = transcript.standing(sems)
    # Linear Algebra and Databases failed once but are being taken again (in progress, not "to repeat", their old
    # attempts replaced); Data Structures failed and has not been retaken: that is what is still "to repeat"
    assert state == {"passed": 8.0, "in_progress": 9.0, "to_repeat": 3.0, "retaken": 2, "courses_passed": 3}, state
    figures = {"record_cgpa": 71.26, "cgpa": 71.3, "passed_hours": 8.0, "record_hours": 14.0}
    lines = transcript.summary_lines(sems, figures, "Regular / Good Standing")
    assert lines == [
        "🎓 *Cumulative GPA:* 71.26% (Portal display: 71.3%)",
        "⏱️ *Passed Hours:* 8 hrs",
        "      └ ⏳ 9h in progress · ❌ 3h to repeat · 🔁 2 retaken",
        "📚 *Total Attempted Hours:* 14 hrs",
        "📌 *Academic Status:* Regular / Good Standing"], lines
    # something failed and not taken again shows up as "to repeat"; zero parts are not listed
    failed_only = [{"term": "T", "average": None, "cumulative": None, "courses": [
        {"code": "1", "name": "X", "hours": 3, "mark": 40, "passed": False, "repeated": False},
        {"code": "2", "name": "Y", "hours": 3, "mark": 80, "passed": True, "repeated": False}]}]
    assert "      └ ❌ 3h to repeat" in transcript.summary_lines(failed_only, {"passed_hours": 3.0})
    assert not any("└" in l for l in transcript.summary_lines(sems[1:2], {"passed_hours": 6.0}))
    # a guest: the major first, the hours left after the status block
    guest = transcript.summary_lines(sems, figures, None, major="نظم المعلومات", remaining=124.0)
    assert guest[0].startswith("📘 *Major:*") and "نظم المعلومات" in guest[0] and guest[-1] == "🏁 *Remaining Hours:* 124 hrs"
    assert transcript.summary_lines(sems, figures, None, remaining=-5.0)[-1] == "🏁 *Remaining Hours:* 0 hrs"
    assert transcript.summary_lines([], {}) == []

    text = "\n".join(transcript.blocks(sems))
    order = [text.index(m) for m in ("In Progress", "Second semester 2025/2026", "First semester 2024/2025")]
    assert order == sorted(order), "the semester in progress first"
    assert text.startswith("⏳ *First semester 2026/2027* · In Progress")
    assert "📅 *Second semester 2025/2026* · average 76 · cumulative 71.3" in text
    for fragment in ("✅ *88* ·", "✅ *74* ·", "🔁 *35* · \u2068جبر خطي (1)\u2069 · 3h", "🔁 *35* · \u2068قواعد البيانات\u2069 · 3h",
                     "❌ *30* · \u2068تراكيب البيانات\u2069 · 3h", "✅ *95* ·"):
        assert fragment in text, fragment       # Databases: failed, replaced by this term's retake (the portal did not flag it)
    assert text.count("⏳ \u2068") == 3, "the three courses of the term in progress"
    assert text.count("🔁 *") == 2 and text.count("❌ *") == 1
    # an Arabic name cannot reorder its neighbours: every course line starts with a left-to-right mark, the name is isolated
    for line in text.splitlines():
        if "· 3h" in line or "· 2h" in line:
            assert line.startswith("\u200e") and "\u2068" in line and "\u2069" in line, repr(line)
    assert "Big\\_Data \\*2\\*" in text and _legacy_markdown_ok(text), "Markdown characters in a name are escaped"
    assert text.rstrip().endswith("⏳ in progress_") and transcript.blocks([]) == []
    # marks the results page already shows for the term in progress
    published = "\n".join(transcript.blocks(sems, {"0135220": 74.0, "0101221": 40.0}))
    assert "✅ *74* · \u2068قواعد البيانات\u2069" in published and "❌ *40* · \u2068جبر خطي (1)\u2069" in published
    assert "🔁 *35*" in published, "a replaced attempt is still marked as such"
    # a full record fits one message
    many = sems * 3
    assert len("\n".join(transcript.blocks(many))) < notifier.MESSAGE_LIMIT


def test_portal_answer_is_one_message():
    """A portal answer is ONE message (the screenshot with its caption); a balance holds everything
    in it, only a transcript's semesters follow as a single text. No "retrieved" status message, no
    "details" heading, and the hub it came from returns under the answer."""
    from uninotifier import portal, maintenance
    bot = _package("uninotifier.bot")

    sems = _sample_transcript()
    sent_files, sent_msgs, screens, edits, api_calls = [], [], [], [], []
    keep = (bot.notifier.send_png_file, bot.notifier.send_message, bot.notifier.send_long_message, bot.send_view,
            bot.edit_view, bot.api, bot.portal.fetch_many, bot.sync_base, bot.sync_retakes, bot.lock_grades,
            bot.SWR_SECTIONS)
    bot.SWR_SECTIONS = ()          # every request here reads afresh (test_portal_answers_from_the_saved_copy has the cache)
    bot.notifier.send_png_file = lambda tok, chat, png, name, caption="", url=None, rows=None: (
        sent_files.append((name, caption, rows)) or True)
    bot.notifier.send_message = lambda tok, chat, text, url=None, rows=None: sent_msgs.append(("msg", text)) or True
    bot.notifier.send_long_message = lambda tok, chat, text, url=None, rows=None: sent_msgs.append(("long", text)) or True
    bot.send_view = lambda text, rows: screens.append(text) or {"message_id": 1}
    bot.edit_view = lambda mid, text, rows: edits.append((mid, text, rows))
    bot.api = lambda method, **d: api_calls.append(method) or True
    bot.sync_base = bot.sync_retakes = lambda *a: None
    bot.lock_grades = lambda *a: None

    def reset():
        sent_files.clear(); sent_msgs.clear(); screens.clear(); edits.clear()

    try:
        # a transcript: the summary in the caption of the screenshot, the semesters in one text message
        transcript_result = portal.Result("transcript", True, b"png", [], {"record_cgpa": 71.26, "cgpa": 71.3, "passed_hours": 8.0,
                                          "record_hours": 14.0, "base_cgpa": 70.0, "base_hours": 14.0},
                                          [], [], {"semesters": sems, "status": "منتظم", "retakes": {}})
        bot.send_portal_result(transcript_result)
        assert len(sent_files) == 1 and [k for k, _ in sent_msgs] == ["long"], (sent_files, sent_msgs)
        caption = sent_files[0][1]
        assert caption.startswith("📜 *ZUJ Portal — Transcript*") and "⏱️ *Passed Hours:* 8 hrs" in caption
        assert "      └ ⏳ 9h in progress · ❌ 3h to repeat · 🔁 2 retaken" in caption and "Total Attempted Hours" in caption \
            and len(caption) < 1024
        assert sent_msgs[0][1].startswith("⏳ *First semester 2026/2027* · In Progress"), "no \"details\" heading"
        assert "details" not in sent_msgs[0][1] and "Retrieved" not in sent_msgs[0][1]

        # a balance: everything in the one message
        reset()
        info = {"total": 1250.0, "rows": [("First semester 2026/2027", 18.0, 1250.0), ("Second semester 2025/2026", 12.0, 0.0)]}
        balance = portal.Result("balance", True, b"png", [], {"balance": 1250.0}, [], portal.balance_lines(info))
        bot.send_portal_result(balance)
        assert len(sent_files) == 1 and sent_msgs == [], "a balance is one message"
        assert "Total balance:* 1,250 JD" in sent_files[0][1] and "First semester 2026/2027 · 18h · 1,250 JD" in sent_files[0][1]
        # ... unless it cannot fit a caption: the total stays there, the rest follows once
        reset()
        long_rows = [(f"Semester {n}", 15.0, 100.0 + n) for n in range(60)]
        big = portal.Result("balance", True, b"png", [], {"balance": 9.0}, [], portal.balance_lines({"total": 9000.0, "rows": long_rows}))
        bot.send_portal_result(big)
        assert len(sent_files) == 1 and len(sent_files[0][1]) <= notifier.CAPTION_LIMIT and "Total balance" in sent_files[0][1]
        assert [k for k, _ in sent_msgs] == ["long"] and "Semester 59" in sent_msgs[0][1] and "Total balance" not in sent_msgs[0][1]

        # the request: a button's own message shows progress; a typed command shows none; the hub returns under the answer
        bot.portal.fetch_many = lambda sections, ask: {"transcript": transcript_result}
        for message_id in (5, None):
            reset()
            api_calls.clear()
            assert bot._portal_lock.acquire(blocking=False)
            bot.run_portal("transcript", message_id)
            assert not bot._portal_lock.locked()
            assert screens and screens[-1].startswith("📜 *Academic Record & Balance*"), screens
            assert not any("retrieved" in t for t in screens) and len(screens) == 1, "no status message, just the hub"
            if message_id:
                assert [e[0] for e in edits] == [5] and "signing in" in edits[0][1] and edits[0][2] == []
            else:
                assert edits == [] and screens == [screens[-1]], "typed: nothing but the answer and the hub"
            assert "sendChatAction" in api_calls and len(sent_files) == 1

        # a failure is said in the message that was pressed, with a way to retry; typed: as the menu itself
        for exc in (portal.PortalError("The portal refused the sign-in."), RuntimeError("boom")):
            for message_id in (5, None):
                reset()
                bot.portal.fetch_many = lambda sections, ask, exc=exc: (_ for _ in ()).throw(exc)
                assert bot._portal_lock.acquire(blocking=False)
                bot.run_portal("balance", message_id)
                shown = edits[-1][1:] if message_id else (screens[-1], None)
                assert shown[0].startswith("❌ *Tuition Balance*") and "refused" in shown[0] or "couldn't be read" in shown[0]
                rows = edits[-1][2] if message_id else None
                if rows:
                    assert rows[0][0] == {"text": "🔄 Try Again", "callback_data": "portal:balance"}
                assert sent_files == [] and not bot._portal_lock.locked()
        # maintenance: the screenshot says why, and the hub returns under it
        reset()
        bot.portal.fetch_many = lambda sections, ask: (_ for _ in ()).throw(maintenance.PortalUnavailable(
            "registration portal", "the sign-in page shows a maintenance notice", "قيد الصيانة", b"PNG"))
        assert bot._portal_lock.acquire(blocking=False)
        bot.run_portal("transcript", None)
        assert len(sent_files) == 1 and sent_files[0][2] == [] and screens[-1].startswith("📜 *Academic Record & Balance*")
    finally:
        (bot.notifier.send_png_file, bot.notifier.send_message, bot.notifier.send_long_message, bot.send_view,
         bot.edit_view, bot.api, bot.portal.fetch_many, bot.sync_base, bot.sync_retakes, bot.lock_grades,
         bot.SWR_SECTIONS) = keep


def test_archive():
    assert archive.course_folder_name("20261-Big Data-AllSec-3-") == "Big Data"
    assert archive.course_folder_name("20261-Linear Algebra (1)-AllSec-3-F") == "Linear Algebra (1)"
    assert archive.course_folder_name("20261-الاتصال والتواصل باللغة العربية-ش31-3-ك") == "الاتصال والتواصل باللغة العربية"
    assert archive.categorize("quiz", "x") == archive.EXAMS
    assert archive.categorize("resource", "Midterm 2025") == archive.EXAMS
    assert archive.categorize("resource", "Course Syllabus") == archive.LECTURES
    assert archive.categorize("resource", "Chapter 3 slides") == archive.LECTURES
    assert archive.categorize("assign", "HW1") == archive.ASSIGNMENTS


def test_archive_layout():
    """English course folders with exactly three numbered subfolders. The semester
    root's 00_Schedules is the student's own: nothing is filed there."""
    assert (archive.GENERAL, archive.LECTURES, archive.ASSIGNMENTS, archive.EXAMS) == (
        "00_Schedules", "01_Lectures", "02_Assignments", "03_Exams")
    assert archive.SUBFOLDERS == (archive.LECTURES, archive.ASSIGNMENTS, archive.EXAMS)

    # portal names -> the English folder, whatever language or spelling the portal uses
    for portal_name, english in (
            ("20261-الاتصال والتواصل باللغة الإنجليزية-ش27-3-و", "English Communication Skills"),
            ("20261-الاتصال والتواصل باللغة العربية-ش31-3-ك", "Arabic Communication Skills"),
            ("20261-Databases-AllSec-3-F", "Database Systems"),
            ("20261-Numerical Analysis(1)-AllSec-3-", "Numerical Analysis (1)"),
            ("20261-Big Data-AllSec-3-", "Big Data"),
            ("20261-Some New Course-AllSec-3-", "Some New Course")):
        assert archive.english_course_name(portal_name) == english, portal_name
    assert archive.english_course_name("الاتصال والتواصل باللغة العربية") == "Arabic Communication Skills", \
        "a folder name already on disk maps the same way"

    # routing: what goes where
    assert archive.categorize("resource", "Exam Schedule") == archive.EXAMS
    assert archive.categorize("resource", "جدول الامتحانات النهائية") == archive.EXAMS
    assert archive.categorize("resource", "Midterm exam dates") == archive.EXAMS
    assert archive.categorize("resource", "Midterm 2025 solutions") == archive.EXAMS
    assert archive.categorize("quiz", "Exam schedule quiz") == archive.EXAMS
    for plan in ("Course Syllabus", "Big Data Course Plan for Bachelor program 0135321 (1)", "خطة المادة",
                 "Course Plan - Projects", "Final exam syllabus"):
        assert archive.categorize("resource", plan) == archive.LECTURES, plan
    assert archive.categorize("resource", "Lecture 3 slides") == archive.LECTURES
    assert archive.categorize("resource", "Homework 2") == archive.ASSIGNMENTS

    root = archive.ARCHIVE_ROOT, archive._ARCHIVE_ROOT_RESOLVED
    with tempfile.TemporaryDirectory() as tmp:
        archive.ARCHIVE_ROOT = Path(tmp) / "01-First_Semester"
        archive._ARCHIVE_ROOT_RESOLVED = archive.ARCHIVE_ROOT.resolve()
        try:
            arabic = "20261-الاتصال والتواصل باللغة العربية-ش31-3-ك"
            base = archive.ensure_course_tree(arabic)
            assert base.name == "Arabic Communication Skills"
            assert sorted(p.name for p in base.iterdir()) == [
                "01_Lectures", "02_Assignments", "03_Exams"], "strictly three"
            assert sorted(p.name for p in archive.ARCHIVE_ROOT.iterdir()) == ["Arabic Communication Skills"]
            assert not (base / "00_Schedules").exists() and not list(base.glob("*chedules*"))

            assert archive.target(arabic, "resource", "Lecture 1") == base / archive.LECTURES
            assert archive.target(arabic, "resource", "Course Syllabus") == base / archive.LECTURES
            assert archive.target(arabic, "resource", "Exam dates") == base / archive.EXAMS
            assert archive.target("20261-Databases-AllSec-3-F", "assign", "Project 1") == \
                archive.ARCHIVE_ROOT / "Database Systems" / archive.ASSIGNMENTS
            assert not (archive.ARCHIVE_ROOT / archive.GENERAL).exists(), "the personal folder is never created or used"
            try:
                archive.ensure_inside_archive(Path(tmp) / "02-Other_Semester" / "x")
            except archive.ArchiveBoundaryError:
                pass
            else:
                raise AssertionError("the boundary guard must still hold")
        finally:
            archive.ARCHIVE_ROOT, archive._ARCHIVE_ROOT_RESOLVED = root


def test_resource_page_file_link():
    """A resource's intermediate page lists the theme's favicon through
    pluginfile.php before the real file; the file is the mod_resource one."""
    page = ('<link rel="shortcut icon" href="https://x.jo/pluginfile.php/2/theme_mb2nl/favicon/0/favicon.ico">'
            '<a href="https://x.jo/pluginfile.php/204704/mod_resource/content/4/The%20Bisection%20Method.pdf?forcedownload=1&amp;a=b">file</a>')
    assert archive._file_link(page) == "https://x.jo/pluginfile.php/204704/mod_resource/content/4/The%20Bisection%20Method.pdf?forcedownload=1&a=b"
    # no mod_resource link: still never the theme's icon
    other = ('<a href="https://x.jo/pluginfile.php/2/theme_mb2nl/favicon/0/favicon.ico">i</a>'
             '<a href="https://x.jo/pluginfile.php/9/mod_folder/content/0/a.pdf">f</a>')
    assert archive._file_link(other) == "https://x.jo/pluginfile.php/9/mod_folder/content/0/a.pdf"
    assert archive._file_link('<a href="https://x.jo/pluginfile.php/2/theme_x/favicon/0/favicon.ico">i</a>') is None
    assert archive._file_link("<p>nothing</p>") is None


def test_migration_helpers():
    from uninotifier.tools import migrate_archive as m

    assert m.FAVICON_NAME.match("favicon.ico") and m.FAVICON_NAME.match("favicon (13).ico")
    assert not m.FAVICON_NAME.match("favicon-lecture.ico") and not m.FAVICON_NAME.match("slides.pdf")
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        (tmp / "a.pdf").write_bytes(b"1")
        assert m.free_path(tmp / "b.pdf") == tmp / "b.pdf" and m.free_path(tmp / "a.pdf") == tmp / "a (1).pdf"
        # favicon detection: only identical image copies of recorded resources count
        icon = b"\x89PNG\r\n\x1a\nrest"
        for n in ("favicon.ico", "favicon (1).ico"):
            (tmp / n).write_bytes(icon)
        state = {"files": {"1-1": {"path": str(tmp / "favicon.ico")}, "1-2": {"path": str(tmp / "favicon (1).ico")},
                           "1-3": {"path": str(tmp / "a.pdf")}},
                 "content": {"1-1": {"type": "resource"}, "1-2": {"type": "resource"}, "1-3": {"type": "resource"}}}
        assert sorted(m.find_favicon_junk(state)) == ["1-1", "1-2"]
        (tmp / "favicon (1).ico").write_bytes(icon + b"different")
        assert m.find_favicon_junk(state) == {}, "copies that differ are not touched"
        (tmp / "favicon (1).ico").write_bytes(icon)
        state["content"]["1-2"]["type"] = "assign"
        assert sorted(m.find_favicon_junk(state)) == ["1-1"], "only downloaded resources are considered"


def test_downloader_dedup_and_routing():
    class Resp:
        def __init__(self, body, url="https://x/pluginfile.php/1/mod_resource/content/1/Lecture.pdf", ctype="application/pdf"):
            self._body, self.url, self.ok = body, url, True
            self.headers = {"content-type": ctype, "last-modified": "Wed, 01 Jan 2025 00:00:00 GMT"}

        def body(self):
            return self._body

        def text(self):
            return ""

    class Ctx:
        def __init__(self, body):
            self.body = body
            self.request = self

        def get(self, url, timeout=0):
            return Resp(self.body)

    with tempfile.TemporaryDirectory() as tmp:
        folder = Path(tmp) / "Big Data" / "01_Lectures"
        folder.mkdir(parents=True)
        old_root = archive.ARCHIVE_ROOT, archive._ARCHIVE_ROOT_RESOLVED
        archive.ARCHIVE_ROOT, archive._ARCHIVE_ROOT_RESOLVED = Path(tmp), Path(tmp).resolve()
        try:
            first = archive.download_resource(Ctx(b"PDF-A" * 100), "https://x/mod/resource/view.php?id=1", folder)
            assert first["path"].name == "Lecture.pdf" and not first.get("duplicate")
            # a second Moodle item with the very same bytes: nothing is written
            again = archive.download_resource(Ctx(b"PDF-A" * 100), "https://x/mod/resource/view.php?id=2", folder)
            assert again["duplicate"] and again["path"] == first["path"] and again["sha256"] == first["sha256"]
            assert [p.name for p in folder.iterdir()] == ["Lecture.pdf"], "no 'Lecture (1).pdf'"
            # different bytes under the same name: a real new file, numbered
            other = archive.download_resource(Ctx(b"PDF-B" * 100), "https://x/mod/resource/view.php?id=3", folder)
            assert not other.get("duplicate") and other["path"].name == "Lecture (1).pdf"
            # ...and its exact copy is recognised too, preferring the clean name when both match
            assert archive.download_resource(Ctx(b"PDF-B" * 100), "u", folder)["path"].name == "Lecture (1).pdf"
            # an instructor's re-upload (tag set) is never swallowed: the caller compares versions itself
            updated = archive.download_resource(Ctx(b"PDF-A" * 100), "u", folder, tag="updated 2026-10-12")
            assert not updated.get("duplicate") and "updated" in updated["path"].name

            # the shared general folder: only the same course's files count
            general = Path(tmp) / archive.GENERAL
            general.mkdir()
            a = archive.download_resource(Ctx(b"SYL"), "u", general, prefix="Big Data - ")
            b = archive.download_resource(Ctx(b"SYL"), "u", general, prefix="Linear Algebra (1) - ")
            c = archive.download_resource(Ctx(b"SYL"), "u", general, prefix="Big Data - ")
            assert not a.get("duplicate") and not b.get("duplicate") and c["duplicate"] and c["path"] == a["path"]
        finally:
            archive.ARCHIVE_ROOT, archive._ARCHIVE_ROOT_RESOLVED = old_root

    # practice questions / problem sheets are assignment material, not lecture notes
    for name in ("اسئلة مقترحة", "أسئلة مقترحة", "Practice questions", "Suggested Questions", "Exercises 3",
                 "Problem sheet 2", "Worksheet", "تمارين الفصل الأول"):
        assert archive.categorize("resource", name) == archive.ASSIGNMENTS, name
    assert archive.categorize("resource", "Midterm questions") == archive.EXAMS, "exam words still win"
    assert archive.categorize("resource", "Chapter 2 lecture notes") == archive.LECTURES
    assert archive.categorize("resource", "Review") == archive.LECTURES and archive.categorize("resource", "ملخص المادة") == archive.LECTURES
    assert archive.categorize("resource", "Course Syllabus") == archive.LECTURES


def test_tidy_archive():
    from uninotifier.tools import tidy_archive as t

    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp) / "01-First_Semester"
        old = archive.ARCHIVE_ROOT, archive._ARCHIVE_ROOT_RESOLVED
        archive.ARCHIVE_ROOT, archive._ARCHIVE_ROOT_RESOLVED = root, root.resolve()
        try:
            lectures = root / "Big Data" / archive.LECTURES
            general = root / archive.GENERAL
            lectures.mkdir(parents=True)
            general.mkdir()
            (lectures / "Notes.pdf").write_bytes(b"same")
            (lectures / "Notes (1).pdf").write_bytes(b"same")
            (lectures / "Notes (2).pdf").write_bytes(b"same")
            (lectures / "Derivative.docx").write_bytes(b"version one")
            (lectures / "Derivative (1).docx").write_bytes(b"a different document")   # NOT a duplicate
            (lectures / "Sheet.docx").write_bytes(b"questions")
            (general / "Big Data - Syllabus.docx").write_bytes(b"syl")
            (general / "Big Data - Syllabus (1).docx").write_bytes(b"syl")
            (general / "My timetable.pdf").write_bytes(b"mine")
            (general / "My timetable (1).pdf").write_bytes(b"mine")

            assert (t.base_name(Path("a (2).pdf")), t.base_name(Path("a.pdf")), t.base_name(Path("b (c).pdf"))) == \
                ("a.pdf", "a.pdf", "b (c).pdf")
            dupes = {k.name: sorted(c.name for c in cs) for k, cs in t.find_duplicates(root)}
            # the personal folder is never de-duplicated, whatever it holds
            assert dupes == {"Notes.pdf": ["Notes (1).pdf", "Notes (2).pdf"]}, dupes
            # the keeper is the clean name even when a numbered copy is older
            os.utime(lectures / "Notes.pdf", (2e9, 2e9))
            assert [k.name for k, _ in t.find_duplicates(root) if k.parent == lectures] == ["Notes.pdf"]

            course = "20261-Big Data-AllSec-3-"
            state = {"files": {"1": {"path": str(lectures / "Sheet.docx")},
                               "2": {"path": str(general / "Big Data - Syllabus (1).docx")},
                               "3": {"path": str(lectures / "Notes.pdf")}},
                     "content": {"1": {"type": "resource", "name": "اسئلة مقترحة", "course": course},
                                 "2": {"type": "resource", "name": "Course Syllabus", "course": course},
                                 "3": {"type": "resource", "name": "Lecture notes", "course": course}}}
            moves = t.misfiled(state, {})
            # a syllabus the bot once put in the shared folder goes back to its course, without the prefix;
            # the student's own files there (not in the state) are not touched
            assert moves == [(lectures / "Sheet.docx", root / "Big Data" / archive.ASSIGNMENTS / "Sheet.docx"),
                             (general / "Big Data - Syllabus (1).docx", lectures / "Syllabus (1).docx")], moves
        finally:
            archive.ARCHIVE_ROOT, archive._ARCHIVE_ROOT_RESOLVED = old


def test_auto_done_only_for_saved_files_and_links():
    """"Mark as done" is ticked automatically only on static material that really
    landed; graded work, forums, pages and labels are never ticked."""
    monitor = _package("uninotifier.monitor")
    ticked = []
    orig = monitor.moodle.set_completion
    monitor.moodle.set_completion = lambda ctx, key, cmid, done: ticked.append(cmid) or True
    try:
        with tempfile.TemporaryDirectory() as tmp:
            saved = Path(tmp) / "Lecture 1.pdf"
            saved.write_bytes(b"%PDF")

            def tick(kind, artifact, n):
                item = {"type": kind, "cmid": str(n), "manual": True, "done": False}
                monitor.auto_complete(None, "k", item, artifact)
                return item["done"]

            assert tick("resource", {"kind": "file", "path": str(saved)}, 1)
            assert tick("folder", {"kind": "file", "path": tmp}, 2)
            assert tick("url", {"kind": "link", "target": "https://youtu.be/x"}, 3)
            assert not tick("resource", {"kind": "file", "path": str(Path(tmp) / "gone.pdf")}, 4)
            assert not tick("resource", {"kind": "unknown"}, 5)          # download failed
            assert not tick("url", {"kind": "unknown"}, 6)               # link not resolved
            for n, kind in enumerate(("assign", "quiz", "forum", "page", "label", "book"), 10):
                assert not tick(kind, {"kind": "file", "path": str(saved)}, n), kind
            assert ticked == ["1", "2", "3"], ticked
    finally:
        monitor.moodle.set_completion = orig
    assert media.CONTAINER_TYPES == {"forum"}


def test_announce_first_run_and_duplicate_files():
    """monitor.py --announce reports an empty state's first scan; an empty forum
    stays silent, a second Moodle item for an already saved file is not attached again."""
    monitor = _package("uninotifier.monitor")
    course = "20261-Big Data-AllSec-3-"
    items = {"1-1": {"name": "Lecture 1", "type": "resource", "course": course, "cmid": "1", "url": "https://x/mod/resource/view.php?id=1"},
             "1-2": {"name": "Lecture 1 (again)", "type": "resource", "course": course, "cmid": "2", "url": "https://x/mod/resource/view.php?id=2"},
             "1-3": {"name": "Announcements", "type": "forum", "course": course, "cmid": "3", "url": "https://x/mod/forum/view.php?id=3"}}
    seen = {}

    def fake_scan(saved, notify_new, mfa_handler=None):
        seen["notify_new"] = notify_new
        return {"courses": {"1": {"name": course, "url": "https://x/c"}}, "content": dict(items),
                "deadlines": {}, "artifacts": {}, "updates": {}, "edits": {}, "posts": {},
                "forum_posts": {"1-3": []}, "submissions": {}, "link_changes": [], "seconds": 1.0,
                "browser_used": False}

    with tempfile.TemporaryDirectory() as tmp:
        keep = {k: getattr(monitor, k) for k in ("DATA_FILE", "USERNAME", "PASSWORD", "BOT_TOKEN", "CHAT_ID",
                                                 "run_scan", "notify_new_item", "send_schedule_messages",
                                                 "mfa_allowed", "notify_new_course")}
        monitor.DATA_FILE = Path(tmp) / "saved_items.json"
        monitor.USERNAME = monitor.PASSWORD = monitor.BOT_TOKEN = monitor.CHAT_ID = "x"
        monitor.send_schedule_messages = lambda saved: None
        monitor.mfa_allowed = lambda saved: False
        monitor.notify_new_course = lambda course: True
        monitor.run_scan = fake_scan
        sent = []
        monitor.notify_new_item = lambda it, deadline, art, submitted=None: sent.append(it["name"]) or True
        try:
            monitor.run()                                   # plain first run: silent baseline
            assert seen["notify_new"] is False and sent == []
            monitor.DATA_FILE.unlink()
            monitor.run(announce=True)                      # --announce: everything is reported
            assert seen["notify_new"] is True and sent == ["Lecture 1", "Lecture 1 (again)"], sent
            state = json.loads(monitor.DATA_FILE.read_text(encoding="utf-8"))
            assert set(state["content"]) == set(items)      # the forum is known, just not announced
        finally:
            for k, v in keep.items():
                setattr(monitor, k, v)

    docs = []
    orig = monitor.deliver
    monitor.deliver = lambda header, body="", url=None, rows=None, shots=(), shot_name="", document=None: \
        docs.append((header, document)) or True
    try:
        with tempfile.TemporaryDirectory() as tmp:
            pdf = Path(tmp) / "Lecture 1.pdf"
            pdf.write_bytes(b"%PDF")
            item = items["1-1"]
            monitor.notify_new_item(item, None, {"kind": "file", "path": str(pdf), "size": 4})
            monitor.notify_new_item(item, None, {"kind": "file", "path": str(pdf), "size": 4, "duplicate": True})
    finally:
        monitor.deliver = orig
    assert docs[0][1] == pdf and docs[1][1] is None and "not attached again" in docs[1][0]


def test_alerts_are_recorded_as_they_are_sent():
    """A run that dies half-way must not make the next one repeat what was already
    delivered: every ID is written down the moment its message goes out."""
    monitor = _package("uninotifier.monitor")

    course = "20261-Big Data-AllSec-3-"

    def item(n, kind="label"):
        return {"name": f"Item {n}", "type": kind, "course": course, "course_id": "1", "cmid": str(n),
                "url": f"https://x/mod/{kind}/view.php?id={n}"}

    content = {f"1-{n}": item(n) for n in (10, 11, 12)}
    content["1-20"] = item(20, "forum")
    posts = [{"id": str(n), "discussion": "1", "subject": f"S{n}", "author": "Dr", "time": n, "html": "x",
              "is_reply": False, "url": f"https://x/d#p{n}", "attachments": []} for n in (101, 102, 103)]

    def result():
        return {"courses": {"1": {"name": course, "url": "https://x/c"}}, "content": dict(content),
                "deadlines": {}, "artifacts": {}, "updates": {}, "edits": {},
                "posts": {"1-20": list(posts)}, "forum_posts": {"1-20": list(posts)},
                "submissions": {}, "link_changes": [], "seconds": 1.0, "browser_used": False}

    class Crash(BaseException):
        pass

    with tempfile.TemporaryDirectory() as tmp:
        keep = {k: getattr(monitor, k) for k in ("DATA_FILE", "USERNAME", "PASSWORD", "BOT_TOKEN", "CHAT_ID",
                                                 "run_scan", "notify_new_item", "notify_forum_post",
                                                 "send_schedule_messages", "mfa_allowed", "notify_new_course")}
        monitor.DATA_FILE = Path(tmp) / "saved_items.json"
        monitor.USERNAME = monitor.PASSWORD = monitor.BOT_TOKEN = monitor.CHAT_ID = "x"
        monitor.send_schedule_messages = lambda saved: None
        monitor.mfa_allowed = lambda saved: False
        monitor.notify_new_course = lambda course: True
        def fake_scan(saved, notify_new, mfa_handler=None):
            r = result()          # the portal's view; "posts" is filtered against what was already seen, as run_scan does
            r["posts"] = monitor.new_forum_posts(saved, r["forum_posts"], notify_new)
            return r
        monitor.run_scan = fake_scan
        # a state that already knows the course and the forum (so this is not a "first run")
        monitor.save_data({**monitor.load_saved_data(), "courses": {"1": {"name": course, "url": "https://x/c"}},
                           "content": {"1-0": item(0)}, "forums": {"1-20": ["100"]}})
        sent_items, sent_posts = [], []
        try:
            # run 1: dies while announcing the third new item
            def flaky_item(it, deadline, art, submitted=None):
                if it["name"] == "Item 12":
                    raise Crash("process killed")
                sent_items.append(it["name"])
                return True
            monitor.notify_new_item = flaky_item
            monitor.notify_forum_post = lambda forum, post: sent_posts.append(post["id"]) or True
            try:
                monitor.run()
            except Crash:
                pass
            assert sent_items == ["Item 10", "Item 11"], sent_items
            on_disk = json.loads(monitor.DATA_FILE.read_text(encoding="utf-8"))
            assert {"1-10", "1-11"} <= set(on_disk["content"]) and "1-12" not in on_disk["content"], \
                "what was delivered is already on disk, what wasn't is not"

            # run 2: only what was never delivered goes out — nothing is announced twice
            monitor.notify_new_item = lambda it, deadline, art, submitted=None: sent_items.append(it["name"]) or True
            monitor.run()
            # the new forum itself is no news (it is only a container) — its posts are
            assert sent_items == ["Item 10", "Item 11", "Item 12"], sent_items
            assert sent_posts == ["101", "102", "103"]
            assert "1-20" in json.loads(monitor.DATA_FILE.read_text(encoding="utf-8"))["content"]

            # run 3 finds nothing new at all
            monitor.run()
            assert sent_items == ["Item 10", "Item 11", "Item 12"] and sent_posts == ["101", "102", "103"], (sent_items, sent_posts)

            # forum posts: dying after the first post, the other two are still owed — and only those
            monitor.save_data({**monitor.load_saved_data(), "forums": {"1-20": ["100"]}})
            sent_posts.clear()

            def flaky_post(forum, post):
                if post["id"] == "102":
                    raise Crash("killed")
                sent_posts.append(post["id"])
                return True
            monitor.notify_forum_post = flaky_post
            try:
                monitor.run()
            except Crash:
                pass
            assert sent_posts == ["101"]
            assert "101" in json.loads(monitor.DATA_FILE.read_text(encoding="utf-8"))["forums"]["1-20"]
            monitor.notify_forum_post = lambda forum, post: sent_posts.append(post["id"]) or True
            monitor.run()
            assert sent_posts == ["101", "102", "103"], sent_posts

            # Telegram refusing one item run after run: it is owed every run, and reported once it is stuck
            _fresh_failsafe()
            monitor.save_data({**monitor.load_saved_data(),
                               "content": {k: v for k, v in monitor.load_saved_data()["content"].items() if k != "1-12"}})
            monitor.notify_new_item = lambda it, deadline, art, submitted=None: False
            for _ in range(monitor.UNDELIVERED_ALERT_RUNS):
                monitor.run()
            failsafe.wait_idle()
            owed = json.loads(monitor.DATA_FILE.read_text(encoding="utf-8"))["undelivered"]
            assert owed["item:1-12"]["runs"] == monitor.UNDELIVERED_ALERT_RUNS, owed
            assert len([t for t in OUTBOX if "could not be sent" in t and "Item 12" in t]) == 1, OUTBOX
            monitor.notify_new_item = lambda it, deadline, art, submitted=None: sent_items.append(it["name"]) or True
            monitor.run()
            assert json.loads(monitor.DATA_FILE.read_text(encoding="utf-8"))["undelivered"] == {}
            assert any("have gone out" in t for t in OUTBOX), OUTBOX
        finally:
            for k, v in keep.items():
                setattr(monitor, k, v)

    # a reminder is saved the moment it is sent, and writing state keeps the run lock alive
    flags = []
    orig = monitor.send_telegram
    monitor.send_telegram = lambda text, url=None: True
    try:
        due = {"assign-1": {"type": "assign", "course": course, "name": "HW", "url": "u",
                            "timestamp": time.time() + 1800}}
        monitor.send_deadline_reminders(due, on_sent=lambda: flags.append(dict(due["assign-1"])))
    finally:
        monitor.send_telegram = orig
    assert len(flags) == 2 and flags[0].get("reminded_24h") and not flags[0].get("reminded_1h") and flags[1].get("reminded_1h")
    with tempfile.TemporaryDirectory() as tmp:
        lock = Path(tmp) / "monitor.lock"
        lock.write_text("1")
        os.utime(lock, (1, 1))
        orig = monitor.LOCK_FILE, monitor.DATA_FILE
        monitor.LOCK_FILE, monitor.DATA_FILE = lock, Path(tmp) / "s.json"
        try:
            monitor.save_data({"a": 1})
            assert lock.stat().st_mtime > 1e9, "saving state refreshes the lock"
        finally:
            monitor.LOCK_FILE, monitor.DATA_FILE = orig


def test_archive_boundary():
    root = archive.ARCHIVE_ROOT
    assert archive.ensure_inside_archive(root / "Big Data" / "x.pdf")
    for bad in (root.parent / "Other_Semester", root / ".." / "02-Second_Semester",
                Path("C:/Windows/x.txt")):
        try:
            archive.ensure_inside_archive(bad)
        except archive.ArchiveBoundaryError:
            continue
        raise AssertionError(f"{bad} should be rejected")
    assert archive._unique_path(root / "Big Data", "a.pdf", "updated 2026-10-10").name ==         "a (updated 2026-10-10).pdf"


def _approx(a, b, tol=1e-6):
    return a is not None and abs(a - b) < tol


def test_gpa_formulas():
    from uninotifier import gpa

    def course(hours=3, grade=None, retake=False, old=None):
        return {**gpa.new_course(), "hours": hours, "grade": grade, "retake": retake, "old_grade": old}

    courses = {"A": course(3, 80), "B": course(3, 60), "C": course(2, 90), "D": course()}
    assert _approx(gpa.semester_average(courses), 75.0), "ungraded courses are left out"
    assert gpa.semester_average({"D": course()}) is None
    assert gpa.DEFAULT_HOURS == 3 and gpa.PASS_MARK == 50
    assert not hasattr(gpa, "SCHOLARSHIP_MARK") and not hasattr(gpa, "standing")

    # CGPA = (CGPA * passed + semester average * semester hours) / (passed + semester hours)
    assert _approx(gpa.cumulative(courses, 70, 30), (70 * 30 + 75 * 8) / 38)
    assert gpa.cumulative(courses, None, None) is None and gpa.cumulative(courses, None, 30) is None
    assert _approx(gpa.cumulative(courses, None, 0), 75.0), "a first semester needs no CGPA"

    # a retake replaces the old grade in the points and adds no hours
    retake = {"R": course(3, 80, True, 55)}
    assert _approx(gpa.cumulative(retake, 70, 30), (70 * 30 - 55 * 3 + 80 * 3) / 30)
    mixed = {**retake, "A": course(3, 80)}
    assert _approx(gpa.cumulative(mixed, 70, 30), (70 * 30 - 55 * 3 + 80 * 3 + 80 * 3) / 33)
    assert _approx(gpa.semester_average(mixed), 80.0), "retakes still count in the semester average"

    # grade from marks once every part has one
    c = {**gpa.new_course(), "syllabus": [["Midterm", 30.0], ["Final", 70.0]], "marks": {"Midterm": 24.0}}
    assert gpa.course_grade(c) is None
    c["marks"]["Final"] = 50.0
    assert gpa.course_grade(c) == 74.0
    c["grade"] = 80.0
    assert gpa.course_grade(c) == 80.0, "a grade typed by hand wins"


def test_gpa_parsing_and_need():
    from uninotifier import gpa

    assert gpa.parse_number(" 85 % ") == 85 and gpa.parse_number("62.5") == 62.5
    for bad in ("abc", "101", "-1", ""):
        try:
            gpa.parse_number(bad)
        except gpa.InputError:
            continue
        raise AssertionError(bad)
    assert gpa.parse_cgpa("72.5 33") == (72.5, 33.0) and gpa.parse_cgpa("72.5, 33") == (72.5, 33.0)
    for bad in ("72", "101 30", "70 400"):
        try:
            gpa.parse_cgpa(bad)
        except gpa.InputError:
            continue
        raise AssertionError(bad)

    syl = gpa.parse_syllabus("Midterm 30, Assignments 10, Project: 20, Final Exam 40")
    assert syl == [["Midterm", 30.0], ["Assignments", 10.0], ["Project", 20.0], ["Final Exam", 40.0]]
    for bad, why in (("Midterm 30, Final 50", "add up"), ("Midterm 60, Quiz 40", "Final"),
                     ("Midterm 30, Midterm 30, Final 40", "twice"), ("Final 100", "at least two"),
                     ("Midterm, Final 40", "Couldn't read")):
        try:
            gpa.parse_syllabus(bad)
        except gpa.InputError as e:
            assert why in str(e), (bad, str(e))
            continue
        raise AssertionError(bad)

    marks = gpa.parse_marks("mid 24, Assignments 8/10, proj 15", syl)
    assert marks == {"Midterm": 24.0, "Assignments": 8.0, "Project": 15.0}, marks
    assert gpa.parse_marks("Midterm 12/15", syl) == {"Midterm": 24.0}, "scaled to the part's weight"
    for bad in ("Midterm 31", "Quiz 5", "", "Midterm"):
        try:
            gpa.parse_marks(bad, syl)
        except gpa.InputError:
            continue
        raise AssertionError(bad)

    course = {**gpa.new_course(), "syllabus": syl, "marks": marks}      # 47 of 60 before the final
    need = gpa.need_for(course, 60)
    assert need["earned"] == 47 and need["required"] == 13 and need["weight"] == 40 and need["missing"] == []
    assert gpa.need_for(course, 90)["required"] == 43, "above the final's weight: out of reach"
    assert gpa.need_for(gpa.new_course(), 60) is None
    partial = {**course, "marks": {"Midterm": 24.0}}
    assert gpa.need_for(partial, 60)["missing"] == ["Assignments", "Project"]


def test_gpa_screens_and_flow():
    bot = _package("uninotifier.bot")
    from uninotifier import gpa
    from uninotifier import gpa_ui

    names = ["Arabic Communication Skills", "Big Data", "Database Systems"]
    with tempfile.TemporaryDirectory() as tmp:
        gpa.STATE_FILE = Path(tmp) / "gpa.json"

        text, rows = gpa_ui.view_home(names)
        assert "Course Adjustments" in text and "Big Data" in text and "set your current CGPA" in text
        assert [r[0]["callback_data"] for r in rows[:3]] == ["gpa:c:0", "gpa:c:1", "gpa:c:2"]
        assert all(len(b["callback_data"].encode()) <= 64 for r in rows for b in r)
        assert gpa.load()["courses"]["Big Data"]["hours"] == 3, "3 credit hours by default"

        # quick toggles: hours
        text, rows = gpa_ui.on_callback(names, "gpa:c:0")
        assert "Credit hours:* 3" in text
        text, rows = gpa_ui.on_callback(names, "gpa:h:0:2")
        assert "Credit hours:* 2" in text and any(b["text"] == "✅ 2h" for b in rows[0])

        # grade entry through a question
        prompt = gpa_ui.on_callback(names, "gpa:g:0")
        assert isinstance(prompt, gpa_ui.Prompt) and prompt.kind == "grade"
        err = gpa_ui.on_text(names, prompt, "150")
        assert isinstance(err, str) and "between 0 and 100" in err
        text, rows = gpa_ui.on_text(names, prompt, "85")
        assert "Grade:* 85" in text
        gpa_ui.on_text(names, gpa_ui.on_callback(names, "gpa:g:1"), "40")
        text, _ = gpa_ui.view_home(names)
        assert "✅ 85" in text and "❌ 40" in text and "Below the 50% pass mark:* Big Data" in text
        assert "Semester Average:* 58%" in text and "60%" not in text   # (85*2 + 40*3) / 5

        # retake: asks for the old grade first; turning it off clears it
        prompt = gpa_ui.on_callback(names, "gpa:r:1")
        assert prompt.kind == "old_grade"
        assert not gpa.load()["courses"]["Database Systems"]["retake"], "not a retake until the old grade is given"
        gpa_ui.on_text(names, prompt, "45")
        c = gpa.load()["courses"]["Big Data"]
        assert c["retake"] and c["old_grade"] == 45
        assert "retake, was 45" in gpa_ui.view_home(names)[0]
        assert "previous attempt was a fail" in gpa_ui.view_course(names, 1)[0]
        gpa_ui.on_callback(names, "gpa:r:1")
        assert gpa.load()["courses"]["Big Data"]["retake"] is False

        # CGPA and passed hours -> projected CGPA
        prompt = gpa_ui.on_callback(names, "gpa:cg")
        gpa_ui.on_text(names, prompt, "70 30")
        text, _ = gpa_ui.view_home(names)
        assert "Projected CGPA:* 68.29%" in text, text           # (70*30 + 85*2 + 40*3) / 35

        # syllabus -> marks -> /need
        prompt = gpa_ui.on_callback(names, "gpa:m:0")
        assert "Set the syllabus" in prompt[0], "marks need a syllabus first"
        prompt = gpa_ui.on_callback(names, "gpa:s:0")
        assert prompt.kind == "syllabus"
        err = gpa_ui.on_text(names, prompt, "Midterm 30, Final 50")
        assert isinstance(err, str) and "add up to 80" in err
        gpa_ui.on_text(names, prompt, "Midterm 30, Assignments 30, Final 40")
        prompt = gpa_ui.on_callback(names, "gpa:m:0")
        assert prompt.kind == "marks"
        text, _ = gpa_ui.on_text(names, prompt, "Midterm 24, Assignments 17")
        assert "Midterm 24/30" in text and "Assignments 17/30" in text and "Final 40" in text

        # Grade Needed: scope -> target -> plan; graded courses are locked, the rest share the target
        text, rows = gpa_ui.on_callback(names, "gpa:need")
        assert "What are you aiming for?" in text and "/need" not in text
        assert [r[0]["callback_data"] for r in rows[:3]] == ["gpa:t:s", "gpa:t:c", "gpa:home"]
        text, rows = gpa_ui.on_callback(names, "gpa:t:c")
        assert [b["text"] for r in rows[:3] for b in r] == ["50", "60", "70", "75", "80", "85",
                                                            "90", "✏️ Custom Value"]
        assert rows[0][1]["callback_data"] == "gpa:n:c:60"
        gpa_ui.on_callback(names, "gpa:x:0")                     # Arabic: no grade any more, its syllabus marks stay
        text, rows = gpa_ui.on_callback(names, "gpa:n:s:60")
        # open: Arabic (2h, marks 41/60) and Database Systems (3h); locked: Big Data 40 (3h)
        # semester: (60*8 - 40*3) / 5 = 72 on each open course
        assert "🔒 *Big Data* · 3h · 40 (entered)" in text and "Needed on each remaining course:* 72%" in text, text
        assert "Final needed: *31 / 40*" in text, "Arabic: 72 - 41 earned before the final"
        assert _legacy_markdown_ok(text)
        assert "Out of reach" in gpa_ui.view_plan(names, "s", 95)[0]
        assert "Already secured" in gpa_ui.view_plan(names, "s", 10)[0]
        # cumulative (70% over 30 h): (65*38 - 70*30 - 40*3) / 5 = 50
        text, rows = gpa_ui.view_plan(names, "c", 65)
        assert "Needed on each remaining course:* 50%" in text and "CGPA at that pace:* 65%" in text, text
        assert rows[0][1]["callback_data"] == "gpa:n:s:65", "the same target as a semester one"
        assert gpa_ui.on_callback(names, "gpa:n:60")[0].startswith("📊 *Target Semester GPA: 60%*"), "old buttons"
        # a grade published on the portal is locked and wins over a typed one
        assert gpa.record_portal_grades({"Database Systems": 90, "Not A Course": 50}) == ["Database Systems"]
        text, _ = gpa_ui.view_plan(names, "s", 60)
        assert "🔒 *Database Systems* · 3h · 90 (portal)" in text and "Needed on each remaining course:* 45%" in text and "(and 50 to pass)" in text
        assert gpa.record_portal_grades({"Database Systems": 90}) == [], "unchanged: nothing to save"
        state = gpa.load()
        state["courses"]["Database Systems"]["portal_grade"] = None
        gpa.save(state)

        # the bot: commands, typed answers, cancel
        sent = []
        orig = (bot.send_view, bot.edit_view, bot.course_names)
        bot.send_view = lambda text, rows: sent.append((text, rows))
        bot.edit_view = lambda mid, text, rows: sent.append((f"edit:{mid}", text, rows))
        bot.course_names = lambda: names
        try:
            bot.handle_command("/gpa")
            assert "GPA & Grade Hub" in sent[-1][0]
            bot.handle_command("/need")
            assert "Grade Needed" in sent[-1][0] and "What are you aiming for?" in sent[-1][0]
            bot.handle_command("/need 60")
            assert "What are you aiming for?" in sent[-1][0], "no syntax: always the buttons"

            bot.handle_gpa_callback("gpa:v:c", 7)             # ✏️ Custom Value asks for the number
            assert "Send the target" in sent[-1][0] and bot._pending_prompt.kind == "target_c"
            bot.handle_command("77.5")
            assert sent[-1][0].startswith("🎓 *Target Cumulative CGPA: 77.5%*") and bot._pending_prompt is None

            bot.handle_gpa_callback("gpa:g:2", 7)             # asks, does not edit
            assert "Send the grade" in sent[-1][0] and sent[-1][1] == [gpa_ui.CANCEL_ROW]
            assert bot._pending_prompt is not None
            bot.handle_command("not a number")                # an unusable answer keeps the question open
            assert "Send just a number" in sent[-1][0] and bot._pending_prompt is not None
            bot.handle_command("92")
            assert bot._pending_prompt is None and "Grade:* 92" in sent[-1][0]

            bot.handle_gpa_callback("gpa:c:2", 7)             # a tapped button edits its message
            assert sent[-1][0] == "edit:7"
            bot.handle_gpa_callback("gpa:g:2", 7)
            bot.handle_gpa_callback("gpa:cancel", 7)
            assert bot._pending_prompt is None and sent[-1][0] == "edit:7" and "Course Adjustments" in sent[-1][1]
            bot.handle_gpa_callback("gpa:g:2", 7)
            bot.handle_command("/gpa")                        # any command drops the open question
            assert bot._pending_prompt is None

            # passed hours / CGPA taken from the transcript button arrive as a new message
            bot.handle_gpa_callback("gpa:use:33:72.5", 9)
            assert not sent[-1][0].startswith("edit:") and "72.5%" in sent[-1][0] and "33 record hours" in sent[-1][0]
        finally:
            bot.send_view, bot.edit_view, bot.course_names = orig
            bot._pending_prompt = None
        assert gpa.load()["cgpa"] == 72.5 and gpa.load()["passed_hours"] == 33

        # the GPA & Grade Hub: Grade Needed and Finals Watch only (Course Adjustments lives inside Grade Needed,
        # this term's grades inside Finals Watch)
        buttons = [b["callback_data"] for row in bot.view_academic()[1] for b in row if "callback_data" in b]
        assert buttons == ["gpa:need", "finals:status", "menu"], buttons
        from uninotifier import gpa_ui as _gpa_ui
        need_buttons = [b["callback_data"] for row in _gpa_ui.view_need_picker()[1] for b in row if "callback_data" in b]
        assert "gpa:home" in need_buttons, "Course Adjustments sits inside Grade Needed"
        hub_text = bot.view_academic()[0]
        assert "This Term" not in hub_text and "Course Adjustments" not in hub_text.split("Grade Needed")[0]
        record = [b["callback_data"] for row in bot.view_record()[1] for b in row if "callback_data" in b]
        assert record == ["portal:transcript", "portal:balance", "guest:start", "menu"], "no combined button"
        record_text = bot.view_record()[0]
        assert "verification code" not in record_text and "re-syncs" not in record_text
        assert {"gpa", "need", "grades", "transcript", "balance"} <= {c for c, _ in bot.COMMANDS}


def test_portal():
    bot = _package("uninotifier.bot")
    monitor = _package("uninotifier.monitor")
    from uninotifier import portal

    # on demand only: the scheduled monitor never touches the portal
    assert not hasattr(monitor, "portal")
    assert not re.search(r"^\s*(import|from)\s+portal\b", Path("monitor.py").read_text(encoding="utf-8"), re.M)

    raw = [[["Course", "Mark"], ["Big Data", " 85 "], ["", ""]],
           [["menu"], ["x"]],                         # a one-column table is layout, not data
           [["Code", "Name", "Hours"], ["1", "Databases", "3"]]]
    tables = portal.clean_tables(raw)
    assert tables == [[["Course", "Mark"], ["Big Data", "85"]], [["Code", "Name", "Hours"], ["1", "Databases", "3"]]]
    lines = portal.table_lines(tables)
    assert lines[0] == "*Course | Mark*" and lines[1] == "Big Data | 85"
    assert "more rows" in portal.table_lines(tables, limit=2)[-1]

    text = "الساعات المنجزة: 33\nالمعدل التراكمي 72.5\nالرصيد: -1,250.50"
    assert portal.find_figures(text) == {"passed_hours": 33.0, "cgpa": 72.5, "balance": -1250.5}
    assert portal.find_figures("Passed credit hours 45 | Cumulative GPA 68.4 | Balance 120") == {
        "passed_hours": 45.0, "cgpa": 68.4, "balance": 120.0}
    assert portal.find_figures("nothing useful") == {}

    # credentials: from .env, never asked for in chat
    with tempfile.TemporaryDirectory() as tmp:
        portal.ENV_FILE = Path(tmp) / ".env"
        saved_env = {k: os.environ.pop(k, None) for k in ("PORTAL_USERNAME", "PORTAL_PASSWORD")}
        try:
            assert portal.credentials() is None
            portal.ENV_FILE.write_text("PORTAL_USERNAME=209900001\nPORTAL_PASSWORD=secret\n", encoding="utf-8")
            assert portal.credentials() == ("209900001", "secret")
        finally:
            for k, v in saved_env.items():
                if v is not None:
                    os.environ[k] = v
    text = bot.portal_setup_text()
    assert "PORTAL_USERNAME" in text and "PORTAL_PASSWORD" in text and "never send" in text

    # the portal's verification code is relayed through the chat
    replies = []
    orig = bot.send_view
    bot.send_view = lambda text, rows: replies.append(text)
    try:
        assert not bot.handle_portal_code("123456"), "nothing is waiting for a code"
        import queue
        bot._portal_code = waiting = queue.Queue()
        assert not bot.handle_portal_code("what is this")
        assert bot.handle_portal_code("481 516") and waiting.get_nowait() == "481516"
        assert "Code received" in replies[-1]
    finally:
        bot.send_view, bot._portal_code = orig, None

    # results: screenshot + figures + the full text; the transcript offers "use in /gpa"
    sent_files, sent_msgs = [], []
    orig = (bot.notifier.send_png_file, bot.notifier.send_message, bot.notifier.send_long_message)
    bot.notifier.send_png_file = lambda tok, chat, png, name, caption="", url=None, rows=None: (
        sent_files.append((name, caption, rows)) or True)
    bot.notifier.send_message = lambda tok, chat, text, url=None, rows=None: sent_msgs.append((text, rows)) or True
    bot.notifier.send_long_message = lambda tok, chat, text, url=None, rows=None: sent_msgs.append((text, rows)) or True
    try:
        from uninotifier import gpa
        keep_gpa = gpa.STATE_FILE
        gpa_tmp = tempfile.TemporaryDirectory()
        gpa.STATE_FILE = Path(gpa_tmp.name) / "gpa.json"
        gpa.save({**gpa.load(), "courses": {"Linear Algebra (1)": gpa.new_course(), "Big Data": gpa.new_course()}})
        bot.send_portal_result(portal.Result("transcript", True, b"png", tables,
                                             {"passed_hours": 33.0, "cgpa": 72.5, "record_hours": 36.0, "record_cgpa": 71.25,
                                              "base_cgpa": 71.25, "base_hours": 36.0}, [], [],
                                             {"status": "منتظم", "retakes": {"0101221": 35.0}}))
        name, caption, rows = sent_files[-1]
        assert name == "ZUJ Portal - Transcript.png"
        assert "🎓 *Cumulative GPA:* 71.25% (Portal display: 72.5%)" in caption, caption
        assert "⏱️ *Passed Hours:* 33 hrs" in caption and "📌 *Academic Status:* Regular / Good Standing" in caption
        assert "📚 *Total Attempted Hours:* 36 hrs" in caption
        assert "🏁 *Remaining Hours:* 99 hrs" in caption, "your own transcript shows the hours left too (132 − 33 passed)"
        body = caption.split("Retrieved")[1].split("\n", 1)[1]        # without the clock: "6:18:35 PM" holds a 35
        for internal in ("35", "attempted", "synced", "record"):
            assert internal not in body, (internal, caption)
        assert rows == [[{"text": "🏛️ Open Registration Portal ↗️", "url": portal.LOGIN_URL}]], "no button to press"
        # under the hood: the 66-hour-style base and the retake are synced silently
        assert (gpa.load()["cgpa"], gpa.load()["passed_hours"], gpa.load()["base_source"]) == (71.25, 36.0, "portal")
        la = gpa.load()["courses"]["Linear Algebra (1)"]
        assert (la["retake"], la["old_grade"], la["retake_source"]) == (True, 35.0, "portal")
        assert not gpa.load()["courses"]["Big Data"]["retake"]
        # a transcript whose record couldn't be read leaves the retakes alone
        bot.send_portal_result(portal.Result("transcript", True, None, [], {"passed_hours": 33.0}))
        assert gpa.load()["courses"]["Linear Algebra (1)"]["retake"], "no retake data: nothing cleared"
        gpa.STATE_FILE = keep_gpa
        gpa_tmp.cleanup()
        assert any("Big Data | 85" in m[0] for m in sent_msgs), sent_msgs
        bot.send_portal_result(portal.Result("grades", False, b"png", [], {}, ["Home", "Student Services"]))
        assert "couldn't find the Grades page" in sent_files[-1][1] and "Student Services" in sent_files[-1][1]
        bot.send_portal_result(portal.Result("balance", True, None, [], {"balance": 120.0}, [],
                                             ["💰 *Total balance:* 120 JD — owed", "• First semester 2026/2027 · 120 JD"]))
        assert "Total balance:* 120" in sent_msgs[-1][0] and sent_msgs[-1][0].endswith("• First semester 2026/2027 · 120 JD"),             "no screenshot: the caption, with the details in it, goes out as ONE text message"
        bot.send_portal_result(portal.Result("transcript", True, b"png", [], {"passed_hours": 63.0, "cgpa": 71.3, "record_hours": 72.0, "record_cgpa": 71.26}, [],
                                             ["📅 *First semester 2026/2027*"]))
        assert "⏱️ *Passed Hours:* 63 hrs" in sent_files[-1][1], sent_files[-1][1]
    finally:
        bot.notifier.send_png_file, bot.notifier.send_message, bot.notifier.send_long_message = orig

    # no credentials: the bot explains .env and does not start a request
    replies = []
    orig = (bot.send_view, portal.credentials)
    bot.send_view = lambda text, rows: replies.append(text)
    portal.credentials = lambda: None
    try:
        assert bot.start_portal("grades") and "not configured" in replies[-1]
    finally:
        bot.send_view, portal.credentials = orig
    assert not bot._portal_lock.locked()


def test_day_separator_and_daily_recap():
    """The first automatic message of a day gets a date line above it; a checked day
    with nothing new ends with one recap — at 11:59 PM, or from the first check after
    midnight when the PC slept through it, unless that check finds something."""
    import threading
    from datetime import timedelta

    bot = _package("uninotifier.bot")
    from uninotifier import daily
    monitor = _package("uninotifier.monitor")

    assert daily.day_name(date(2026, 10, 13)) == "Tuesday, 13 Oct 2026"
    assert daily.separator_text(date(2026, 10, 13)) == "━━━━━━━━ 📅 Tuesday, 13 Oct 2026 ━━━━━━━━"
    assert daily.recap_text(date(2026, 10, 13), date(2026, 10, 13)) == (
        "✨ *Daily Recap:* No new files, assignments, or announcements posted today across your courses.")
    late = daily.recap_text(date(2026, 10, 12), date(2026, 10, 13))
    assert late.startswith("✨ *Daily Recap (Monday, 12 Oct):*") and "posted yesterday" in late, late

    with tempfile.TemporaryDirectory() as tmp:
        orig = (daily.CLAIM_DIR, notifier.send_message, telegram.post, notifier.before_send,
                daily.ensure_separator, daily.catch_up, daily.send_recap)
        keep = {k: getattr(monitor, k) for k in ("DATA_FILE", "USERNAME", "PASSWORD", "BOT_TOKEN", "CHAT_ID",
                                                 "run_scan", "notify_new_item", "send_schedule_messages",
                                                 "mfa_allowed", "notify_new_course")}
        daily.CLAIM_DIR = Path(tmp) / "claims"
        try:
            # through the real notifier: the separator goes first, once a day, silently
            posted = []

            class Ok:
                ok, status_code = True, 200
            telegram.post = lambda url, json=None, data=None, files=None, timeout=None: posted.append(data) or Ok()
            notifier.before_send = lambda: daily.ensure_separator("t", "c")
            assert notifier.send_message("t", "c", "first") and notifier.send_message("t", "c", "second")
            assert [d["text"] for d in posted] == [daily.separator_text(date.today()), "first", "second"], posted
            assert posted[0].get("disable_notification") and "disable_notification" not in posted[1]
            notifier.before_send = None

            sent, ok = [], {"v": False}
            notifier.send_message = lambda token, chat, text, url=None, rows=None, silent=False: \
                sent.append((text, silent)) or ok["v"]
            tue = datetime(2030, 10, 15, 8, 0)                  # never "today": the check above claimed today
            daily.ensure_separator("t", "c", tue)                # undelivered: the next message tries again
            ok["v"] = True
            daily.ensure_separator("t", "c", tue)
            daily.ensure_separator("t", "c", tue.replace(hour=22))
            daily.ensure_separator("t", "c", tue + timedelta(days=1))
            assert sent == [(daily.separator_text(date(2030, 10, 15)), True)] * 2 + [
                (daily.separator_text(date(2030, 10, 16)), True)], sent

            # the day record: checked / found something new
            sent.clear()
            days = {}
            daily.note_check(days, False, datetime(2026, 10, 12, 9, 0))
            daily.note_check(days, False, datetime(2026, 10, 13, 9, 0))
            daily.note_check(days, True, datetime(2026, 10, 13, 15, 0))
            daily.note_check(days, False, datetime(2026, 10, 13, 16, 0))   # a later quiet check undoes nothing
            assert days == {"checked": ["2026-10-12", "2026-10-13"], "active": ["2026-10-13"]}, days

            # 11:59 PM: once, only for a quiet day
            assert not daily.send_recap("t", "c", days, datetime(2026, 10, 12, 23, 58))
            assert daily.send_recap("t", "c", days, datetime(2026, 10, 12, 23, 59))
            assert not daily.send_recap("t", "c", days, datetime(2026, 10, 12, 23, 59, 40)), "bot + monitor: once"
            assert not daily.catch_up("t", "c", days, False, datetime(2026, 10, 13, 7, 0)), "already sent"
            assert not daily.send_recap("t", "c", days, datetime(2026, 10, 13, 23, 59)), "a busy day"
            assert sent == [(daily.recap_text(date(2026, 10, 12), date(2026, 10, 12)), False)], sent

            # asleep at 11:59 PM: the first check of the morning sends yesterday's
            daily.note_check(days, False, datetime(2026, 10, 14, 10, 0))
            assert daily.catch_up("t", "c", days, False, datetime(2026, 10, 15, 7, 5))
            assert not daily.catch_up("t", "c", days, False, datetime(2026, 10, 15, 7, 20))
            assert sent[-1] == (daily.recap_text(date(2026, 10, 14), date(2026, 10, 15)), False), sent
            # ... unless that check found new items (posted the evening before?)
            daily.note_check(days, False, datetime(2026, 10, 15, 9, 0))
            assert not daily.catch_up("t", "c", days, True, datetime(2026, 10, 16, 7, 0))
            assert not daily.catch_up("t", "c", days, False, datetime(2026, 10, 16, 7, 15))
            # a day with no successful check: nothing can be said about it
            assert not daily.catch_up("t", "c", days, False, datetime(2026, 10, 18, 7, 0))
            assert len(sent) == 2, sent
            # an undelivered recap is retried
            daily.note_check(days, False, datetime(2026, 10, 18, 9, 0))
            ok["v"] = False
            assert not daily.send_recap("t", "c", days, datetime(2026, 10, 18, 23, 59))
            ok["v"] = True
            assert daily.send_recap("t", "c", days, datetime(2026, 10, 18, 23, 59, 20))
            # only the last week is kept
            for n in range(1, 11):
                daily.note_check(days, True, datetime(2026, 11, n, 9, 0))
            assert len(days["checked"]) == daily.KEEP_DAYS and days["active"][0] == "2026-11-04", days

            # the bot: separators for its clock threads only, never for replies to commands
            threads = []
            daily.ensure_separator = lambda token, chat, now=None: threads.append(threading.current_thread().name)
            bot.before_send()
            t = threading.Thread(target=bot.before_send, name="lectures")
            t.start()
            t.join()
            assert threads == ["lectures"], threads

            # monitor.run: what counts as "something new" for the day
            course = "20261-Big Data-AllSec-3-"
            content = {"1-1": {"name": "Lecture 1", "type": "resource", "course": course, "cmid": "1",
                               "url": "https://x/mod/resource/view.php?id=1"}}
            monitor.DATA_FILE = Path(tmp) / "saved_items.json"
            monitor.USERNAME = monitor.PASSWORD = monitor.BOT_TOKEN = monitor.CHAT_ID = "x"
            monitor.send_schedule_messages = lambda saved: None
            monitor.mfa_allowed = lambda saved: False
            monitor.notify_new_course = lambda course: True
            monitor.notify_new_item = lambda it, deadline, art, submitted=None: True
            monitor.run_scan = lambda saved, notify_new, mfa_handler=None: {
                "courses": {"1": {"name": course, "url": "https://x/c"}}, "content": dict(content),
                "deadlines": {}, "artifacts": {}, "updates": {}, "edits": {}, "posts": {},
                "forum_posts": {}, "submissions": {}, "link_changes": [], "seconds": 1.0, "browser_used": False}
            found = []
            daily.catch_up = lambda token, chat, days, found_new, now=None: found.append(found_new) or False
            daily.send_recap = lambda token, chat, days, now=None: False
            monitor.run()                                        # the silent baseline is not news
            content["1-2"] = {**content["1-1"], "name": "Announcements", "type": "forum", "cmid": "2"}
            monitor.run()                                        # nor is an empty forum
            content["1-3"] = {**content["1-1"], "name": "Lecture 2", "cmid": "3"}
            monitor.run()
            assert found == [False, False, True], found
            today = str(date.today())
            state = json.loads(monitor.DATA_FILE.read_text(encoding="utf-8"))
            assert state["days"] == {"checked": [today], "active": [today]}, state["days"]
        finally:
            (daily.CLAIM_DIR, notifier.send_message, telegram.post, notifier.before_send,
             daily.ensure_separator, daily.catch_up, daily.send_recap) = orig
            for k, v in keep.items():
                setattr(monitor, k, v)


def _legacy_markdown_ok(text: str) -> bool:
    """Telegram's legacy Markdown rejects a message with an unclosed * _ or `. Read left to right the
    way Telegram does: a backslash escapes _ * ` [ outside a code span, inside one it is a letter."""
    text = re.sub(r"(?<!\\)\[[^\]]*\]\([^)]*\)", "", text)            # [label](url): its underscores are the url's
    i, in_code, stars, unders = 0, False, 0, 0
    while i < len(text):
        c = text[i]
        if in_code:
            in_code = c != "`"
        elif c == "\\" and i + 1 < len(text) and text[i + 1] in "_*`[":
            i += 1
        elif c == "`":
            in_code = True
        elif c == "*":
            stars += 1
        elif c == "_":
            unders += 1
        i += 1
    return not in_code and stars % 2 == 0 and unders % 2 == 0


def test_help_hub_dashboard_and_campus():
    """/help is a small site: every page is reachable, valid Markdown, fits one message
    and leads back to the hub. The dashboard tells the truth about the monitor."""
    bot = _package("uninotifier.bot")
    from uninotifier import guide

    # walk every page reachable from the hub
    seen, queue_ = {}, ["help"]
    while queue_:
        action = queue_.pop()
        if action in seen:
            continue
        seen[action] = text, rows = guide.view(action.partition(":")[2], "🟢 *Live Monitor Active*")
        assert _legacy_markdown_ok(text) and len(text) < notifier.MESSAGE_LIMIT, (action, text)
        buttons = [b for row in rows for b in row]
        assert all(len(b["callback_data"].encode()) <= 64 for b in buttons if "callback_data" in b)
        if action != "help":
            assert guide.HUB_BUTTON in rows[-1], f"{action}: no way back to the hub"
        for b in buttons:
            if b.get("callback_data", "").startswith("help"):
                assert b["callback_data"] == "help" or b["callback_data"][5:] in guide.PAGES, b
                queue_.append(b["callback_data"])
    assert set(seen) == {"help"} | {f"help:{k}" for k in guide.PAGES}, "every page is linked"
    hub_text, hub_rows = seen["help"]
    assert [r[0]["text"] for r in hub_rows] == [
        "📚 Schedule & Moodle Monitoring", "⏳ Tasks & Custom Reminders Engine",
        "📊 Smart GPA, Target Allocation & Finals Watch", "📜 Academic Records & Financials",
        "⚙️ Architecture, Healthchecks & Stealth Engine", "🏠 Main Menu"], "one section per hub, no Command Reference"
    assert "commands" not in guide.PAGES
    assert "🟢" in hub_text and "never saved twice" in seen["help:schedule"][0]
    # one self-contained message per section: no sub-pages, no carousel, only the way back
    assert set(guide.PAGES) == set(guide.SECTIONS) and len(guide.SECTIONS) == 5
    for key in guide.SECTIONS:
        text, rows = seen[f"help:{key}"]
        assert rows == [[guide.HUB_BUTTON, guide.MENU_BUTTON]], (key, rows)
        assert not re.search(r"Previous|Next|◀️|▶️", json.dumps(rows, ensure_ascii=False)), key
    # the buttons of the old two-level guide open the section they belonged to
    for old_key, section in guide.OLD_PAGES.items():
        assert guide.view(old_key)[0] == guide.view(section)[0], old_key
    assert "Check Grades Now" in seen["help:gpa"][0] and "sits inside Grade Needed" in seen["help:gpa"][0]
    assert "Sunday to Wednesday" in seen["help:tasks"][0] and "replaces" in seen["help:gpa"][0]
    assert guide.view("no-such-page")[0] == guide.view_hub()[0]
    assert guide.view("commands")[0] == guide.view_hub()[0], "an old Command Reference button lands on the hub"
    assert {"help", "remind", "need"} <= {c for c, _ in bot.COMMANDS}

    # dispatch: /help opens the hub; a button edits the page in place; /courses is the submenu
    sent, edited = [], []
    orig = (bot.send_view, bot.edit_view, bot.api)
    bot.send_view = lambda text, rows: sent.append(text)
    bot.edit_view = lambda message_id, text, rows: edited.append(text)
    bot.api = lambda method, **data: None
    try:
        bot.handle_command("/help")
        bot.handle_callback({"id": "1", "data": "help:tasks", "message": {"message_id": 5}})
        bot.handle_callback({"id": "2", "data": "help", "message": {"message_id": 5}})
        bot.handle_command("/courses")
        bot.handle_callback({"id": "3", "data": "courselist", "message": {"message_id": 6}})
        bot.handle_command("/campus")
    finally:
        bot.send_view, bot.edit_view, bot.api = orig
    assert sent[0].startswith("📖") and "Your own reminders" in edited[0] and edited[1].startswith("📖")
    assert sent[1].startswith("📚 *My Courses & Schedule*") and "Section 03" in edited[2]
    assert sent[2].startswith("❓ Unknown command"), "Campus Guide was removed"
    for text in (bot.view_courses()[0], bot.view_course_list()[0], bot.view_tasks()[0]):
        assert _legacy_markdown_ok(text), text

    # the dashboard's monitor line
    orig = bot.load_last_run
    try:
        for run, expected in (({}, "⚪"), ({"finished": time.time() - 300, "ok": True}, "🟢 *Live Monitor Active*"),
                              ({"finished": time.time() - 3 * 3600, "ok": True}, "🟡 *Monitor paused* — last check 3 h ago"),
                              ({"finished": time.time() - 60, "ok": False, "error": "exit code 1"}, "🟠")):
            bot.load_last_run = lambda run=run: run
            assert bot.monitor_status().startswith(expected), (run, bot.monitor_status())
            assert _legacy_markdown_ok(bot.view_menu()[0])
    finally:
        bot.load_last_run = orig


OWNER_ID = "424242"


def _tg(user_id=OWNER_ID, chat_id=OWNER_ID, chat_type="private", text=None, data=None, is_bot=False, **user_extra):
    """An update the way Telegram sends it: a message (text) or a button press (data)."""
    user = {"id": int(user_id), "is_bot": is_bot, "first_name": "T", **user_extra}
    chat = {"id": int(chat_id), "type": chat_type}
    if data is not None:
        return {"update_id": 1, "callback_query": {"id": "cb", "from": user, "data": data,
                                                    "message": {"message_id": 5, "chat": chat}}}
    return {"update_id": 1, "message": {"message_id": 1, "from": user, "chat": chat, "text": text}}


def _as_owner(bot):
    """The bot and the access gate believe the owner is OWNER_ID; returns what to restore."""
    keep = (bot.CHAT_ID, access.ADMIN_USER_ID)
    bot.CHAT_ID, access.ADMIN_USER_ID = OWNER_ID, OWNER_ID
    return keep


def test_update_routing_from_telegram():
    """handle_update is what every Telegram message and button press goes through: the owner's
    own are routed (a local variable once shadowed the chat module and made every update fail),
    everyone else's is dropped."""
    bot = _package("uninotifier.bot")
    keep_owner = _as_owner(bot)
    seen, answered = [], []
    orig = (bot.handle_command, bot.handle_callback, bot.api)
    bot.handle_command = lambda text, message_id=None: seen.append(("command", text))
    bot.handle_callback = lambda cb: seen.append(("callback", cb["data"]))
    bot.api = lambda method, **data: answered.append(method)
    try:
        bot.handle_update(_tg(text="/start"))
        bot.handle_update(_tg(data="menu"))
        bot.handle_update(_tg(user_id="999", chat_id="999", text="/start"))
        bot.handle_update(_tg(user_id="999", chat_id="999", data="menu"))
    finally:
        bot.handle_command, bot.handle_callback, bot.api = orig
        bot.CHAT_ID, access.ADMIN_USER_ID = keep_owner
    assert seen == [("command", "/start"), ("callback", "menu")], seen
    assert answered == [], "a stranger's button press is not even acknowledged"


def test_access_lock():
    """Only the owner is served: the id must match on `from` AND on a private chat of that same id,
    anything else is dropped silently (no reply, no button answer, no API call at all), the owner is
    told once per unknown id (never what was typed) and at most a few times a day, and a missing
    or malformed ADMIN_USER_ID fails closed."""
    bot = _package("uninotifier.bot")
    keep_owner = _as_owner(bot)
    calls, handled = [], []
    orig = (bot.handle_command, bot.handle_callback, bot.api)
    bot.api = lambda method, **data: calls.append(method)
    bot.handle_command = lambda text, message_id=None: handled.append(text)
    bot.handle_callback = lambda cb: handled.append(cb["data"])
    _fresh_failsafe()
    access._logged.clear()
    access._handled.clear()
    access._alerts.clear()
    try:
        # the owner, in every shape that is legitimately theirs
        assert access.is_owner(_tg(text="hi")) and access.is_owner(_tg(data="menu"))
        # ... and everything that is not
        strangers = {
            "another user, own private chat": _tg(user_id="777", chat_id="777", text="/start"),
            "another user writing in the owner's chat id": _tg(user_id="777", chat_id=OWNER_ID, text="/start"),
            "the owner's id, but from a group": _tg(chat_type="group", chat_id="-100123", text="/start"),
            "the owner's id in a supergroup": _tg(chat_type="supergroup", text="/start"),
            "a bot with the owner's id": _tg(is_bot=True, text="/start"),
            "the owner's user in someone else's chat": _tg(chat_id="777", text="/start"),
            "a channel post (no sender)": {"update_id": 2, "message": {"chat": {"id": -1001, "type": "channel"}, "text": "x"}},
            "no sender, no chat": {"update_id": 3, "message": {"text": "x"}},
            "an update of another kind": {"update_id": 4, "edited_message": {"from": {"id": int(OWNER_ID)}, "chat": {"id": int(OWNER_ID), "type": "private"}, "text": "x"}},
            "a stranger's button press": _tg(user_id="777", chat_id="777", data="menu"),
            "a button of the owner's chat pressed by someone else": _tg(user_id="777", chat_id=OWNER_ID, data="menu"),
        }
        for name, update in strangers.items():
            assert not access.is_owner(update), name
            bot.handle_update(update)
        assert handled == [] and calls == [], ("nothing is handled and nothing is sent to anyone", handled, calls)
        failsafe.wait_idle()
        # the owner is told, once per id: 777, -100123/owner-group, bot, unknown ... each distinct sender id once
        told = [t for t in OUTBOX if "not you" in t]
        assert told and all("Someone who is not you" in t for t in told)
        assert len(told) == 3, ("one alert per distinct sender id: 777, the owner's id from a group/bot, and 'unknown'", told)
        first = [t for t in told if "id 777" in t]
        assert len(first) == 1, "once per unknown id"
        assert "ignored it" in first[0] and "/start" not in first[0] and "hunter2" not in first[0], "never what was typed"
        assert _legacy_markdown_ok(first[0])

        # a flood of different strangers cannot become a flood of alerts
        OUTBOX.clear()
        failsafe.STATE_FILE.unlink(missing_ok=True)
        access._alerts.clear()
        now = 1_000_000.0
        for n in range(30):
            access.deny(_tg(user_id=str(5000 + n), chat_id=str(5000 + n), text="spam"), now=now + n)
        failsafe.wait_idle()
        assert len([t for t in OUTBOX if "not you" in t]) == access.MAX_ALERTS_PER_DAY
        # a name that tries to break the message does not
        OUTBOX.clear()
        failsafe.STATE_FILE.unlink(missing_ok=True)
        access._alerts.clear()
        access.deny(_tg(user_id="888", chat_id="888", text="x", username="a_b*c`d", first_name="[x](http://evil)"), now=now + 99999)
        failsafe.wait_idle()
        assert len(OUTBOX) == 1 and _legacy_markdown_ok(OUTBOX[0]), OUTBOX

        # fails closed: no / malformed ADMIN_USER_ID means nobody is the owner, and the bot will not start
        for bad in ("", "abc", "-5", "0", "12 34", "1" * 30):
            access.ADMIN_USER_ID = bad if access._ID.fullmatch(bad) else ""
            assert not access.is_owner(_tg(text="hi")) and access.problems(OWNER_ID), bad
        access.ADMIN_USER_ID = OWNER_ID
        assert access.problems(OWNER_ID) == [] and access.problems("555"), "the answering chat must be the owner's"
        assert access.problems("")
    finally:
        bot.handle_command, bot.handle_callback, bot.api = orig
        bot.CHAT_ID, access.ADMIN_USER_ID = keep_owner

    # configuration: ADMIN_USER_ID wins, TELEGRAM_CHAT_ID is the fallback (and is flagged), junk is refused
    saved_env = {k: os.environ.get(k) for k in ("ADMIN_USER_ID", "TELEGRAM_CHAT_ID")}
    try:
        os.environ["ADMIN_USER_ID"], os.environ["TELEGRAM_CHAT_ID"] = " 123456 ", "999"
        assert access._admin_from_env() == "123456" and access.from_chat_id is False
        os.environ["ADMIN_USER_ID"] = ""
        assert access._admin_from_env() == "999" and access.from_chat_id is True
        os.environ["ADMIN_USER_ID"] = "not-a-number"
        assert access._admin_from_env() == "", "junk fails closed instead of falling back"
    finally:
        for k, v in saved_env.items():
            os.environ.pop(k, None)
            if v is not None:
                os.environ[k] = v
        access._admin_from_env()
        access.ADMIN_USER_ID = keep_owner[1]


def test_modules_import_quietly_and_names_resolve():
    """Importing any module is side-effect free (a tool once sent a Telegram message on
    import), and every `module.name` the code uses exists — the checks a split or a
    rename can silently break (the showcase tool lost monitor.deliver that way)."""
    import ast
    import importlib
    import subprocess

    probe = (
        "import importlib, pkgutil, socket, sys\n"
        "def blocked(*a, **k): raise AssertionError('network used while importing')\n"
        "socket.socket.connect = blocked\n"
        "socket.create_connection = blocked\n"
        "import uninotifier\n"
        "for m in pkgutil.walk_packages(uninotifier.__path__, 'uninotifier.'):\n"
        "    importlib.import_module(m.name)\n"
    )
    # ... and touches no file: the child imports everything into an empty state folder, which must
    # stay empty (views.py once read schedule.json at import and wrote schedule.last_good.json)
    with tempfile.TemporaryDirectory() as empty:
        done = subprocess.run([sys.executable, "-c", probe], capture_output=True, text=True,
                              encoding="utf-8", errors="replace", timeout=120,
                              env={**os.environ, "UNINOTIFIER_DATA": empty})
        assert done.returncode == 0 and not done.stdout.strip(), (done.stdout + done.stderr)[-600:]
        written = sorted(p.name for p in Path(empty).rglob("*"))
        assert written == [], ("importing wrote into the state folder", written)

    root = Path(importlib.import_module("uninotifier").__file__).parent
    missing = []
    for path in sorted(root.rglob("*.py")):
        tree = ast.parse(path.read_text(encoding="utf-8"))
        aliases = {}
        for node in ast.walk(tree):
            if isinstance(node, ast.ImportFrom) and node.level == 0 and (node.module or "").startswith("uninotifier"):
                for a in node.names:
                    try:
                        mod = importlib.import_module(f"{node.module}.{a.name}")
                    except ImportError:
                        continue                       # a plain name, not a submodule
                    aliases[a.asname or a.name] = mod
        for node in ast.walk(tree):
            if (isinstance(node, ast.Attribute) and isinstance(node.value, ast.Name)
                    and node.value.id in aliases and not hasattr(aliases[node.value.id], node.attr)):
                missing.append(f"{path.relative_to(root)}:{node.lineno} {node.value.id}.{node.attr}")
    assert not missing, missing


def _fresh_failsafe():
    failsafe.wait_idle()          # a delivery an earlier test started in the background may still be writing the file
    failsafe.STATE_FILE.unlink(missing_ok=True)
    failsafe._inflight.clear()
    OUTBOX.clear()


def test_failsafe_reports_once_and_resolves():
    """A problem reaches you once, not once per occurrence; a different problem, or
    the same one after the cool-down, reaches you again; "fine again" follows once;
    nothing in it can raise, leak the token or break Markdown."""
    _fresh_failsafe()
    try:
        raise ValueError("a_b *c* `d` 123456789:AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA")
    except ValueError as e:
        exc = e
    assert failsafe.report("k", "Something *broke*", detail="Fix _this_", exc=exc) is True
    assert len(OUTBOX) == 1 and "needs attention" in OUTBOX[0] and "ValueError" in OUTBOX[0]
    assert "AAAAAAAAAAAAAAAAAAAA" not in OUTBOX[0] and _legacy_markdown_ok(OUTBOX[0]), OUTBOX[0]
    assert not failsafe.report("k", "Something broke", exc=exc) and not failsafe.report("k", "Something broke")
    assert len(OUTBOX) == 1, "held back during the cool-down"
    assert failsafe.report("k", "Something broke", fingerprint="other") and len(OUTBOX) == 2, "a different instance"
    state = json.loads(failsafe.STATE_FILE.read_text(encoding="utf-8"))
    state["k"]["at"] = 0
    state["k"]["suppressed"] = 2
    failsafe.STATE_FILE.write_text(json.dumps(state), encoding="utf-8")
    assert failsafe.report("k", "Again", fingerprint="other") and "2 similar" in OUTBOX[-1], "after the cool-down"

    assert failsafe.resolve("k", "✅ fine") and OUTBOX[-1] == "✅ fine"
    assert not failsafe.resolve("k", "✅ fine") and not failsafe.resolve("never-reported", "x"), "once, and only if open"

    # background delivery, and a failing door out
    assert failsafe.report("bg", "In the background", wait=False)
    failsafe.wait_idle()
    assert OUTBOX[-1].count("In the background") == 1
    orig = failsafe.send
    failsafe.send = lambda *a, **k: (_ for _ in ()).throw(OSError("network down"))
    try:
        assert failsafe.report("down", "Network is down") is False, "never raises"
        assert "down" not in json.loads(failsafe.STATE_FILE.read_text(encoding="utf-8")), "a failed send is retried next time"
    finally:
        failsafe.send = orig
    assert failsafe.report("down", "Network is back") is True

    # bug signatures: the same bug, the same signature
    assert failsafe.signature(exc).startswith("ValueError@test_architecture.py:")

    # whatever escapes a thread is reported (and then still printed by the default hook)
    import threading
    hooks = (sys.excepthook, threading.excepthook)
    err = sys.stderr
    sys.stderr = open(os.devnull, "w")
    try:
        failsafe.install_excepthooks("bot")
        t = threading.Thread(target=lambda: 1 / 0, name="worker")
        t.start()
        t.join()
        failsafe.wait_idle()
    finally:
        sys.stderr.close()
        sys.stderr = err
        sys.excepthook, threading.excepthook = hooks
    assert "background thread" in OUTBOX[-1] and "worker" in OUTBOX[-1] and "ZeroDivisionError" in OUTBOX[-1]


def test_bot_contains_and_reports_failures():
    """One failing button or message neither stops the bot nor goes unseen: the
    person is answered, the bug is reported (once), and what they typed is not logged."""
    _fresh_failsafe()
    bot = _package("uninotifier.bot")
    keep_owner = _as_owner(bot)
    api_calls, records = [], []
    orig = (bot.handle_command, bot.handle_callback, bot.api)
    bot.api = lambda method, **data: api_calls.append((method, data))
    bot.handle_command = lambda text, message_id=None: 1 / 0
    bot.handle_callback = lambda cb: {}["missing"]
    handler = logging.Handler()
    handler.emit = lambda record: records.append(record.getMessage())
    logging.getLogger("uninotifier").addHandler(handler)
    try:
        bot.dispatch({**_tg(text="hunter2-secret"), "update_id": 7})
        bot.dispatch({**_tg(data="x"), "update_id": 8, "callback_query": {**_tg(data="x")["callback_query"], "id": "cb1"}})
        bot.dispatch({**_tg(user_id="999", chat_id="999", text="hi"), "update_id": 9})   # a stranger: dropped before any handler
        failsafe.wait_idle()
    finally:
        logging.getLogger("uninotifier").removeHandler(handler)
        bot.handle_command, bot.handle_callback, bot.api = orig
        bot.CHAT_ID, access.ADMIN_USER_ID = keep_owner
    methods = [m for m, _ in api_calls]
    assert methods == ["sendMessage", "answerCallbackQuery"], methods
    assert "Something went wrong" in api_calls[0][1]["text"] and api_calls[1][1]["callback_query_id"] == "cb1"
    failures = [t for t in OUTBOX if "failed to handle" in t]
    assert len(failures) == 2 and "message" in failures[0] and "button press" in failures[1], OUTBOX
    assert len([t for t in OUTBOX if "not you" in t]) == 1, "the stranger was dropped, and the owner told once"
    assert not any("hunter2" in r for r in records), "the update's text must not be logged"
    assert all(_legacy_markdown_ok(t) for t in OUTBOX)


def test_bot_heartbeat_and_monitor_watch():
    """The bot's pulse is written as its loops turn; the scheduled monitor notices when
    one stops (not right after Windows starts), says so at most once, and says when it is back.
    The bot also has its own healthchecks.io check."""
    from uninotifier import healthcheck
    from uninotifier.monitor import alerts, state as mstate

    heartbeat.FILE.unlink(missing_ok=True)
    heartbeat._state.pop("poll", None)
    heartbeat._state.pop("ticker", None)
    heartbeat._last_write = 0
    assert heartbeat.read() == {} and not heartbeat.fresh_in_memory()
    heartbeat.beat("poll", now=1000)
    heartbeat.beat("ticker", now=1001)                 # inside the write throttle: memory only
    assert heartbeat.read()["poll"] == 1000 and "ticker" not in heartbeat.read()
    heartbeat._last_write = 0
    heartbeat.beat("ticker", now=1002)
    assert heartbeat.stale_parts(heartbeat.read(), now=1100) == []
    assert heartbeat.stale_parts(heartbeat.read(), now=1000 + heartbeat.STALE_SECONDS + 5) == ["poll", "ticker"]

    sent = []
    keep = (alerts.send_telegram, mstate.DATA_FILE, mstate.LOCK_FILE, alerts.selfheal.restart_bot)
    alerts.send_telegram = lambda text, url=None: sent.append(text) or True
    alerts.selfheal.restart_bot = lambda pid: False          # this test is about the alerts: a restart that cannot happen
    tmp = Path(tempfile.mkdtemp(prefix="uninotifier-watch-"))
    mstate.DATA_FILE, mstate.LOCK_FILE = tmp / "saved_items.json", tmp / "monitor.lock"
    saved = mstate.load_saved_data()
    long_up = 99999
    try:
        now = 5000.0
        heartbeat.FILE.unlink()
        alerts.check_bot(saved, now, long_up)
        assert sent == [], "no pulse on disk: the bot never ran here"
        heartbeat.FILE.write_text(json.dumps({"poll": now - 30, "ticker": now - 20}), encoding="utf-8")
        alerts.check_bot(saved, now, long_up)
        assert sent == [] and not saved["watch"], "fresh"
        heartbeat.FILE.write_text(json.dumps({"poll": now - 30, "ticker": now - 3600}), encoding="utf-8")
        alerts.check_bot(saved, now, uptime=120)
        assert sent == [], "just after Windows started: give the bot time"
        alerts.check_bot(saved, now, long_up)
        assert len(sent) == 1 and "isn't responding" in sent[0] and "lecture clock" in sent[0] and _legacy_markdown_ok(sent[0])
        alerts.check_bot(saved, now + 900, long_up)
        assert len(sent) == 1, "once per cool-down"
        heartbeat.FILE.write_text(json.dumps({"poll": now + 1000, "ticker": now + 1000}), encoding="utf-8")
        alerts.check_bot(saved, now + 1010, long_up)
        assert len(sent) == 2 and "bot is back" in sent[1] and not saved["watch"]
        alerts.check_bot(saved, now + 1020, long_up)
        assert len(sent) == 2, "back is said once"
    finally:
        alerts.send_telegram, mstate.DATA_FILE, mstate.LOCK_FILE, alerts.selfheal.restart_bot = keep
        heartbeat.FILE.unlink(missing_ok=True)

    # the bot's own healthchecks.io check, independent of the monitor's
    calls = []
    orig_post = healthcheck.requests.post
    healthcheck.requests.post = lambda url, data=None, timeout=None: calls.append(url) or type(
        "R", (), {"raise_for_status": lambda self: None})()
    saved_env = {k: os.environ.pop(k, None) for k in (healthcheck.ENV_NAME, healthcheck.BOT_ENV_NAME)}
    try:
        assert healthcheck.ping_bot() is False and calls == [], "unset = disabled"
        os.environ[healthcheck.ENV_NAME] = "https://hc-ping.com/monitor-uuid"
        assert healthcheck.ping_bot() is False and calls == [], "the monitor's URL is not the bot's"
        os.environ[healthcheck.BOT_ENV_NAME] = "https://hc-ping.com/bot-uuid/"
        assert healthcheck.ping_bot() is True and calls == ["https://hc-ping.com/bot-uuid"]
    finally:
        healthcheck.requests.post = orig_post
        for k, v in saved_env.items():
            os.environ.pop(k, None)
            if v is not None:
                os.environ[k] = v


def test_bot_self_heal():
    """A bot that stops beating is restarted by the next scheduled check (a few times,
    spaced out, only when it is really this project's bot.py), you hear about it only if
    that does not help, and a recovery says it was done for you."""
    from uninotifier import selfheal
    from uninotifier.monitor import alerts, state as mstate

    # the restart itself: only a process that really is this project's bot.py is ended
    commands, listing = [], {"out": ""}

    class Done:
        def __init__(self, out="", code=0):
            self.stdout, self.stderr, self.returncode = out, "", code

    def fake_run(cmd, timeout):
        commands.append(cmd[0].lower())
        return Done(listing["out"]) if cmd[0].lower().startswith("powershell") else Done()

    keep_run, keep_sleep = selfheal._run, selfheal.time.sleep
    selfheal._run, selfheal.time.sleep = fake_run, lambda s: None
    try:
        root = str(selfheal.paths.ROOT)
        ours = f'"C:\\py\\pythonw.exe" "{root}\\bot.py"'
        listing["out"] = ours
        assert selfheal.is_our_bot(4242) and not selfheal.is_our_bot(0) and not selfheal.is_our_bot(None)
        assert not selfheal.is_our_bot(True) and not selfheal.is_our_bot("12")
        listing["out"] = '"C:\\Windows\\notepad.exe" bot.py'
        assert not selfheal.is_our_bot(4242), "a recycled pid: same name is not enough"
        listing["out"] = ""
        assert not selfheal.is_our_bot(4242), "no such process"

        commands.clear()
        listing["out"] = ours
        assert selfheal.restart_bot(4242) is True
        assert commands == ["powershell.exe", "taskkill", "schtasks"], commands
        commands.clear()
        listing["out"] = ""
        assert selfheal.restart_bot(4242) is True and "taskkill" not in commands, "nothing to end: only start the task"
    finally:
        selfheal._run, selfheal.time.sleep = keep_run, keep_sleep

    # the policy in the scheduled check
    sent, restarts = [], []
    keep = (alerts.send_telegram, mstate.DATA_FILE, mstate.LOCK_FILE, alerts.selfheal.restart_bot)
    alerts.send_telegram = lambda text, url=None: sent.append(text) or True
    alerts.selfheal.restart_bot = lambda pid: restarts.append(pid) or True
    tmp = Path(tempfile.mkdtemp(prefix="uninotifier-heal-"))
    mstate.DATA_FILE, mstate.LOCK_FILE = tmp / "saved_items.json", tmp / "monitor.lock"
    saved = mstate.load_saved_data()
    long_up, t0 = 99999, 10000.0
    stuck = {"poll": t0 - 3600, "ticker": t0 - 3600, "pid": 4242}
    try:
        heartbeat.FILE.write_text(json.dumps(stuck), encoding="utf-8")
        alerts.check_bot(saved, t0, uptime=60)
        assert restarts == [], "just after Windows started"
        alerts.check_bot(saved, t0, long_up)
        assert restarts == [4242] and sent == [], "restarted quietly, judged on the next check"
        alerts.check_bot(saved, t0 + 300, long_up)
        assert restarts == [4242] and sent == [], "the restart is given its chance"
        heartbeat.FILE.write_text(json.dumps({"poll": t0 + 890, "ticker": t0 + 890, "pid": 5}), encoding="utf-8")
        alerts.check_bot(saved, t0 + 900, long_up)
        assert len(sent) == 1 and "restarted it" in sent[0] and "bot is back" not in sent[0] and not saved["watch"], sent
        assert _legacy_markdown_ok(sent[0])

        # it keeps failing: three restarts, spaced out, then a human is told
        sent.clear(); restarts.clear()
        heartbeat.FILE.write_text(json.dumps(stuck), encoding="utf-8")
        for step in range(3):
            alerts.check_bot(saved, t0 + 2000 + step * 900, long_up)
        assert restarts == [4242] * 3 and sent == []
        alerts.check_bot(saved, t0 + 2000 + 3 * 900, long_up)
        assert len(sent) == 1 and "restarted it 3 time(s)" in sent[0] and _legacy_markdown_ok(sent[0]), sent
        alerts.check_bot(saved, t0 + 2000 + 4 * 900, long_up)
        assert len(sent) == 1 and len(restarts) == 3, "no more restarts, no repeat of the alert"
        heartbeat.FILE.write_text(json.dumps({"poll": t0 + 9000, "ticker": t0 + 9000, "pid": 5}), encoding="utf-8")
        alerts.check_bot(saved, t0 + 9010, long_up)
        assert "bot is back" in sent[-1] and not saved["watch"]

        # a check the bot started itself never restarts the bot
        sent.clear(); restarts.clear()
        heartbeat.FILE.write_text(json.dumps(stuck), encoding="utf-8")
        from uninotifier.monitor import config as mconfig
        mconfig.MANUAL_RUN = True
        try:
            alerts.check_bot(saved, t0 + 20000, long_up)
        finally:
            mconfig.MANUAL_RUN = False
        assert restarts == [] and len(sent) == 1 and "isn't responding" in sent[0]
    finally:
        alerts.send_telegram, mstate.DATA_FILE, mstate.LOCK_FILE, alerts.selfheal.restart_bot = keep
        heartbeat.FILE.unlink(missing_ok=True)


def test_schedule_validation_and_last_good():
    """schedule.json is checked as strictly as it is used: every mistake is named,
    and a broken file never switches the lecture reminders off — the last good copy
    stays in force until it is fixed, and you are told once per version of the mistake."""
    import copy
    good = json.loads(paths.SCHEDULE_FILE.read_text(encoding="utf-8"))
    assert schedule.validate(good) == [], schedule.validate(good)

    def broken(mutate):
        data = copy.deepcopy(good)
        mutate(data)
        return schedule.validate(data)

    cases = {
        "unknown setting": lambda d: d.update(lecture_reminder_minute=30),
        "unknown day": lambda d: d["days"].update(mond=[]),
        "missing end": lambda d: d["days"]["wed"][0].pop("end"),
        "bad time": lambda d: d["days"]["wed"][0].update(start="7:00 PM"),
        "ends before it starts": lambda d: d["days"]["wed"][0].update(start="21:00"),
        "bad mode": lambda d: d["days"]["wed"][0].update(mode="Hybrid"),
        "bad minutes": lambda d: d.update(lecture_reminder_minutes=0),
        "bad hour": lambda d: d.update(morning_summary_hour=25),
        "bad remind_modes": lambda d: d.update(remind_modes=["Zoom"]),
        "teachers": lambda d: d.update(teachers={"Big Data": 5}),
        "finals dates": lambda d: d.update(finals={"from": "2027-01-04"}),
        "finals order": lambda d: d.update(finals={"from": "2027-02-01", "to": "2027-01-04"}),
        "days not an object": lambda d: d.update(days=[]),
    }
    for name, mutate in cases.items():
        assert broken(mutate), f"should be rejected: {name}"
    assert schedule.validate([1]) and schedule.validate("x")
    ok = copy.deepcopy(good)
    ok["finals"] = {"from": "2027-01-04", "to": "2027-01-31", "interval_minutes": 30}
    ok["_note"] = "comments are fine"
    assert schedule.validate(ok) == []

    _fresh_failsafe()
    tmp = Path(tempfile.mkdtemp(prefix="uninotifier-schedule-"))
    keep = (schedule.SCHEDULE_FILE, schedule.LAST_GOOD_FILE, schedule._good, schedule._broken, dict(schedule._read))
    schedule.SCHEDULE_FILE, schedule.LAST_GOOD_FILE = tmp / "schedule.json", tmp / "last_good.json"
    schedule._good, schedule._broken = None, False
    schedule._read.update(path=None, text=None, data=None)
    try:
        # no file and no earlier copy: empty, and told
        assert schedule.load_schedule() == {"days": {}}
        failsafe.wait_idle()
        assert len(OUTBOX) == 1 and "OFF" in OUTBOX[0] and "copy schedule.example.json to schedule.json" in OUTBOX[0], OUTBOX
        assert schedule.load_schedule() == {"days": {}}
        failsafe.wait_idle()
        assert len(OUTBOX) == 1, "a file that stays missing is reported once"

        schedule.SCHEDULE_FILE.write_text(json.dumps(good), encoding="utf-8")
        assert schedule.load_schedule() == good and schedule.LAST_GOOD_FILE.exists()
        failsafe.wait_idle()
        assert "valid again" in OUTBOX[-1] and len(OUTBOX) == 2

        trailing = json.dumps(good)[:-1] + ",}"                 # the classic comma mistake
        schedule.SCHEDULE_FILE.write_text(trailing, encoding="utf-8")
        assert schedule.load_schedule() == good, "the last good copy stays in force"
        failsafe.wait_idle()
        assert len(OUTBOX) == 3 and "needs fixing" in OUTBOX[-1] and "not valid JSON" in OUTBOX[-1]
        assert "keep using the last good copy" in OUTBOX[-1] and _legacy_markdown_ok(OUTBOX[-1])
        assert schedule.load_schedule() == good
        failsafe.wait_idle()
        assert len(OUTBOX) == 3, "the same mistake is not repeated"

        worse = copy.deepcopy(good)
        worse["days"]["wed"][0]["mode"] = "Hybrid"
        schedule.SCHEDULE_FILE.write_text(json.dumps(worse), encoding="utf-8")
        assert schedule.load_schedule() == good
        failsafe.wait_idle()
        assert len(OUTBOX) == 4 and "Hybrid" in OUTBOX[-1], "another mistake is a new report"

        # a restart while it is still broken still falls back to the copy on disk
        schedule._good, schedule._broken = None, False
        schedule._read.update(path=None, text=None, data=None)
        assert schedule.load_schedule() == good

        schedule.SCHEDULE_FILE.write_text(json.dumps(ok), encoding="utf-8")
        assert schedule.load_schedule() == ok
        failsafe.wait_idle()
        assert "valid again" in OUTBOX[-1]
        edited = copy.deepcopy(ok)
        edited["lecture_reminder_minutes"] = 15
        schedule.SCHEDULE_FILE.write_text(json.dumps(edited), encoding="utf-8")
        assert schedule.load_schedule()["lecture_reminder_minutes"] == 15, "an edit takes effect on the next call"
        sent = len(OUTBOX)
        edited["lecture_reminder_minutes"] = 20                   # saved by Notepad as "UTF-8 with BOM"
        schedule.SCHEDULE_FILE.write_text("﻿" + json.dumps(edited), encoding="utf-8")
        assert schedule.load_schedule()["lecture_reminder_minutes"] == 20, "a byte-order mark is not a mistake"
        failsafe.wait_idle()
        assert len(OUTBOX) == sent, OUTBOX[sent:]
    finally:
        schedule.SCHEDULE_FILE, schedule.LAST_GOOD_FILE, schedule._good, schedule._broken = keep[:4]
        schedule._read.update(keep[4])
        schedule._arabic_to_english.cache_clear()


def test_state_backup_and_recovery():
    """saved_items.json: backed up after a good run, kept to 14 days, and an unreadable
    file is replaced by the newest good backup (and reported) instead of silently
    starting over. An empty file is NOT corruption (it is how --announce is used)."""
    from uninotifier import fileio
    from uninotifier.monitor import state as mstate

    _fresh_failsafe()
    tmp = Path(tempfile.mkdtemp(prefix="uninotifier-state-"))
    keep = (mstate.DATA_FILE, mstate.LOCK_FILE)
    mstate.DATA_FILE, mstate.LOCK_FILE = tmp / "saved_items.json", tmp / "monitor.lock"
    try:
        content = {"courses": {"1": {"name": "Big Data"}}, "content": {"1-1": {"name": "Slides"}}}
        fileio.write_json_atomic(mstate.DATA_FILE, content)
        assert not list(tmp.glob("*.tmp")), "the temporary file is gone"
        mstate.backup_state()
        folder = mstate.backup_dir()
        today = time.strftime("%Y-%m-%d")
        assert json.loads((folder / "saved_items.latest.json").read_text(encoding="utf-8")) == content
        assert (folder / f"saved_items.{today}.json").exists()

        for n in range(20):                                   # old days: only the newest 14 days stay
            (folder / f"saved_items.2020-01-{n + 1:02d}.json").write_text("{}", encoding="utf-8")
        mstate.backup_state()
        days = sorted(p.name for p in folder.glob("saved_items.????-??-??.json"))
        assert len(days) == mstate.BACKUP_KEEP_DAYS and f"saved_items.{today}.json" in days, days

        # an empty baseline or an unreadable file never replaces a good backup
        mstate.DATA_FILE.write_text("{}", encoding="utf-8")
        mstate.backup_state()
        mstate.DATA_FILE.write_text("{ not json", encoding="utf-8")
        mstate.backup_state()
        assert json.loads((folder / "saved_items.latest.json").read_text(encoding="utf-8")) == content

        # empty is a valid state (--announce): loaded as it is, nobody alarmed
        mstate.DATA_FILE.write_text("{}", encoding="utf-8")
        loaded = mstate.load_saved_data()
        assert loaded["courses"] == {} and set(mstate.STATE_KEYS) <= set(loaded) and OUTBOX == []

        # unreadable (also: valid JSON that is not an object) -> restored from the backup, reported
        for damage in ("{ not json", "[1, 2]", ""):
            OUTBOX.clear()
            failsafe.STATE_FILE.unlink(missing_ok=True)
            mstate.DATA_FILE.write_text(damage, encoding="utf-8")
            loaded = mstate.load_saved_data()
            assert loaded["courses"] == content["courses"] and loaded["content"] == content["content"], damage
            assert json.loads(mstate.DATA_FILE.read_text(encoding="utf-8")) == content, "the file itself is repaired"
            assert len(OUTBOX) == 1 and "restored from a backup" in OUTBOX[0] and _legacy_markdown_ok(OUTBOX[0]), OUTBOX
            assert list(folder.glob("saved_items.corrupt-*.json")), "the damaged file is kept"

        # no backup at all: a fresh baseline, said out loud
        for p in folder.glob("saved_items.*.json"):
            if "corrupt" not in p.name:
                p.unlink()
        OUTBOX.clear()
        failsafe.STATE_FILE.unlink(missing_ok=True)
        mstate.DATA_FILE.write_text("{ not json", encoding="utf-8")
        loaded = mstate.load_saved_data()
        assert loaded["courses"] == {} and len(OUTBOX) == 1 and "no backup" in OUTBOX[0]
        assert len(list(folder.glob("saved_items.corrupt-*.json"))) <= mstate.CORRUPT_KEEP
    finally:
        mstate.DATA_FILE, mstate.LOCK_FILE = keep


def test_telegram_transport():
    """The one door to Telegram: connections are reused, a dead pooled connection is replaced
    (but a send is never repeated), sends are paced, a 429 is honoured once, and every call
    is counted so a connection problem can be measured."""
    import threading
    from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

    import requests

    class Reply:
        def __init__(self, status=200, body=None):
            self.status_code, self._body = status, body or {"ok": True}
            self.ok = status < 400

        def json(self):
            return self._body

    def reset_stats():
        telegram._stats.clear()
        telegram.STATS_FILE.unlink(missing_ok=True)

    sleeps, calls = [], []
    keep = (telegram.post, telegram.time.sleep, telegram.MIN_SEND_INTERVAL)
    telegram.time.sleep = sleeps.append
    try:
        # ok / bookkeeping: polls and other calls are counted apart
        reset_stats()
        telegram.post = lambda url, **kw: calls.append((url, kw)) or Reply()
        assert telegram.request("T", "getMe").ok and telegram.request("T", "getUpdates", json={"timeout": 5}, timeout=15).ok
        assert calls[0][0].endswith("/botT/getMe") and calls[1][1]["timeout"] == (telegram.CONNECT_TIMEOUT, 15)
        hour = time.strftime("%Y-%m-%dT%H")
        assert telegram._stats[hour] == {"api:ok": 1, "poll:ok": 1}, telegram._stats

        # a pooled connection that died: a call that is safe to repeat is repeated once, on a fresh connection
        calls.clear()
        outcomes = [requests.ConnectionError("Connection aborted: ConnectionResetError(10054, 'forcibly closed')"), Reply()]

        def flaky(url, **kw):
            calls.append(url)
            out = outcomes.pop(0)
            if isinstance(out, Exception):
                raise out
            return out
        telegram.post = flaky
        assert telegram.request("T", "answerCallbackQuery", json={"callback_query_id": "1"}).ok and len(calls) == 2
        assert telegram._stats[hour]["api:ConnectionReset"] == 1

        # ... but a send is never repeated by the transport (a duplicate message is worse than a late one)
        calls.clear()
        outcomes[:] = [requests.ConnectionError("boom"), Reply()]
        try:
            telegram.request("T", "sendMessage", json={"chat_id": 1, "text": "x"})
            raise AssertionError("must raise")
        except requests.ConnectionError:
            pass
        assert len(calls) == 1, "sendMessage was repeated"
        # a read timeout is not repeated either (the call may have been processed)
        calls.clear()
        outcomes[:] = [requests.ReadTimeout("slow"), Reply()]
        try:
            telegram.request("T", "getUpdates", json={})
            raise AssertionError("must raise")
        except requests.ReadTimeout:
            pass
        assert len(calls) == 1 and telegram._stats[hour]["poll:ReadTimeout"] == 1
        # ... while a timeout on a call that is harmless to repeat (often the handshake of a new connection) is repeated once
        calls.clear()
        outcomes[:] = [requests.ReadTimeout("handshake"), Reply()]
        assert telegram.request("T", "setMyCommands", json={"commands": []}).ok and len(calls) == 2
        calls.clear()
        outcomes[:] = [requests.ReadTimeout("a"), requests.ReadTimeout("b")]
        try:
            telegram.request("T", "setMyCommands", json={})
            raise AssertionError("must raise")
        except requests.ReadTimeout:
            pass
        assert len(calls) == 2, "once more, not forever"
        calls.clear()
        outcomes[:] = [requests.ReadTimeout("slow"), Reply()]
        try:
            telegram.request("T", "sendMessage", json={"chat_id": 1})
            raise AssertionError("must raise")
        except requests.ReadTimeout:
            pass
        assert len(calls) == 1, "a send is never repeated"
        # ... unless it never left: a stalled handshake (a ReadTimeout whose limit is the CONNECT timeout, not
        # the call's own) or a connect timeout means nothing was sent, so even a send goes again - once
        calls.clear()
        hs = requests.ReadTimeout(f"HTTPSConnectionPool(host='api.telegram.org', port=443): Read timed out. (read timeout={telegram.CONNECT_TIMEOUT})")
        outcomes[:] = [hs, Reply()]
        assert telegram.request("T", "sendMessage", json={"chat_id": 1}, timeout=20).ok and len(calls) == 2
        assert telegram._stats[hour]["api:HandshakeTimeout"] == 1
        calls.clear()
        outcomes[:] = [requests.ConnectTimeout("no route"), Reply()]
        assert telegram.request("T", "sendMessage", json={"chat_id": 1}).ok and len(calls) == 2
        calls.clear()
        outcomes[:] = [hs, hs, Reply()]
        try:
            telegram.request("T", "sendMessage", json={"chat_id": 1}, timeout=20)
            raise AssertionError("must raise")
        except requests.ReadTimeout:
            pass
        assert len(calls) == 2, "once more, not forever"
        calls.clear()
        outcomes[:] = [requests.ReadTimeout("Read timed out. (read timeout=20)"), Reply()]
        try:
            telegram.request("T", "sendMessage", json={"chat_id": 1}, timeout=20)
            raise AssertionError("must raise")
        except requests.ReadTimeout:
            pass
        assert len(calls) == 1, "a read timeout of the call itself may mean it was processed: never repeated"
        assert telegram.stalled_before_sending(hs, 20) and not telegram.stalled_before_sending(hs, telegram.CONNECT_TIMEOUT)

        # 429: "retry after N" is honoured once (and not for an upload, whose file is already read)
        sleeps.clear()
        outcomes[:] = [Reply(429, {"ok": False, "parameters": {"retry_after": 3}}), Reply()]
        telegram.post = flaky
        assert telegram.request("T", "sendMessage", json={"chat_id": 1}).ok and sleeps == [3.5]
        sleeps.clear()
        outcomes[:] = [Reply(429, {"ok": False, "parameters": {"retry_after": 99}})]
        assert telegram.request("T", "sendMessage", json={"chat_id": 1}).status_code == 429 and sleeps == [], "too long: left to the caller"
        outcomes[:] = [Reply(429, {"ok": False, "parameters": {"retry_after": 2}}), Reply()]
        assert telegram.request("T", "sendDocument", data={"chat_id": 1}, files={"f": b"x"}).status_code == 429

        # pacing: back-to-back sends to one chat are spaced; another chat, or a call that sends nothing, is not
        sleeps.clear()
        telegram.MIN_SEND_INTERVAL = 0.4
        telegram._next_slot.clear()
        telegram.post = lambda url, **kw: Reply()
        telegram.request("T", "sendMessage", json={"chat_id": 7, "text": "a"})
        telegram.request("T", "sendMessage", json={"chat_id": 7, "text": "b"})
        telegram.request("T", "sendMessage", json={"chat_id": 8, "text": "c"})
        telegram.request("T", "answerCallbackQuery", json={"callback_query_id": "1"})
        # Windows' monotonic clock ticks every ~16 ms: two sends in one tick wait the full 0.4 s, give or take a float's last bit
        assert len(sleeps) == 1 and 0.2 < sleeps[0] <= 0.4 + 1e-6, sleeps
    finally:
        telegram.post, telegram.time.sleep, telegram.MIN_SEND_INTERVAL = keep

    # backoff: 0.5, 1, 2, 4, then 8 s (with jitter), never more
    waits = [telegram.backoff_seconds(n) for n in (1, 2, 3, 4, 5, 20)]
    assert waits == sorted(waits) or waits[4] <= waits[5] * 1.5
    assert 0.39 <= waits[0] <= 0.61 and all(w <= 9.7 for w in waits) and waits[5] >= 6.3

    # statistics: summed over the window, old hours dropped
    reset_stats()
    telegram._stats[time.strftime("%Y-%m-%dT%H")] = {"poll:ok": 5, "poll:ReadTimeout": 1}
    telegram.STATS_FILE.write_text(json.dumps({"2020-01-01T00": {"poll:ok": 99}}), encoding="utf-8")
    assert telegram.summary(24) == {"poll:ok": 5, "poll:ReadTimeout": 1}
    assert "2020-01-01T00" not in json.loads(telegram.STATS_FILE.read_text(encoding="utf-8"))
    reset_stats()

    # the real session: ONE shared by every thread, keep-alive probes on, and ONE TCP connection
    # for many requests — also when each request comes from a thread of its own
    assert (socket_opts := dict(((lvl, opt), val) for lvl, opt, val in telegram._keepalive_options()))
    import socket
    assert socket_opts[(socket.SOL_SOCKET, socket.SO_KEEPALIVE)] == 1
    seen = {}
    t = threading.Thread(target=lambda: seen.update(other=telegram._session()))
    t.start()
    t.join()
    mine = telegram._session()
    assert mine is telegram._session() and seen["other"] is mine, "every thread shares the one session"
    telegram.reset_session()
    assert telegram._session() is not mine

    connections, slow = [], threading.Event()

    class Handler(BaseHTTPRequestHandler):
        protocol_version = "HTTP/1.1"

        def setup(self):
            connections.append(1)
            super().setup()

        def do_POST(self):
            self.rfile.read(int(self.headers.get("Content-Length", 0)))
            if self.path.endswith("/slow"):
                slow.wait(5)                          # held open until the test has reset the pool
            body = b'{"ok": true, "result": true}'
            self.send_response(200)
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)

        def log_message(self, *a):
            pass

    server = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
    threading.Thread(target=server.serve_forever, daemon=True).start()
    try:
        url = f"http://127.0.0.1:{server.server_address[1]}/botT/getMe"
        telegram.reset_session()
        for _ in range(5):
            assert _REAL_POST(url, json={}, timeout=5).ok
        assert len(connections) == 1, f"5 requests opened {len(connections)} connections"
        # five short-lived threads (chat actions, workers, alerts): still that one warm connection
        for _ in range(5):
            worker = threading.Thread(target=lambda: _REAL_POST(url, json={}, timeout=5))
            worker.start()
            worker.join()
        assert len(connections) == 1, f"5 threads opened {len(connections) - 1} new connection(s)"
        connections.clear()
        for _ in range(3):
            requests.post(url, json={}, timeout=5)          # what the old code did
        assert len(connections) == 3
        # a reset while another thread's call is in flight: that call still finishes normally
        outcome = {}
        busy = threading.Thread(target=lambda: outcome.update(ok=_REAL_POST(url[:-len("getMe")] + "slow", json={}, timeout=10).ok))
        busy.start()
        time.sleep(0.3)
        telegram.reset_session()
        slow.set()
        busy.join(10)
        assert outcome == {"ok": True}, outcome
    finally:
        server.shutdown()
        telegram.reset_session()


def test_buttons_are_answered_first():
    """A button press is answered once and at once, a double tap is ignored, a slow
    "Mark as done" says "updating…" and then shows its outcome on the button, and a
    long job shows a chat action instead of a message."""
    import threading

    bot = _package("uninotifier.bot")
    calls = []
    keep = (bot.api, bot.send_view, bot.edit_view, bot.DEBOUNCE_SECONDS, bot.ACTION_REFRESH_SECONDS,
            moodle.http_set_completion, bot.set_item_done, bot.DONE_WAIT, notifier.send_document_group)
    bot.api = lambda method, **data: calls.append((method, data))
    bot.send_view = lambda text, rows: calls.append(("sendMessage", {"text": text}))
    bot.edit_view = lambda message_id, text, rows: calls.append(("editMessageText", {"text": text}))
    bot.set_item_done = lambda cmid, done: calls.append(("state", {"cmid": cmid, "done": done}))
    real_start_check = bot.checks.start_check
    bot.checks.start_check = lambda message_id: True        # a "check" press must never start the real monitor

    def answers():
        return [d for m, d in calls if m == "answerCallbackQuery"]

    def press(cid, data, mid=9, keyboard=None):
        message = {"message_id": mid}
        if keyboard:
            message["reply_markup"] = {"inline_keyboard": keyboard}
        bot.handle_callback({"id": cid, "data": data, "message": message})

    try:
        # answered exactly once, whatever the route did — and the first answer wins
        assert bot.answer("a1", "first") is True and bot.answer("a1", "second") is False
        assert [d.get("text") for d in answers()] == ["first"]
        calls.clear()
        press("a2", "menu")
        press("a3", "no-such-screen")
        assert [d["callback_query_id"] for d in answers()] == ["a2", "a3"] and all("text" not in d for d in answers())
        calls.clear()
        press("a5", "check")                                 # the route's own words are kept (never a real check)
        assert [d.get("text") for d in answers()][-1] == "🔄 Checking…"

        # a double tap (same button, same message, within the window) is answered but not handled twice
        calls.clear()
        bot.DEBOUNCE_SECONDS = 0.2
        bot._last_press.clear()
        press("b1", "tasks")
        press("b2", "tasks")
        press("b3", "tasks", mid=10)                         # another message: its own press
        press("b4", "courses")                               # another button
        assert [m for m, _ in calls if m in ("editMessageText", "sendMessage")] != []
        assert len([1 for m, _ in calls if m == "editMessageText"]) == 3, ("tasks twice would be 4", calls)
        assert [d["callback_query_id"] for d in answers()] == ["b1", "b2", "b3", "b4"]
        time.sleep(0.25)
        calls.clear()
        press("b5", "tasks")
        assert len([1 for m, _ in calls if m == "editMessageText"]) == 1, "after the window it is a new press"
        bot.DEBOUNCE_SECONDS = 0

        # Mark as done: fast portal -> the outcome is the toast, the button flips
        keyboard = [[{"text": "✅ Mark as Done", "callback_data": "done:7:1"}], [{"text": "🔙", "callback_data": "menu"}]]
        moodle.http_set_completion = lambda cmid, done: True
        calls.clear()
        press("d1", "done:7:1", keyboard=keyboard)
        assert answers() == [{"callback_query_id": "d1", "text": "✅ Marked as done on the portal"}], answers()
        edit = [d for m, d in calls if m == "editMessageReplyMarkup"][-1]
        assert edit["reply_markup"]["inline_keyboard"][0][0]["callback_data"] == "done:7:0"
        assert edit["reply_markup"]["inline_keyboard"][1] == keyboard[1], "the other buttons stay"
        assert ("state", {"cmid": "7", "done": True}) in calls

        # fast failures: only the toast, the keyboard stays; an expired session is a dialog
        moodle.http_set_completion = lambda cmid, done: False
        calls.clear()
        press("d2", "done:7:1", keyboard=keyboard)
        assert "didn't accept" in answers()[0]["text"] and not [1 for m, _ in calls if m == "editMessageReplyMarkup"]

        def expired(cmid, done):
            raise moodle.SessionExpiredError("gone")
        moodle.http_set_completion = expired
        calls.clear()
        press("d3", "done:7:1", keyboard=keyboard)
        assert answers()[0]["show_alert"] is True and "expired" in answers()[0]["text"]

        # a slow portal: answered "updating…" in time, the outcome lands on the button
        bot.DONE_WAIT = 0.05
        for outcome, expected in ((True, "done:7:0"), (False, "done:7:1")):
            gate = threading.Event()
            moodle.http_set_completion = lambda cmid, done, gate=gate, outcome=outcome: gate.wait(5) and outcome
            calls.clear()
            started = time.monotonic()
            press(f"s{outcome}", "done:7:1", keyboard=keyboard)
            assert time.monotonic() - started < 1.0 and [d["text"] for d in answers()] == ["⏳ Updating the portal…"]
            gate.set()
            for t in threading.enumerate():
                if t.name == "done-toggle":
                    t.join(5)
            assert len(answers()) == 1, "the press is answered once"
            edit = [d for m, d in calls if m == "editMessageReplyMarkup"][-1]
            row = edit["reply_markup"]["inline_keyboard"][0][0]
            assert row["callback_data"] == expected, (outcome, row)
            if not outcome:
                assert "tap to retry" in row["text"], "a late failure is shown on the button"

        # a long job shows the chat action (renewed while it runs, stopped after), not a message
        calls.clear()
        bot.ACTION_REFRESH_SECONDS = 0.02
        with bot.working("upload_document"):
            time.sleep(0.12)
        n = len([1 for m, _ in calls if m == "sendChatAction"])
        assert n >= 3 and all(d["action"] == "upload_document" for m, d in calls if m == "sendChatAction"), calls
        time.sleep(0.1)
        assert len([1 for m, _ in calls if m == "sendChatAction"]) == n, "stops with the block"
        assert bot.send_schedule_files.__name__ == "send_schedule_files"
        calls.clear()
        notifier.send_document_group = lambda *a, **k: True
        bot.send_schedule_files()
        assert calls[0][0] == "sendChatAction" and calls[0][1]["action"] == "upload_document"
        assert not [1 for m, d in calls if m == "sendMessage" and "Fetching" in d.get("text", "")], "no loading message"
    finally:
        (bot.api, bot.send_view, bot.edit_view, bot.DEBOUNCE_SECONDS, bot.ACTION_REFRESH_SECONDS,
         moodle.http_set_completion, bot.set_item_done, bot.DONE_WAIT, notifier.send_document_group) = keep
        bot.checks.start_check = real_start_check


def test_files_telegram_already_holds_are_not_uploaded_again():
    """A document is uploaded once; afterwards its file_id (remembered by content) sends it with no
    upload at all. A changed file is a new file, a file_id Telegram refuses is dropped and the file
    uploaded again, a network failure neither uploads nor forgets, and the timetable album is covered."""
    import requests

    tmp = Path(tempfile.mkdtemp(prefix="uninotifier-fileids-"))
    pdf, png = tmp / "plan.pdf", tmp / "plan.png"
    pdf.write_bytes(b"%PDF-" + b"a" * 5000)
    png.write_bytes(b"\x89PNG" + b"b" * 9000)

    class Reply:
        def __init__(self, status=200, body=None, text=""):
            self.status_code, self._body, self.text = status, body or {"ok": True, "result": True}, text
            self.ok = status < 400

        def json(self):
            return self._body

    sent, state = [], {"mode": "ok", "ids": iter(range(1, 10 ** 6))}

    def post(url, json=None, data=None, files=None, timeout=None):
        method = url.rsplit("/", 1)[1]
        uploaded = sum(len(v[1]) if isinstance(v[1], (bytes, bytearray)) else os.fstat(v[1].fileno()).st_size
                       for v in (files or {}).values())          # an album holds bytes, a single file an open handle
        sent.append({"method": method, "data": data, "files": files, "bytes": uploaded})
        if state["mode"] == "down":
            raise requests.ConnectionError("network down")
        if state["mode"] == "stale":
            refs = [m["media"] for m in __import__("json").loads(data["media"])] if method == "sendMediaGroup" else [data.get("document", "attach")]
            if any(not str(r).startswith("attach") and r != "attach" for r in refs):
                return Reply(400, {"ok": False}, '{"ok":false,"description":"Bad Request: wrong file identifier/HTTP URL specified"}')
        if method == "sendMediaGroup":
            n = len(__import__("json").loads(data["media"]))
            return Reply(body={"ok": True, "result": [{"document": {"file_id": f"ID{next(state['ids'])}"}} for _ in range(n)]})
        return Reply(body={"ok": True, "result": {"document": {"file_id": f"ID{next(state['ids'])}"}}})

    keep = (telegram.post, notifier.time.sleep)
    telegram.post, notifier.time.sleep = post, lambda s: None
    fileids.STORE.unlink(missing_ok=True)
    try:
        def album(items=None):
            sent.clear()
            return notifier.send_document_group("T", "1", items or [(pdf, "📎 PDF"), (png, "🖼️ PNG")])

        # first time: both are uploaded and remembered by content
        assert album() and sent[0]["bytes"] == 5005 + 9004 and set(sent[0]["files"]) == {"f0", "f1"}
        assert len(json.loads(fileids.STORE.read_text(encoding="utf-8"))) == 2
        # second time: not a byte uploaded, the album refers to the file_ids
        assert album() and sent[0]["bytes"] == 0 and sent[0]["files"] is None
        media = json.loads(sent[0]["data"]["media"])
        assert [m["media"] for m in media] == ["ID1", "ID2"] and [m["caption"] for m in media] == ["📎 PDF", "🖼️ PNG"]
        # the same content under another name and path is the same file
        copy = tmp / "other-name.pdf"
        copy.write_bytes(pdf.read_bytes())
        assert fileids.fingerprint(copy) == fileids.fingerprint(pdf) and fileids.fingerprint(tmp / "nope") is None
        # a changed file is uploaded again, the unchanged one still is not
        png.write_bytes(b"\x89PNG" + b"c" * 9000)
        assert album() and sent[0]["bytes"] == 9004 and set(sent[0]["files"]) == {"f1"}
        assert [m["media"] for m in json.loads(sent[0]["data"]["media"])] == ["ID1", "attach://f1"]
        assert album() and sent[0]["bytes"] == 0, "and its new file_id is remembered too"

        # Telegram refuses a file_id: they are forgotten and uploaded afresh, once
        state["mode"] = "stale"
        assert album() and [s["bytes"] for s in sent] == [0, 5005 + 9004], [s["bytes"] for s in sent]
        state["mode"] = "ok"
        assert album() and sent[0]["bytes"] == 0, "the fresh ids work"

        # the network is down: nothing is uploaded and nothing is forgotten
        state["mode"] = "down"
        assert not album() and all(s["bytes"] == 0 for s in sent) and len(sent) == notifier.MAX_ATTEMPTS
        state["mode"] = "ok"
        assert album() and sent[0]["bytes"] == 0

        # one document
        sent.clear()
        assert notifier.send_document("T", "1", pdf, "cap", rows=[[{"text": "x", "callback_data": "y"}]])
        assert sent[0]["data"]["document"] == fileids.lookup(fileids.fingerprint(pdf)) and sent[0]["bytes"] == 0
        assert "reply_markup" in sent[0]["data"]
        fileids.STORE.unlink()
        sent.clear()
        assert notifier.send_document("T", "1", pdf, "cap") and sent[0]["bytes"] == 5005 and "document" not in sent[0]["data"]
        assert notifier.send_document("T", "1", pdf, "cap") and sent[1]["bytes"] == 0
        state["mode"] = "stale"
        sent.clear()
        assert notifier.send_document("T", "1", pdf, "cap") and [s["bytes"] for s in sent] == [0, 5005]
        state["mode"] = "ok"

        # the timetable: the second request uploads nothing
        bot = _package("uninotifier.bot")
        keep_api = bot.api
        bot.api = lambda method, **d: {"message_id": 1} if method == "sendMessage" else True
        fileids.STORE.unlink(missing_ok=True)
        try:
            sent.clear()
            bot.send_schedule_files()
            first = sum(s["bytes"] for s in sent if s["method"] == "sendMediaGroup")
            sent.clear()
            bot.send_schedule_files()
            again = sum(s["bytes"] for s in sent if s["method"] == "sendMediaGroup")
        finally:
            bot.api = keep_api
        assert first == paths.SCHEDULE_PDF.stat().st_size + paths.SCHEDULE_IMAGE.stat().st_size and again == 0, (first, again)

        # the memory itself: bounded, tolerant of damage
        fileids.MAX_ENTRIES = 3
        for n in range(6):
            fileids.remember(f"fp{n}", f"id{n}", "x")
        assert len(json.loads(fileids.STORE.read_text(encoding="utf-8"))) == 3
        fileids.STORE.write_text("{ not json", encoding="utf-8")
        assert fileids.lookup("fp5") is None and fileids.ids_from_reply(None) == [] and fileids.ids_from_reply({"result": True}) == []
        assert fileids.ids_from_reply({"result": [{"document": {"file_id": "a"}}, {"photo": []}]}) == ["a", None]
    finally:
        telegram.post, notifier.time.sleep = keep
        fileids.MAX_ENTRIES = 300
        fileids.STORE.unlink(missing_ok=True)


def test_monitor_pruning_and_background_pings():
    """A run no longer asks the portal about every archived file (each is asked once per two hours,
    spread evenly over the runs; /check asks about all of them), and its two healthchecks.io pings
    go out beside the work, in order, without the process ending before the last one is out."""
    import threading
    from uninotifier import healthcheck
    from uninotifier.monitor import config as mconfig, scan as mscan

    # --- each file's turn
    ids = [f"{course}-{n}" for course in (101, 102, 103, 104) for n in range(8)]
    runs = round(mconfig.FILE_PROBE_SECONDS / mconfig.SCAN_SLOT_SECONDS)
    assert runs == 8
    base = 1_800_000_000 // mconfig.SCAN_SLOT_SECONDS * mconfig.SCAN_SLOT_SECONDS
    per_slot = []
    for slot in range(runs):
        now = base + slot * mconfig.SCAN_SLOT_SECONDS + 44          # a run lands anywhere inside its slot
        per_slot.append({i for i in ids if mscan.probe_due(i, now)})
    for i in ids:
        assert sum(i in due for due in per_slot) == 1, f"{i}: exactly one turn per {runs} runs"
    assert sum(len(d) for d in per_slot) == len(ids) and max(len(d) for d in per_slot) <= 3 * len(ids) // runs, [len(d) for d in per_slot]
    assert all(mscan.probe_due(i, base + 7, manual=True) for i in ids), "/check asks about every file"
    again = {i for i in ids if mscan.probe_due(i, base + 3 * mconfig.SCAN_SLOT_SECONDS + 9 * runs * mconfig.SCAN_SLOT_SECONDS)}
    assert again == per_slot[3], "the same file has the same turn every cycle"
    keep_every = mconfig.FILE_PROBE_SECONDS
    mconfig.FILE_PROBE_SECONDS = mconfig.SCAN_SLOT_SECONDS
    try:
        assert all(mscan.probe_due(i, base + 5) for i in ids), "probing every run is a setting away"
    finally:
        mconfig.FILE_PROBE_SECONDS = keep_every

    # --- the pings
    calls, gate = [], threading.Event()

    def fake_post(url, data=None, timeout=None):
        gate.wait(5)
        calls.append(url)
        if url.endswith("/boom"):
            raise ValueError("not a requests error")
        return type("R", (), {"raise_for_status": lambda self: None})()

    keep = (healthcheck.requests.post, os.environ.get(healthcheck.ENV_NAME))
    healthcheck.requests.post = fake_post
    os.environ[healthcheck.ENV_NAME] = "https://hc-ping.com/uuid-9"
    try:
        started = time.monotonic()
        healthcheck.start_bg()
        healthcheck.report_bg(0, "6 courses")
        assert time.monotonic() - started < 0.3, "the pings must not hold the run"
        assert healthcheck.wait(0.05) is False and calls == [], "still being sent"
        gate.set()
        assert healthcheck.wait(5) is True
        assert calls == ["https://hc-ping.com/uuid-9/start", "https://hc-ping.com/uuid-9"], ("start before the result", calls)
        # a ping that blows up does not stop the next ones
        calls.clear()
        os.environ[healthcheck.ENV_NAME] = "https://hc-ping.com/boom"
        healthcheck.report_bg(1, "x")
        os.environ[healthcheck.ENV_NAME] = "https://hc-ping.com/uuid-9"
        healthcheck.report_bg(0, "y")
        assert healthcheck.wait(5) and calls[-1] == "https://hc-ping.com/uuid-9", calls
    finally:
        gate.set()
        healthcheck.requests.post = keep[0]
        os.environ.pop(healthcheck.ENV_NAME, None)
        if keep[1] is not None:
            os.environ[healthcheck.ENV_NAME] = keep[1]


def test_answers_go_out_beside_the_edit():
    """The answer to a press is sent by a standing worker while the caller goes on with the edit,
    so the two Telegram round trips overlap instead of adding up; it is still sent once, a failing
    send does not stop the workers, and the workers are long-lived (a warm connection each)."""
    import threading

    bot = _package("uninotifier.bot")
    calls, gate = [], threading.Event()
    names = {}

    def fake(method, **data):
        names.setdefault(method, set()).add(threading.current_thread().name)
        if method == "answerCallbackQuery":
            gate.wait(5)                                  # a slow Telegram: 5 s unless released
            if data.get("text") == "boom":
                raise OSError("network")
        calls.append((method, data))
        return True

    keep = (bot.api, bot.ANSWER_ASYNC)
    bot.api = fake
    bot.ANSWER_ASYNC = True
    try:
        started = time.monotonic()
        assert bot.answer("w1", "⏳ Fetching…") is True
        assert time.monotonic() - started < 0.5, "answer() must not wait for Telegram"
        bot.edit_view(7, "the screen", [[{"text": "x", "callback_data": "menu"}]])      # the caller carries on at once
        assert any(m == "editMessageText" for m, _ in calls) and not any(m == "answerCallbackQuery" for m, _ in calls)
        assert bot.answer("w1", "again") is False
        gate.set()
        assert bot.flush_answers(5)
        sent = [d for m, d in calls if m == "answerCallbackQuery"]
        assert sent == [{"callback_query_id": "w1", "text": "⏳ Fetching…"}], sent
        assert all(n.startswith("answer-") for n in names["answerCallbackQuery"]), names

        # a send that fails neither raises into the caller nor kills the worker
        assert bot.answer("w2", "boom") is True
        assert bot.flush_answers(5)
        assert bot.answer("w3", "still alive") is True and bot.flush_answers(5)
        assert [d["text"] for m, d in calls if m == "answerCallbackQuery"][-1] == "still alive"
        workers = [t for t in threading.enumerate() if t.name.startswith("answer-")]
        assert len(workers) == bot.ANSWER_WORKERS and all(t.daemon for t in workers)
        assert "getMe" in names, "each worker opens its connection before the first press"
    finally:
        gate.set()
        bot.api, bot.ANSWER_ASYNC = keep


def test_one_menu_anchored_at_the_bottom():
    """There is one menu and it is the last thing in the chat: a new screen replaces the old
    one, an old screen that is pressed is moved to the bottom, files are followed by the menu
    again, a status line is not a menu, alerts' buttons are untouched, and it all survives a restart."""
    bot = _package("uninotifier.bot")
    log_, ids, fail = [], iter(range(100, 10 ** 6)), {"delete": False, "send": False}

    def fake(method, **data):
        log_.append((method, data))
        if method == "sendMessage":
            return None if fail["send"] else {"message_id": next(ids)}
        if method == "deleteMessage":
            return None if fail["delete"] else True
        return True

    def calls(*names):
        return [(m, d) for m, d in log_ if m in names]

    keep = (bot.api, notifier.send_document_group, bot.DONE_WAIT)
    bot.api = fake
    rows_a, rows_b = [[{"text": "📚 Courses", "callback_data": "courses"}]], [[{"text": "🔙 Main Menu", "callback_data": "menu"}]]
    try:
        # the first screen with buttons is the menu
        bot.reset_menu()
        first = bot.send_view("Main", rows_a)["message_id"]
        assert bot.current_menu_id() == first and not calls("deleteMessage")
        # a status line / a question being asked is a plain message: the menu stays
        bot.send_view("⏳ Checking…", [])
        assert bot.current_menu_id() == first and not calls("deleteMessage")
        # a new screen replaces the menu: sent first, the old one removed after
        log_.clear()
        second = bot.send_view("Courses", rows_b)["message_id"]
        assert [m for m, _ in log_] == ["sendMessage", "deleteMessage"] and log_[1][1]["message_id"] == first
        assert bot.current_menu_id() == second
        # navigation inside the menu edits it in place
        log_.clear()
        bot.edit_view(second, "Courses (page 2)", rows_b)
        assert [m for m, _ in log_] == ["editMessageText"] and bot.current_menu_id() == second and bot._menu["text"] == "Courses (page 2)"

        # a button pressed on an OLD screen: the answer goes to the bottom, the old screen goes away
        log_.clear()
        bot.edit_view(first, "Tasks", rows_a)
        assert [m for m, _ in log_] == ["sendMessage", "deleteMessage", "deleteMessage"], log_
        third = bot.current_menu_id()
        assert third not in (first, second) and not calls("editMessageText")
        assert {d["message_id"] for _, d in calls("deleteMessage")} == {first, second}, "the menu it replaces AND the old screen"

        # a finished status line that gets the result's buttons becomes the menu at the bottom
        status = bot.send_view("🔄 Checking the portal…", [])["message_id"]
        log_.clear()
        bot.edit_view(status, "✅ Check Complete", rows_b)
        assert [m for m, _ in log_] == ["sendMessage", "deleteMessage", "deleteMessage"]
        assert {d["message_id"] for m, d in calls("deleteMessage")} == {third, status}

        # files were posted below the menu: the menu returns to the bottom, silently, old one removed
        menu = bot.current_menu_id()
        notifier.send_document_group = lambda *a, **k: True
        log_.clear()
        bot.send_schedule_files()
        names = [m for m, _ in log_]
        assert names[0] == "sendChatAction" and names[-2:] == ["sendMessage", "deleteMessage"], names
        assert calls("sendMessage")[-1][1]["disable_notification"] is True
        assert calls("sendMessage")[-1][1]["text"] == "✅ Check Complete" and calls("deleteMessage")[-1][1]["message_id"] == menu

        # Telegram will not delete it (too old): its buttons are removed instead
        fail["delete"] = True
        log_.clear()
        old = bot.current_menu_id()
        bot.send_view("Next", rows_a)
        assert [m for m, _ in log_] == ["sendMessage", "deleteMessage", "editMessageReplyMarkup"]
        assert log_[2][1]["message_id"] == old and log_[2][1]["reply_markup"] == {"inline_keyboard": []}
        fail["delete"] = False

        # a send that fails changes nothing: the old menu stays (and is not deleted)
        before = bot.current_menu_id()
        fail["send"] = True
        log_.clear()
        assert bot.send_view("Lost", rows_b) is None
        assert bot.current_menu_id() == before and not calls("deleteMessage")
        fail["send"] = False

        # a restart: the menu is remembered
        remembered = bot.current_menu_id()
        bot._menu.update(message_id=None, text=None, rows=None)
        bot._menu_loaded = False
        assert bot.current_menu_id() == remembered and bot._menu["rows"] == rows_a
        log_.clear()
        bot.edit_view(remembered - 1, "Old one from before the restart", rows_b)
        assert [m for m, _ in log_][:2] == ["sendMessage", "deleteMessage"]

        # with no menu at all, nothing is re-sent
        bot.reset_menu()
        log_.clear()
        bot.reanchor()
        assert log_ == []
        # ... and the very first edit of an unknown message with buttons adopts it as the menu
        bot.edit_view(555, "Adopted", rows_a)
        assert bot.current_menu_id() == 555 and [m for m, _ in log_] == ["editMessageText"]

        # an alert's buttons are not a menu: notifier messages never register
        log_.clear()
        telegram_keep = telegram.post
        telegram.post = lambda url, **kw: type("R", (), {"ok": True, "status_code": 200, "json": lambda self: {"ok": True}})()
        try:
            assert notifier.send_message("T", "1", "🎥 Lecture", rows=[[{"text": "Join", "callback_data": "x"}]])
        finally:
            telegram.post = telegram_keep
        assert bot.current_menu_id() == 555 and log_ == []
    finally:
        bot.api, notifier.send_document_group, bot.DONE_WAIT = keep
        bot.reset_menu()


def test_course_list_is_a_timetable():
    """The course list is ordered like the week: grouped by the days a course meets (Sun · Tue together,
    Mon · Wed together), each group by the time of its lectures, a course that meets at the same time on
    both days says it once and only spells out what differs, and the rest stays short."""
    from uninotifier import timetable
    bot = _package("uninotifier.bot")

    def lec(course, start, end, mode="In-Person", label="In-Person Lecture", place=None):
        d = {"start": start, "end": end, "course": course, "course_ar": course, "section": "01", "mode": mode, "label": label}
        if place:
            d["place"] = place
        return d

    days = {
        "sat": [lec("Big Data", "08:00", "10:00", label="Lab")],
        "sun": [lec("Late Course", "15:30", "17:00", place="Room 5"), lec("Early Course", "09:30", "11:00", place="Room 1"),
                lec("Big Data", "11:00", "12:30", place="Room 60"), lec("Same Time Other Room", "13:00", "14:00", place="A")],
        "mon": [lec("Evening", "19:00", "20:30", mode="Activity", label="Activity / Assignment")],
        "tue": [lec("Early Course", "09:30", "11:00", place="Room 1"), lec("Late Course", "15:30", "17:00", place="Room 5"),
                lec("Big Data", "11:00", "12:30", mode="Activity", label="Activity / Self-Study"),
                lec("Same Time Other Room", "13:00", "14:00", place="B")],
        "wed": [lec("Evening", "19:00", "20:30", mode="Online", label="Online Lecture"),
                lec("Odd Times", "08:00", "09:00")],
        "thu": [lec("Odd Times", "10:00", "11:00", place="Hall")],
    }
    names = ["Late Course", "Early Course", "Big Data", "Same Time Other Room", "Evening", "Odd Times", "No Timetable"]
    state = {"courses": {str(n): {"name": f"20261-{name}-AllSec-3-", "url": f"https://x/c{n}"} for n, name in enumerate(names)},
             "teachers": {}}
    keep = (schedule.load_schedule, bot.load_state)
    schedule.load_schedule = lambda: {"days": days, "semester": "S"}
    schedule._arabic_to_english.cache_clear()
    bot.load_state = lambda: state
    try:
        text = bot.view_course_list()[0]
        titles = [l for l in text.splitlines() if l.startswith("📅")]
        assert titles == ["📅 *Sun · Tue*", "📅 *Mon · Wed*", "📅 *Wed · Thu*", "📅 *No timetable entry*"], titles
        listed = [l.split("[")[1].split("]")[0] for l in text.splitlines() if l[:2].rstrip(".").isdigit() or l[:3].rstrip(".").isdigit()]
        # Sun · Tue: by the time of the lectures, earliest first (9:30, 11:00, 13:00, 15:30); then the other patterns
        assert listed[:4] == ["Early Course", "Big Data", "Same Time Other Room", "Late Course"], listed
        assert listed.index("Evening") > listed.index("Late Course") and listed[-1] == "No Timetable"
        assert [l for l in text.splitlines() if l.startswith("📅")][0] == "📅 *Sun · Tue*"
        N = "\u00a0"                       # every indent is a non-breaking space: it cannot be trimmed
        assert not any(l.startswith(" ") for l in text.splitlines()), "no ordinary leading spaces"
        # a course block: name, instructor, section, the hour right under it; the kind of session and its room
        # on lines of their own, the room hanging from the kind
        assert ("1. 📘 [Early Course](https://x/c1)\n" + N * 4 + "🔢 Section 01\n" + N * 4 + "🕘 Sun · Tue · 9:30–11:00 AM\n"
                + N * 7 + "└ 🏫 In-Person\n" + N * 9 + "└ 📍 Room 1") in text
        assert "\n\n" + N * 4 + "🕘" not in text, "no blank line between the section and the hour"
        assert text.count("9:30–11:00 AM") == 1, "the same hour on both days is said once, alike days once"
        assert not any("In-Person · 📍" in line or "Activity · 📍" in line for line in text.splitlines())
        # the same hour but another kind of session / another room: the hour once, a branch per day
        assert (N * 4 + "🕘 Sun · Tue · 11:00 AM–12:30 PM\n" + N * 7 + "├ Sun — 🏫 In-Person\n" + N * 7 + "│" + N * 2 + "└ 📍 Room 60\n"
                + N * 7 + "└ Tue — 🧩 Activity") in text
        assert (N * 4 + "🕘 Sun · Tue · 1:00–2:00 PM\n" + N * 7 + "├ Sun — 🏫 In-Person\n" + N * 7 + "│" + N * 2 + "└ 📍 A\n"
                + N * 7 + "└ Tue — 🏫 In-Person\n" + N * 10 + "└ 📍 B") in text
        assert N * 4 + "🕘 Mon · Wed · 7:00–8:30 PM\n" + N * 7 + "├ Mon — 🧩 Activity\n" + N * 7 + "└ Wed — 💻 Online" in text
        # short words only; a day outside the campus days is its own entry when there is one; different hours: an entry each
        assert not any(w in text for w in ("In-Person Lecture", "Online Lecture", "Self-Study", "Assignment"))
        assert N * 4 + "🕘 Wed · 8:00–9:00 AM\n" + N * 7 + "└ 🏫 In-Person\n" in text
        assert N * 4 + "🕘 Thu · 10:00–11:00 AM\n" + N * 7 + "└ 🏫 In-Person\n" + N * 9 + "└ 📍 Hall" in text
        assert N * 4 + "🕘 Sat · 8:00–10:00 AM\n" + N * 7 + "└ 🏫 Lab" in text, "a day outside Sun–Thu is shown when the schedule has one"
        assert "──" not in text and _legacy_markdown_ok(text) and len(text) < notifier.MESSAGE_LIMIT
    finally:
        schedule.load_schedule, bot.load_state = keep
        schedule._arabic_to_english.cache_clear()

    # the pure helpers
    slots = [("sun", lec("A", "09:00", "10:00")), ("tue", lec("A", "09:00", "10:00"))]
    assert timetable.pattern(slots) == ("sun", "tue") and timetable.group_title(("sun", "tue")) == "Sun · Tue"
    assert timetable.group_title(()) == "Other days"
    N = "\u00a0"
    assert timetable.describe(slots) == ["🕘 Sun · Tue · 9:00–10:00 AM", N * 3 + "└ 🏫 In-Person"]
    assert timetable.time_range({"start": "09:30", "end": "11:00"}) == "9:30–11:00 AM"
    assert timetable.time_range({"start": "11:00", "end": "12:30"}) == "11:00 AM–12:30 PM"
    assert timetable.time_range({"start": "12:30", "end": "14:00"}) == "12:30–2:00 PM"
    two_a_day = [("sun", lec("A", "09:00", "10:00")), ("sun", lec("A", "12:00", "13:00"))]
    assert timetable.describe(two_a_day) == ["🕘 Sun · 9:00–10:00 AM", N * 3 + "└ 🏫 In-Person", "🕘 Sun · 12:00–1:00 PM", N * 3 + "└ 🏫 In-Person"], "two sessions on one day: an entry each"
    assert timetable.order_key([]) > timetable.order_key([("sat", lec("A", "08:00", "09:00"))]) > timetable.order_key(slots)

    # the same hierarchy in the lecture alerts and the morning summary
    one = lec("Arabic Communication Skills", "19:00", "20:30", mode="Online", label="Online Lecture", place="Hall 3")
    assert timetable.lecture_lines(one) == ["🕘 7:00–8:30 PM", N * 3 + "└ 💻 Online"], "online: no room, even if one is listed"
    live = lec("Big Data", "09:00", "10:00", place="Building 5, Room 2147")
    assert timetable.lecture_lines(live) == ["🕘 9:00–10:00 AM", N * 3 + "└ 🏫 In-Person", N * 5 + "└ 📍 Room 2147"]
    sched = {"lecture_reminder_minutes": 60, "morning_summary_hour": 7, "remind_modes": ["Online"], "days": {"wed": [{**one, "section": "31"}]}}
    wed = datetime(2026, 10, 21, 18, 30)                      # a Wednesday, half an hour before the hour
    by_kind = {kind: text for _, kind, text, _ in schedule.due_messages(sched, wed)}
    assert by_kind["pre"].endswith("\n🕘 7:00–8:30 PM\n" + N * 3 + "└ 💻 Online"), by_kind["pre"]
    summary = {kind: text for _, kind, text, _ in schedule.due_messages(sched, datetime(2026, 10, 21, 8, 0))}["summary"]
    assert ("🕘 7:00–8:30 PM · *Arabic Communication Skills* · Section 31\n" + N * 3 + "└ 💻 Online") in summary
    start = {kind: text for _, kind, text, _ in schedule.due_messages(sched, datetime(2026, 10, 21, 19, 2))}["start"]
    assert start.startswith("🎥 *Lecture starting now:*") and "\n🕘 7:00–8:30 PM\n" in start


def test_rooms_are_always_short():
    """A location is "📍 Room <number>" and nothing more — no building or faculty name, no block prefix —
    and a session that is not in a room (Activity, Online) has none, so a location line can never wrap."""
    from uninotifier import timetable
    N = "\u00a0"
    for raw, short in (("Building 5, Room 2147", "Room 2147"), ("Building 5, Room 2135", "Room 2135"),
                       ("Engineering College, Room 32123", "Room 32123"), ("Room 60-2104", "Room 2104"),
                       ("Room 449", "Room 449"), ("room 2147", "Room 2147"), ("Faculty of IT Room 12-345", "Room 345"),
                       ("Room A-12", "Room A-12"), ("Hall 3", "Hall 3"), ("  ", ""), (None, ""), ("", "")):
        assert timetable.room(raw) == short, (raw, timetable.room(raw))

    def lec(mode, place):
        return {"start": "09:30", "end": "11:00", "course": "X", "course_ar": "x", "section": "01", "mode": mode,
                "label": mode, "place": place}
    for mode in ("Activity", "Online"):
        lines = timetable.describe([("sun", lec(mode, "Building 5, Room 2147")), ("tue", lec(mode, "Room 60-2104"))])
        assert not any("📍" in line for line in lines), (mode, lines)
    # two days, one in a room and one remote: the room only under its day, in the one format
    lines = timetable.describe([("sun", lec("In-Person", "Room 60-2104")), ("tue", lec("Activity", "Room 60-2104"))])
    assert lines == ["🕘 Sun · Tue · 9:30–11:00 AM", N * 3 + "├ Sun — 🏫 In-Person", N * 3 + "│" + N * 2 + "└ 📍 Room 2104",
                     N * 3 + "└ Tue — 🧩 Activity"], lines

    # the real schedule: every room is already in that form, nothing remote has one, and the course list shows
    # nothing but "📍 Room <number>"
    import json
    schedule_data = json.loads(paths.SCHEDULE_FILE.read_text(encoding="utf-8"))
    for day, entries in schedule_data["days"].items():
        for item in entries:
            assert timetable.room(item.get("place")) == item.get("place", ""), (day, item["course"], item.get("place"))
            assert not (item["mode"] in timetable.REMOTE_MODES and item.get("place")), (day, item["course"])
    bot = _package("uninotifier.bot")
    text = bot.view_course_list()[0]
    places = [line.strip(N + " ") for line in text.splitlines() if "📍" in line]
    assert places and all(re.fullmatch("(│\u00a0*)?└ 📍 Room \\d+", p) for p in places), places
    assert "Building" not in text and "College" not in text and "60-2104" not in text and max(len(p) for p in places) <= 16


def test_poll_outage_tracking_and_quiet_logs():
    """One failed long poll is routine and silent; an outage is logged once when it starts and
    once when it ends, with a pause that grows; a failed poll is not a WARNING line each time."""
    import requests
    bot = _package("uninotifier.bot")
    records = []
    handler = logging.Handler()
    handler.emit = lambda record: records.append((record.levelno, record.getMessage()))
    logger = logging.getLogger("uninotifier")
    logger.addHandler(handler)
    try:
        tracker = bot.OutageTracker()
        pauses = [tracker.failed(now=100.0 + n) for n in range(3)]
        assert records == [] and pauses[0] < pauses[2] + 1
        tracker.failed(now=103.0)
        assert [m for lvl, m in records if lvl == logging.WARNING] == ["Telegram is unreachable: 4 polls in a row failed"]
        tracker.failed(now=104.0)
        assert len(records) == 1, "an outage is reported once"
        tracker.ok(now=160.0)
        assert tracker.failures == 0 and not tracker.reported
        tracker.failed(now=200.0)
        assert tracker.failures == 1 and tracker.since == 200.0

        # call_api: a failed getUpdates is a debug note, any other failed call still a warning
        records.clear()
        keep = telegram.post
        telegram.post = lambda url, **kw: (_ for _ in ()).throw(requests.ReadTimeout("slow"))
        try:
            assert notifier.call_api("T", "getUpdates", {}, timeout=1) is None
            assert [m for lvl, m in records if lvl >= logging.WARNING] == []
            assert notifier.call_api("T", "sendMessage", {"chat_id": 1}, timeout=1) is None
            assert len([m for lvl, m in records if lvl == logging.WARNING]) == 1
        finally:
            telegram.post = keep
    finally:
        logger.removeHandler(handler)
    assert 5 <= bot.POLL_TIMEOUT <= 50 and bot.POLL_SLACK >= 5


def test_network_resilience():
    """A dropped connection is retried (per request, then per check); a stalled git
    push can't outlive its time limit; a run killed holding the lock doesn't block
    the next ones; a retried check still attaches the files it downloaded."""
    import subprocess
    import sys

    import requests

    from uninotifier import gitsync
    from uninotifier import httpctx
    monitor = _package("uninotifier.monitor")

    logging.disable(logging.WARNING)      # the simulated failures stay out of monitor.log
    try:
        _resilience_checks(subprocess, sys, requests, gitsync, httpctx, monitor)
    finally:
        logging.disable(logging.INFO)


def _resilience_checks(subprocess, sys, requests, gitsync, httpctx, monitor):
    # 1. per request: connection resets / timeouts / 502-504 are retried with backoff
    class Resp:
        def __init__(self, code):
            self.status_code, self.url, self.headers, self.text = code, "u", {}, ""
    calls = []

    def flaky(outcomes):
        def request(method, url, timeout=None, allow_redirects=True, **kw):
            calls.append(url)
            out = outcomes.pop(0)
            if isinstance(out, Exception):
                raise out
            return Resp(out)
        return request
    keep = httpctx.BACKOFF_SECONDS
    httpctx.BACKOFF_SECONDS = 0
    try:
        http = httpctx.HttpContext(requests.cookies.RequestsCookieJar())
        session = http._session()
        session.request = flaky([requests.ConnectionError("Connection aborted. 10054"),
                                 requests.ReadTimeout("handshake timed out"), 200])
        assert http.get("https://x/a?sesskey=s").status == 200 and len(calls) == 3
        session.request = flaky([503, 200])
        assert http.get("https://x/b").status == 200
        session.request = flaky([404])
        assert http.get("https://x/c").status == 404, "a real answer is not retried"
        session.request = flaky([requests.ConnectionError("down")] * 3)
        try:
            http.get("https://x/d")
            raise AssertionError("should have raised after the retries")
        except requests.ConnectionError:
            pass
        session.request = flaky([502, 502, 502])
        assert http.get("https://x/e").status == 502, "after the retries the answer is returned as is"
    finally:
        httpctx.BACKOFF_SECONDS = keep

    # 2. per check: a network failure is retried, sign-in and other errors are not
    keep = (monitor.run_scan, monitor.SCAN_RETRY_SECONDS)
    monitor.SCAN_RETRY_SECONDS = (0, 0)
    attempts = []

    def scan(outcomes):
        def run_scan(saved, notify_new, mfa_handler=None):
            attempts.append(1)
            out = outcomes.pop(0)
            if isinstance(out, BaseException):
                raise out
            return out
        return run_scan
    try:
        monitor.run_scan = scan([requests.ConnectionError("Connection aborted."), {"ok": 1}])
        assert monitor.scan_with_retries({}, True) == {"ok": 1} and monitor.SUMMARY["attempts"] == 2
        for error, tries in ((monitor.LoginError("bad password"), 1), (ValueError("parser broke"), 1),
                             (requests.ReadTimeout("Read timed out"), 3)):
            attempts.clear()
            monitor.run_scan = scan([error] * 3)
            try:
                monitor.scan_with_retries({}, True)
                raise AssertionError("should have raised")
            except type(error):
                pass
            assert len(attempts) == tries, (error, attempts)
    finally:
        monitor.run_scan, monitor.SCAN_RETRY_SECONDS = keep

    # 3. a bounded command really ends — even when its own child holds the output open
    with tempfile.TemporaryDirectory() as tmp:
        marker = Path(tmp) / "grandchild.pid"
        hold = ("import subprocess, sys, time\n"
                "p = subprocess.Popen([sys.executable, '-c', 'import time; time.sleep(120)'])\n"
                f"open(r'{marker}', 'w').write(str(p.pid))\n"
                "time.sleep(120)\n")
        started = time.monotonic()
        try:
            gitsync.run_bounded([sys.executable, "-c", hold], timeout=3)
            raise AssertionError("should have timed out")
        except subprocess.TimeoutExpired:
            pass
        assert time.monotonic() - started < 20, time.monotonic() - started
        assert not monitor.pid_alive(int(marker.read_text())), "the whole tree is killed"
    assert gitsync.run_bounded([sys.executable, "-c", "print('hi')"], timeout=30).stdout.strip() == "hi"

    # 4. the run lock: a dead holder's lock is taken over at once; a live holder's is respected
    with tempfile.TemporaryDirectory() as tmp:
        keep = monitor.LOCK_FILE
        monitor.LOCK_FILE = Path(tmp) / "monitor.lock"
        try:
            dead = subprocess.Popen([sys.executable, "-c", "pass"])
            dead.wait()
            monitor.LOCK_FILE.write_text(str(dead.pid))
            assert monitor.acquire_lock() and monitor.LOCK_FILE.read_text() == str(os.getpid())
            alive = subprocess.Popen([sys.executable, "-c", "import time; time.sleep(30)"])
            try:
                monitor.LOCK_FILE.write_text(str(alive.pid))
                assert not monitor.acquire_lock()
            finally:
                alive.kill()
            monitor.LOCK_FILE.write_text("")                 # being written: its age decides
            assert not monitor.acquire_lock()
        finally:
            monitor.LOCK_FILE = keep

    # 5. a retried check finds its own download "identical" — still a new file, attached
    with tempfile.TemporaryDirectory() as tmp:
        saved = Path(tmp) / "Lecture 3.pdf"
        saved.write_bytes(b"%PDF")
        first = {"v": True}

        def download(context, url, dest):
            dup, first["v"] = not first["v"], False
            return {"path": saved, "sha256": "h", "modified": None, **({"duplicate": True} if dup else {})}
        keep = (monitor.archive.download_resource, monitor.archive.target)
        monitor.archive.download_resource = download
        monitor.archive.target = lambda course, kind, name: Path(tmp)
        try:
            item = {"name": "Lecture 3", "type": "resource", "course": "20261-Big Data-AllSec-3-",
                    "url": "https://x/mod/resource/view.php?id=901"}
            assert not monitor.process_item(None, "s", item).get("duplicate")     # attempt 1
            assert not monitor.process_item(None, "s", item)["duplicate"], "attempt 2: its own copy"
            other = {**item, "name": "Lecture 3 (again)", "url": "https://x/mod/resource/view.php?id=902"}
            assert monitor.process_item(None, "s", other)["duplicate"], "another item sharing the file"
        finally:
            monitor.archive.download_resource, monitor.archive.target = keep


class _FakePage:
    """Just enough of a Playwright page for the maintenance checks."""

    def __init__(self, text="", title="", png=b"PNG", broken=False, url="about:blank"):
        self.text, self._title, self.png, self.broken, self.url = text, title, png, broken, url

    def title(self):
        if self.broken:
            raise RuntimeError("page closed")
        return self._title

    def inner_text(self, selector, timeout=None):
        if self.broken:
            raise RuntimeError("page closed")
        return self.text

    def screenshot(self, type="png", timeout=None):
        if self.broken:
            raise RuntimeError("page closed")
        return self.png

    def goto(self, url, wait_until=None, timeout=None):
        pass


def test_sign_in_needs_no_form_when_microsoft_still_remembers_you():
    """Moodle's own session can lapse while Microsoft's has not: the OIDC redirect then comes straight
    back to the dashboard and no e-mail form exists. That is signed in (nothing is filled — it used to
    wait 30 s for the field and fail every run); a real Microsoft page still gets the e-mail filled."""
    from uninotifier import zuj_auth

    class Page:
        def __init__(self, lands_on):
            self.url, self.lands_on, self.filled = "about:blank", lands_on, []

        def goto(self, url, **kwargs):
            self.url = self.lands_on

        def fill(self, selector, text):
            self.filled.append((selector, text))
            raise zuj_auth.LoginError("stop here: the form was reached")

        click = wait_for_load_state = wait_for_timeout = fill

    class Handler:
        done = 0

        def succeeded(self):
            self.done += 1

    for landed in ("https://elearning.zuj.edu.jo/my/", "https://elearning.zuj.edu.jo/"):
        page, handler = Page(landed), Handler()
        zuj_auth.login_via_sso(page, "me@std-zuj.edu.jo", "pw", mfa_handler=handler)
        assert page.filled == [] and handler.done == 1, (landed, page.filled, handler.done)
    for form in ("https://login.microsoftonline.com/common/oauth2/authorize",
                 "https://elearning.zuj.edu.jo/auth/oidc/?source=loginpage",
                 "https://elearning.zuj.edu.jo/login/index.php"):
        page = Page(form)
        try:
            zuj_auth.login_via_sso(page, "me@std-zuj.edu.jo", "pw")
            raise AssertionError(f"{form} needs the form")
        except zuj_auth.LoginError:
            pass
        assert page.filled == [('input[type="email"]', "me@std-zuj.edu.jo")], (form, page.filled)


def test_maintenance_and_finals_pause():
    """A maintenance page is recognised (Arabic or English) and reported with a
    screenshot — the run ends cleanly, no exit code 1; an unexpected sign-in failure
    gets the screenshot too; a wrong password keeps its own message. Finals Watch
    pauses itself once every grade is published."""
    from datetime import timedelta

    bot = _package("uninotifier.bot")
    from uninotifier import finals
    from uninotifier import maintenance
    monitor = _package("uninotifier.monitor")
    from uninotifier import portal
    from uninotifier import zuj_auth

    # what counts as a maintenance notice
    for text, phrase in (("الموقع قيد الصيانة حالياً، يرجى المحاولة لاحقاً", "قيد الصيانة"),
                         ("نعتذر، النظام تحت  الصيانة", "تحت الصيانة"),
                         ("This site is undergoing maintenance and is currently not available", "undergoing maintenance"),
                         ("Under Maintenance", "Under Maintenance"),
                         ("503 Service Unavailable", "Service Unavailable"),
                         ("502 Bad Gateway — nginx", "Bad Gateway")):
        assert maintenance.detect(text) == phrase, (text, maintenance.detect(text))
    for text in ("Dashboard · Big Data · Database Systems", "اسم المستخدم / كلمة المرور غير صحيحة",
                 "Incorrect username or password", "", None):
        assert maintenance.detect(text) is None, text

    with tempfile.TemporaryDirectory() as tmp:
        keep_snapshot = maintenance.SNAPSHOT
        maintenance.SNAPSHOT = Path(tmp) / "maintenance_snapshot.png"
        logging.disable(logging.CRITICAL)          # the simulated failures stay out of the real logs
        try:
            # unavailable(): reads the page, keeps the screenshot, never raises itself
            e = maintenance.unavailable(_FakePage("نعتذر، البوابة قيد الصيانة"), "registration portal",
                                        TimeoutError("Timeout 30000ms exceeded"))
            assert e.phrase == "قيد الصيانة" and e.png == b"PNG" and maintenance.SNAPSHOT.read_bytes() == b"PNG"
            text = maintenance.alert_text(e)
            assert text.startswith("⚠️ *تعذر الدخول للبوابة، مرفق لقطة للشاشة الحالية للتحقق من الحالة*"), text
            assert "Under maintenance" in text and "قيد الصيانة" in text and _legacy_markdown_ok(text)
            e = maintenance.unavailable(_FakePage(broken=True), "e-learning portal", "the form never appeared")
            assert e.phrase is None and e.png is None and e.detail == "the form never appeared"
            text = maintenance.alert_text(e, "Try again later.")
            assert text.splitlines()[0] == maintenance.HEADLINE and "not a wrong password" in text
            assert "E-learning portal" in text and "مرفق لقطة" not in text, "no picture: no mention of one"
            assert _legacy_markdown_ok(text)

            # a failure on a page that is already inside the portal says so, instead of a bare timeout
            timeout = RuntimeError('Page.fill: Timeout 30000ms exceeded. Call log: - waiting for locator("input[type=\\"email\\"]")')
            e = maintenance.unavailable(_FakePage("Dashboard · Course overview"), "e-learning portal", timeout, signed_in=True)
            assert e.phrase is None and e.signed_in and "already signed in" in str(e)
            text = maintenance.alert_text(e)
            assert "Your password is fine" in text and "already inside the portal" in text, text
            assert "uninotifier.tools.login_refresh" in text and "next check tries again by itself" in text, text
            assert "not a wrong password" not in text and "Timeout 30000ms" in text and _legacy_markdown_ok(text), text
            e = maintenance.unavailable(_FakePage("Service Unavailable"), "e-learning portal", timeout, signed_in=True)
            assert "Under maintenance" in maintenance.alert_text(e), "a maintenance notice still wins"

            # the Moodle sign-in sorts its failures
            browser = monitor.Browser()
            keep_auth = monitor.is_authenticated
            try:
                for error, page_text, expected in (
                        (zuj_auth.InvalidCredentialsError("Your password is incorrect"), "", zuj_auth.InvalidCredentialsError),
                        (zuj_auth.MFARequiredError("2FA"), "", zuj_auth.MFARequiredError),
                        (RuntimeError("net::ERR_CONNECTION_RESET at https://login.microsoftonline.com"), "", RuntimeError),
                        (RuntimeError("net::ERR_CONNECTION_RESET"), "Service Unavailable", maintenance.PortalUnavailable),
                        (RuntimeError("Page.fill: Timeout 30000ms exceeded.\nCall log:\n  - waiting for locator('input')"),
                         "", maintenance.PortalUnavailable),
                        (zuj_auth.LoginError("Sign-in flow did not reach the portal"), "", maintenance.PortalUnavailable)):
                    browser._page = _FakePage(page_text)
                    monitor.is_authenticated = lambda page, error=error: (_ for _ in ()).throw(error)
                    try:
                        browser.sign_in()
                        raise AssertionError("should have raised")
                    except Exception as raised:
                        assert type(raised) is expected, (error, raised)
                # the same timeout, with the browser on the dashboard / on Microsoft's page / on the OIDC redirect
                for url, inside in (("https://elearning.zuj.edu.jo/my/", True),
                                    ("https://login.microsoftonline.com/common/oauth2/authorize", False),
                                    ("https://elearning.zuj.edu.jo/auth/oidc/?source=loginpage", False)):
                    browser._page = _FakePage("Dashboard", url=url)
                    monitor.is_authenticated = lambda page: (_ for _ in ()).throw(RuntimeError(
                        'Page.fill: Timeout 30000ms exceeded.\nCall log:\n  - waiting for locator("input[type=\\"email\\"]")'))
                    try:
                        browser.sign_in()
                        raise AssertionError("should have raised")
                    except maintenance.PortalUnavailable as raised:
                        assert raised.signed_in is inside, (url, raised.signed_in)
            finally:
                monitor.is_authenticated = keep_auth
                browser._page = None

            # the registration portal: a maintenance page instead of the sign-in form
            try:
                portal.login(_FakePage("Service Unavailable"), "u", "p", lambda *a: None)
                raise AssertionError("should have raised")
            except maintenance.PortalUnavailable as raised:
                assert raised.site == "registration portal" and raised.phrase == "Service Unavailable"

            # monitor.run: maintenance ends cleanly (exit 0) with the screenshot; then "back to normal"
            keep = {k: getattr(monitor, k) for k in ("DATA_FILE", "USERNAME", "PASSWORD", "BOT_TOKEN", "CHAT_ID",
                                                     "run_scan", "send_schedule_messages", "mfa_allowed",
                                                     "send_telegram", "SCAN_RETRY_SECONDS")}
            keep_png, keep_summary = monitor.notifier.send_png_file, dict(monitor.SUMMARY)
            monitor.DATA_FILE = Path(tmp) / "saved_items.json"
            monitor.USERNAME = monitor.PASSWORD = monitor.BOT_TOKEN = monitor.CHAT_ID = "x"
            monitor.send_schedule_messages = lambda saved: None
            monitor.mfa_allowed = lambda saved: False
            monitor.SCAN_RETRY_SECONDS = (0, 0)
            pngs, texts, scans = [], [], []
            monitor.notifier.send_png_file = lambda token, chat, png, name, caption="", url=None, rows=None: \
                pngs.append((name, caption)) or True
            monitor.send_telegram = lambda text, url=None: texts.append(text) or True
            course = {"1": {"name": "20261-Big Data-AllSec-3-", "url": "https://x/c"}}
            item = {"1-1": {"name": "Lecture 1", "type": "label", "course": "20261-Big Data-AllSec-3-", "cmid": "1",
                            "url": "https://x/mod/label/view.php?id=1", "hash": "h"}}
            monitor.save_data({**monitor.load_saved_data(), "courses": dict(course), "content": dict(item)})
            before = json.loads(monitor.DATA_FILE.read_text(encoding="utf-8"))

            def failing(error):
                def run_scan(saved, notify_new, mfa_handler=None):
                    scans.append(1)
                    raise error
                return run_scan

            def exit_code():
                try:
                    monitor.run()
                except SystemExit as ex:
                    return ex.code
                return 0
            try:
                monitor.run_scan = failing(maintenance.PortalUnavailable(
                    "e-learning portal", "the session was refused", "Under Maintenance", b"PNG"))
                assert exit_code() == 0 and len(scans) == 1, "no crash, and no pointless retries"
                assert pngs[-1][0] == "maintenance_snapshot.png" and "Under maintenance" in pngs[-1][1]
                assert "tell you when it's back" in pngs[-1][1] and monitor.SUMMARY["maintenance"] == "Under Maintenance"
                after = json.loads(monitor.DATA_FILE.read_text(encoding="utf-8"))
                assert after["content"] == before["content"] and after["courses"] == before["courses"], "state untouched"
                assert exit_code() == 0 and len(pngs) == 1, "a second maintenance run: no repeat alert (cooldown)"

                monitor.run_scan = failing(maintenance.PortalUnavailable(
                    "e-learning portal", "the form never appeared", None, b"PNG"))
                assert exit_code() == 1 and "the form never appeared (not a wrong password)" in pngs[-1][1]

                monitor.run_scan = failing(zuj_auth.InvalidCredentialsError("Your password is incorrect"))
                assert exit_code() == 1 and "Portal login failed" in texts[-1] and len(pngs) == 2, "no screenshot for a wrong password"

                monitor.SUMMARY.clear()
                monitor.run_scan = lambda saved, notify_new, mfa_handler=None: {
                    "courses": dict(course), "content": dict(item), "deadlines": {}, "artifacts": {}, "updates": {},
                    "edits": {}, "posts": {}, "forum_posts": {}, "submissions": {}, "link_changes": [],
                    "seconds": 1.0, "browser_used": False}
                keep_recap = (monitor.daily.catch_up, monitor.daily.send_recap)
                monitor.daily.catch_up = monitor.daily.send_recap = lambda *a, **k: False
                try:
                    assert exit_code() == 0 and any("Back to normal" in t for t in texts), texts
                finally:
                    monitor.daily.catch_up, monitor.daily.send_recap = keep_recap
            finally:
                for k, v in keep.items():
                    setattr(monitor, k, v)
                monitor.notifier.send_png_file = keep_png
                monitor.SUMMARY.clear()
                monitor.SUMMARY.update(keep_summary)

            # the bot: an on-demand portal request that meets maintenance
            sent_png, views = [], []
            orig = (bot.portal.fetch_many, bot.notifier.send_png_file, bot.send_view, bot.edit_view, bot.load_last_run)
            bot.portal.fetch_many = lambda sections, ask: (_ for _ in ()).throw(maintenance.PortalUnavailable(
                "registration portal", "the sign-in page shows a maintenance notice", "قيد الصيانة", b"PNG"))
            bot.notifier.send_png_file = lambda token, chat, png, name, caption="", url=None, rows=None: \
                sent_png.append(caption) or True
            bot.send_view = lambda text, rows: views.append(text) or {"message_id": 1}
            bot.edit_view = lambda message_id, text, rows: views.append(text)
            try:
                assert bot._portal_lock.acquire(blocking=False)
                bot.run_portal("grades", None)
                assert "Registration portal" in sent_png[-1] and "Under maintenance" in sent_png[-1]
                assert sent_png[-1].startswith("⚠️ *تعذر الدخول للبوابة") and "Try again" in sent_png[-1]
                assert views[-1].startswith("📊 *GPA & Grade Hub*") and not bot._portal_lock.locked(),                     "the screenshot says why; the hub it came from returns under it (no extra status message)"
                assert not any("retrieved" in v or "maintenance." in v for v in views)
                bot.load_last_run = lambda: {"finished": time.time() - 60, "ok": False, "maintenance": "Under Maintenance"}
                assert bot.monitor_status().startswith("🛠️ *Portal under maintenance*")
            finally:
                (bot.portal.fetch_many, bot.notifier.send_png_file, bot.send_view, bot.edit_view, bot.load_last_run) = orig
        finally:
            maintenance.SNAPSHOT = keep_snapshot
            logging.disable(logging.INFO)

    # Finals Watch pauses itself when every grade is out
    with tempfile.TemporaryDirectory() as tmp:
        keep = (finals.STATE_FILE, finals.evaluate)
        finals.STATE_FILE = Path(tmp) / "finals.json"
        published = {"n": 2}

        def evaluate(state, grades, transcript):
            state.update(term="First semester 2026/2027", published=published["n"], total=3,
                         complete=published["n"] == 3)
            return [{"text": "🎓 *Grade Released*", "rows": []}]
        finals.evaluate = evaluate
        try:
            today = date.today()
            sch = {"finals": {"from": str(today - timedelta(days=1)), "to": str(today + timedelta(days=20))}}
            state = finals.load()
            assert finals.active(state=state, sch=sch), "inside the window: on by itself"
            msgs, state = finals.check(lambda sections: {"transcript": None}, state, sch)
            assert len(msgs) == 1 and finals.active(state=state, sch=sch), "2 of 3: keeps watching"
            published["n"] = 3
            msgs, state = finals.check(lambda sections: {"transcript": None}, state, sch)
            assert [m["text"].splitlines()[0] for m in msgs] == ["🎓 *Grade Released*", "⏸️ *Finals Watch Paused*"]
            assert msgs[-1]["rows"][0][0]["callback_data"] == "finals:on"
            assert not finals.active(state=finals.load(), sch=sch) and not finals.due(state=finals.load(), sch=sch)
            assert "paused automatically" in finals.status_text(finals.load(), sch)
            msgs, _ = finals.check(lambda sections: {"transcript": None}, finals.load(), sch)
            assert len(msgs) == 1, "a manual check while paused doesn't announce the pause again"
            # /finals on resumes; the next complete reading pauses again
            state = finals.load()
            state["enabled"] = True
            assert finals.active(state=state, sch=sch)
            msgs, state = finals.check(lambda sections: {"transcript": None}, state, sch)
            assert not state["enabled"] and "Paused" in msgs[-1]["text"]
            # next season's window arms it again
            later = {"finals": {"from": str(today), "to": str(today + timedelta(days=20))}}
            assert finals.active(state=state, sch=later)
            for m in msgs:
                assert _legacy_markdown_ok(m["text"])
        finally:
            finals.STATE_FILE, finals.evaluate = keep


TEAMS_PAGE = """<html><body><div role="main"><h2>رابط الدخول إلى المحاضرة</h2>
<table class="generaltable w-auto mb-4"><tbody>
<tr><td class="cell c0">Meeting type</td><td class="cell c1 lastcol">Recurring meeting</td></tr>
<tr><td class="cell c0">Repeat on</td><td class="cell c1 lastcol">Monday, Wednesday</td></tr>
<tr><td class="cell c0">Meeting time</td><td class="cell c1 lastcol">7:00 PM – 8:30 PM</td></tr>
<tr class="lastrow"><td class="cell c0">Organizer</td><td class="cell c1 lastcol">Dr. Someone</td></tr>
</tbody></table>
<input type="text" id="teamscheduler-meetingurl" class="form-control"
 value="https://teams.microsoft.com/l/meetup-join/19%3ameeting_ABC%40thread.v2/0?context=%7b%22Tid%22%3a%22x%22%7d" readonly>
<form method="get" action="https://teams.microsoft.com/l/meetup-join/19%3ameeting_ABC%40thread.v2/0">
<button type="submit">Join Teams Meeting</button></form></div></body></html>"""

ATTENDANCE_PAGE = """<html><body><div role="main"><table class="generaltable attwidth"><thead><tr>
<th class="header c0 datecol">Date</th><th class="header c1 desccol">Description</th>
<th class="header c2 statuscol">Status</th><th class="header c3 pointscol">Points</th></tr></thead><tbody>
<tr><td class="datecol cell c0"><nobr>Mon 12 Oct 2026</nobr> <nobr>7PM - 8:30PM</nobr></td>
 <td class="desccol cell c1"><div class="text_to_html">Regular class session</div></td>
 <td class="statuscol cell c2">Present</td><td class="pointscol cell c3">2 / 2</td></tr>
<tr><td class="datecol cell c0"><nobr>Wed 14 Oct 2026</nobr> <nobr>7PM - 8:30PM</nobr></td>
 <td class="desccol cell c1"><div class="text_to_html">Regular class session</div></td>
 <td class="statuscol cell c2"><a href="https://elearning.zuj.edu.jo/mod/attendance/attendance.php?sessid=55&amp;sesskey=SECRET">Submit attendance</a></td>
 <td class="pointscol cell c3">? / 2</td></tr>
<tr><td class="datecol cell c0"><nobr>Mon 19 Oct 2026</nobr> <nobr>7PM - 8:30PM</nobr></td>
 <td class="desccol cell c1"><div class="text_to_html">Regular class session</div></td>
 <td class="statuscol cell c2">?</td><td class="pointscol cell c3">? / 2</td></tr>
</tbody></table></div></body></html>"""

CHECK_IN_PAGE = """<html><body><form action="attendance.php" method="post">
<input type="hidden" name="sessid" value="55"><input type="hidden" name="sesskey" value="SECRET">
<input type="radio" name="status" id="s1" value="101"><label for="s1">Present</label>
<input type="radio" name="status" id="s2" value="102"><label for="s2">Late</label>
<input type="submit" name="submitbutton" value="Save changes"></form></body></html>"""


def test_teams_meeting_and_attendance_items():
    """teamscheduler and attendance are no longer "Unclassified": a Teams meeting comes
    with its details and a 🔗 Join Teams Lecture button (its link also feeds the lecture
    reminders); attendance with the next session and ✍️ Check In. Self check-in finds the
    open session and plans the POST, but never submits until verified."""
    from uninotifier import activities
    monitor = _package("uninotifier.monitor")

    course = "20261-الاتصال والتواصل باللغة العربية-ش31-3-ك"
    teams = {"name": "رابط الدخول إلى المحاضرة", "type": "teamscheduler", "course": course, "cmid": "76459",
             "url": "https://elearning.zuj.edu.jo/mod/teamscheduler/view.php?id=76459", "section": "روابط هامة ورئيسية"}
    attend = {**teams, "name": "الحضور", "type": "attendance", "cmid": "75482",
              "url": "https://elearning.zuj.edu.jo/mod/attendance/view.php?id=75482",
              "_body": "<p>الرجاء تأكيد الحضور بالدخول إلى هذا الرابط .</p>"}
    assert media.classify(teams) == "meeting" and media.classify(attend) == "attendance"

    meeting = activities.parse_meeting(TEAMS_PAGE)
    assert meeting["url"].startswith("https://teams.microsoft.com/l/meetup-join/19%3ameeting_ABC")
    assert ("Repeat on", "Monday, Wednesday") in meeting["details"] and len(meeting["details"]) == 4
    assert activities.parse_meeting("<html><body>nothing</body></html>") == {"url": None, "details": []}

    sessions = activities.parse_sessions(ATTENDANCE_PAGE)
    assert [s["date"] for s in sessions] == ["Mon 12 Oct 2026", "Wed 14 Oct 2026", "Mon 19 Oct 2026"]
    assert sessions[0]["start"] == datetime(2026, 10, 12, 19, 0) and sessions[0]["status"] == "Present"
    nxt = activities.next_session(sessions, datetime(2026, 10, 13, 9, 0))
    assert nxt["date"] == "Wed 14 Oct 2026"
    assert activities.next_session(sessions, datetime(2026, 10, 14, 20, 0))["date"] == "Wed 14 Oct 2026", "in progress"
    assert activities.open_session(sessions)["submit_url"].endswith("sessid=55&sesskey=SECRET")

    # the messages: no raw URL, a button instead; nothing leaks the session key
    text = media.build_meeting_message(teams, meeting)
    assert text.startswith("💻 *New Online Lecture*") and "Microsoft Teams Meeting" in text
    assert "📆 *Repeat on:* Monday, Wednesday" in text and "teams.microsoft.com" not in text
    assert media.meeting_rows(teams, meeting) == [[{"text": "🔗 Join Teams Lecture", "url": meeting["url"]}]]
    assert media.meeting_rows(teams, None)[0][0]["url"] == teams["url"], "no link read: the portal page"
    text = media.build_attendance_message(attend, sessions, datetime(2026, 10, 13, 9, 0))
    assert text.startswith("📝 *Attendance Check-in*") and "⏱️ *Type:* Attendance Check-in" in text
    assert "Next session:* Wed 14 Oct 2026 · 7PM - 8:30PM" in text and "SECRET" not in text
    assert media.attendance_rows(attend) == [[{"text": "✍️ Check In / تأكيد الحضور", "url": attend["url"]}]]
    for t in (text, media.build_meeting_message(teams, meeting), media.build_meeting_message(teams, None)):
        assert _legacy_markdown_ok(t), t

    # the monitor: reads the page, keeps the join link as the item's target, sends the right message
    class Ctx:
        def __init__(self):
            self.request = self

        def get(self, url, timeout=None, max_redirects=None):
            class R:
                ok, status = True, 200

                def text(self_inner):
                    return TEAMS_PAGE if "teamscheduler" in url else ATTENDANCE_PAGE
            return R()
    art = monitor.process_item(Ctx(), "s", teams)
    assert art["kind"] == "meeting" and art["target"] == meeting["url"] and not monitor.needs_screenshot(art)
    art2 = monitor.process_item(Ctx(), "s", attend)
    assert art2["kind"] == "attendance" and len(art2["sessions"]) == 3 and not monitor.needs_screenshot(art2)
    sent = []
    orig = monitor.deliver
    monitor.deliver = lambda header, body="", url=None, rows=None, shots=(), shot_name="", document=None: \
        sent.append((header, body, rows)) or True
    try:
        assert monitor.notify_new_item(teams, None, art) and monitor.notify_new_item(attend, None, art2)
    finally:
        monitor.deliver = orig
    assert sent[0][0].startswith("💻") and sent[0][2][0][0]["text"] == "🔗 Join Teams Lecture"
    assert sent[1][0].startswith("📝") and sent[1][2][0][0]["text"].startswith("✍️ Check In")
    assert "الرجاء تأكيد الحضور" in sent[1][1], "the instructor's note follows as the body"

    # self check-in: plans the POST for "Present", never sends it until verified
    posted = []

    class Http:
        def get(self, url, timeout=None):
            class R:
                ok, status = True, 200

                def text(self_inner):
                    return CHECK_IN_PAGE
            return R()

        def post(self, url, data=None, headers=None):
            posted.append((url, data))
    result = activities.check_in(Http(), sessions)
    assert result["reason"] == "dry run" and result["plan"]["url"].endswith("/mod/attendance/attendance.php")
    assert result["plan"]["data"] == {"sessid": "55", "sesskey": "SECRET", "status": "101",
                                      "submitbutton": "Save changes"}, result["plan"]
    assert activities.check_in(Http(), sessions, dry_run=False)["reason"] == "submitting is not verified yet"
    assert posted == [], "nothing is ever submitted before SUBMIT_VERIFIED"
    assert activities.check_in(Http(), sessions[:1])["reason"] == "no session is open for check-in"


def test_tasks_and_reminders_planner():
    """Upcoming Tasks & Reminders: deadlines with 🔗 Open Task, personal reminders
    (➕ typed, 📌 from a post) scheduled by buttons, delivered once at the minute;
    Monday stays an activity — no lecture alerts."""
    from datetime import timedelta

    bot = _package("uninotifier.bot")
    monitor = _package("uninotifier.monitor")
    from uninotifier import reminders

    sch = schedule.load_schedule()
    # Monday is an activity: no lecture alert or Join button, Wednesday's online lecture keeps them
    for hm in ((18, 0), (19, 0)):
        assert schedule.due_messages(sch, datetime(2026, 10, 19, *hm)) == [], "Monday"
    assert [k for _, k, _, _ in schedule.due_messages(sch, datetime(2026, 10, 21, 19, 0))] == ["start"]

    # when: next campus day, before the next lecture
    tue = datetime(2026, 10, 13, 18, 0)
    # campus days are Sunday–Wednesday only: Saturday's labs and Wednesday's online lecture don't count
    assert reminders.next_campus_day(tue, sch) == datetime(2026, 10, 18, 8, 30), "Sunday, not Saturday's lab"
    assert reminders.next_campus_day(datetime(2026, 10, 17, 7, 30), sch) == datetime(2026, 10, 18, 8, 30)
    assert reminders.next_campus_day(datetime(2026, 10, 18, 9, 0), sch) == datetime(2026, 10, 20, 8, 30), "then Tuesday"
    # before next lecture: in-person lectures only (Tuesday's Big Data slot is an activity)
    assert reminders.before_next_lecture("Big Data", tue, sch) == (datetime(2026, 10, 18, 10, 30), "Big Data")
    assert reminders.before_next_lecture("Big Data", datetime(2026, 10, 18, 12, 0), sch)[0] == \
        datetime(2026, 10, 25, 10, 30), "not Tuesday's activity slot"
    assert reminders.before_next_lecture("Arabic Communication Skills", tue, sch) == \
        (datetime(2026, 10, 18, 9, 0), "Linear Algebra (1)"), "an online-only course: the next class on campus"
    assert reminders.before_next_lecture(None, tue, sch)[0] == datetime(2026, 10, 18, 9, 0), "any course"
    bot._reminder_draft.update(note="x", course="Arabic Communication Skills")
    assert "has no in-person lecture" in bot.view_when(bot._reminder_draft)[0]
    bot._reminder_draft.clear()
    for text, hm in (("18:45", (18, 45)), ("6:45 pm", (18, 45)), ("6PM", (18, 0)), ("12 am", (0, 0)), ("9", (9, 0))):
        assert reminders.parse_time(text) == hm, text
    for bad in ("25:00", "abc", "6:75"):
        try:
            reminders.parse_time(bad)
            raise AssertionError(bad)
        except ValueError:
            pass
    labels = [label for label, _ in reminders.day_choices(datetime(2026, 10, 13, 10, 0))]
    assert labels[:3] == ["Today", "Tomorrow", "Next Sunday"] and "Thu 15" in labels and len(labels) == 9
    assert reminders.day_choices(datetime(2026, 10, 13, 23, 0))[0][0] == "Tomorrow"
    assert reminders.when_text(datetime(2026, 10, 14, 18, 0).timestamp()) == "Wed 14 Oct · 6:00 PM"

    with tempfile.TemporaryDirectory() as tmp:
        keep = (reminders.STATE_FILE, reminders.SOURCES_FILE, reminders.CLAIM_DIR)
        reminders.STATE_FILE, reminders.SOURCES_FILE = Path(tmp) / "r.json", Path(tmp) / "s.json"
        reminders.CLAIM_DIR = Path(tmp) / "sent"
        sent, views, edits, answers = [], [], [], []
        orig = (reminders.notifier.send_message, bot.send_view, bot.edit_view, bot.api, bot.load_state)
        replies = []
        reminders.notifier.send_message = lambda token, chat, text, url=None, rows=None, silent=False, reply_to=None: \
            sent.append((text, rows)) or replies.append(reply_to) or True
        bot.send_view = lambda text, rows: views.append((text, rows)) or {"message_id": 1}
        bot.edit_view = lambda mid, text, rows: edits.append((text, rows))
        api_calls, api_ok = [], {"v": False}

        def fake_api(method, **data):
            answers.append((method, data.get("text")))
            api_calls.append((method, data))
            return {"message_id": 99} if api_ok["v"] and method == "sendMessage" else None
        bot.api = fake_api
        try:
            # delivery: once, at the minute; snooze makes it a new delivery
            item = reminders.add("Buy the textbook", time.time() - 5, "Big Data", "https://x/post")
            assert reminders.send_due("t", "c") == 1 and reminders.send_due("t", "c") == 0
            text, rows = sent[-1]
            assert text.startswith("⏰ *Reminder*") and "Buy the textbook" in text and _legacy_markdown_ok(text)
            assert rows[0][0] == {"text": "🌐 Open on Moodle", "url": "https://x/post"}
            assert replies[-1] is None, "a typed reminder quotes nothing"
            assert [b["callback_data"] for b in rows[1]] == [f"rem:z:{item['id']}", f"rem:done:{item['id']}"]
            assert reminders.pending() == [], "delivered: no longer upcoming"
            reminders.snooze(item["id"], now=time.time() - 3600 - 60)   # due a minute ago: a new delivery
            assert reminders.send_due("t", "c") == 1, "snoozed: sent again"

            # 📌 under a forum post: the monitor files the source, the button carries a short key
            delivered = []
            keep_deliver = monitor.deliver
            monitor.deliver = lambda header, body="", url=None, rows=None, shots=(), shot_name="", document=None: \
                delivered.append(rows) or True
            try:
                forum = {"name": "Announcements", "course": "20261-Big Data-AllSec-3-", "url": "https://x/f"}
                post = {"id": "101", "subject": "Install Hadoop before Sunday", "author": "Dr", "time": 1,
                        "html": "<p>x</p>", "is_reply": False, "url": "https://x/d#p101", "attachments": []}
                assert monitor.notify_forum_post(forum, post)
            finally:
                monitor.deliver = keep_deliver
            assert delivered[-1] == [[{"text": "📌 Add to Reminders", "callback_data": "rem:s:p101"}]]
            assert reminders.source("p101") == {**reminders.source("p101"), "note": "Install Hadoop before Sunday",
                                                "course": "Big Data", "url": "https://x/d#p101"}

            tomorrow = (datetime.now() + timedelta(days=1)).strftime("%Y%m%d")
            press = iter(range(10 ** 6))                       # every press has its own id, as on Telegram
            cb = lambda data, mid=5: bot.handle_callback({"id": f"p{next(press)}", "data": data, "message": {"message_id": mid}})
            cb("rem:s:p101")
            text, rows = views[-1]
            assert text.startswith("📌 *Add to Reminders*") and "Install Hadoop" in text and "Big Data" in text
            assert [r[0]["callback_data"] for r in rows] == ["rem:q:campus", "rem:q:lecture", "rem:q:custom", "rem:cancel"]
            assert rows[0][0]["text"].startswith("🏫 Next Campus Day · ") and rows[1][0]["text"].startswith("⏰ Before Next Lecture · ")
            cb("rem:q:custom")
            assert edits[-1][0].startswith("🗓️ *Pick a day*")
            cb(f"rem:day:{tomorrow}")
            assert edits[-1][0].startswith("🕐 *Pick a time") and edits[-1][1][0][0]["callback_data"] == f"rem:hr:{tomorrow}:0800"
            cb(f"rem:hr:{tomorrow}:1800")
            assert edits[-1][0].startswith("✅ *Reminder set*")
            mine = reminders.pending()
            assert len(mine) == 1 and mine[0]["note"] == "Install Hadoop before Sunday" and mine[0]["url"] == "https://x/d#p101"
            assert mine[0]["message_id"] == 5, "remembers the chat message the 📌 sat under"
            assert datetime.fromtimestamp(mine[0]["due"]).strftime("%Y%m%d %H:%M") == f"{tomorrow} 18:00"
            cb("rem:s:p101")
            cb("rem:q:lecture")
            assert edits[-1][0].startswith("✅ *Reminder set*") and len(reminders.pending()) == 2
            cb("rem:q:campus")
            assert answers[-1] == ("answerCallbackQuery", "⌛ That reminder draft expired — start again."), "used up"
            cb("rem:s:p999")
            assert "too old" in answers[-1][1]

            # ➕ typed note -> typed time
            bot.handle_command("/remind")
            assert views[-1][0].startswith("➕ *New Custom Reminder*") and bot._reminder_wait == "note"
            bot.handle_command("اشتري كتاب كلية الآداب طابق 1")
            assert views[-1][0].startswith("📌 *Add to Reminders*") and "كلية الآداب" in views[-1][0]
            cb("rem:q:custom")
            cb(f"rem:type:{tomorrow}")
            assert bot._reminder_wait == "time"
            bot.handle_command("25:99")
            assert views[-1][0].startswith("⚠️") and bot._reminder_wait == "time", "a bad time keeps the question"
            bot.handle_command("6:45 pm")
            assert views[-1][0].startswith("✅ *Reminder set*") and bot._reminder_wait is None
            assert len(reminders.pending()) == 3
            bot.handle_command("/remind")
            cb("rem:cancel")
            assert bot._reminder_wait is None and edits[-1][0].startswith("✖️")

            # the planner view: deadlines with Open Task, reminders, ➕ and 🗑️
            bot.load_state = lambda: {"deadlines": {"assign-1": {
                "name": "Assignment 2 — ER Diagram", "type": "assign", "course": "20261-Databases-AllSec-3-F",
                "url": "https://x/mod/assign/view.php?id=1", "timestamp": time.time() + 86400 * 2}}}
            text, rows = bot.view_tasks()
            assert "Course deadlines" in text and "Assignment 2" in text and "Your reminders" in text
            assert "كلية الآداب" in text and _legacy_markdown_ok(text)
            assert rows[0][0]["text"].startswith("🔗 Open Task 1 · Assignment 2") and rows[0][0]["url"].endswith("id=1")
            views_btns = [r[0] for r in rows if r[0].get("callback_data", "").startswith("rem:view:")]
            pinned = [r for r in reminders.pending() if r.get("message_id")]
            assert len(views_btns) == len(pinned) == 2, "one per pinned reminder; the typed one has none"
            assert views_btns[0]["text"].startswith("🔗 View Original Post · Install Hadoop")
            assert [b["callback_data"] for b in rows[3]] == ["rem:new", "rem:list"]

            # 🔗 View Original Post: a reply quoting the pinned message (tap the quote to jump)
            api_ok["v"] = True
            cb(views_btns[0]["callback_data"])
            method, data = api_calls[-2]
            assert method == "sendMessage" and data["reply_parameters"] == {"message_id": 5}
            assert data["text"].startswith("📌 *Original post*") and "Install Hadoop" in data["text"]
            # ...and if that message was deleted, the Moodle link instead
            api_ok["v"] = False
            cb(views_btns[0]["callback_data"])
            assert views[-1][0].startswith("⚠️ The original message was deleted")
            assert views[-1][1][0][0] == {"text": "🌐 Open on Moodle", "url": "https://x/d#p101"}
            typed = next(r for r in reminders.pending() if not r.get("message_id"))
            cb(f"rem:view:{typed['id']}")
            assert answers[-1][1] == "That reminder has no original post."

            cb("rem:list")
            assert edits[-1][0].startswith("🗑️ *Manage Reminders*") and len(edits[-1][1]) == 4
            first = reminders.pending()[0]["id"]
            cb(f"rem:del:{first}")
            assert len(reminders.pending()) == 2 and len(edits[-1][1]) == 3
            # the reminder itself arrives as a reply to the original post
            state = reminders.load()
            for r in state["items"]:
                if r.get("message_id"):
                    r["due"] = time.time() - 1
            reminders.save(state)
            assert reminders.send_due("t", "c") == 1 and replies[-1] == 5
        finally:
            (reminders.notifier.send_message, bot.send_view, bot.edit_view, bot.api, bot.load_state) = orig
            reminders.STATE_FILE, reminders.SOURCES_FILE, reminders.CLAIM_DIR = keep
            bot._reminder_draft.clear()
            bot._reminder_wait = None


def test_pin_rule_and_forum_announcements():
    """📌 Add to Reminders on every notification except assignments and quizzes (those are
    tracked by themselves); a forum whose description carries an announcement is announced,
    an empty forum stays silent."""
    monitor = _package("uninotifier.monitor")
    from uninotifier import reminders

    course = "20261-الاتصال والتواصل باللغة الإنجليزية-ش27-3-و"
    notice = ("<h3>إعلان هام لطلبة المساق</h3><p>يرجى من جميع الطلبة شراء الكتاب يومي الثلاثاء والأربعاء، "
              "من كلية الآداب، الطابق الأول.</p>")

    def item(kind, cmid, **extra):
        return {"name": f"{kind} {cmid}", "type": kind, "course": course, "cmid": str(cmid),
                "url": f"https://x/mod/{kind}/view.php?id={cmid}", **extra}
    forum = item("forum", 66689, name="اعلان هام", _body=notice)
    assert not media.silent_container(forum) and media.silent_container(item("forum", 14908))
    assert not media.silent_container(item("page", 1))

    with tempfile.TemporaryDirectory() as tmp:
        keep = (reminders.SOURCES_FILE, monitor.deliver)
        reminders.SOURCES_FILE = Path(tmp) / "s.json"
        sent = []
        monitor.deliver = lambda header, body="", url=None, rows=None, shots=(), shot_name="", document=None: \
            sent.append((header, body, rows or [])) or True
        try:
            def pins(rows):
                return [b["callback_data"] for row in rows for b in row if b.get("callback_data", "").startswith("rem:s:")]
            pdf = Path(tmp) / "Chapter 1.pdf"
            pdf.write_bytes(b"%PDF")
            cases = [(item("assign", 1), {"kind": "activity"}, []),
                     (item("quiz", 2), {"kind": "activity"}, []),
                     (item("resource", 3), {"kind": "file", "path": str(pdf), "size": 4}, ["rem:s:i3"]),
                     (item("url", 4), {"kind": "link", "target": "https://youtu.be/x"}, ["rem:s:i4"]),
                     (item("page", 5, _body="<p>Install Outlook</p>"), {"kind": "text"}, ["rem:s:i5"]),
                     (item("book", 6), {"kind": "activity"}, ["rem:s:i6"]),
                     (forum, {"kind": "activity"}, ["rem:s:i66689"]),
                     (item("teamscheduler", 7), {"kind": "meeting", "meeting": {"url": "https://teams.microsoft.com/l/x",
                                                                               "details": []}}, ["rem:s:i7"])]
            for it, art, expected in cases:
                assert monitor.notify_new_item(it, None, art)
                assert pins(sent[-1][2]) == expected, (it["type"], sent[-1][2])
            header, body, rows = sent[-2]                       # the forum announcement
            assert header.startswith("📢 *New Announcement*") and "شراء الكتاب" in body
            assert reminders.source("i66689")["course"] == "English Communication Skills"
            assert sent[-1][2][0][0]["text"] == "🔗 Join Teams Lecture", "the meeting keeps its Join button first"

            # updates: a re-uploaded file and an edited page get 📌 too, under their own keys
            monitor.notify_file_update(item("resource", 3), {"path": str(pdf), "size": 4, "modified": 2,
                                                             "old_modified": 1})
            assert pins(sent[-1][2]) == ["rem:s:u3"]
            monitor.notify_content_update(item("page", 5, _body="<p>v2</p>"), {"shots": []})
            assert pins(sent[-1][2]) == ["rem:s:e5"]
        finally:
            reminders.SOURCES_FILE, monitor.deliver = keep

    # monitor.run: the forum with an announcement in its description is reported, the empty one is not
    with tempfile.TemporaryDirectory() as tmp:
        keep = {k: getattr(monitor, k) for k in ("DATA_FILE", "USERNAME", "PASSWORD", "BOT_TOKEN", "CHAT_ID",
                                                 "run_scan", "notify_new_item", "send_schedule_messages",
                                                 "mfa_allowed", "notify_new_course")}
        monitor.DATA_FILE = Path(tmp) / "saved_items.json"
        monitor.USERNAME = monitor.PASSWORD = monitor.BOT_TOKEN = monitor.CHAT_ID = "x"
        monitor.send_schedule_messages = lambda saved: None
        monitor.mfa_allowed = lambda saved: False
        monitor.notify_new_course = lambda c: True
        announced = []
        monitor.notify_new_item = lambda it, deadline, art, submitted=None: announced.append(it["name"]) or True
        content = {"8806-66689": forum, "8806-14908": item("forum", 14908, name="Announcements")}
        monitor.run_scan = lambda saved, notify_new, mfa_handler=None: {
            "courses": {"8806": {"name": course, "url": "https://x/c"}}, "content": dict(content), "deadlines": {},
            "artifacts": {}, "updates": {}, "edits": {}, "posts": {}, "forum_posts": {}, "submissions": {},
            "link_changes": [], "seconds": 1.0, "browser_used": False}
        keep_daily = (monitor.daily.catch_up, monitor.daily.send_recap)
        monitor.daily.catch_up = monitor.daily.send_recap = lambda *a, **k: False
        try:
            monitor.save_data({**monitor.load_saved_data(), "courses": {"8806": {"name": course, "url": "https://x/c"}},
                               "content": {"8806-1": item("page", 1)}})
            monitor.run()
            assert announced == ["اعلان هام"], announced
            assert {"8806-66689", "8806-14908"} <= set(json.loads(monitor.DATA_FILE.read_text(encoding="utf-8"))["content"])
        finally:
            for k, v in keep.items():
                setattr(monitor, k, v)
            monitor.daily.catch_up, monitor.daily.send_recap = keep_daily


def test_retake_replacement_in_plan():
    """A retaken course's old (recorded) grade leaves the weighted sum and its hours
    count once; retakes found on the portal are synced, hand-set ones left alone."""
    from uninotifier import gpa

    with tempfile.TemporaryDirectory() as tmp:
        keep = gpa.STATE_FILE
        gpa.STATE_FILE = Path(tmp) / "gpa.json"
        try:
            gpa.save({**gpa.load(), "courses": {"Calculus": gpa.new_course(), "Physics": gpa.new_course(),
                                                "Ethics": dict(gpa.new_course(), retake=True, old_grade=55.0)}})
            assert gpa.sync_retakes({"Calculus": 35.0, "Not here": 40.0}) == ["Calculus"]
            courses = gpa.load()["courses"]
            assert (courses["Calculus"]["retake"], courses["Calculus"]["old_grade"]) == (True, 35.0)
            assert courses["Ethics"]["retake"] and courses["Ethics"].get("retake_source") is None, "hand-set: untouched"
            # base: 70% over 30 h, which includes Calculus's 35 and Ethics's 55 (3 h each).
            # target CGPA 72: hours = 30 + 3 (Physics only) = 33; points = 72*33 = 2376
            # = 70*30 - 35*3 - 55*3 + 9x  ->  x = (2376 - 2100 + 105 + 165) / 9 = 60.67
            result = gpa.plan(courses, 72, "cgpa", 70.0, 30.0)
            assert abs(result["needed"] - 546 / 9) < 1e-9 and abs(result["cgpa"] - 72) < 1e-9, result
            # once Calculus's new mark is published it replaces the 35
            gpa.record_portal_grades({"Calculus": 80.0})
            result = gpa.plan(gpa.load()["courses"], 72, "cgpa", 70.0, 30.0)
            assert abs(result["needed"] - (2376 - 2100 + 105 + 165 - 240) / 6) < 1e-9, result
            assert gpa.sync_retakes({}) == ["Calculus"] and not gpa.load()["courses"]["Calculus"]["retake"], \
                "no longer listed on the portal: cleared"
        finally:
            gpa.STATE_FILE = keep


def _guest_sem(term, rows, average=None):
    return {"term": term, "average": average, "cumulative": None,
            "courses": [{"code": c, "name": n, "hours": h, "mark": m,
                         "passed": None if m is None else m >= 50, "repeated": False} for c, n, h, m in rows]}


def test_guest_lookup():
    """guest.py reads someone else's record and touches nothing of the owner's: the
    report lists passed and owed courses (a renumbered retake counts), the status,
    this term's retakes and the averages CGPA targets need."""
    from uninotifier import guest
    from uninotifier import portal

    semesters = [
        _guest_sem("First semester 2026/2027", [("0101221", "جبر خطي (1)", 3, None), ("0135220", "قواعد البيانات", 3, None),
                                           ("0135321", "البيانات الضخمة", 3, None)]),
        _guest_sem("First semester 2025/2026", [("0112212", "تراكيب البيانات", 3, 30), ("0135323", "تنقيب البيانات", 3, 72)], 51),
        _guest_sem("First semester 2024/2025", [("0101221", "جبر خطي (1)", 3, 35), ("0142251", "قواعد البيانات", 3, 20),
                                           ("0112220", "برمجة كينونية", 3, 75)], 43),
    ]
    result = portal.Result("transcript", True, None, [], {"passed_hours": 6.0, "cgpa": 52.0, "record_cgpa": 52.33,
                                                          "record_hours": 15.0}, [], [],
                           {"semesters": semesters, "status": "منتظم"})
    text = guest.build_report("202400001", result)
    assert "Cumulative GPA:       52.33%  (the portal shows 52%)" in text and "Passed credit hours:  6" in text
    assert "Academic status:      Regular / Good Standing" in text and "COURSES TAKEN — 2 passed" in text
    owed = text.split("FAILED AND NOT YET PASSED")[1].split("THIS TERM")[0]
    assert "0112212" in owed and "being retaken" not in owed.split("0112212")[1].split("\n")[0]
    assert "0142251" in owed and "→ being retaken this term" in owed.split("0142251")[1].split("\n")[0], \
        "renumbered Databases: same name, so it counts as retaken"
    assert "0101221  جبر خطي (1)  · 3h  🔁 retake — replaces 35" in text
    assert "0135220  قواعد البيانات  · 3h  🔁 retake — replaces 35" in text, "20 is recorded as 35"
    # targets: base = record before this term (3 attempts + 2 = 15 h); retakes replace, no extra hours
    plans = dict(guest.targets(semesters))
    base, hours = (30 * 0 + 35 * 3 + 35 * 3 + 75 * 3 + 35 * 3 + 72 * 3) / 15, 15     # failures floored at 35
    assert abs(plans[60]["needed"] - (60 * (hours + 3) - base * hours + 35 * 3 + 35 * 3) / 9) < 1e-9
    assert "TO REACH A CUMULATIVE GPA OF" in text and "60%  →" in text
    assert guest.build_report("x", portal.Result("transcript", True)).endswith("no courses could be read from it.")

    # the portal read: the guest's credentials, and nothing written to portal_debug/
    import inspect
    params = inspect.signature(portal.fetch_many).parameters
    assert "creds" in params and params["keep_debug"].default is True
    src = inspect.getsource(guest)
    for forbidden in (r"\bnotifier\b", r"gpa\.save", r"save_data", r"send_message", r"\breminders\b", r"set_base"):
        assert not re.search(forbidden, src), forbidden
    assert "keep_debug=False" in src and "creds=(username, password)" in src and "getpass" in src


def test_guest_summary_and_target_math():
    """The guest's major comes from the financial page's column (nothing else of that page is used), the
    hours left from `degree_hours` when the plan is known (never invented), and the target calculator
    follows the university's rules and agrees with the existing planner for the term in progress."""
    from uninotifier import guest, portal
    bot = _package("uninotifier.bot")

    # --- the major
    table = [["", "الفصل الدراسي", "الحالة", "التخصص", "عدد الساعات", "الرصيد"],
             ["", "الفصل الأول 2026/2027", "منتظم", "نظم المعلومات", "18", "د.أ1250"],
             ["", "الفصل الثاني 2025/2026", "منتظم", "علم الحاسوب", "12", "د.أ0"]]
    assert portal.parse_major([table]) == "نظم المعلومات", "the newest semester's"
    assert portal.parse_major([[["الفصل الدراسي", "الرصيد"], ["x", "1"]]]) is None
    assert portal.parse_major([[["التخصص"], [""], ["12"], ["علم الحاسوب"]]]) == "علم الحاسوب", "empty and numeric cells skipped"
    assert portal.parse_major([]) is None

    # --- degree_hours is a setting like the others
    from uninotifier import schedule
    base = {"days": {}, "degree_hours": 132}
    assert schedule.validate(base) == [] and schedule.validate({"days": {}, "degree_hours": 132.5}) == []
    assert schedule.validate({"days": {}, "degree_hours": None}) == [], "null = not set (the example's value)"
    for bad in ("132", 10, 500, True):
        assert schedule.validate({"days": {}, "degree_hours": bad}), bad
    assert "degree_hours" in schedule.KNOWN_KEYS

    # --- the trajectory: the same answer as the planner for this term, and the formula for more terms
    semesters = [
        _guest_sem("First semester 2026/2027", [("0101221", "جبر خطي (1)", 3, None), ("0135220", "قواعد البيانات", 3, None),
                                           ("0135321", "البيانات الضخمة", 3, None)]),
        _guest_sem("First semester 2025/2026", [("0112212", "تراكيب البيانات", 3, 30), ("0135323", "تنقيب البيانات", 3, 72)], 51),
        _guest_sem("First semester 2024/2025", [("0101221", "جبر خطي (1)", 3, 35), ("0142251", "قواعد البيانات", 3, 20),
                                           ("0112220", "برمجة كينونية", 3, 75)], 43),
    ]
    plans = dict(guest.targets(semesters))
    for target in guest.TARGETS:
        plan = guest.trajectory(semesters, target)
        n, hours, needed, verdict = plan["rows"][0]
        assert n == 1 and hours == 9.0 and abs(needed - plans[target]["needed"]) < 1e-9, (target, needed, plans[target])
    plan = guest.trajectory(semesters, 70)
    assert plan["in_progress"] and plan["term_hours"] == 9.0 and plan["base"][1] == 9.0, plan["base"]   # the two retaken courses left the record: 3 courses, 9 h
    two = plan["rows"][1]
    base_cgpa, base_hours = plan["base"]
    assert abs(two[2] - (70 * (base_hours + 18) - base_cgpa * base_hours) / 18) < 1e-9 and two[1] == 18.0
    assert plan["rows"][0][2] > plan["rows"][1][2] > plan["rows"][2][2], "more terms: a lower average is enough"

    # no term in progress: a full load (18 h) is assumed, and the verdicts
    done = [_guest_sem("First semester 2025/2026", [("1", "A", 3, 60)], 60)]
    nothing = guest.trajectory(done, 70)
    assert not nothing["in_progress"] and nothing["term_hours"] == 18 and abs(nothing["rows"][0][2] - (70 * 21 - 180) / 18) < 1e-9
    assert [guest.trajectory(done, t)["rows"][0][3] for t in (5, 40, 70, 100)] == ["secured", "pass", "ok", "out"]
    assert guest.trajectory([], 80)["rows"][0][2] == 80, "no record yet: the target itself"

    # --- the summary carries the hours left only when the plan's size is known
    results = {"transcript": portal.Result("transcript", True, b"PNG", [], {"passed_hours": 33.0, "cgpa": 60.0},
                                           [], [], {"semesters": semesters, "status": "منتظم"}),
               "balance": portal.Result("balance", True, b"PNG", [table], {}, [], [])}
    captions, texts = [], []
    keep = (bot.api, bot.notifier.upload, schedule.load_schedule)
    bot.api = lambda method, **d: texts.append(d.get("text", "")) or {"message_id": 1}
    bot.notifier.upload = lambda token, method, data, field, name, payload: captions.append(data["caption"]) or {"message_id": 2}
    bot._guest.clear()
    try:
        schedule.load_schedule = lambda: {"days": {}}
        bot.send_guest_results("2024", results)
        assert "Remaining Hours" not in captions[-1] and "📘 *Major:*" in captions[-1]
        schedule.load_schedule = lambda: {"days": {}, "degree_hours": 132}
        bot.send_guest_results("2024", results)
        assert "🏁 *Remaining Hours:* 99 hrs" in captions[-1], captions[-1]
        assert captions[-1].index("📘 *Major:*") < captions[-1].index("🎓 *Cumulative GPA:*") < captions[-1].index("🏁 *Remaining Hours:*")
        assert bot._guest["record"]["figures"]["passed_hours"] == 33.0
        # the financial page is only for the major: a page that failed to load simply leaves it out
        results["balance"] = portal.Result("balance", False)
        bot.send_guest_results("2024", results)
        assert "📘 *Major:*" not in captions[-1]
    finally:
        bot.api, bot.notifier.upload, schedule.load_schedule = keep
        bot._guest.clear()


def test_attendance_on_online_lecture_alert():
    """✍️ Check In rides on an online lecture's start alert — for whichever course it is,
    one button per attendance activity — and never on the hour-before heads-up."""
    import json as _json

    course = "Arabic Communication Skills"
    portal_name = "20261-الاتصال والتواصل باللغة العربية-ش31-3-ك"
    with tempfile.TemporaryDirectory() as tmp:
        keep = (lectures.STATE_FILE, lectures.CLAIM_DIR, lectures.notifier.send_message,
                lectures.notifier.send_png_file, lecturelinks.LINKS_FILE)
        lectures.STATE_FILE = Path(tmp) / "saved_items.json"
        lectures.CLAIM_DIR = Path(tmp) / "claims"
        lecturelinks.LINKS_FILE = Path(tmp) / "links.json"
        sent = []
        lectures.notifier.send_message = lambda token, chat, text, url=None, rows=None, **kw: sent.append(rows) or True
        lectures.notifier.send_png_file = lambda token, chat, png, name, caption="", url=None, rows=None: \
            sent.append(rows) or True
        try:
            def state(*cmids):
                lectures.STATE_FILE.write_text(_json.dumps({"content": {
                    f"8805-{c}": {"name": f"الحضور {c}", "type": "attendance", "course": portal_name, "cmid": str(c),
                                  "url": f"https://x/mod/attendance/view.php?id={c}"} for c in cmids}}), encoding="utf-8")
            state(75482)
            assert lectures.attendance_rows(course) == [[{"text": "✍️ Check In / تأكيد الحضور",
                                                          "url": "https://x/mod/attendance/view.php?id=75482"}]]
            assert lectures.attendance_rows("Big Data") == [], "only the course's own activities"
            # no meeting link yet: the alert still carries Check In
            assert lectures.send_due("t", "c", datetime(2026, 10, 14, 19, 0)) == 1
            assert sent[-1] == lectures.attendance_rows(course)
            # with a link: Join Meeting first, then Check In; the heads-up an hour before has neither
            lecturelinks.set_link(course, "https://teams.microsoft.com/l/meetup-join/19%3ax", "portal", "post")
            state(75482, 70299)
            assert lectures.send_due("t", "c", datetime(2026, 10, 21, 18, 0)) == 1 and sent[-1] is None
            assert lectures.send_due("t", "c", datetime(2026, 10, 21, 19, 0)) == 1
            texts = [b["text"] for row in sent[-1] for b in row]
            assert texts[0] == "🎥 Join Meeting" and [t for t in texts if t.startswith("✍️")] == [
                "✍️ Check In / تأكيد الحضور · الحضور 75482", "✍️ Check In / تأكيد الحضور · الحضور 70299"], texts
            lectures.STATE_FILE.unlink()
            assert lectures.attendance_rows(course) == [], "no saved state: no button, no crash"
        finally:
            (lectures.STATE_FILE, lectures.CLAIM_DIR, lectures.notifier.send_message,
             lectures.notifier.send_png_file, lecturelinks.LINKS_FILE) = keep


def test_guest_lookup_in_telegram():
    """/guest: student ID, then the password (deleted from the chat at once), their
    verification code relayed, their record sent one message after another — and none
    of the owner's data touched; 🚪 Sign Out removes it all from the chat."""
    bot = _package("uninotifier.bot")
    from uninotifier import portal

    semesters = [_guest_sem("First semester 2026/2027", [("0101221", "جبر خطي (1)", 3, None)]),
                 _guest_sem("First semester 2025/2026", [("0101221", "جبر خطي (1)", 3, 40),
                                                          ("0135323", "تنقيب البيانات", 3, 72)], 56)]
    results = {
        "transcript": portal.Result("transcript", True, b"PNG", [], {"passed_hours": 3.0, "cgpa": 56.0,
                                                                     "base_cgpa": 56.0, "base_hours": 6.0},
                                    [], [], {"semesters": semesters, "status": "منتظم", "retakes": {"0101221": 40.0}}),
        "grades": portal.Result("grades", True, b"PNG", [], {}, [], ["ℹ️ No results are published yet."]),
        "balance": portal.Result("balance", True, b"PNG-FINANCE", [[
            ["", "الفصل الدراسي", "الحالة", "التخصص", "عدد الساعات", "الرصيد"],
            ["", "الفصل الأول 2026/2027", "منتظم", "نظم المعلومات", "18", "د.أ1250"]]],
            {"balance": 1250.0}, [], ["💰 *Total balance:* 1,250 JD — owed to the university (debit)"]),
    }
    calls, files, fetched, threads = [], [], [], []

    def fake_api(method, **data):
        calls.append((method, data))
        return {"message_id": 100 + len(calls)} if method == "sendMessage" else True

    class Inline:                                    # run the worker in this thread
        def __init__(self, target, args=(), daemon=None):
            self.target, self.args = target, args

        def start(self):
            threads.append(1)
            self.target(*self.args)
    gpa_before = gpa.STATE_FILE.read_text(encoding="utf-8") if gpa.STATE_FILE.exists() else None
    keep = (bot.api, bot.send_view, bot.notifier.upload, bot.portal.fetch_many, bot.threading.Thread)
    bot.api = fake_api
    bot.send_view = lambda text, rows: calls.append(("view", {"text": text})) or {"message_id": 1}
    bot.notifier.upload = lambda token, method, data, field, name, payload: files.append((name, data["caption"])) \
        or {"message_id": 500 + len(files)}

    def fake_fetch(sections, ask_code, creds=None, keep_debug=True):
        fetched.append((sections, creds, keep_debug))
        return results
    bot.portal.fetch_many = fake_fetch
    bot.threading.Thread = Inline
    try:
        bot.handle_command("/guest", 10)
        assert bot._guest["step"] == "user" and "student ID" in calls[-1][1]["text"]
        bot.handle_command("not an id!", 11)
        assert "doesn't look like a student ID" in calls[-1][1]["text"] and bot._guest["step"] == "user"
        bot.handle_command("202400001", 12)
        assert bot._guest["step"] == "password"
        bot.handle_command("s3cret-Pass", 13)
        assert ("deleteMessage", {"chat_id": bot.CHAT_ID, "message_id": 13}) in calls, "password deleted at once"
        assert fetched == [(["transcript", "grades", "balance"], ("202400001", "s3cret-Pass"), False)], \
            "the record, this term's marks, and the financial page for the major only"
        texts = [d.get("text", "") for m, d in calls if m == "sendMessage"]
        assert not any("s3cret" in t for t in texts + [c for _, c in files]), "the password is never echoed"
        # the owner's transcript layout: the screenshot carries the executive summary (with the major) ...
        assert [n for n, _ in files] == ["Guest 202400001 - Transcript.png"], "one file: no grades or finance screenshot"
        caption = files[0][1]
        assert caption.startswith("👤 *Guest 202400001* — 📜 Transcript") and "📘 *Major:* \u2068نظم المعلومات\u2069" in caption
        assert "⏱️ *Passed Hours:* 3 hrs" in caption and "🎓 *Cumulative GPA:*" in caption
        assert "📌 *Academic Status:* Regular / Good Standing" in caption
        assert "🏁 *Remaining Hours:* 129 hrs" in caption, "the degree plan (132 h in schedule.json) minus the 3 passed"
        # ... then the semesters
        blocks = [t for t in texts if "In Progress" in t]
        assert len(blocks) == 1 and "🔁 *40*" in blocks[0] and "✅ *72*" in blocks[0], texts
        assert not any("1,250" in t or "1250" in t or "debit" in t for t in texts + [c for _, c in files]), \
            "nothing of the financial page is shown"
        assert not any("Schedule" in t for t in texts), "no class or exam schedule"
        assert bot._guest["step"] == "done" and not bot._portal_lock.locked()
        assert bot._guest["record"]["semesters"] and "balance" not in bot._guest["record"]
        final = [d for m, d in calls if m == "sendMessage"][-1]
        assert [r[0]["callback_data"] for r in final["reply_markup"]["inline_keyboard"]] == ["guest:target", "guest:wipe", "guest:out"]
        panel = bot._guest["panel"]

        # 🎯 Target GPA: the panel itself becomes the picker, then the answer; typed values; a command keeps the session
        def press(data, mid=None):
            calls.clear()
            bot.handle_callback({"id": f"g{len(calls)}{data}", "data": data, "message": {"message_id": mid or panel}})
            return [d for m, d in calls if m == "editMessageText"]

        edit = press("guest:target")[-1]
        assert edit["message_id"] == panel and "Target GPA" in edit["text"]
        picks = [b["callback_data"] for row in edit["reply_markup"]["inline_keyboard"] for b in row]
        assert picks[:7] == [f"guest:t:{t}" for t in (60, 65, 70, 75, 80, 85, 90)] and picks[-2:] == ["guest:tcustom", "guest:back"]
        edit = press("guest:t:75")[-1]
        assert "Target GPA — Guest 202400001" in edit["text"] and "To reach *75%* cumulative" in edit["text"]
        assert "This term (3h)" in edit["text"] and "+ the next one (6h)" in edit["text"] and "+ the next two (9h)" in edit["text"]
        assert "🎓 *Now:* " in edit["text"] and "university's own rules" in edit["text"] and _legacy_markdown_ok(edit["text"])
        buttons = [b["callback_data"] for row in edit["reply_markup"]["inline_keyboard"] for b in row]
        assert buttons == ["guest:target", "guest:wipe", "guest:out"], "another target, or sign out"
        assert press("guest:back")[-1]["text"].startswith("👤 *Guest session finished*")
        press("guest:tcustom")
        assert bot._guest["step"] == "target"
        calls.clear()
        bot.handle_command("not a number", 30)
        assert "between 40 and 100" in calls[-1][1]["text"] and bot._guest["step"] == "target"
        calls.clear()
        bot.handle_command("78", 31)
        edit = [d for m, d in calls if m == "editMessageText"][-1]
        assert "To reach *78%* cumulative" in edit["text"] and edit["message_id"] == panel and bot._guest["step"] == "done"
        assert 30 in bot._guest["messages"] and 31 in bot._guest["messages"], "typed answers go with the session"
        press("guest:tcustom")
        bot.handle_command("/start", 32)                                   # a command drops the prompt, not the session
        assert bot._guest["step"] == "done" and bot._guest.get("record")
        # nothing of the owner's: no GPA base, no retakes, no portal_debug dump (keep_debug=False above)
        now = gpa.STATE_FILE.read_text(encoding="utf-8") if gpa.STATE_FILE.exists() else None
        assert now == gpa_before
        # 🚪 Sign Out & Delete Messages: every message of the session goes, the student ID included
        recorded = list(bot._guest["messages"])
        assert 12 in recorded and 13 not in recorded
        bot.handle_callback({"id": "1", "data": "guest:wipe", "message": {"message_id": recorded[-1]}})
        deleted = {d["message_id"] for m, d in calls if m == "deleteMessage"}
        assert set(recorded) <= deleted and bot._guest == {}
        assert "Signed out of the guest account" in calls[-1][1]["text"]
        late = press("guest:t:75")[-1]                                     # a stale button after sign-out
        assert late["text"] == "🚪 This guest session has ended." and late["reply_markup"]["inline_keyboard"] == []

        # a command while it waits for the password drops the prompt; so does the timeout
        bot.handle_command("/guest", 20)
        bot.handle_command("202400001", 21)
        bot.handle_command("/start", 22)
        assert bot._guest == {}
        bot.handle_command("/guest", 23)
        bot._guest["at"] -= bot.GUEST_TIMEOUT + 1
        bot.handle_command("202400001", 24)
        assert bot._guest == {}
    finally:
        (bot.api, bot.send_view, bot.notifier.upload, bot.portal.fetch_many, bot.threading.Thread) = keep
        bot._guest.clear()


def test_no_personal_values_in_the_code():
    """Anyone can run it: paths, student ID and semester come from .env / schedule.json."""
    from uninotifier import archive
    bot = _package("uninotifier.bot")
    project = Path(__file__).resolve().parents[1]
    code = list((project / "src").rglob("*.py")) + list(project.glob("*.py"))
    # your own IDs, read from .env (so they are never written in this file either); a value still
    # copied from .env.example (what install.bat leaves before you fill it in) is nobody's
    from dotenv import dotenv_values
    from uninotifier import envcheck
    env = dotenv_values(project / ".env") if (project / ".env").is_file() else {}
    own_ids = {v.strip() for v in (env.get("STUDENT_ID"), env.get("PORTAL_USERNAME"),
                                   (env.get("ZUJ_USERNAME") or "").split("@")[0])
               if v and len(v.strip()) >= 6 and not envcheck._PLACEHOLDER.search(v.strip())}
    assert len(code) > 30, code
    for path in code:
        text = path.read_text(encoding="utf-8")
        for personal in ("Humam", "Year_04", *own_ids):
            assert personal not in text, (path.name, personal)
    for path in [*(project / "scripts").glob("*.ps1"), project / ".env.example", project / "schedule.example.json",
                 *project.glob("*.md"), *(project / "docs").glob("*.md")]:      # what a new student reads and runs
        text = path.read_text(encoding="utf-8").replace(REPO_PATH, "")      # the project's own public address is fine
        for personal in ("Humam", "Taibeh", "Year_04", *own_ids):
            assert personal not in text, (path.name, personal)
    for path in (project / "tests").glob("*.py"):           # the tests use a made-up ID, never the real one
        text = path.read_text(encoding="utf-8")
        for personal in own_ids:
            assert personal not in text, (path.name, "the real student ID")
    example = (project / ".env.example").read_text(encoding="utf-8")
    assert "\nARCHIVE_ROOT=\n" in example and "\nSTUDENT_ID=\n" in example, "optional settings are empty: off until filled in"
    assert archive.ARCHIVE_ROOT.is_absolute() and bot.semester() == schedule.load_schedule().get("semester", "")
    assert set(portal_sections()) == {"transcript", "grades", "balance"}


def test_layout_watchdog():
    """A page that suddenly reads short (the site changed) keeps its known items and is
    reported once; ✅ when it reads normally again; a shortfall that lasts 3 days is
    accepted as the new normal; an empty course list fails the run without touching state."""
    monitor = _package("uninotifier.monitor")
    from uninotifier import watchdog

    course = "20261-Numerical Analysis(1)-AllSec-3-"

    def items(n, start=1, text=False):
        return {f"8802-{i}": {"name": f"Item {i}", "type": "resource", "course": course, "cmid": str(i),
                              "url": f"https://x/mod/resource/view.php?id={i}", **({"hash": "h"} if text else {})}
                for i in range(start, start + n)}
    courses = {"8802": {"name": course, "url": "https://x/c"}}

    gone, findings = watchdog.check_courses(courses, items(2), items(46))
    assert gone == ["8802"] and findings == ["Numerical Analysis (1): 2 item(s) read, 46 known"], findings
    assert watchdog.check_courses(courses, items(30), items(46)) == ([], []), "a normal change"
    assert watchdog.check_courses(courses, items(1), items(3)) == ([], []), "too small to judge"
    assert watchdog.check_courses(courses, items(5), items(5, text=True))[1] == [
        "Numerical Analysis (1): the course page's text could not be read"]
    assert watchdog.check_forums({"a": None, "b": None}) == ["none of the 2 forums could be read"]
    assert watchdog.check_forums({"a": None, "b": []}) == []
    assert _legacy_markdown_ok(watchdog.alert_text(findings))

    with tempfile.TemporaryDirectory() as tmp:
        keep = {k: getattr(monitor, k) for k in ("DATA_FILE", "USERNAME", "PASSWORD", "BOT_TOKEN", "CHAT_ID",
                                                 "run_scan", "notify_new_item", "send_schedule_messages",
                                                 "mfa_allowed", "notify_new_course", "send_telegram")}
        keep_daily = (monitor.daily.catch_up, monitor.daily.send_recap)
        monitor.daily.catch_up = monitor.daily.send_recap = lambda *a, **k: False
        monitor.DATA_FILE = Path(tmp) / "saved_items.json"
        monitor.USERNAME = monitor.PASSWORD = monitor.BOT_TOKEN = monitor.CHAT_ID = "x"
        monitor.send_schedule_messages = lambda saved: None
        monitor.mfa_allowed = lambda saved: False
        monitor.notify_new_course = lambda c: True
        announced, said = [], []
        monitor.notify_new_item = lambda it, deadline, art, submitted=None: announced.append(it["name"]) or True
        monitor.send_telegram = lambda text, url=None: said.append(text) or True
        reading = {"courses": dict(courses), "content": items(2)}

        def fake_scan(saved, notify_new, mfa_handler=None):
            # the same steps run_scan takes, on a fixed reading
            content, failed = dict(reading["content"]), []
            gone, anomalies = watchdog.check_courses(reading["courses"], content, saved["content"])
            protect, accepted = watchdog.review(saved["layout"], gone, anomalies, reading["courses"], clock["now"])
            for cid in protect:
                failed.append(cid)
                content = {i: it for i, it in content.items() if watchdog.course_id(i) != cid}
            for iid, it in saved["content"].items():
                if watchdog.course_id(iid) in failed:
                    content.setdefault(iid, it)
            return {"courses": reading["courses"], "content": content, "deadlines": {}, "artifacts": {},
                    "updates": {}, "edits": {}, "posts": {}, "forum_posts": {}, "submissions": {},
                    "link_changes": [], "seconds": 1.0, "browser_used": False,
                    "anomalies": anomalies, "accepted": accepted}
        monitor.run_scan = fake_scan
        clock = {"now": time.time()}
        keep_time = monitor.time.time
        monitor.time.time = lambda: clock["now"]
        try:
            monitor.save_data({**monitor.load_saved_data(), "courses": dict(courses), "content": items(46)})

            def run():
                try:
                    monitor.run()
                except SystemExit as e:
                    return e.code
                return 0
            assert run() == 0
            saved = json.loads(monitor.DATA_FILE.read_text(encoding="utf-8"))
            assert len(saved["content"]) == 46 and announced == [], "known items kept, nothing announced"
            assert said[-1].startswith("🧩 *Moodle looks different*") and "2 item(s) read, 46 known" in said[-1]
            n = len(said)
            clock["now"] += 3600
            run()
            assert len(said) == n, "not repeated within 6 hours"
            clock["now"] += watchdog.REPEAT_SECONDS
            run()
            assert len(said) == n + 1, "repeated after 6 hours"

            reading["content"] = items(46)                     # the page reads normally again
            clock["now"] += 900
            run()
            assert said[-1] == watchdog.RECOVERED_TEXT and announced == []
            run()
            assert said[-1] == watchdog.RECOVERED_TEXT and said.count(watchdog.RECOVERED_TEXT) == 1

            reading["content"] = items(10)                     # the instructor really removed 36 items
            for _ in range(13):                                 # 3 days of checks, 6 h apart
                clock["now"] += 6 * 3600
                run()
            assert any(t.startswith("ℹ️ *Accepted as the new normal:* Numerical Analysis (1)") for t in said)
            assert len(json.loads(monitor.DATA_FILE.read_text(encoding="utf-8"))["content"]) == 10

            before = monitor.DATA_FILE.read_text(encoding="utf-8")
            reading["courses"] = {}

            def empty_scan(saved, notify_new, mfa_handler=None):
                if not reading["courses"] and saved["courses"]:
                    raise watchdog.LayoutChanged("the list of your enrolled courses came back empty")
            monitor.run_scan = empty_scan
            assert run() == 1
            after = json.loads(monitor.DATA_FILE.read_text(encoding="utf-8"))
            assert after["content"] == json.loads(before)["content"], "state untouched"
            assert after["failures"]["kind"] == "layout" and after["failures"]["count"] == 1, "alert after 2 runs"
        finally:
            for k, v in keep.items():
                setattr(monitor, k, v)
            monitor.daily.catch_up, monitor.daily.send_recap = keep_daily
            monitor.time.time = keep_time


def test_recorded_real_pages():
    """Your real Moodle, replayed: the pages recorded by tools.record_pages are read
    again by today's code and must give exactly what they gave when recorded — every
    course, item, name, text fingerprint, forum post, deadline, Teams link and
    attendance session. Skipped when nothing has been recorded on this PC."""
    from uninotifier.tools import record_pages

    if not record_pages.MOODLE_FILE.exists():
        print("  (skipped: no recorded pages — run python -m uninotifier.tools.record_pages)")
        return
    recording = json.loads(record_pages.MOODLE_FILE.read_text(encoding="utf-8"))
    replay = record_pages.Replay(recording["calls"])
    got = record_pages.read_everything(replay, record_pages.PLACEHOLDER)
    expected = recording["expected"]
    assert got["courses"] == expected["courses"]
    for cid, items in expected["content"].items():
        assert got["content"][cid] == items, f"course {expected['courses'][cid]} reads differently"
    for key in ("deadlines", "forums", "meetings", "sessions"):
        assert got[key] == expected[key], key
    assert sum(len(c) for c in got["content"].values()) > 0


def test_upload_tried_again_sends_the_whole_file():
    """A send whose TLS handshake stalled is tried again (nothing was sent), and the second try
    must carry the whole file: requests reads a file object while it builds the request, so the
    second try used to send it empty. The fake door reads the body exactly as requests does."""
    import requests

    class Reply:
        status_code, ok = 200, True

        def json(self):
            return {"ok": True, "result": {"message_id": 5, "document": {"file_id": "F", "file_unique_id": "U"}}}

    stall = requests.ReadTimeout("HTTPSConnectionPool(host='api.telegram.org', port=443): Read timed out. "
                                 f"(read timeout={telegram.CONNECT_TIMEOUT})")
    seen = []

    def door(url, files=None, **kw):
        body = {field: value[1] if isinstance(value[1], bytes) else value[1].read()
                for field, value in (files or {}).items()}
        seen.append({field: len(data) for field, data in body.items()})
        if len(seen) == 1:
            raise stall
        return Reply()

    keep = telegram.post
    telegram.post = door
    try:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "Lecture 1.pdf"
            path.write_bytes(b"%PDF" + b"x" * 5000)
            fileids.forget(fileids.fingerprint(path))
            assert notifier.send_document("T", "42", path, "📎 Lecture 1")
            assert seen == [{"document": 5004}, {"document": 5004}], seen
    finally:
        telegram.post = keep


def test_refused_markdown_goes_out_as_plain_text():
    """Telegram refusing a message's Markdown is final — the same message would be refused every
    run. It goes out once more as plain text (links and addresses whole), the owner is told once,
    and a notification that still cannot be sent for several checks in a row is reported."""
    import importlib

    assert notifier.plain_text("*Lecture 3* — [Open](https://x.edu/mod_resource/a_b.pdf) snake\\_case "
                               "https://y.edu/z_1 `code`") == \
        "Lecture 3 — Open (https://x.edu/mod_resource/a_b.pdf) snake_case https://y.edu/z_1 code"
    assert notifier.without_markdown({"chat_id": 1, "text": "x"}) is None, "nothing to drop: not tried again"

    class Reply:
        def __init__(self, status, body):
            self.status_code, self.ok, self._body = status, status < 400, body
            self.text = json.dumps(body)

        def json(self):
            return self._body

    refused = {"ok": False, "error_code": 400, "description": "Bad Request: can't parse entities: "
                                                              "Can't find end of the entity starting at byte offset 7"}
    sent = []

    def door(url, json=None, data=None, files=None, **kw):
        body = dict(json or data or {})
        sent.append((url.rsplit("/", 1)[1], body))
        if "parse_mode" in body or "parse_mode" in str(body.get("media", "")):
            return Reply(400, refused)
        return Reply(200, {"ok": True, "result": {"message_id": 9}})

    keep = telegram.post
    telegram.post = door
    _fresh_failsafe()
    try:
        assert notifier.send_message("T", "42", "*Big_Data* lecture")
        assert [m for m, _ in sent] == ["sendMessage", "sendMessage"]
        assert "parse_mode" not in sent[1][1] and sent[1][1]["text"] == "BigData lecture", sent[1]
        # an album's captions too
        sent.clear()
        media = json.dumps([{"type": "document", "media": "F1", "caption": "*x_y*", "parse_mode": "Markdown"}])
        assert notifier._request("T", "sendMediaGroup", {"chat_id": "42", "media": media})
        assert json.loads(sent[1][1]["media"]) == [{"type": "document", "media": "F1", "caption": "xy"}]
        # ... and the bot's own screens (call_api)
        sent.clear()
        assert notifier.call_api("T", "editMessageText", {"chat_id": 42, "message_id": 3, "text": "*a_b*",
                                                          "parse_mode": "Markdown"}) == {"message_id": 9}
        assert sent[1][1]["text"] == "ab" and "parse_mode" not in sent[1][1]
        # a refusal that is not about Markdown is not retried
        sent.clear()
        telegram.post = lambda url, **kw: sent.append(1) or Reply(400, {"ok": False, "description": "Bad Request: chat not found"})
        assert not notifier.send_message("T", "42", "*x*") and len(sent) == 1
        failsafe.wait_idle()
        told = [t for t in OUTBOX if "without its formatting" in t]
        assert len(told) == 1, ("told once, not once per message", OUTBOX)
    finally:
        telegram.post = keep

    # the monitor: a notification that keeps failing is reported after UNDELIVERED_ALERT_RUNS checks
    alerts = importlib.import_module("uninotifier.monitor.alerts")
    mstate = importlib.import_module("uninotifier.monitor.state")
    mconfig = importlib.import_module("uninotifier.monitor.config")
    _fresh_failsafe()
    with tempfile.TemporaryDirectory() as tmp:
        keep = (mstate.DATA_FILE, mstate.LOCK_FILE)
        mstate.DATA_FILE, mstate.LOCK_FILE = Path(tmp) / "saved_items.json", Path(tmp) / "monitor.lock"
        try:
            saved = {"undelivered": {}}
            stuck = {"item:1-5": "Lecture 3 — Big Data"}
            for run in range(1, mconfig.UNDELIVERED_ALERT_RUNS):
                alerts.track_undelivered(saved, stuck, now=run)
            assert not OUTBOX and saved["undelivered"]["item:1-5"]["runs"] == mconfig.UNDELIVERED_ALERT_RUNS - 1
            alerts.track_undelivered(saved, stuck, now=10)
            assert len(OUTBOX) == 1 and "Lecture 3" in OUTBOX[0] and "could not be sent" in OUTBOX[0], OUTBOX
            alerts.track_undelivered(saved, stuck, now=11)
            assert len(OUTBOX) == 1, "the same stuck set is not reported every run"
            alerts.track_undelivered(saved, {}, now=12)                    # it went out
            assert saved["undelivered"] == {} and len(OUTBOX) == 2 and "have gone out" in OUTBOX[1], OUTBOX
            alerts.track_undelivered(saved, {"item:1-6": "x"}, now=13)     # a blip: one run, never reported
            alerts.track_undelivered(saved, {}, now=14)
            assert len(OUTBOX) == 2 and saved["undelivered"] == {}
        finally:
            mstate.DATA_FILE, mstate.LOCK_FILE = keep


def test_your_own_data_survives_a_damaged_file():
    """reminders.json and gpa.json hold what you typed in — nothing can rebuild them. A damaged
    file used to read as empty and the next save wiped everything; now it is kept aside, the
    newest backup is put back, and you are told. Every save leaves a backup."""
    from uninotifier import userdata

    keep = (reminders.STATE_FILE, gpa.STATE_FILE)
    with tempfile.TemporaryDirectory() as tmp:
        reminders.STATE_FILE, gpa.STATE_FILE = Path(tmp) / "reminders.json", Path(tmp) / "gpa.json"
        backups = Path(tmp) / "backups"
        _fresh_failsafe()
        try:
            assert reminders.load()["items"] == [] and not OUTBOX, "a file that does not exist yet is simply empty"
            first = reminders.add("Bring the lab report", time.time() + 3600)
            reminders.add("Office hours", time.time() + 7200)
            assert (backups / "reminders.latest.json").is_file()
            assert len(list(backups.glob("reminders.????-??-??.json"))) == 1, "one backup per day"

            # a damaged file (a power cut, a bad edit): restored from the backup, never read as empty
            reminders.STATE_FILE.write_text('{"next_id": 3, "items": [{"id": 1, "no', encoding="utf-8")
            restored = reminders.load()
            assert [r["note"] for r in restored["items"]] == ["Bring the lab report", "Office hours"], restored
            assert reminders.get(first["id"])["note"] == "Bring the lab report", "the file itself is repaired"
            assert len(list(backups.glob("reminders.corrupt-*.json"))) == 1, "the damaged file is kept to look at"
            failsafe.wait_idle()
            assert len(OUTBOX) == 1 and "restored from a backup" in OUTBOX[0] and "reminders" in OUTBOX[0], OUTBOX
            # adding after the repair keeps what was there
            reminders.add("Third", time.time() + 9000)
            assert len(reminders.load()["items"]) == 3

            # not a JSON object at all, and no backup left: empty — and said so
            for old in backups.glob("reminders.*"):
                old.unlink()
            _fresh_failsafe()
            reminders.STATE_FILE.write_text("[1, 2]", encoding="utf-8")
            assert reminders.load()["items"] == []
            failsafe.wait_idle()
            assert len(OUTBOX) == 1 and "no backup" in OUTBOX[0], OUTBOX

            # a file that exists but cannot be opened is not "empty" either: it raises, nothing is saved over it
            assert not reminders.STATE_FILE.exists(), "the damaged file was moved aside"
            reminders.STATE_FILE.mkdir()
            try:
                reminders.load()
                raise AssertionError("an unopenable file must not read as empty")
            except OSError:
                pass
            reminders.STATE_FILE.rmdir()

            # the Grade Needed plan the same way
            _fresh_failsafe()
            state = gpa.load()
            state["courses"] = {"Database Systems": {"hours": 3, "items": []}}
            gpa.save(state)
            gpa.STATE_FILE.write_text("\x00\x00", encoding="utf-8")
            assert "Database Systems" in gpa.load()["courses"]
            failsafe.wait_idle()
            assert len(OUTBOX) == 1 and "Grade Needed plan" in OUTBOX[0], OUTBOX
            assert userdata.backup_dir(gpa.STATE_FILE) == backups
        finally:
            reminders.STATE_FILE, gpa.STATE_FILE = keep


def test_mark_as_done_never_writes_back_an_older_state():
    """The bot's "Mark as done" changes one field of saved_items.json while a check may be saving
    it. Without coordination the bot could write back the copy it read before the monitor's save —
    and the next check would announce again what it had just announced. Both take the file's
    write lock: the monitor's save waits for the bot's few milliseconds, and lands after it."""
    import threading

    from uninotifier import fileio

    bot = _package("uninotifier.bot")
    mstate = importlib.import_module("uninotifier.monitor.state")
    with tempfile.TemporaryDirectory() as tmp:
        data = Path(tmp) / "saved_items.json"
        keep = (bot.DATA_FILE, bot.load_state, mstate.DATA_FILE, mstate.LOCK_FILE)
        bot.DATA_FILE = mstate.DATA_FILE = data
        mstate.LOCK_FILE = Path(tmp) / "monitor.lock"
        real_load = keep[1]
        fileio.write_json_atomic(data, {"courses": {"1": {}}, "content": {"1-5": {"cmid": "5", "done": False}}})
        newer = {"courses": {"1": {}}, "content": {"1-5": {"cmid": "5", "done": False},
                                                   "1-6": {"cmid": "6", "done": False, "name": "just announced"}}}
        saver = threading.Thread(target=mstate.save_data, args=(newer,))
        blocked = []

        def load_then_let_the_monitor_try():
            state = real_load()
            saver.start()                     # the monitor saves while the bot holds what it read
            saver.join(0.3)
            blocked.append(saver.is_alive())
            return state
        bot.load_state = load_then_let_the_monitor_try
        try:
            bot.set_item_done("5", True)
            saver.join(5)
            assert blocked == [True], "the monitor's save waited for the bot's read-modify-write"
            final = json.loads(data.read_text(encoding="utf-8"))
            assert "1-6" in final["content"], "the monitor's newer save is not overwritten by an older copy"
            assert not data.with_name("saved_items.json.lock").exists(), "the lock is released"

            # a lock left by a process that died holding it is taken over; a live one makes a writer wait, then give up
            lock = data.with_name("saved_items.json.lock")
            lock.touch()
            old = time.time() - fileio.LOCK_STALE_SECONDS - 5
            os.utime(lock, (old, old))
            with fileio.write_lock(data):
                pass
            lock.touch()
            try:
                with fileio.write_lock(data, wait=0.1):
                    raise AssertionError("a held lock must not be entered")
            except TimeoutError:
                pass
            lock.unlink()
        finally:
            bot.DATA_FILE, bot.load_state, mstate.DATA_FILE, mstate.LOCK_FILE = keep


def test_catch_alls_tell_the_owner():
    """A catch-all keeps the bot and the check running — but a bug it swallows must not leave a
    feature switched off with only a log line to show for it: the owner is told (failsafe.bug,
    once per bug). /check never stays on "Checking the portal…". And no new silent catch-all can
    slip into the modules that run unattended."""
    import ast
    import subprocess

    bot = _package("uninotifier.bot")
    checks = importlib.import_module("uninotifier.bot.checks")

    def boom(*a, **k):
        raise KeyError("text")

    keep = (schedule.due_messages, notifier.send_message, subprocess.run, bot.edit_view)
    _fresh_failsafe()
    try:
        # a lecture alert whose sending crashes
        schedule.due_messages = lambda sched, now: [("2099-01-01-test-before", "before", "⏰ text", "Big Data")]
        notifier.send_message = boom
        assert lectures.send_due("T", "42") == 0
        lectures.release("2099-01-01-test-before")
        # a reminder whose sending crashes
        reminders.add("Lab report", time.time() - 60)
        assert reminders.send_due("T", "42") == 0
        failsafe.wait_idle()
        assert any("lecture alert could not be sent" in t for t in OUTBOX), OUTBOX
        assert any("reminder could not be sent" in t for t in OUTBOX), OUTBOX

        # /check crashing: reported, and the screen says so instead of "Checking…" forever
        screens = []
        bot.edit_view = lambda message_id, text, rows: screens.append(text)
        subprocess.run = boom
        checks.run_check(7)
        failsafe.wait_idle()
        assert screens[0].startswith("🔄") and "Check Failed" in screens[-1], screens
        assert any("/check failed" in t for t in OUTBOX), OUTBOX
        assert not checks._check_lock.locked()
    finally:
        schedule.due_messages, notifier.send_message, subprocess.run, bot.edit_view = keep
        reminders.STATE_FILE.unlink(missing_ok=True)

    # the guard: in the unattended code, a handler that logs an exception also tells the owner —
    # or is listed here with the reason it need not
    told_elsewhere = {
        ("bot/checks.py", "Could not show that the check failed"),            # the failure itself is reported
        ("bot/finals_flow.py", "Finals check failed"),                       # shown on the Finals screen / notice
        ("bot/portal_flow.py", "Portal request failed"),                     # shown as the answer, Try Again
        ("monitor/run.py", "%s failure while talking to the portal%s"),      # alerts.record_failure
        ("monitor/scan.py", "Failed scanning course %s — skipping it this run"),   # read in the browser; watchdog
        ("monitor/scan.py", "Could not keep the screenshot of the link source"),   # the link itself is kept
        ("monitor/notify.py", "Could not forward attachment %s"),            # one file; the post still goes out
        ("monitor/notify.py", "Could not download forum attachment %s"),     # one file; the post still goes out
        ("monitor/state.py", "Could not set the unreadable state file aside"),     # the recovery reports right after
        ("monitor/state.py", "Could not put the backup back"),                     # ... likewise
        ("userdata.py", "Could not set the unreadable %s aside"),            # the recovery reports right after
        ("userdata.py", "Could not put the backup of %s back"),              # ... likewise
    }
    root = Path(lectures.__file__).parent
    guarded = ["monitor/run.py", "monitor/scan.py", "monitor/notify.py", "monitor/alerts.py", "monitor/state.py",
               "lectures.py", "reminders.py", "userdata.py", "gitsync.py", "bot/checks.py", "bot/finals_flow.py",
               "bot/portal_flow.py", "bot/guest_flow.py", "probe.py", "cambridge.py"]
    silent = []
    for name in guarded:
        for handler in ast.walk(ast.parse((root / name).read_text(encoding="utf-8"))):
            if not isinstance(handler, ast.ExceptHandler):
                continue
            body = "\n".join(ast.unparse(statement) for statement in handler.body)
            logged = [n.args[0].value for statement in handler.body for n in ast.walk(statement)
                      if isinstance(n, ast.Call) and getattr(n.func, "attr", "") == "exception"
                      and n.args and isinstance(n.args[0], ast.Constant)]
            if logged and "failsafe." not in body and (name, logged[0]) not in told_elsewhere:
                silent.append(f"{name}: {logged[0]}")
    assert not silent, ("catch-alls that only log — report with failsafe.bug, or list the reason", silent)


def test_state_files_are_written_through_fileio():
    """Every state file is replaced in one step after its data is flushed to the disk (fileio), so a
    power cut cannot leave one empty. A hand-made "write a .tmp, then rename" skips the flush — none
    may come back."""
    root = Path(lectures.__file__).parent
    pattern = re.compile(r"""with_suffix\(\s*["'][^"']*tmp["']\s*\)|\btmp\.replace\(""")
    found = [f"{path.relative_to(root).as_posix()}:{n}"
             for path in sorted(root.rglob("*.py")) if path.name != "fileio.py"
             for n, line in enumerate(path.read_text(encoding="utf-8").splitlines(), 1) if pattern.search(line)]
    assert not found, ("write state files with fileio.write_json_atomic / write_text_atomic", found)


def test_forwarded_messages_are_never_acted_on():
    """A forward carries someone else's words: a forwarded "123456" would pass for a verification
    code, a forwarded "/guest" for a command. No forward is acted on — the owner's own is ignored
    without an alert (the owner is no intruder), a stranger's is dropped and reported as usual."""
    bot = _package("uninotifier.bot")
    keep_owner = _as_owner(bot)
    calls, handled = [], []
    orig = (bot.handle_command, bot.handle_callback, bot.api)
    bot.api = lambda method, **data: calls.append(method)
    bot.handle_command = lambda text, message_id=None: handled.append(text)
    bot.handle_callback = lambda cb: handled.append(cb["data"])
    _fresh_failsafe()
    access._logged.clear()
    access._handled.clear()
    access._alerts.clear()

    def forwarded(update, **marker):
        update["message"].update(marker)
        return update
    try:
        owner_forwards = [
            forwarded(_tg(text="123456"), forward_origin={"type": "user", "date": 1, "sender_user": {"id": 99}}),
            forwarded(_tg(text="/guest"), forward_origin={"type": "hidden_user", "date": 1, "sender_user_name": "X"}),
            forwarded(_tg(text="/start"), forward_from={"id": 99}, forward_date=1),          # older Bot API fields
            forwarded(_tg(text="/start"), forward_date=1),
            forwarded(_tg(text="/start"), is_automatic_forward=True),
            # sent from the owner's account through an inline bot: that bot wrote the text
            forwarded(_tg(text="123456"), via_bot={"id": 777000, "is_bot": True, "username": "somebot"}),
        ]
        for update in owner_forwards:
            assert access.is_forwarded(update) and not access.is_owner(update), update
            bot.handle_update(update)
        assert handled == [] and calls == [], ("a forward is never acted on", handled, calls)
        failsafe.wait_idle()
        assert OUTBOX == [], "the owner's own forward raises no alert"

        stranger = forwarded(_tg(user_id="777", chat_id="777", text="/start"), forward_origin={"type": "user", "date": 1})
        bot.handle_update(stranger)
        failsafe.wait_idle()
        assert handled == [] and len(OUTBOX) == 1 and "forwarded message" in OUTBOX[0], OUTBOX

        # the owner's ordinary messages and buttons are untouched
        assert access.is_owner(_tg(text="/start")) and access.is_owner(_tg(data="menu"))
        bot.handle_update(_tg(text="/start"))
        assert handled == ["/start"]
    finally:
        bot.handle_command, bot.handle_callback, bot.api = orig
        bot.CHAT_ID, access.ADMIN_USER_ID = keep_owner


def test_secrets_never_reach_logs_or_alerts():
    """Every .env value (also URL-encoded, as inside a request), the bot token, Moodle's session key
    and cookie and a healthchecks.io address are masked in log lines, tracebacks and alerts. .env is
    read again when it changes; short settings (BOT_POLL_TIMEOUT=20) are not secrets."""
    from uninotifier import maintenance

    password = "Hunter2!x y_z"
    token = "123456789:AAH" + "x" * 32
    keep = (notifier.SECRETS_FILE, notifier.SECRETS_RECHECK_SECONDS)
    with tempfile.TemporaryDirectory() as tmp:
        env = Path(tmp) / ".env"
        env.write_text(f'ZUJ_PASSWORD="{password}"\nTELEGRAM_BOT_TOKEN={token}\nPORTAL_USERNAME=student4711\n'
                       "BOT_POLL_TIMEOUT=20\nHEALTHCHECK_URL=https://hc-ping.com/0f1e2d3c-aaaa-bbbb\n", encoding="utf-8")
        notifier.SECRETS_FILE, notifier.SECRETS_RECHECK_SECONDS = env, 0
        notifier.reset_secrets()
        _fresh_failsafe()
        try:
            leaky = (f"POST /login password={password} user=student4711 encoded=Hunter2%21x%20y_z plus=Hunter2%21x+y_z "
                     f"https://api.telegram.org/bot{token}/sendMessage "
                     "HTTPSConnectionPool(host='elearning.zuj.edu.jo'): Max retries exceeded with url: "
                     "/lib/ajax/service.php?sesskey=AbC123xyz&info=core_course "
                     '{"sesskey":"Zz9Yy8"} Cookie: MoodleSession=k3y5ecret; MoodleSessionzuj=0ther '
                     "{'MoodleSession': 'pyrepr'} ping https://hc-ping.com/5a5a5a5a-1111-2222/fail "
                     "timeout 20 s")
            clean = notifier.redact(leaky)
            for secret in (password, "Hunter2%21x%20y_z", "Hunter2%21x+y_z", "student4711", token, "AbC123xyz",
                           "Zz9Yy8", "k3y5ecret", "0ther", "pyrepr", "5a5a5a5a", "0f1e2d3c"):
                assert secret not in clean, (secret, clean)
            assert "<ZUJ_PASSWORD>" in clean and "<PORTAL_USERNAME>" in clean and "timeout 20 s" in clean, clean
            assert "sesskey=<redacted>" in clean and "hc-ping.com/<redacted>" in clean

            # a traceback in the log file
            try:
                raise ConnectionError(f"login as student4711 with {password} refused")
            except ConnectionError:
                record = logging.LogRecord("uninotifier", logging.ERROR, __file__, 1, "Sign-in failed for %s",
                                           ("student4711",), sys.exc_info())
            line = notifier.RedactingFormatter("%(message)s").format(record)
            assert "Traceback" in line and password not in line and "student4711" not in line, line

            # an alert: headline, detail and the exception line
            failsafe.report("test:secret", "Sign-in as student4711 failed", detail=f"used {password}",
                            exc=RuntimeError(f"bot{token} sesskey=AbC123xyz"))
            assert len(OUTBOX) == 1 and all(s not in OUTBOX[0] for s in ("student4711", password, token, "AbC123xyz"))
            # the portal's "can't sign in" reason shown on Telegram
            error = maintenance.unavailable(None, "portal", TimeoutError("page.goto timed out: ...?u=student4711"))
            assert "student4711" not in error.detail

            # .env changed: the new value is masked from then on
            env.write_text(env.read_text(encoding="utf-8") + "PORTAL_PASSWORD=Brand-New-Secret\n", encoding="utf-8")
            os.utime(env, (time.time() + 5, time.time() + 5))
            assert "Brand-New-Secret" not in notifier.redact("pw Brand-New-Secret")
        finally:
            notifier.SECRETS_FILE, notifier.SECRETS_RECHECK_SECONDS = keep
            notifier.reset_secrets()


def test_guest_password_never_stays_in_the_chat():
    """The guest's password message is deleted — tried again when Telegram does not confirm it, and
    the owner asked to delete it by hand when it still stands. A reply after the password prompt
    expired is taken for the password: deleted, not read as anything else (a reminder note)."""
    bot = _package("uninotifier.bot")
    guest_flow = importlib.import_module("uninotifier.bot.guest_flow")
    reminder_flow = importlib.import_module("uninotifier.bot.reminder_flow")
    calls, views, outcomes = [], [], []

    def api(method, **data):
        calls.append((method, data))
        return outcomes.pop(0) if method == "deleteMessage" and outcomes else {"message_id": 99}

    keep = (bot.api, bot.send_view, guest_flow.DELETE_RETRY_SECONDS, reminder_flow._reminder_wait)
    bot.api, bot.send_view = api, lambda text, rows: views.append(text)
    guest_flow.DELETE_RETRY_SECONDS = (0, 0)
    try:
        # Telegram confirms only the third try: three calls, no warning
        outcomes[:] = [None, None, True]
        assert guest_flow.delete_password_message(41)
        assert [m for m, _ in calls] == ["deleteMessage"] * 3
        # never confirmed: the owner is asked to delete it, in a reply to that very message
        calls.clear()
        outcomes[:] = [None, None, None]
        assert not guest_flow.delete_password_message(42)
        warning = calls[-1]
        assert warning[0] == "sendMessage" and "delete it yourself" in warning[1]["text"]
        assert warning[1]["reply_parameters"]["message_id"] == 42

        # the password prompt expired, and a reminder note is being asked for: the reply is deleted, not used
        calls.clear()
        outcomes[:] = []
        noted = []
        keep_text = reminder_flow.handle_reminder_text
        reminder_flow.handle_reminder_text = noted.append
        reminder_flow._reminder_wait = "note"
        guest_flow._guest.clear()
        guest_flow._guest.update(step="password", user="20230001", at=time.time() - guest_flow.GUEST_TIMEOUT - 1, messages=[])
        try:
            bot.handle_command("S3cret-pass", 55)
        finally:
            reminder_flow.handle_reminder_text = keep_text
        assert ("deleteMessage", {"chat_id": bot.CHAT_ID, "message_id": 55}) in calls
        assert noted == [] and "timed out" in views[-1] and guest_flow._guest == {}

        # ... but a command after the expiry is a command, and a late student ID is not deleted
        calls.clear()
        guest_flow._guest.update(step="password", at=0, messages=[])
        bot.handle_command("/help", 56)
        assert not any(m == "deleteMessage" for m, _ in calls) and guest_flow._guest == {}
        calls.clear()
        guest_flow._guest.update(step="user", at=0, messages=[])
        reminder_flow._reminder_wait = None
        bot.handle_command("20230001", 57)
        assert not any(m == "deleteMessage" for m, _ in calls)
    finally:
        bot.api, bot.send_view, guest_flow.DELETE_RETRY_SECONDS, reminder_flow._reminder_wait = keep
        guest_flow._guest.clear()


def test_stranger_floods_cost_no_disk_access():
    """A stranger hammering the bot: the first update of an id goes to failsafe (which reads and
    writes its file); every further one is turned away in memory. Many ids at once: at most
    MAX_ALERTS_PER_DAY alerts, then nothing touches the disk; the memory stays bounded; stale
    entries are pruned, and expired access:* records leave failsafe.json once a day."""
    disk = {"load": 0, "save": 0}
    real_load, real_save = failsafe._load, failsafe._save

    def counted_load():
        disk["load"] += 1
        return real_load()

    def counted_save(state):
        disk["save"] += 1
        return real_save(state)

    keep = (access.ADMIN_USER_ID, access.MAX_TRACKED)
    access.ADMIN_USER_ID, access.MAX_TRACKED = OWNER_ID, 50
    failsafe._load, failsafe._save = counted_load, counted_save
    for table in (access._logged, access._handled, access._alerts):
        table.clear()
    access._pruned.update(memory=0.0, file=0.0)
    _fresh_failsafe()
    t0 = 1_000_000.0
    try:
        stranger = _tg(user_id="777", chat_id="777", text="/start")
        access.deny(stranger, now=t0)
        failsafe.wait_idle()
        after_first = dict(disk)
        for n in range(1, 500):
            access.deny(stranger, now=t0 + n)
        assert disk == after_first, ("repeats of a known id never touch the disk", after_first, disk)
        assert len(OUTBOX) == 1

        # a flood of different ids: a few alerts, then nothing reaches the disk; memory stays bounded
        for n in range(3000):
            access.deny(_tg(user_id=str(10_000 + n), chat_id=str(10_000 + n), text="x"), now=t0 + 600 + n * 0.01)
        failsafe.wait_idle()
        assert len(OUTBOX) == access.MAX_ALERTS_PER_DAY, len(OUTBOX)
        settled = dict(disk)
        for n in range(3000, 3500):
            access.deny(_tg(user_id=str(10_000 + n), chat_id=str(10_000 + n), text="x"), now=t0 + 700 + n * 0.01)
        assert disk == settled, "past the daily cap, a flood costs no disk access"
        assert len(access._handled) <= access.MAX_TRACKED and len(access._logged) <= access.MAX_TRACKED

        # a day later: stale entries are pruned, the id may alert again, and failsafe.json is cleaned up
        later = t0 + access.ALERT_COOLDOWN + 3600
        record = real_load()                       # failsafe holds back on the wall clock: a day passes there too
        record["access:777"]["at"] = time.time() - access.ALERT_COOLDOWN - 1
        real_save(record)
        access.deny(stranger, now=later)
        failsafe.wait_idle()
        assert all(later - t < access.ALERT_COOLDOWN for t in access._handled.values())
        assert len(OUTBOX) == access.MAX_ALERTS_PER_DAY + 1, "after its window, the id is told about again"

        # failsafe.forget drops only the expired records of the given prefix
        real_save({"access:1": {"at": 0}, "access:2": {"at": later}, "delivery": {"at": 0, "active": True}})
        assert failsafe.forget("access:", access.ALERT_COOLDOWN, now=later) == 1
        assert set(real_load()) == {"access:2", "delivery"}
    finally:
        failsafe._load, failsafe._save = real_load, real_save
        access.ADMIN_USER_ID, access.MAX_TRACKED = keep
        for table in (access._logged, access._handled, access._alerts):
            table.clear()


def test_the_run_is_isolated_from_your_files():
    """The run (and every process it starts) reads the made-up .env and writes the throw-away data
    folder — never your .env or data/. UNINOTIFIER_ENV / UNINOTIFIER_DATA are what make it so."""
    import subprocess

    bot = _package("uninotifier.bot")
    real_root = Path(__file__).resolve().parents[1]
    assert paths.ENV_FILE == _HOME / ".env" and paths.DATA == _HOME / "data"
    assert not str(paths.DATA).startswith(str(real_root)), "state lives outside the project"
    assert bot.BOT_TOKEN == "123456:test-token" and bot.CHAT_ID == "42", "only the made-up credentials are loaded"
    assert notifier.SECRETS_FILE == paths.ENV_FILE
    from uninotifier import gitsync
    assert not gitsync.enabled() and gitsync.STATE_REPO == paths.DATA, "a throw-away state folder is never a backup repository"
    child = subprocess.run(
        [sys.executable, "-c", "from uninotifier import paths; print(paths.DATA); print(paths.ENV_FILE)"],
        capture_output=True, text=True, timeout=60, env={**os.environ, "PYTHONPATH": str(real_root / "src")})
    assert child.stdout.split() == [str(_HOME / "data"), str(_HOME / ".env")], (child.stdout, child.stderr[-300:])


def test_the_guards_block_the_network_and_the_browser():
    """The suite-wide guards: a real request to any outside host is refused and recorded (even when
    the code would have swallowed the error), so is starting a browser; loopback keeps working for
    the tests' own servers. The project's own paths are guarded too — the monitor's Browser class."""
    import requests
    from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
    import threading

    scan = importlib.import_module("uninotifier.monitor.scan")
    try:
        for attempt in (lambda: requests.get("https://elearning.zuj.edu.jo/my/", timeout=5),
                        lambda: socket.create_connection(("93.184.216.34", 443), timeout=5)):
            try:
                attempt()
                raise AssertionError("an outside connection got through")
            except (requests.ConnectionError, OSError):
                pass
        assert any("elearning.zuj.edu.jo" in h for h in GUARD_HITS) and any("93.184.216.34" in h for h in GUARD_HITS), GUARD_HITS
        GUARD_HITS.clear()
        try:
            scan.Browser().page                                  # what the monitor does to start Chromium
            raise AssertionError("a browser got started")
        except RuntimeError:
            pass
        assert GUARD_HITS and "browser" in GUARD_HITS[0]
    finally:
        GUARD_HITS.clear()

    class Hello(BaseHTTPRequestHandler):
        def do_GET(self):
            self.send_response(200)
            self.end_headers()
            self.wfile.write(b"hi")

        def log_message(self, *args):
            pass

    server = ThreadingHTTPServer(("127.0.0.1", 0), Hello)
    threading.Thread(target=server.serve_forever, daemon=True).start()
    try:
        assert requests.get(f"http://127.0.0.1:{server.server_port}/", timeout=5).text == "hi"
        assert not GUARD_HITS, "loopback is allowed"
    finally:
        server.shutdown()


def test_uploads_use_their_own_send_buffer():
    """Calls that carry files go through a session of their own: its connections connect within
    CONNECT_TIMEOUT and write within UPLOAD_WRITE_TIMEOUT, and its sockets may get a larger send buffer
    (UPLOAD_SEND_BUFFER, 0 = default). Everything else uses the shared pool; a reset starts both afresh."""
    keep = telegram.UPLOAD_SEND_BUFFER
    try:
        telegram.reset_session()
        telegram.UPLOAD_SEND_BUFFER = 0
        plain = telegram._upload_session()
        assert plain is not telegram._session(), "uploads always have their own pool"
        plain_adapter = plain.get_adapter("https://api.telegram.org")
        assert not any(opt == socket.SO_SNDBUF for _, opt, _ in plain_adapter.poolmanager.connection_pool_kw["socket_options"])
        assert plain_adapter.poolmanager.pool_classes_by_scheme["https"] is telegram._UploadPool
        shared_pools = telegram._session().get_adapter("https://api.telegram.org").poolmanager.pool_classes_by_scheme
        assert shared_pools["https"] is not telegram._UploadPool, "only uploads get the long write timeout"
        telegram.UPLOAD_SEND_BUFFER = 1 << 20
        upload = telegram._upload_session()
        assert upload is not telegram._session() and upload is telegram._upload_session()
        options = upload.get_adapter("https://api.telegram.org").poolmanager.connection_pool_kw["socket_options"]
        assert (socket.SOL_SOCKET, socket.SO_SNDBUF, 1 << 20) in options
        assert (socket.SOL_SOCKET, socket.SO_KEEPALIVE, 1) in options, "keep-alive probes stay on"
        shared = telegram._session().get_adapter("https://api.telegram.org").poolmanager.connection_pool_kw["socket_options"]
        assert not any(opt == socket.SO_SNDBUF for _, opt, _ in shared), "the long poll and buttons keep the default"

        # the door picks the session by whether the call carries files
        picked = []
        real = (telegram._session, telegram._upload_session)

        class Probe:
            def __init__(self, name):
                self.name = name

            def post(self, url, **kw):
                picked.append(self.name)

        telegram._session, telegram._upload_session = (lambda: Probe("shared")), (lambda size=None: Probe("upload"))
        try:
            _REAL_POST("https://api.telegram.org/botT/sendDocument", data={}, files={"document": ("a", b"x")})
            _REAL_POST("https://api.telegram.org/botT/sendMessage", json={})
        finally:
            telegram._session, telegram._upload_session = real
        assert picked == ["upload", "shared"], picked

        telegram.reset_session()
        assert telegram._upload_session() is not upload, "a reset starts the upload pool afresh too"

        # every call connects within CONNECT_TIMEOUT — an outage is noticed in 6 s, uploads included
        given, keep_post = [], telegram.post

        class Ok:
            status_code, ok = 200, True

        telegram.post = lambda url, **kw: given.append(kw["timeout"]) or Ok()
        try:
            telegram.request("T", "sendDocument", data={"chat_id": 1}, files={"document": ("a", b"x")}, timeout=60)
            telegram.request("T", "sendMessage", json={"chat_id": 1}, timeout=20)
        finally:
            telegram.post = keep_post
        assert given == [(telegram.CONNECT_TIMEOUT, 60), (telegram.CONNECT_TIMEOUT, 20)], given
        assert telegram.UPLOAD_WRITE_TIMEOUT >= 30

        # ... and an upload connection really connects under that, writes under UPLOAD_WRITE_TIMEOUT and
        # reads under the read timeout (urllib3 would write under the connect timeout: 6 s)
        import urllib3
        from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
        import threading

        seen = []

        class WriteProbe(telegram._WriteTimeout, urllib3.connection.HTTPConnection):
            def connect(self):
                seen.append(("connect", self.timeout))
                super().connect()

            def send(self, data):
                seen.append(("write", self.sock.gettimeout()))
                super().send(data)

        class Pool(urllib3.HTTPConnectionPool):
            ConnectionCls = WriteProbe

        class Sink(BaseHTTPRequestHandler):
            def do_POST(self):
                self.rfile.read(int(self.headers["Content-Length"]))
                self.send_response(200)
                self.send_header("Content-Length", "2")
                self.end_headers()
                self.wfile.write(b"ok")

            def log_message(self, *a):
                pass

        server = ThreadingHTTPServer(("127.0.0.1", 0), Sink)
        threading.Thread(target=server.serve_forever, daemon=True).start()
        try:
            pool = Pool("127.0.0.1", server.server_port)
            reply = pool.urlopen("POST", "/botT/sendDocument", body=b"x" * 50000,
                                 timeout=urllib3.Timeout(connect=telegram.CONNECT_TIMEOUT, read=60))
            assert reply.status == 200
            connection = pool.pool.queue[-1] or pool.pool.queue[0]
            read_timeout = connection.sock.gettimeout() if connection and connection.sock else 60
        finally:
            server.shutdown()
        assert seen[0] == ("connect", telegram.CONNECT_TIMEOUT), seen
        writes = {t for kind, t in seen if kind == "write"}
        assert writes == {telegram.UPLOAD_WRITE_TIMEOUT}, seen
        assert read_timeout == 60, read_timeout
    finally:
        telegram.UPLOAD_SEND_BUFFER = keep
        telegram.reset_session()


def test_schedule_files_never_hold_up_the_bot():
    """/schedule uploads files, which can take minutes on a bad connection. It runs beside the
    polling loop: the command returns at once, a second tap meanwhile starts nothing, and a crash
    is reported and frees it for the next tap."""
    import threading

    bot = _package("uninotifier.bot")
    views = importlib.import_module("uninotifier.bot.views")
    release, runs = threading.Event(), []

    def slow_upload():
        runs.append(1)
        release.wait(5)

    def worker():
        return next((t for t in threading.enumerate() if t.name == "schedule"), None)

    keep = views.send_schedule_files
    bot.send_schedule_files = slow_upload
    bot.api = lambda method, **data: {"message_id": 1}          # the button's answer (restored by the runner)
    _fresh_failsafe()
    try:
        started = time.monotonic()
        bot.handle_command("/schedule")
        bot.route_callback({"id": "cb1", "message": {"message_id": 3}}, "schedule", 3)   # a second tap meanwhile
        assert time.monotonic() - started < 1, "the polling loop is never held up by the upload"
        time.sleep(0.1)
        assert runs == [1], "one upload at a time"
        release.set()
        worker() and worker().join(5)

        def broken():
            raise OSError("disk gone")
        bot.send_schedule_files = broken
        assert views.start_schedule_files()
        worker() and worker().join(5)
        failsafe.wait_idle()
        assert any("timetable files could not be sent" in t for t in OUTBOX), OUTBOX
        assert views.start_schedule_files(), "a crash frees it for the next tap"
        worker() and worker().join(5)
    finally:
        bot.send_schedule_files = keep


def test_a_restart_never_replays_an_update():
    """The offset is saved before each update is handled: a bot killed while handling one (the
    monitor restarts a stuck bot) does not get the same update again — and cannot get stuck on it
    again. A saved offset of another bot (another token) is ignored."""
    bot = _package("uninotifier.bot")

    class Killed(BaseException):
        pass

    asked, handled, sleeps = [], [], []
    queue = [{"update_id": 10}, {"update_id": 11}, {"update_id": 12}]

    def get_updates(token, method, params=None, timeout=20, **kw):
        asked.append(params.get("offset"))
        return [u for u in queue if params.get("offset") is None or u["update_id"] >= params["offset"]]

    def dispatch(update):
        handled.append((update["update_id"], bot.load_offset()))
        if update["update_id"] == 11:
            raise Killed()                      # the process dies while handling update 11

    keep = (notifier.call_api, bot.dispatch, bot.OFFSET_FILE, bot.BOT_TOKEN, bot.time.sleep)
    with tempfile.TemporaryDirectory() as tmp:
        notifier.call_api, bot.dispatch, bot.OFFSET_FILE = get_updates, dispatch, Path(tmp) / "offset.json"
        try:
            assert bot.load_offset() is None
            try:
                bot.poll_once(bot.load_offset(), bot.OutageTracker())
            except Killed:
                pass
            assert handled == [(10, 11), (11, 12)], ("each offset is on disk before its update is handled", handled)
            # ... the restart: it asks past update 11, so 11 is not handled again; 12 is
            assert bot.poll_once(bot.load_offset(), bot.OutageTracker()) == 13
            assert asked == [None, 12] and [u for u, _ in handled] == [10, 11, 12]
            # another bot's offset is not used
            bot.BOT_TOKEN = "999999:other-token"
            assert bot.load_offset() is None
            bot.BOT_TOKEN = keep[3]
            # Telegram unreachable: the offset stays, the loop backs off
            notifier.call_api = lambda *a, **k: None
            bot.time.sleep = sleeps.append
            assert bot.poll_once(13, bot.OutageTracker()) == 13 and len(sleeps) == 1
        finally:
            notifier.call_api, bot.dispatch, bot.OFFSET_FILE, bot.BOT_TOKEN, bot.time.sleep = keep


def test_a_long_retry_after_is_not_waited_out():
    """HTTP 429: a "retry after" up to MAX_RETRY_WAIT is waited out once; a longer one gives the send
    up at once instead of holding the thread (on the bot: every button) for minutes. A malformed
    one is treated as no delay given."""
    class Reply:
        def __init__(self, status, body):
            self.status_code, self.ok, self._body = status, status < 400, body
            self.text = json.dumps(body)

        def json(self):
            return self._body

    def throttled(seconds):
        return Reply(429, {"ok": False, "error_code": 429, "parameters": {"retry_after": seconds}})

    replies, calls, sleeps = [], [], []

    def door(url, **kw):
        calls.append(url)
        return replies.pop(0)

    keep = (telegram.post, time.sleep)
    telegram.post, time.sleep = door, sleeps.append
    try:
        replies[:] = [throttled(300)]
        started = time.monotonic()
        assert not notifier.send_message("T", "42", "hello")
        assert len(calls) == 1 and sleeps == [] and time.monotonic() - started < 1, (calls, sleeps)

        calls.clear()
        replies[:] = [throttled(45), Reply(200, {"ok": True, "result": {}})]
        assert notifier.send_message("T", "42", "hello")
        assert sleeps == [45.0] and len(calls) == 2, (sleeps, calls)

        calls.clear()
        sleeps.clear()
        replies[:] = [throttled("soon"), Reply(200, {"ok": True, "result": {}})]
        assert notifier.send_message("T", "42", "hello") and sleeps == [notifier.BACKOFF_SECONDS[0]], sleeps
    finally:
        telegram.post, time.sleep = keep


def test_windows_device_names_are_saved_as_files():
    """A lecturer's file called "CON", "nul.pdf" or "Com1.txt" would be written to a Windows device
    (or fail) instead of the archive: such names get a "_" in front; ordinary names are untouched,
    and no name can still climb out of its folder."""
    for reserved in ("CON", "nul.pdf", "Com1.txt", "LPT9", "aux .doc", "PRN.tar.gz", "CONOUT$"):
        assert archive.safe_name(reserved) == "_" + reserved, reserved
    for ordinary in ("console.pdf", "Lecture CON.pdf", "auxiliary.docx", "com10.txt", "nullable.py", "Big Data"):
        assert archive.safe_name(ordinary) == ordinary, ordinary
    assert archive.safe_name("../../x") == "_.._x" and archive.safe_name("..") == "item"
    assert archive.safe_name("a" * 300) == "a" * 120
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / archive.safe_name("CON")
        path.write_bytes(b"lecture")
        assert path.read_bytes() == b"lecture" and [p.name for p in Path(tmp).iterdir()] == ["_CON"]


def test_the_poll_slows_down_when_telegram_asks():
    """A 429 on getUpdates: the loop waits what Telegram asked for (it used to try again within 8 s),
    but never longer than MAX_POLL_WAIT, so its heartbeat never looks stuck."""
    bot = _package("uninotifier.bot")

    class Reply:
        status_code, ok = 429, False

        def __init__(self, seconds):
            self._body = {"ok": False, "error_code": 429, "description": "Too Many Requests",
                          "parameters": {"retry_after": seconds}}

        def json(self):
            return self._body

    # call_api tells the caller how long Telegram asked to wait
    keep = (telegram.post, time.sleep, notifier.call_api)
    sleeps = []
    telegram.post, time.sleep = (lambda url, **kw: Reply(45)), sleeps.append
    try:
        told = {}
        assert notifier.call_api("T", "getUpdates", {}, outcome=told) is None and told == {"retry_after": 45.0}
        # ... and the loop waits it out — up to MAX_POLL_WAIT
        for asked, waited in ((45, 45), (900, bot.MAX_POLL_WAIT), (None, None)):
            sleeps.clear()

            def refused(token, method, data=None, timeout=20, outcome=None, asked=asked):
                if asked and outcome is not None:
                    outcome["retry_after"] = asked
                return None
            notifier.call_api = refused
            assert bot.poll_once(7, bot.OutageTracker()) == 7
            if waited:
                assert sleeps == [waited], (asked, sleeps)
            else:
                assert len(sleeps) == 1 and sleeps[0] <= 1, ("no 429: the usual short back-off", sleeps)
        assert bot.MAX_POLL_WAIT < heartbeat.STALE_SECONDS / 2
    finally:
        telegram.post, time.sleep, notifier.call_api = keep


def test_an_odd_getupdates_answer_never_stops_the_loop():
    """getUpdates answering something that is not a list of updates (call_api gives True for an
    "ok" without a result) or a list with malformed entries: skipped and logged — the polling
    loop never crashes on it."""
    bot = _package("uninotifier.bot")
    handled, sleeps = [], []
    keep = (notifier.call_api, bot.dispatch, bot.OFFSET_FILE, time.sleep)
    with tempfile.TemporaryDirectory() as tmp:
        bot.dispatch, bot.OFFSET_FILE, time.sleep = (lambda u: handled.append(u["update_id"])), Path(tmp) / "o.json", sleeps.append
        try:
            for odd in (True, {"result": []}, "x"):
                notifier.call_api = lambda *a, odd=odd, **k: odd
                assert bot.poll_once(4, bot.OutageTracker()) == 4
            assert handled == [] and len(sleeps) == 3, "each odd answer counts as a failed poll"
            notifier.call_api = lambda *a, **k: [None, {"message": {}}, {"update_id": "7"}, {"update_id": True},
                                                 {"update_id": 5}]
            assert bot.poll_once(4, bot.OutageTracker()) == 6 and handled == [5]
        finally:
            notifier.call_api, bot.dispatch, bot.OFFSET_FILE, time.sleep = keep


def test_the_statistics_do_not_grow_forever():
    """The always-on bot counts its Telegram calls per hour in memory; hours past STATS_KEEP_HOURS
    leave memory at the next flush, as they leave the file."""
    keep = dict(telegram._stats)
    try:
        telegram._stats.clear()
        old = time.strftime("%Y-%m-%dT%H", time.localtime(time.time() - (telegram.STATS_KEEP_HOURS + 5) * 3600))
        now = time.strftime("%Y-%m-%dT%H")
        telegram._stats.update({"2020-01-01T00": {"api:ok": 1}, old: {"poll:ok": 3}, now: {"api:ok": 2}})
        telegram.flush_stats()
        assert set(telegram._stats) == {now}, telegram._stats
        assert set(json.loads(telegram.STATS_FILE.read_text(encoding="utf-8"))) >= {now}
    finally:
        telegram._stats.clear()
        telegram._stats.update(keep)
        telegram.STATS_FILE.unlink(missing_ok=True)


def test_the_moodle_probe():
    """The 2-minute look: one batched request (updates since the last look + the calendar). The
    first look is the baseline; a new change starts ONE check; our own Mark as done (completion)
    is no change; a moved deadline is; a refused key or session is read afresh, an expired session
    waits for the monitor; nothing starts while a check runs, nor more than once per MIN_TRIGGER_GAP."""
    import requests
    from uninotifier import probe

    def updates(*pairs):
        return {"data": {"instances": [{"contextlevel": "module", "id": cm, "updates": [{"name": area, "timeupdated": t}]}
                                       for cm, area, t in pairs], "warnings": []}}

    courses = probe.course_ids()
    assert len(courses) == 6, courses                         # from the fixture state
    quiet = [updates() for _ in courses]
    calendar = {"data": {"events": [{"id": 1, "timesort": 100, "name": "HW1"}]}}

    # reading one answer
    found, fp, refused = probe.signals([updates((5, "configuration", 9), (5, "completion", 9))] + quiet[1:] + [calendar], courses)
    assert found == {f"{courses[0]}:5:configuration:9"} and fp and not refused
    assert probe.signals(quiet[:3], courses)[2], "a short answer (Moodle stopped at an error) is refused"
    assert probe.signals(quiet + [{"error": True, "exception": {"errorcode": "invalidsesskey"}}], courses)[2]

    class FakeHttp:
        def __init__(self, key="k"):
            self.key = key

        def sesskey(self):
            return self.key

    answers, launched = [], []
    keep = (probe.HttpContext.from_state_file, moodle.ajax_batch, probe.launch_check, probe.MONITOR_LOCK)
    probe.HttpContext.from_state_file = staticmethod(lambda path: FakeHttp())
    moodle.ajax_batch = lambda http, key, calls: answers.pop(0)
    probe.launch_check = lambda: launched.append(1) or type("P", (), {"poll": lambda self: 0})()
    with tempfile.TemporaryDirectory() as tmp:
        probe.MONITOR_LOCK = Path(tmp) / "monitor.lock"
        try:
            p, t = probe.Probe(), 1_000_000.0
            answers[:] = [quiet + [calendar]]
            assert not probe.look(p, t) and p.since == t, "the first look is the baseline"
            answers[:] = [[updates((7, "contentfiles", 50))] + quiet[1:] + [calendar]]
            assert probe.look(p, t + 120), "a new file"
            answers[:] = [[updates((7, "contentfiles", 50))] + quiet[1:] + [calendar]]
            assert not probe.look(p, t + 240), "the same change (the overlap asks back) counts once"
            answers[:] = [[updates((7, "completion", 300))] + quiet[1:] + [calendar]]
            assert not probe.look(p, t + 360), "our own Mark as done is not a change"
            answers[:] = [quiet + [{"data": {"events": [{"id": 1, "timesort": 200, "name": "HW1"}]}}]]
            assert probe.look(p, t + 480), "a deadline moved"
            answers[:] = [quiet + [{"error": True}]]
            assert not probe.look(p, t + 600) and p.sesskey is None and p.http is None, "refused: read afresh"
            probe.HttpContext.from_state_file = staticmethod(lambda path: FakeHttp(None))
            assert not probe.look(p, t + 720) and p.wait_until == t + 720 + probe.SESSION_RETRY, \
                "an expired session waits for the monitor, never signs in"

            # tick: only when due, never while a check runs, one check per MIN_TRIGGER_GAP
            probe.HttpContext.from_state_file = staticmethod(lambda path: FakeHttp())
            q, t = probe.Probe(), 2_000_000.0
            answers[:] = [quiet + [calendar]]
            assert not probe.tick(q, t)                                   # baseline
            assert not probe.tick(q, t + 30) and not answers, "not due yet: no request at all"
            probe.MONITOR_LOCK.write_text("1")
            assert not probe.tick(q, t + probe.PROBE_SECONDS), "a check is running: no look"
            probe.MONITOR_LOCK.unlink()
            for n, at in enumerate((2, 3, 4), 1):
                answers[:] = [[updates((9, "configuration", n))] + quiet[1:] + [calendar]]
                probe.tick(q, t + at * probe.PROBE_SECONDS)
            assert launched == [1], "three changes 2 min apart: one check (MIN_TRIGGER_GAP)"
            answers[:] = [[updates((9, "configuration", 99))] + quiet[1:] + [calendar]]
            later = t + 2 * probe.PROBE_SECONDS + probe.MIN_TRIGGER_GAP + probe.PROBE_SECONDS   # due, and past the gap
            assert probe.tick(q, later) and launched == [1, 1]

            # no network: logged, nothing raised, the next look reads afresh
            def offline(*a):
                raise requests.ConnectionError("down")
            moodle.ajax_batch = offline
            assert not probe.tick(q, t + 100 * probe.PROBE_SECONDS) and q.http is None
        finally:
            probe.HttpContext.from_state_file, moodle.ajax_batch, probe.launch_check, probe.MONITOR_LOCK = keep


def test_cambridge_one_assignments():
    """Cambridge One (temporary): the feed's address comes from the page's; the watch needs
    CAMBRIDGE_UNTIL and stops after it. The parser finds an assignment's fields under the usual
    names (or one level down), dates as ms / s / ISO, and falls back to a fingerprint id. The
    first feed is a baseline; a new assignment is a 🚨 alert with the screenshot and 🔗 Open in
    Cambridge One, seen only once it reached Telegram; a moved due date and a removal (missing
    twice) are told; the 24 h reminder goes once per due date and only for "not submitted". A
    refused session is renewed once in the browser, else the owner is told once to --login and
    nothing is retried until the session file changes; network blips stay quiet for an hour.
    Its session cookies and any JWT are masked in logs."""
    import requests
    from uninotifier import cambridge

    page = "https://www.cambridgeone.org/class/learner/org_test/class/c-1/assignments"
    assert cambridge.feed_url(page) == "https://www.cambridgeone.org/class/apigateway/org_test/class/c-1/learner/assignments"
    for bad in ("https://www.cambridgeone.org/dashboard/learner/dashboard", "https://example.invalid/x"):
        try:
            cambridge.feed_url(bad)
            raise AssertionError(bad)
        except ValueError:
            pass

    recorded = json.loads((Path(__file__).resolve().parent / "fixtures" / "cambridge_feed.json").read_text(encoding="utf-8"))
    assert cambridge.parse(recorded) == [] and cambridge.class_title(recorded).startswith("Teacher. B2")

    t = 1_800_000_000.0
    ms = lambda seconds: int(seconds * 1000)          # noqa: E731
    a1 = {"uuid": "a1", "title": "Unit 1 Writing_task", "due_date": ms(t + 3 * 86400), "status": "not_started"}
    a2 = {"assignment": {"name": "Unit 2 Reading"}, "endDate": "2027-01-10T21:00:00Z", "submitted": True}
    a3 = {"id": 7, "name": "Unit 3 Quiz", "state": "something new"}
    p1, p2, p3 = cambridge.parse({"assignmentsData": {"entities": [a1, a2, a3, "junk"]}})
    assert (p1["id"], p1["due"], p1["submitted"]) == ("a1", t + 3 * 86400, False)
    assert p2["title"] == "Unit 2 Reading" and p2["id"].startswith("fp-") and p2["submitted"] is True
    assert p2["due"] == datetime.fromisoformat("2027-01-10T21:00:00+00:00").timestamp()
    assert (p3["id"], p3["due"], p3["submitted"]) == ("7", None, None)
    assert cambridge.to_epoch(1_800_000_000) == 1_800_000_000 and cambridge.to_epoch("soon") is None
    assert "fields" not in p1 and p1["activities"] == [], "no raw field dump any more"

    # the real assignment (2026-10-21): its id, a short title, the activities, and the submission rule
    real = json.loads((Path(__file__).resolve().parent / "fixtures" / "cambridge_assignment.json").read_text(encoding="utf-8"))
    entity = real["assignmentsData"]["entities"][0]
    [a] = cambridge.parse(real)
    assert a["id"] == entity["assigned_path_id"] and a["title"] == "Assignment 1"
    assert a["full_title"] == "Assignment 1, Ross. B2. sec. 27 . 3:30-5:00"
    assert (a["start"], a["due"]) == (entity["startdate"] / 1000, entity["duedate"] / 1000)
    assert len(a["activities"]) == 15 and a["activities"][0]["name"] == "Vocabulary: Character adjectives 1"
    assert a["submitted"] is False, "not started: not submitted"
    text = cambridge.new_message(a, cambridge.class_title(real), 1_792_540_000)
    lines = text.splitlines()
    assert lines[:6] == ["🚨 *New Assignment — Cambridge One*", "📘 *Course:* English Communication Skills",
                         "🏫 *Class:* Ross. B2. sec. 27 . 3:30-5:00", "📌 *Title:* Assignment 1",
                         f"🗓️ *Opens:* {notifier.format_timestamp(a['start'])}",
                         f"🕐 *Due Date:* {notifier.format_timestamp(a['due'])}"], lines[:6]
    assert lines[6:8] == ["", "📋 *Activities (15):*"] and lines[8] == "• Vocabulary: Character adjectives 1"
    assert lines[22] == "• Writing: An article" and lines[-1].startswith("🕒 *Detected:*")
    for raw in ("rule_id", "sub-type", "tags", "assigned_path_id", "__analytics", "•  "):
        assert raw not in text, raw
    assert len(text) <= notifier.CAPTION_LIMIT, "fits the screenshot's caption: one message"

    def progressed(*done_flags):
        e = json.loads(json.dumps(entity))
        for item, (submitted, status) in zip(e["items"], done_flags):
            item["__analytics"].update(submittedOnce=submitted, status=status)
        return cambridge.parse({**real, "assignmentsData": {"entities": [e]}})[0]

    all_done = [(True, "completed")] * 15
    assert progressed(*all_done)["submitted"] is True, "every activity submitted and completed"
    assert progressed(*all_done[:14], (True, "in_progress"))["submitted"] is False, "one not completed"
    assert progressed(*all_done[:14], (False, "completed"))["submitted"] is False, "one never submitted"
    half = progressed(*all_done[:5])
    assert "10 of 15 activities not done yet" in cambridge.reminder_message(half, "", half["due"] - 3600)
    # the fallback id ignores progress: starting an activity never makes it a new assignment
    no_id = {k: v for k, v in entity.items() if k != "assigned_path_id"}
    started = json.loads(json.dumps(no_id))
    started["items"][0]["__analytics"].update(status="in_progress", launchCount=3)
    assert cambridge.fingerprint(no_id) == cambridge.fingerprint(started)
    # recorded under a fingerprint before its id was known: kept under the id, announced once
    seen = {"fp-0123456789abcdef": {"title": a["full_title"], "due": a["due"], "missing": 0}}
    reminded = {"fp-0123456789abcdef": a["due"]}
    cambridge.adopt_legacy(seen, reminded, [a])
    assert list(seen) == [a["id"]] and list(reminded) == [a["id"]]

    def feed(*entities):
        return {"classDetails": {"title": "Teacher. B2"}, "assignmentsData": {"entities": list(entities)}}

    sent: list[tuple[str, str, object]] = []
    online = [True]

    def png_file(token, chat, png, name, caption="", url=None, rows=None):
        if online[0]:
            sent.append(("png", caption, rows))
        return online[0]

    def long_message(token, chat, text, url=None, rows=None):
        if online[0]:
            sent.append(("text", text, rows))
        return online[0]

    keep_files = (cambridge.SEEN_FILE, cambridge.FEED_FILE, cambridge.STATE_FILE)
    keep_calls = (notifier.send_png_file, notifier.send_long_message, cambridge.fetch, cambridge.visit)
    keep_env = {k: os.environ.get(k) for k in (cambridge.ENV_URL, cambridge.ENV_UNTIL)}
    shots = []
    with tempfile.TemporaryDirectory() as tmp:
        cambridge.SEEN_FILE, cambridge.FEED_FILE = Path(tmp) / "cambridge.json", Path(tmp) / "cambridge_feed.json"
        cambridge.STATE_FILE = Path(tmp) / "cambridge_state.json"
        notifier.send_png_file, notifier.send_long_message = png_file, long_message
        shoot = lambda: shots.append(1) or b"png"     # noqa: E731
        _fresh_failsafe()
        try:
            # baseline, then a new assignment due in 10 h (the alert says so: no reminder on top)
            assert cambridge.apply(feed(a1), t, page, shoot)["baseline"] and not sent and not shots
            a4 = {"uuid": "a4", "title": "Unit 4", "due_date": ms(t + 10 * 3600), "status": "in_progress"}
            counts = cambridge.apply(feed(a1, a4), t + 60, page, shoot)
            assert counts["new"] == 1 and counts["reminded"] == 0 and shots == [1], counts
            kind, caption, rows = sent.pop()
            assert kind == "png" and caption.startswith("🚨 *New Assignment — Cambridge One*") and "Unit 4" in caption
            assert "📘 *Course:* English Communication Skills" in caption and "uuid" not in caption
            assert rows == [[{"text": cambridge.BUTTON_TEXT, "url": page}]] and not sent

            # Telegram unreachable: not seen, so the next check sends it
            a5 = {"uuid": "a5", "title": "Unit 5", "due_date": ms(t + 5 * 86400)}
            online[0] = False
            assert cambridge.apply(feed(a1, a4, a5), t + 120, page, shoot)["new"] == 0
            online[0] = True
            assert cambridge.apply(feed(a1, a4, a5), t + 180, page, shoot)["new"] == 1 and len(sent) == 1
            sent.clear()
            assert cambridge.apply(feed(a1, a4, a5), t + 240, page, shoot) == {
                "assignments": 3, "new": 0, "moved": 0, "gone": 0, "reminded": 0, "baseline": False}

            # a1's due date moves to 20 h away: told, and (not submitted) the reminder follows — once
            a1_moved = {**a1, "due_date": ms(t + 20 * 3600)}
            counts = cambridge.apply(feed(a1_moved, a4, a5), t + 300, page, shoot)
            assert counts["moved"] == 1 and counts["reminded"] == 1, counts
            assert "Due Date Changed" in sent[0][1] and "↩️ *Was:*" in sent[0][1]
            assert sent[1][1].startswith("🚨 *Reminder — Cambridge One*") and "not submitted yet" in sent[1][1]
            assert sent[1][2] == [[{"text": cambridge.BUTTON_TEXT, "url": page}]]
            sent.clear()
            assert cambridge.apply(feed(a1_moved, a4, a5), t + 360, page, shoot)["reminded"] == 0 and not sent

            # submitted, or an unknown status: no reminder
            a6 = {"uuid": "a6", "title": "Unit 6", "due_date": ms(t + 30 * 3600), "status": "weird"}
            a7 = {"uuid": "a7", "title": "Unit 7", "due_date": ms(t + 30 * 3600), "status": "submitted"}
            cambridge.apply(feed(a1_moved, a4, a5, a6, a7), t + 400, page, shoot)
            sent.clear()
            assert cambridge.apply(feed(a1_moved, a4, a5, a6, a7), t + 8 * 3600, page, shoot)["reminded"] == 0

            # gone only once it is missing twice in a row (one odd answer is not a removal)
            assert cambridge.apply(feed(a1_moved, a5, a6, a7), t + 9 * 3600, page, shoot)["gone"] == 0
            assert cambridge.apply(feed(a1_moved, a5, a6, a7), t + 10 * 3600, page, shoot)["gone"] == 1
            assert "Assignment Removed" in sent[-1][1] and "Unit 4" in sent[-1][1]
            assert json.loads(cambridge.FEED_FILE.read_text(encoding="utf-8"))["entities"][0]["uuid"] == "a1"

            # an unreadable record: set aside, told, a new baseline
            cambridge.SEEN_FILE.write_text("{broken", encoding="utf-8")
            sent.clear()
            assert cambridge.apply(feed(a1_moved), t + 11 * 3600, page, shoot)["baseline"] and not sent
            assert list(Path(tmp).glob("cambridge.corrupt-*.json"))
            assert any("Cambridge One record was unreadable" in m for m in OUTBOX), OUTBOX

            # the thread's tick: settings, renewal, waiting for --login, quiet network blips
            os.environ[cambridge.ENV_URL] = page
            os.environ.pop(cambridge.ENV_UNTIL, None)
            assert not cambridge.active(), "no end date: not watching"
            os.environ[cambridge.ENV_UNTIL] = "2027-01-31"
            assert cambridge.active(date(2027, 1, 31)) and not cambridge.active(date(2027, 2, 1))
            w = cambridge.Watcher()
            assert cambridge.tick(w, t, date(2027, 2, 1)) is None and w.ended, "past CAMBRIDGE_UNTIL the thread ends"

            cambridge.STATE_FILE.write_text(json.dumps({"cookies": [
                {"name": "c1_sid", "value": "s1", "domain": ".cambridgeone.org", "path": "/"}]}), encoding="utf-8")
            answers, renewals = [], []

            def fake_fetch(http, url):
                assert http.headers["cambridgeone-app-version"] == "v2" and http.cookies.get("c1_sid") == "s1"
                answer = answers.pop(0)
                if isinstance(answer, Exception):
                    raise answer
                return answer

            cambridge.fetch = fake_fetch
            cambridge.visit = lambda shot: renewals.append(shot) or (renewal_ok[0], None)
            renewal_ok = [True]
            today = date(2026, 11, 1)
            _fresh_failsafe()
            w = cambridge.Watcher()
            answers[:] = [cambridge.Expired("HTTP 401"), feed(a1_moved)]
            assert cambridge.tick(w, t, today) is not None and renewals == [False], "renewed once, then read"
            assert cambridge.tick(w, t + 60, today) is None and not answers, "not due yet: no request"

            renewal_ok[0] = False
            answers[:] = [cambridge.Expired("HTTP 403")]
            assert cambridge.tick(w, t + cambridge.CHECK_SECONDS, today) is None and w.blocked_at is not None
            failsafe.wait_idle()
            assert sum("sign in again" in m for m in OUTBOX) == 1, OUTBOX
            assert cambridge.tick(w, t + 2 * cambridge.CHECK_SECONDS, today) is None and len(renewals) == 2, \
                "waiting for --login: nothing is tried"
            cambridge.STATE_FILE.write_text(cambridge.STATE_FILE.read_text(encoding="utf-8") + " ", encoding="utf-8")
            os.utime(cambridge.STATE_FILE, (t, t))           # --login wrote the session again
            answers[:] = [feed(a1_moved)]
            assert cambridge.tick(w, t + 3 * cambridge.CHECK_SECONDS, today) is not None and w.blocked_at is None
            failsafe.wait_idle()
            assert any("signed in again" in m for m in OUTBOX), OUTBOX

            _fresh_failsafe()
            later = t + 4 * cambridge.CHECK_SECONDS
            for n in range(cambridge.FAILURES_BEFORE_ALERT):
                answers[:] = [requests.ConnectionError("down")]
                assert cambridge.tick(w, later + n * cambridge.RETRY_SECONDS, today) is None
                assert w.next_check == later + n * cambridge.RETRY_SECONDS + cambridge.RETRY_SECONDS
                failsafe.wait_idle()
                assert bool(OUTBOX) == (n == cambridge.FAILURES_BEFORE_ALERT - 1), (n, OUTBOX)
        finally:
            cambridge.SEEN_FILE, cambridge.FEED_FILE, cambridge.STATE_FILE = keep_files
            notifier.send_png_file, notifier.send_long_message, cambridge.fetch, cambridge.visit = keep_calls
            for key, value in keep_env.items():
                if value is None:
                    os.environ.pop(key, None)
                else:
                    os.environ[key] = value

    jwt = "eyJhbGciOiJSUzI1NiJ9.eyJzdWIiOiJ0ZXN0In0.c2lnbmF0dXJlX3Rlc3Q"
    clean = notifier.redact(f"Authorization: Bearer {jwt} Cookie: c1_sid=s3cr3tSID; "
                            "glt_3_AbC-def=GltValue123; gig_loginToken_3_X=TokVal99")
    for secret in (jwt, "s3cr3tSID", "GltValue123", "TokVal99"):
        assert secret not in clean, (secret, clean)
    assert "Bearer <redacted>" in clean and "c1_sid=<redacted>" in clean, clean


def test_cambridge_signs_in_by_itself():
    """Cambridge ends a session about 72 h after the sign-in and its silent renewal does not exist
    (the login token lives 6 h), so the hidden browser signs in with CAMBRIDGE_EMAIL / _PASSWORD
    when the feed refuses the session: a valid session types nothing; an ended one fills Gigya's
    form, presses Enter and saves the new session. Cambridge's own error text makes it
    CredentialsRejected (tried again after a day, not every 6 h: a lockout is worse than waiting);
    no form, or a check only a person can pass, is a plain NeedsLogin; an address that never
    answers is a network blip, not a sign-in alert. Without the two settings nothing is typed and
    the alert says how to stop signing in by hand. The password is in no message or alert."""
    import requests
    from playwright.sync_api import Error
    from uninotifier import cambridge

    email, secret = "student@example.invalid", "fake pass #1 {with} [symbols] & spaces"
    page_url = "https://www.cambridgeone.org/class/learner/org_test/class/c-1/assignments"
    login_url = "https://www.cambridgeone.org/login?returnUrl=%2Fclass"
    scenario: dict = {}
    pages: list = []

    class Locator:
        def __init__(self, page, selector):
            self.page, self.selector = page, selector

        def locator(self, selector):
            return Locator(self.page, f"{self.selector} {selector}")

        def wait_for(self, state=None, timeout=None):
            if scenario.get("no_form"):
                raise Error("Locator.wait_for: Timeout 15000ms exceeded.")

        def fill(self, value, timeout=None):
            self.page.typed.append((self.selector, value))

        def press(self, key, timeout=None):
            self.page.pressed.append(key)

        def all_inner_texts(self):
            return scenario.get("errors", [])

    class Page:
        def __init__(self):
            self.signed, self.typed, self.pressed, self.url = bool(scenario.get("valid")), [], [], "about:blank"

        def goto(self, url, **kwargs):
            self.url = page_url if self.signed else "chrome-error://chromewebdata/" if scenario.get("offline") else login_url

        def expect_response(self, predicate, timeout=None):
            page = self

            class Waiting:
                def __enter__(self):
                    return self

                def __exit__(self, kind, value, tb):
                    if kind is None and not page.signed:
                        raise Error("Page.expect_response: Timeout exceeded.")
                    return False
            return Waiting()

        def locator(self, selector):
            return Locator(self, selector)

        def wait_for_url(self, predicate, timeout=None, wait_until=None):
            assert self.pressed, "nothing was submitted yet"
            if not scenario.get("accepts"):
                raise Error("Page.wait_for_url: Timeout 45000ms exceeded.")
            self.signed, self.url = True, "https://www.cambridgeone.org/"
            assert predicate(self.url) and not predicate(login_url), "leaving /login is what counts"

    class Context:
        def new_page(self):
            pages.append(Page())
            return pages[-1]

        def storage_state(self):
            return {"cookies": [{"name": "c1_sid", "value": "fresh"}], "origins": []}

    class Browser:
        def new_context(self, **kwargs):
            return Context()

        def close(self):
            pass

    class Playwright:
        chromium = type("C", (), {"launch": staticmethod(lambda headless=True: Browser())})()

        def __enter__(self):
            return self

        def __exit__(self, *a):
            return False

    def attempt(**kw):
        scenario.clear()
        scenario.update(kw)
        try:
            return cambridge.visit(False)
        except Exception as e:                       # the point is which one
            return e

    typed_in = lambda: pages[-1].typed               # noqa: E731
    form = "#gigya-login-form"
    keep_env = {k: os.environ.get(k) for k in (cambridge.ENV_URL, cambridge.ENV_UNTIL, cambridge.ENV_EMAIL,
                                               cambridge.ENV_PASSWORD)}
    keep = (_playwright_api.sync_playwright, cambridge.STATE_FILE, cambridge.fetch, cambridge.SEEN_FILE,
            cambridge.FEED_FILE)
    with tempfile.TemporaryDirectory() as tmp:
        cambridge.STATE_FILE = Path(tmp) / "cambridge_state.json"
        cambridge.SEEN_FILE, cambridge.FEED_FILE = Path(tmp) / "cambridge.json", Path(tmp) / "cambridge_feed.json"
        _playwright_api.sync_playwright = lambda: Playwright()

        def old_session():
            cambridge.STATE_FILE.write_text(json.dumps({"cookies": [{"name": "c1_sid", "value": "old"}]}), encoding="utf-8")

        def saved():
            return json.loads(cambridge.STATE_FILE.read_text(encoding="utf-8"))["cookies"][0]["value"]

        try:
            os.environ[cambridge.ENV_URL] = page_url
            os.environ[cambridge.ENV_UNTIL] = "2027-01-31"
            os.environ[cambridge.ENV_EMAIL], os.environ[cambridge.ENV_PASSWORD] = email, secret
            assert cambridge.credentials() == (email, secret)
            for half in (cambridge.ENV_EMAIL, cambridge.ENV_PASSWORD):
                kept = os.environ.pop(half)
                assert cambridge.credentials() is None, f"{half} alone is not enough"
                os.environ[half] = kept

            old_session()
            assert attempt(valid=True) == (True, None) and typed_in() == [] and saved() == "fresh", "valid: nothing typed"

            old_session()
            assert attempt(accepts=True) == (True, None), "ended session: signs in by itself"
            assert typed_in() == [(f"{form} input[name=username]", email), (f"{form} input[name=password]", secret)]
            assert pages[-1].pressed == ["Enter"] and saved() == "fresh", "the new session is saved"

            old_session()
            rejected = attempt(errors=["", "Invalid login or password"])
            assert isinstance(rejected, cambridge.CredentialsRejected), rejected
            assert "Invalid login or password" in str(rejected) and saved() == "old", "a refused sign-in saves nothing"
            stuck = attempt()
            assert type(stuck) is cambridge.NeedsLogin and "check" in str(stuck), stuck
            missing = attempt(no_form=True)
            assert type(missing) is cambridge.NeedsLogin and "could not be filled in" in str(missing), missing
            for failure in (rejected, stuck, missing):
                assert secret not in str(failure) and email not in str(failure), failure
            offline = attempt(offline=True)
            assert isinstance(offline, requests.ConnectionError) and not typed_in(), "no answer is not a sign-in problem"
            assert cambridge.screenshot() is None, "an alert still goes out without its picture"

            # the thread: refused credentials wait a day, a missing form 6 h; the alert says what to do
            t, today = 1_800_000_000.0, date(2026, 11, 1)
            old_session()

            def expired(http, url):
                raise cambridge.Expired("HTTP 401")
            cambridge.fetch = expired
            _fresh_failsafe()
            w = cambridge.Watcher()
            scenario.clear()
            scenario.update(errors=["Invalid login or password"])
            assert cambridge.tick(w, t, today) is None and w.retry_gap == cambridge.CREDENTIALS_RETRY
            failsafe.wait_idle()
            [alert] = [m.replace("\\_", "_") for m in OUTBOX]          # Markdown escapes the underscores
            assert "sign in again" in alert and "automatic sign-in did not work" in alert, alert
            assert "Invalid login or password" in alert and "CAMBRIDGE_PASSWORD" in alert, alert
            assert secret not in alert and email not in alert, alert
            tried = len(pages)
            assert cambridge.tick(w, t + 2 * cambridge.RELOGIN_RETRY, today) is None and len(pages) == tried, \
                "refused credentials are not tried again after 6 h"
            assert cambridge.tick(w, t + cambridge.CREDENTIALS_RETRY + cambridge.CHECK_SECONDS, today) is None
            assert len(pages) == tried + 1, "... but after a day they are"
            old_session()
            os.utime(cambridge.STATE_FILE, (t, t))
            w2 = cambridge.Watcher()
            w2.blocked_at, w2.blocked_mtime, w2.retry_gap = t, t - 100, cambridge.CREDENTIALS_RETRY   # the file has changed since
            cambridge.fetch = lambda http, url: {"classDetails": {}, "assignmentsData": {"entities": []}}
            assert cambridge.tick(w2, t + 60, today) is not None and w2.blocked_at is None, "--login unblocks at once"

            cambridge.fetch = expired
            _fresh_failsafe()
            os.environ.pop(cambridge.ENV_EMAIL), os.environ.pop(cambridge.ENV_PASSWORD)
            w3 = cambridge.Watcher()
            scenario.clear()
            assert cambridge.tick(w3, t, today) is None and w3.retry_gap == cambridge.RELOGIN_RETRY
            assert typed_in() == [], "no settings: nothing is typed"
            failsafe.wait_idle()
            [alert] = [m.replace("\\_", "_") for m in OUTBOX]
            assert "--login" in alert and "CAMBRIDGE_EMAIL" in alert and "CAMBRIDGE_PASSWORD" in alert, alert
        finally:
            (_playwright_api.sync_playwright, cambridge.STATE_FILE, cambridge.fetch, cambridge.SEEN_FILE,
             cambridge.FEED_FILE) = keep
            for key, value in keep_env.items():
                if value is None:
                    os.environ.pop(key, None)
                else:
                    os.environ[key] = value


def test_log_files_stay_small():
    """A log starts a new file at LOG_MAX_BYTES and keeps the last LOG_BACKUPS, newest first
    (monitor.log.1), so a year of running cannot fill the disk. Windows refuses to rename a log
    another process has open: the record being logged still lands, nothing is raised, and the
    roll happens at the next record once the file is free. Both entry points use it."""
    import inspect
    import logging.handlers

    assert "RotatingLogFile" in inspect.getsource(notifier.setup_logging)
    assert notifier.LOG_BACKUPS >= 2 and notifier.LOG_MAX_BYTES >= 1024 * 1024
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "monitor.log"
        handler = notifier.RotatingLogFile(path, maxBytes=400, backupCount=2, encoding="utf-8")
        logger = logging.Logger("rotation-test")
        logger.addHandler(handler)
        real_rename = os.rename
        try:
            for n in range(30):
                logger.warning("line %02d %s", n, "x" * 40)
            names = sorted(p.name for p in Path(tmp).iterdir())
            assert names == ["monitor.log", "monitor.log.1", "monitor.log.2"], names
            assert "line 29" in path.read_text(encoding="utf-8"), "the newest records are in the live file"
            assert "line 00" not in "".join(p.read_text(encoding="utf-8") for p in Path(tmp).iterdir()), "the oldest are gone"

            def busy(src, dst):
                raise PermissionError(13, "The process cannot access the file because it is being used by another process")
            os.rename = busy
            for n in range(30, 40):                      # every roll fails: still no exception, nothing lost
                logger.warning("line %02d %s", n, "x" * 40)
            assert "line 39" in path.read_text(encoding="utf-8")
            os.rename = real_rename
            logger.warning("line 40 %s", "x" * 40)
            assert "line 40" in "".join(p.read_text(encoding="utf-8") for p in Path(tmp).iterdir())
            assert path.stat().st_size < 400 + 100, "rolled again once the file was free"
        finally:
            os.rename = real_rename
            handler.close()


def test_replaced_files_move_to_old_versions():
    """An instructor's new version takes the file's place in its folder; the old one moves, unchanged,
    into _Old_Versions beside it (created only then) as <name>_v<n> (<its date>). The record keeps the
    latest file and every retired one, so neither is downloaded again. A run that repeats (its
    announcement failed) moves and copies nothing twice; the same bytes with a new date change no
    file; a new version under another name retires the old name. The tidy tool leaves the old
    versions alone, and the message says where the previous file went. A file two items share is
    never retired: the updated item's new version goes beside it."""
    from datetime import datetime
    from uninotifier.tools import tidy_archive
    items = importlib.import_module("uninotifier.monitor.items")

    class Resp:
        def __init__(self, body, name):
            self._body, self.ok = body, True
            self.url = f"https://x/pluginfile.php/1/mod_resource/content/1/{name}"
            self.headers = {"content-type": "application/pdf", "last-modified": "Wed, 01 Jan 2025 00:00:00 GMT"}

        def body(self):
            return self._body

        def text(self):
            return ""

    class Ctx:
        def __init__(self, body, name="Lecture.pdf"):
            self._resp = Resp(body, name)
            self.request = self

        def get(self, url, timeout=0):
            return self._resp

    def day(ts):
        return datetime.fromtimestamp(ts).strftime("%Y-%m-%d")

    item = {"course": "20261-Big Data-AllSec-3-F", "type": "resource", "name": "Chapter 1", "url": "https://x/mod/resource/view.php?id=7"}
    t1, t2, t3 = 1_790_000_000, 1_790_500_000, 1_791_000_000
    with tempfile.TemporaryDirectory() as tmp:
        keep = (archive.ARCHIVE_ROOT, archive._ARCHIVE_ROOT_RESOLVED, archive.probe_resource)
        archive.ARCHIVE_ROOT, archive._ARCHIVE_ROOT_RESOLVED = Path(tmp), Path(tmp).resolve()
        probe = {"modified": t2, "size": None}
        archive.probe_resource = lambda context, url: dict(probe)
        try:
            folder = archive.target(item["course"], "resource", item["name"])
            first = archive.download_resource(Ctx(b"V1" * 200), item["url"], folder)
            record = {"path": str(first["path"]), "size": 400, "sha256": first["sha256"], "modified": t1}
            old_versions = folder / archive.OLD_VERSIONS
            assert not old_versions.exists(), "no _Old_Versions until a file is actually replaced"

            # version 2: in the file's place; version 1 retired, unchanged
            art = items.check_file_update(Ctx(b"V2" * 200), item, record)
            main = folder / "Lecture.pdf"
            retired_1 = old_versions / f"Lecture_v1 ({day(t1)}).pdf"
            assert main.read_bytes() == b"V2" * 200 and retired_1.read_bytes() == b"V1" * 200
            assert sorted(p.name for p in folder.iterdir()) == sorted([archive.OLD_VERSIONS, "Lecture.pdf"]), "no stray copies"
            assert art["path"] == str(main.resolve()) and art["old_path"] == str(retired_1.resolve())
            assert art["versions"] == [{"path": str(retired_1.resolve()), "size": 400, "sha256": record["sha256"], "modified": t1}]

            # the announcement failed, the record stayed: the next run finds everything in place
            again = items.check_file_update(Ctx(b"V2" * 200), item, record)
            assert again["path"] == art["path"] and again["versions"] == art["versions"]
            assert sorted(p.name for p in folder.iterdir()) == sorted([archive.OLD_VERSIONS, "Lecture.pdf"])
            assert [p.name for p in old_versions.iterdir()] == [retired_1.name], "nothing moved or copied twice"

            # recorded as run.py does; version 3 retires version 2 beside version 1
            record = {"path": art["path"], "size": art["size"], "sha256": art["sha256"], "modified": art["modified"],
                      "versions": art["versions"]}
            probe["modified"] = t3
            art = items.check_file_update(Ctx(b"V3" * 200), item, record)
            assert main.read_bytes() == b"V3" * 200
            assert sorted(p.name for p in old_versions.iterdir()) == sorted([retired_1.name, f"Lecture_v2 ({day(record['modified'])}).pdf"]), "named after the replaced file's own date"
            assert [v["sha256"] for v in art["versions"]] == [first["sha256"], record["sha256"]], "both old versions tracked"

            # the same bytes with a newer date: no file changes, only the record's date
            record = {**record, "path": art["path"], "sha256": art["sha256"], "modified": t3, "versions": art["versions"]}
            probe["modified"] = t3 + 60
            assert items.check_file_update(Ctx(b"V3" * 200), item, record) is None and record["modified"] == t3 + 60
            assert len(list(folder.iterdir())) == 2 and len(list(old_versions.iterdir())) == 2

            # a new version under another name: the old name is retired, the new one takes the folder
            probe["modified"] = t3 + 3600
            art = items.check_file_update(Ctx(b"V4" * 200, "Lecture-revised.pdf"), item, record)
            assert sorted(p.name for p in folder.iterdir()) == sorted([archive.OLD_VERSIONS, "Lecture-revised.pdf"])
            assert (old_versions / f"Lecture_v3 ({day(record['modified'])}).pdf").read_bytes() == b"V3" * 200
            again = items.check_file_update(Ctx(b"V4" * 200, "Lecture-revised.pdf"), item, record)
            assert again["path"] == art["path"] and sorted(p.name for p in folder.iterdir()) == \
                sorted([archive.OLD_VERSIONS, "Lecture-revised.pdf"]), "repeated: no 'Lecture-revised (1).pdf'"

            # the message: where the previous file went
            text = media.build_update_message(item, t3, t3 + 3600, old_name=Path(art["old_path"]).name,
                                              new_name="Lecture-revised.pdf", where=archive.destination(Path(art["path"])),
                                              retired_to=archive.retired_to(Path(art["old_path"])))
            assert f"`Big Data/01_Lectures/{archive.OLD_VERSIONS}/{Path(art['old_path']).name}`" in text
            assert "takes the previous file's place" in text and "beside the previous file" not in text
            assert "nothing was moved" in media.build_update_message(item, t1, t2, new_name="x.pdf")

            # the tidy tool never "de-duplicates" history
            (old_versions / "Lecture_v1 (copy).pdf").write_bytes(b"V1" * 200)
            assert not [k for k, _ in tidy_archive.find_duplicates(Path(tmp)) if archive.OLD_VERSIONS in k.parts]

            # one file, two items: a new version of one goes beside it; the file stays for the other
            got = archive.download_resource(Ctx(b"S1" * 200, "Shared.pdf"), item["url"], folder)
            shared_file = Path(got["path"])
            rec_a = {"path": str(shared_file), "size": 400, "sha256": got["sha256"], "modified": t1}
            files = {"a": rec_a, "b": dict(rec_a), "c": {"path": art["path"]}}
            shared = archive.shared_paths(files)
            assert shared == {os.path.normcase(str(shared_file))}
            assert archive.is_shared(rec_a, shared) and not archive.is_shared(files["c"], shared)
            probe["modified"] = t3 + 7200
            moved = items.check_file_update(Ctx(b"S2" * 200, "Shared.pdf"), item, rec_a, shared=True)
            assert shared_file.read_bytes() == b"S1" * 200, "the other item's file is untouched"
            new_file = Path(moved["path"])
            assert new_file.parent == folder.resolve() and new_file != shared_file.resolve()
            assert new_file.read_bytes() == b"S2" * 200
            assert moved["versions"] == [] and moved["old_path"] == rec_a["path"]
            assert not [p for p in old_versions.iterdir() if p.name.startswith("Shared")], "nothing retired"
            again = items.check_file_update(Ctx(b"S2" * 200, "Shared.pdf"), item, rec_a, shared=True)
            assert again["path"] == moved["path"]
            assert len([p for p in folder.iterdir() if p.name.startswith("Shared")]) == 2, "repeated: no third copy"
        finally:
            archive.ARCHIVE_ROOT, archive._ARCHIVE_ROOT_RESOLVED, archive.probe_resource = keep


def test_portal_session_is_reused_for_15_minutes():
    """The owner's portal session is kept after a read and reused within REUSE_SECONDS: the browser
    opens with its cookies and no sign-in happens. An ended session (the sign-in form comes back)
    signs in again; an older one is not even tried; a guest's sign-in never uses or touches it;
    a read that fails drops it."""
    from uninotifier import portal

    class Locator:
        def __init__(self, page, selector):
            self.page, self.selector = page, selector

        def count(self):
            return 0 if self.page.signed or "lognForm" not in self.selector else 1

    class Page:
        def __init__(self, signed):
            self.signed, self.url = signed, ""

        def goto(self, url, **kwargs):
            self.url = url

        def locator(self, selector):
            return Locator(self, selector)

        def wait_for_load_state(self, *a, **k):
            pass

        def wait_for_selector(self, *a, **k):
            pass

        def evaluate(self, *a, **k):
            return 1

        def wait_for_timeout(self, *a):
            pass

    server = {"alive": True, "states": [], "fail": False}

    class Context:
        def __init__(self, state):
            self.page = Page(signed=bool(state) and server["alive"])

        def new_page(self):
            return self.page

        def storage_state(self):
            return {"cookies": [{"name": "JSESSIONID", "value": f"s{len(server['states'])}"}], "origins": []}

    class Browser:
        def new_context(self, storage_state=None, **kwargs):
            server["states"].append(storage_state)
            return Context(storage_state)

        def close(self):
            pass

    class Playwright:
        chromium = type("C", (), {"launch": staticmethod(lambda headless=True: Browser())})()

        def __enter__(self):
            return self

        def __exit__(self, *a):
            return False

    sign_ins = []

    def login(page, user, password, ask):
        sign_ins.append(user)
        page.signed = True

    def goto_section(page, spec):
        if server["fail"]:
            raise RuntimeError("the page broke")
        return False

    keep = (_playwright_api.sync_playwright, portal.login, portal._goto_section, portal._read, portal._screenshot,
            portal._keep_debug, portal.credentials, portal.SESSION_FILE)
    with tempfile.TemporaryDirectory() as tmp:
        _playwright_api.sync_playwright = lambda: Playwright()
        portal.login, portal._goto_section = login, goto_section
        portal._read = lambda page, script, default: default
        portal._screenshot = lambda page: b"png"
        portal._keep_debug = lambda *a: None
        portal.credentials = lambda: ("owner", "pw")
        portal.SESSION_FILE = Path(tmp) / "portal_session.json"
        try:
            portal.fetch_many(["transcript"], None)
            assert sign_ins == ["owner"] and server["states"] == [None] and portal.SESSION_FILE.exists()
            saved = json.loads(portal.SESSION_FILE.read_text(encoding="utf-8"))
            assert saved["state"]["cookies"][0]["value"] == "s1"

            portal.fetch_many(["balance"], None)
            assert sign_ins == ["owner"], "within 15 min: the session is reused, no sign-in"
            assert server["states"][-1]["cookies"][0]["value"] == "s1"

            server["alive"] = False                       # the portal ended it: the sign-in form comes back
            portal.fetch_many(["balance"], None)
            assert sign_ins == ["owner", "owner"], "an ended session signs in again, silently"

            server["alive"] = True
            saved = json.loads(portal.SESSION_FILE.read_text(encoding="utf-8"))
            saved["at"] -= portal.REUSE_SECONDS
            portal.SESSION_FILE.write_text(json.dumps(saved), encoding="utf-8")
            assert portal.saved_session() is None
            portal.fetch_many(["balance"], None)
            assert server["states"][-1] is None and len(sign_ins) == 3, "an old session is not even tried"

            before = portal.SESSION_FILE.read_text(encoding="utf-8")
            portal.fetch_many(["transcript"], None, creds=("guest", "gpw"), keep_debug=False)
            assert sign_ins[-1] == "guest" and server["states"][-1] is None, "a guest never uses the owner's session"
            assert portal.SESSION_FILE.read_text(encoding="utf-8") == before, "... nor replaces it"

            server["fail"] = True
            try:
                portal.fetch_many(["balance"], None)
                raise AssertionError("the failure must reach the caller")
            except RuntimeError:
                pass
            assert not portal.SESSION_FILE.exists(), "a failed read drops the session"
        finally:
            (_playwright_api.sync_playwright, portal.login, portal._goto_section, portal._read, portal._screenshot,
             portal._keep_debug, portal.credentials, portal.SESSION_FILE) = keep

    assert "JSESSIONID=<redacted>" in notifier.redact("Cookie: JSESSIONID=AbC123secret; Path=/")


def test_portal_answers_from_the_saved_copy():
    """Transcript and Balance answer at once from the last read ("As of"): younger than
    FRESH_SECONDS, nothing is read; older, the portal is read behind it and the same message is
    updated in place — the stamp only when nothing changed, the screenshot (and the text that
    follows) when something did. A failed refresh leaves the saved copy, saying so. The screenshot
    of a saved answer is sent from the file (by its file_id). Grades are never answered from it."""
    bot = _package("uninotifier.bot")

    sent, calls, screens, reads = [], [], [], []
    ids = iter(range(100, 1000))

    def send_document(tok, chat, path, caption, url=None, rows=None, reply=None):
        mid = next(ids)
        sent.append(("doc", path.name, caption, mid))
        if reply is not None:
            reply["body"] = {"result": {"message_id": mid}}
        return True

    def send_message(tok, chat, text, url=None, rows=None, silent=False, reply_to=None, reply=None):
        mid = next(ids)
        sent.append(("msg", text, mid))
        if reply is not None:
            reply["body"] = {"result": {"message_id": mid}}
        return True

    answers = []

    def fetch_many(sections, ask):
        reads.append(sections)
        answer = answers.pop(0)
        if isinstance(answer, Exception):
            raise answer
        return {sections[0]: answer}

    keep = (bot.notifier.send_png_file, bot.notifier.send_message, bot.notifier.send_long_message,
            bot.notifier.send_document, bot.notifier.upload, bot.send_view, bot.edit_view, bot.api,
            bot.portal.fetch_many, bot.sync_base, bot.sync_retakes, bot.lock_grades)
    bot.notifier.send_png_file = lambda tok, chat, png, name, caption="", url=None, rows=None: (
        sent.append(("png", name, caption)) or True)
    bot.notifier.send_long_message = lambda tok, chat, text, url=None, rows=None: sent.append(("long", text)) or True
    bot.notifier.send_message, bot.notifier.send_document = send_message, send_document
    bot.notifier.upload = lambda tok, method, data, field, name, payload: (
        calls.append((method, data)) or {"message_id": 100, "document": {"file_id": "F2"}})
    bot.send_view = lambda text, rows: screens.append(text) or {"message_id": 1}
    bot.edit_view = lambda mid, text, rows: None
    bot.api = lambda method, **d: calls.append((method, d)) or True
    bot.portal.fetch_many = fetch_many
    bot.sync_base = bot.sync_retakes = bot.lock_grades = lambda *a: None

    def run(section):
        assert bot._portal_lock.acquire(blocking=False)
        bot.run_portal(section, None)
        assert not bot._portal_lock.locked()

    def age(section, seconds):
        meta = bot.CACHE_DIR / f"{section}.json"
        saved = json.loads(meta.read_text(encoding="utf-8"))
        saved["at"] = time.time() - seconds
        meta.write_text(json.dumps(saved), encoding="utf-8")

    def reset():
        sent.clear(); calls.clear(); screens.clear(); reads.clear()

    balance = lambda total: portal_result("balance", {"balance": total}, [f"💰 *Total balance:* {total:,.0f} JD"])  # noqa: E731

    def portal_result(section, figures, lines, data=None):
        from uninotifier import portal
        return portal.Result(section, True, b"png-" + str(figures).encode(), [], figures, [], lines, data or {})

    try:
        # no saved copy: read now (as before), and the answer is kept
        answers[:] = [balance(1250)]
        run("balance")
        assert [s[0] for s in sent] == ["png"] and reads == [["balance"]]
        assert (bot.CACHE_DIR / "balance.json").exists() and (bot.CACHE_DIR / "ZUJ Portal - Tuition Balance.png").exists()

        # a fresh copy: answered from it, nothing read
        reset()
        run("balance")
        assert reads == [] and sent[0][0] == "doc" and sent[0][1] == "ZUJ Portal - Tuition Balance.png"
        assert "🕒 *As of:*" in sent[0][2] and bot.REFRESHING not in sent[0][2] and "1,250 JD" in sent[0][2]
        assert screens and screens[-1].startswith("📜 *Academic Record & Balance*"), "the hub comes back under it"

        # an old copy, nothing changed: answered at once, then only its stamp changes
        reset()
        age("balance", bot.FRESH_SECONDS + 1)
        answers[:] = [balance(1250)]
        run("balance")
        assert bot.REFRESHING in sent[0][2] and reads == [["balance"]]
        assert [c[0] for c in calls] == ["editMessageCaption"], calls
        caption = calls[0][1]["caption"]
        assert "🕒 *Retrieved:*" in caption and bot.UPDATED not in caption
        assert calls[0][1]["message_id"] == sent[0][3], "the message that was sent is the one updated"
        assert json.loads((bot.CACHE_DIR / "balance.json").read_text(encoding="utf-8"))["at"] > time.time() - 60

        # an old copy, the balance changed: the screenshot and caption are replaced in the same message
        reset()
        age("balance", bot.FRESH_SECONDS + 1)
        answers[:] = [balance(900)]
        run("balance")
        assert "1,250 JD" in sent[0][2], "first the saved copy"
        assert [c[0] for c in calls] == ["editMessageMedia"], calls
        media = json.loads(calls[0][1]["media"])
        assert "900 JD" in media["caption"] and bot.UPDATED in media["caption"] and media["type"] == "document"

        # the refresh fails: the saved copy stays, and says so
        from uninotifier import portal
        reset()
        age("balance", bot.FRESH_SECONDS + 1)
        answers[:] = [portal.PortalError("The portal refused the sign-in.")]
        run("balance")
        assert [c[0] for c in calls] == ["editMessageCaption"] and bot.REFRESH_FAILED in calls[0][1]["caption"]
        assert "900 JD" in calls[0][1]["caption"]

        # a transcript: the semesters follow the caption, and are edited in place when they change
        sems = _sample_transcript()
        record = lambda cgpa: portal_result("transcript", {"cgpa": cgpa, "passed_hours": 8.0, "record_cgpa": cgpa, "record_hours": 14.0},  # noqa: E731
                                            [], {"semesters": sems, "status": "منتظم", "retakes": {}})
        reset()
        answers[:] = [record(71.3)]
        run("transcript")
        reset()
        age("transcript", bot.FRESH_SECONDS + 1)
        answers[:] = [record(71.2)]
        run("transcript")
        assert [s[0] for s in sent] == ["doc", "msg"], sent
        assert [c[0] for c in calls] == ["editMessageMedia", "editMessageText"], [c[0] for c in calls]
        assert calls[0][1]["message_id"] == sent[0][3] and calls[1][1]["message_id"] == sent[1][2]
        assert "71.2" in json.loads(calls[0][1]["media"])["caption"], "the new record is in the caption"

        # grades (the Finals Watch check) are never answered from a saved copy
        assert bot.load_cached("grades") is None
    finally:
        (bot.notifier.send_png_file, bot.notifier.send_message, bot.notifier.send_long_message,
         bot.notifier.send_document, bot.notifier.upload, bot.send_view, bot.edit_view, bot.api,
         bot.portal.fetch_many, bot.sync_base, bot.sync_retakes, bot.lock_grades) = keep


def _with_env(values: dict, fn):
    """Run fn with these process variables set ("" = set but empty, None = removed), then put them back."""
    keep = {k: os.environ.get(k) for k in values}
    try:
        for k, v in values.items():
            if v is None:
                os.environ.pop(k, None)
            else:
                os.environ[k] = v
        return fn()
    finally:
        for k, v in keep.items():
            if v is None:
                os.environ.pop(k, None)
            else:
                os.environ[k] = v


def test_settings_are_checked_before_anything_runs():
    """A missing or still-example setting stops a run with a plain list of what to fix (exit code 2):
    in the log, and on Telegram when Telegram itself is set up — never a traceback."""
    from uninotifier import envcheck
    assert envcheck.problems() == [], envcheck.problems()        # the made-up .env is complete

    # values copied unchanged from .env.example count as not set
    example = {k: v for k, v in __import__("dotenv").dotenv_values(paths.ENV_EXAMPLE).items() if v}
    for name in ("ZUJ_USERNAME", "ZUJ_PASSWORD", "TELEGRAM_BOT_TOKEN", "TELEGRAM_CHAT_ID"):
        assert name not in example or _with_env({name: example[name]}, lambda n=name: envcheck.value(n)) == "", name
    for raw in ("your_password", "https://hc-ping.com/your-check-uuid", "https://x.org/class/<class-id>/a",
                "C:\\path\\to\\your\\semester\\folder", "..."):
        assert _with_env({"HEALTHCHECK_URL": raw}, lambda: envcheck.value("HEALTHCHECK_URL")) == "", raw
    assert _with_env({"ZUJ_PASSWORD": " s3cret "}, lambda: envcheck.value("ZUJ_PASSWORD")) == "s3cret"

    # who needs what: the bot runs without the Moodle password, the check doesn't
    missing = _with_env({"ZUJ_PASSWORD": ""}, lambda: (envcheck.problems("bot"), envcheck.problems("monitor")))
    assert missing[0] == [] and len(missing[1]) == 1 and missing[1][0].startswith("ZUJ_PASSWORD is not set"), missing
    shape = _with_env({"TELEGRAM_BOT_TOKEN": "not a token", "ZUJ_USERNAME": "no-at-sign"}, envcheck.problems)
    assert any(p.startswith("TELEGRAM_BOT_TOKEN doesn't look right") for p in shape)
    assert any(p.startswith("ZUJ_USERNAME doesn't look right") for p in shape), shape
    other = _with_env({"ADMIN_USER_ID": "7"}, envcheck.problems)
    assert other == ["TELEGRAM_CHAT_ID must be your private chat with the bot - the same number as ADMIN_USER_ID."], other
    warned = _with_env({"PORTAL_USERNAME": "u1234567", "PORTAL_PASSWORD": "", "CAMBRIDGE_UNTIL": "2020-01-31",
                        "CAMBRIDGE_ASSIGNMENTS_URL": "", "ARCHIVE_ROOT": "relative\\folder"}, envcheck.warnings)
    assert len(warned) == 4 and envcheck.problems() == [], warned

    # require(): the run ends with exit code 2 and the list; Telegram hears about it once
    _fresh_failsafe()

    def stopped(role):
        import contextlib
        import io
        try:
            with contextlib.redirect_stderr(io.StringIO()):      # the console copy of the list
                envcheck.require(role)
        except SystemExit as e:
            return e.code
        return None
    assert stopped("monitor") is None
    assert _with_env({"ZUJ_PASSWORD": ""}, lambda: stopped("monitor")) == envcheck.EXIT_CONFIG
    failsafe.wait_idle()
    assert len(OUTBOX) == 1 and "ZUJ\\_PASSWORD is not set" in OUTBOX[0], OUTBOX
    assert _with_env({"ZUJ_PASSWORD": ""}, lambda: stopped("monitor")) == envcheck.EXIT_CONFIG
    failsafe.wait_idle()
    assert len(OUTBOX) == 1, "repeated runs with the same mistake are held back"
    # without Telegram there is no one to tell: the log and the console only
    _fresh_failsafe()
    assert _with_env({"TELEGRAM_BOT_TOKEN": ""}, lambda: stopped("bot")) == envcheck.EXIT_CONFIG
    failsafe.wait_idle()
    assert OUTBOX == []
    text = envcheck.report(["X is not set - y."])
    assert "  - X is not set - y." in text and str(paths.ENV_FILE) in text and "uninotifier.doctor" in text


def test_archive_root_defaults_to_documents_university():
    """ARCHIVE_ROOT is optional: unset, files go to University in the Documents folder Explorer shows
    (OneDrive's when OneDrive backs up Documents); set, %VARS% and ~ are expanded, so one .env line
    works on any PC."""
    empty = _HOME / "empty.env"
    empty.write_text("", encoding="utf-8")
    keep = paths.ENV_FILE, archive._shell_documents
    paths.ENV_FILE = empty
    try:
        assert archive.DEFAULT_ARCHIVE_ROOT == archive.documents_folder() / "University"
        assert _with_env({"ARCHIVE_ROOT": None}, archive._archive_root) == archive.DEFAULT_ARCHIVE_ROOT
        assert _with_env({"ARCHIVE_ROOT": "  "}, archive._archive_root) == archive.DEFAULT_ARCHIVE_ROOT
        # which Documents: Windows' shell folder first, then OneDrive's, then the home folder's
        shell, onedrive = _HOME / "ShellDocs", _HOME / "OneDrive - ZUJ"
        shell.mkdir(exist_ok=True)
        (onedrive / "Documents").mkdir(parents=True, exist_ok=True)
        archive._shell_documents = lambda: str(shell)
        assert _with_env({"OneDrive": str(onedrive)}, archive.documents_folder) == shell
        archive._shell_documents = lambda: ""
        assert _with_env({"OneDrive": str(onedrive)}, archive.documents_folder) == onedrive / "Documents"
        archive._shell_documents = lambda: str(_HOME / "gone")       # a shell folder that doesn't exist is skipped
        assert _with_env({"OneDrive": None}, archive.documents_folder) == Path.home() / "Documents"
        got = _with_env({"ARCHIVE_ROOT": "%UNINOTIFIER_T%\\Semester_1", "UNINOTIFIER_T": str(_HOME)}, archive._archive_root)
        assert got == _HOME / "Semester_1", got
        assert _with_env({"ARCHIVE_ROOT": "~/Uni"}, archive._archive_root) == Path.home() / "Uni"
    finally:
        paths.ENV_FILE, archive._shell_documents = keep
    assert archive.ARCHIVE_ROOT == _HOME / "archive", "the tests never touch the real archive"


def test_state_backup_lives_in_its_own_repository():
    """The code repository never holds your state. gitsync --init makes data/ a git repository of its
    own that commits only the tracker and your own data — never sessions or caches — and each check
    commits a change (kept locally while there is no remote)."""
    import subprocess
    from uninotifier import gitsync
    repo = Path(tempfile.mkdtemp(prefix="uninotifier-backup-"))
    (repo / "saved_items.json").write_text('{"courses": {}}', encoding="utf-8")
    (repo / "gpa.json").write_text("{}", encoding="utf-8")
    (repo / "auth_state.json").write_text('{"cookies": "secret"}', encoding="utf-8")
    (repo / "logs").mkdir()
    (repo / "logs" / "monitor.log").write_text("x", encoding="utf-8")
    keep = gitsync.STATE_REPO
    gitsync.STATE_REPO = repo
    try:
        assert not gitsync.enabled()
        gitsync.sync_state()                                   # off: does nothing, creates nothing
        assert not (repo / ".git").exists()
        import contextlib
        import io
        with contextlib.redirect_stdout(io.StringIO()):
            assert gitsync.init() == 0 and gitsync.enabled()

        def tracked():
            return sorted(subprocess.run(["git", "ls-files"], cwd=repo, capture_output=True, text=True).stdout.split())
        assert tracked() == [".gitignore", "gpa.json", "saved_items.json"], tracked()
        with contextlib.redirect_stdout(io.StringIO()):
            assert gitsync.init() == 0, "running --init again is harmless"
        (repo / "saved_items.json").write_text('{"courses": {"1": "x"}}', encoding="utf-8")
        (repo / "reminders.json").write_text("[]", encoding="utf-8")
        gitsync.sync_state()
        log_lines = subprocess.run(["git", "log", "--format=%s"], cwd=repo, capture_output=True, text=True).stdout.split("\n")
        assert log_lines[0] == gitsync.COMMIT_MESSAGE and "reminders.json" in tracked() and "auth_state.json" not in tracked()
        status = subprocess.run(["git", "status", "--porcelain"], cwd=repo, capture_output=True, text=True).stdout
        assert status == "", status
    finally:
        gitsync.STATE_REPO = keep
        shutil.rmtree(repo, ignore_errors=True)


def test_doctor_says_what_to_fix():
    """python -m uninotifier.doctor: every check in plain words, exit 0 when nothing blocks a run."""
    import contextlib
    import io
    from uninotifier import doctor
    keep = doctor.check_browser, doctor.check_tasks
    doctor.check_browser = lambda r: r.line(doctor.OK, "Chromium (stubbed in the tests)")
    doctor.check_tasks = lambda r: None

    def run(*argv):
        out = io.StringIO()
        with contextlib.redirect_stdout(out):
            code = doctor.main(list(argv))
        return code, out.getvalue()
    try:
        code, text = run("--offline")
        assert code == 0 and "[ OK ] Settings (.env)" in text and "[ OK ] Timetable: First Semester" in text, text
        assert f"[ OK ] Course files are saved to {_HOME / 'archive'}" in text and "All set." in text, text
        assert "[ -- ] Cambridge One: off" in text, "nothing set: the feature is off and says so"

        # Cambridge One is for the English course: on only when asked, with or without its own sign-in,
        # and a warning when no course of yours looks like an English one
        on = {"CAMBRIDGE_ASSIGNMENTS_URL": "https://www.cambridgeone.org/class/learner/o/class/c/assignments",
              "CAMBRIDGE_UNTIL": "2099-12-31"}
        code, text = _with_env(on, lambda: run("--offline"))
        assert code == 0 and "[ OK ] Cambridge One: watching until 2099-12-31 (the session ends every ~3 days" in text, text
        assert "Cambridge One is on, but" not in text, "the tracker lists an English course"
        code, text = _with_env({**on, "CAMBRIDGE_EMAIL": "me@example.invalid", "CAMBRIDGE_PASSWORD": "pw"},
                               lambda: run("--offline"))
        assert "watching until 2099-12-31, and it signs in again by itself" in text, text
        keep_data = paths.DATA
        paths.DATA = Path(tempfile.mkdtemp(prefix="uninotifier-doctor-data-"))
        try:
            (paths.DATA / "saved_items.json").write_text(json.dumps({"courses": {"1": {"name": "20261-Big Data-AllSec-3-"}}}),
                                                         encoding="utf-8")
            code, text = _with_env(on, lambda: run("--offline"))
            assert code == 0 and "[WARN] Cambridge One is on, but none of your Moodle courses looks like an English course" in text, text
            (paths.DATA / "saved_items.json").write_text(json.dumps({"courses": {"1": "20261-English Communication-1"}}),
                                                         encoding="utf-8")
            assert "Cambridge One is on, but" not in _with_env(on, lambda: run("--offline"))[1]
            (paths.DATA / "saved_items.json").unlink()
            assert "Cambridge One is on, but" not in _with_env(on, lambda: run("--offline"))[1], "no tracker yet: no guess"
        finally:
            shutil.rmtree(paths.DATA, ignore_errors=True)
            paths.DATA = keep_data

        code, text = _with_env({"ZUJ_PASSWORD": ""}, lambda: run("--offline"))
        assert code == 1 and "[FAIL] ZUJ_PASSWORD is not set" in text and "1 problem(s) to fix" in text, text
        keep_schedule = paths.SCHEDULE_FILE
        paths.SCHEDULE_FILE = _HOME / "no-such-schedule.json"
        try:
            code, text = run("--offline")
        finally:
            paths.SCHEDULE_FILE = keep_schedule
        assert code == 1 and "copy schedule.example.json to no-such-schedule.json" in text, text

        # the example left as it was installed (and saved by Notepad with a BOM): fails, so no fake lecture alerts
        tmp = Path(tempfile.mkdtemp(prefix="uninotifier-doctor-"))
        example = json.loads(paths.SCHEDULE_EXAMPLE.read_text(encoding="utf-8"))
        own = json.loads(keep_schedule.read_text(encoding="utf-8"))
        paths.SCHEDULE_FILE = tmp / "schedule.json"
        try:
            paths.SCHEDULE_FILE.write_text("﻿" + json.dumps(example, ensure_ascii=False), encoding="utf-8")
            code, text = run("--offline")
            assert code == 1 and "[FAIL] schedule.json still holds the example classes" in text, text
            assert "not valid JSON" not in text and "[ OK ] Timetable" not in text, text
            own["days"].setdefault("mon", []).append(example["days"]["mon"][0])     # one example class left behind
            paths.SCHEDULE_FILE.write_text(json.dumps(own), encoding="utf-8")
            code, text = run("--offline")
            assert code == 0 and "[ OK ] Timetable" in text and "[WARN] 1 class(es) in schedule.json are still the example's" in text, text
        finally:
            paths.SCHEDULE_FILE = keep_schedule
            shutil.rmtree(tmp, ignore_errors=True)
        assert all(ord(ch) < 128 for line in text.splitlines() if line.startswith("[") for ch in line[:6])
    finally:
        doctor.check_browser, doctor.check_tasks = keep


def test_no_personal_files_are_tracked():
    """A clone gets the code and the examples — never someone's settings, tracker, timetable or session."""
    import subprocess
    project = Path(__file__).resolve().parents[1]
    if not (project / ".git").exists():
        print("  (skipped: not a git checkout)")
        return
    files = subprocess.run(["git", "ls-files"], cwd=project, capture_output=True, text=True).stdout.splitlines()
    assert files, "git ls-files returned nothing"
    personal = [f for f in files if f in (".env", "schedule.json") or f.startswith(("data/", "logs/", "assets/Weekly_"))]
    assert personal == [], personal
    for example in (".env.example", "schedule.example.json", "scripts/setup_tasks.ps1", "scripts/remove_tasks.ps1"):
        assert (project / example).is_file(), example
    ignored = subprocess.run(["git", "check-ignore", ".env", "schedule.json", "data/saved_items.json",
                              "assets/Weekly_Schedule.png"], cwd=project, capture_output=True, text=True).stdout.split()
    assert len(ignored) == 4, ignored
    assert schedule.validate(json.loads((project / "schedule.example.json").read_text(encoding="utf-8"))) == []


def test_the_repository_is_ready_for_students():
    """What a student and a contributor rely on stays in place: exact pins that match pyproject.toml,
    Dependabot for packages and actions, the double-click installer (with Windows line endings), CI that
    runs that installer, issue forms that refuse to be sent without the privacy promise, and the
    security and contributing guides."""
    import subprocess
    project = Path(__file__).resolve().parents[1]

    def lines(name):
        return [l.strip() for l in (project / name).read_text(encoding="utf-8").splitlines()
                if l.strip() and not l.strip().startswith("#")]
    pins = lines("requirements.txt")
    assert all(re.fullmatch(r"[A-Za-z0-9_.-]+==[0-9][\w.]*", p) for p in pins), pins
    assert lines("requirements-dev.txt")[0] == "-r requirements.txt"
    assert all(re.fullmatch(r"[A-Za-z0-9_.-]+==[0-9][\w.]*", p) for p in lines("requirements-dev.txt")[1:])
    declared = re.search(r"^dependencies = \[(.*?)\]", (project / "pyproject.toml").read_text(encoding="utf-8"), re.S | re.M)
    names = {re.split(r"[<>=!~ ]", d.strip().strip('",'))[0].lower() for d in declared.group(1).split(",") if d.strip()}
    assert names == {p.split("==")[0].lower() for p in pins}, ("pyproject.toml and requirements.txt list different packages", names)

    dependabot = (project / ".github" / "dependabot.yml").read_text(encoding="utf-8")
    assert "package-ecosystem: pip" in dependabot and "package-ecosystem: github-actions" in dependabot
    assert dependabot.count("interval: weekly") == 2

    bat = (project / "install.bat").read_bytes()
    assert b"-ExecutionPolicy Bypass" in bat and b'scripts\\install.ps1"' in bat and b"pause" in bat
    attributes = (project / ".gitattributes").read_text(encoding="utf-8")
    assert "*.bat text eol=crlf" in attributes and "*.ps1 text eol=crlf" in attributes
    if (project / ".git").exists():
        checked = subprocess.run(["git", "check-attr", "eol", "install.bat", "scripts/install.ps1"], cwd=project,
                                 capture_output=True, text=True).stdout
        assert checked.count("eol: crlf") == 2, checked
    installer = (project / "scripts" / "install.ps1").read_text(encoding="utf-8")
    assert all(ord(ch) < 128 for ch in installer), "Windows PowerShell 5.1 misreads a BOM-less non-ASCII script"
    for step in ("'3.11'", "'-r', 'requirements.txt'", "'--no-deps', '-e', '.'", "'playwright', 'install', 'chromium'",
                 "'.env.example', '.env'", "'schedule.example.json', 'schedule.json'", "-m uninotifier.doctor"):
        assert step in installer, f"the installer lost: {step}"
    ci = (project / ".github" / "workflows" / "tests.yml").read_text(encoding="utf-8")
    assert "installer:" in ci and "scripts\\install.ps1 -NoEditor" in ci and "pip install --no-deps -e ." in ci

    promise = "I have verified that no passwords, student IDs, or private tokens are included."
    for form in ("bug_report.yml", "feature_request.yml"):
        text = (project / ".github" / "ISSUE_TEMPLATE" / form).read_text(encoding="utf-8")
        assert re.search(rf"- label: {re.escape(promise)}\n\s+required: true", text), form
    config = (project / ".github" / "ISSUE_TEMPLATE" / "config.yml").read_text(encoding="utf-8")
    assert "blank_issues_enabled: false" in config and "/security/advisories/new" in config
    security = (project / "SECURITY.md").read_text(encoding="utf-8")
    assert "/security/advisories/new" in security and "Don't open a public issue" in security
    contributing = (project / "CONTRIBUTING.md").read_text(encoding="utf-8")
    for rule in ("Tests are offline", "No personal data", "Both READMEs change together", "The bot speaks English"):
        assert rule in contributing, rule


REPO_PATH = "github.com/Humam-Taibeh/UniNotifier"      # the project's own public address
PUBLIC_REPO = f"https://{REPO_PATH}.git"


def test_the_readme_translations_say_the_same():
    """README.md and README.ar.md link to each other and stay in step: the same badges, commands,
    diagram and sample output (comments may be translated), every relative link pointing at a file —
    and the tests badge is the live status of the CI workflow that runs this suite."""
    project = Path(__file__).resolve().parents[1]
    en = (project / "README.md").read_text(encoding="utf-8")
    ar = (project / "README.ar.md").read_text(encoding="utf-8")
    assert en.splitlines()[0] == "[العربية](README.ar.md) | **English**", en.splitlines()[0]
    assert ar.splitlines()[0] == "**العربية** | [English](README.md)", ar.splitlines()[0]

    def badges(text):
        head = text.split("\n# UniNotifier\n", 1)[0].splitlines()
        assert head[1:4] == ["", "---", ""] and head[4] == '<div align="center">' and head[-1] == "</div>", head
        return [line for line in head if line.startswith("[![")]
    assert badges(en) == badges(ar) and len(badges(en)) == 5, "the same five badges on both pages"
    labels = " ".join(badges(en))
    for label in ("Python 3.11+", "Platform: Windows", "Code style: pyflakes clean", "License: MIT"):
        assert f"[![{label}](" in labels, label
    workflow = f"https://{REPO_PATH}/actions/workflows/tests.yml"
    assert f"[![Tests]({workflow}/badge.svg)]({workflow})" in labels, "the tests badge is the CI workflow's own"
    ci = (project / ".github" / "workflows" / "tests.yml").read_text(encoding="utf-8")
    for step in ("runs-on: windows-latest", '"3.11"', "python tests/test_architecture.py",
                 "python -m pyflakes src tests monitor.py bot.py", "python -m playwright install chromium"):
        assert step in ci, f"the CI workflow lost: {step}"

    def blocks(text):
        found = []
        for body in re.findall(r"^```\w*\n(.*?)^```", text, re.S | re.M):
            lines = [re.sub(r"\s+#\s.*$", "", line).rstrip() for line in body.splitlines()]
            found.append([line for line in lines if line and not line.lstrip().startswith("# ")])
        return found
    assert len(blocks(en)) >= 8 and blocks(en) == blocks(ar), "a command, the diagram or a sample differs"

    for name, text in (("README.md", en), ("README.ar.md", ar)):
        for target in re.findall(r"\]\(([^)#]+?)\)", text):
            if not target.startswith(("http://", "https://")):
                assert (project / target).exists(), (name, target)
        assert f"git clone {PUBLIC_REPO}\n" in text and "%USERPROFILE%" in text, (name, "clone address and placeholder")
    assert [line.count("|") for line in en.splitlines() if line.startswith("|")] == \
           [line.count("|") for line in ar.splitlines() if line.startswith("|")], "the tables have the same shape"


def portal_sections():
    from uninotifier import portal
    return portal.SECTIONS


if __name__ == "__main__":
    _CHAT_ORIGINALS = {n: getattr(_package("uninotifier.bot"), n) for n in ("send_view", "edit_view", "api")}
    for fn in (test_grades, test_schedule, test_media, test_ui, test_submission_status,
               test_deadline_flag_and_reminders, test_bot_views, test_richtext_and_course_page,
               test_display_names_and_completion_policy, test_forum_messages_and_state,
               test_edit_detection, test_failure_policy, test_notifier_helpers,
               test_mfa_flow, test_mfa_policy_and_bot_relay, test_lecture_links,
               test_gpa_formulas, test_gpa_parsing_and_need, test_gpa_screens_and_flow,
               test_portal, test_portal_parsers, test_http_layer, test_healthcheck,
               test_fast_scan_helpers, test_gradebook_record, test_results_page_parser,
               test_finals_watch, test_finals_commands_and_file_destination, test_finals_check_ends_on_the_grades_screen, test_transcript_layout, test_portal_answer_is_one_message, test_archive,
               test_archive_layout, test_resource_page_file_link, test_migration_helpers,
               test_downloader_dedup_and_routing, test_tidy_archive, test_alerts_are_recorded_as_they_are_sent, test_archive_boundary,
               test_auto_done_only_for_saved_files_and_links, test_announce_first_run_and_duplicate_files,
               test_day_separator_and_daily_recap, test_help_hub_dashboard_and_campus,
               test_update_routing_from_telegram, test_access_lock, test_modules_import_quietly_and_names_resolve,
               test_failsafe_reports_once_and_resolves, test_bot_contains_and_reports_failures,
               test_bot_heartbeat_and_monitor_watch, test_bot_self_heal, test_schedule_validation_and_last_good,
               test_state_backup_and_recovery, test_telegram_transport, test_buttons_are_answered_first,
               test_files_telegram_already_holds_are_not_uploaded_again, test_monitor_pruning_and_background_pings, test_answers_go_out_beside_the_edit, test_course_list_is_a_timetable, test_rooms_are_always_short, test_one_menu_anchored_at_the_bottom, test_poll_outage_tracking_and_quiet_logs,
               test_network_resilience, test_maintenance_and_finals_pause,
               test_teams_meeting_and_attendance_items, test_tasks_and_reminders_planner,
               test_pin_rule_and_forum_announcements, test_retake_replacement_in_plan, test_guest_lookup,
               test_guest_summary_and_target_math,
               test_attendance_on_online_lecture_alert, test_guest_lookup_in_telegram,
               test_no_personal_values_in_the_code, test_layout_watchdog,
               test_recorded_real_pages, test_upload_tried_again_sends_the_whole_file,
               test_refused_markdown_goes_out_as_plain_text, test_your_own_data_survives_a_damaged_file,
               test_mark_as_done_never_writes_back_an_older_state, test_catch_alls_tell_the_owner,
               test_state_files_are_written_through_fileio, test_forwarded_messages_are_never_acted_on,
               test_secrets_never_reach_logs_or_alerts, test_guest_password_never_stays_in_the_chat,
               test_stranger_floods_cost_no_disk_access, test_the_run_is_isolated_from_your_files,
               test_the_guards_block_the_network_and_the_browser, test_uploads_use_their_own_send_buffer,
               test_schedule_files_never_hold_up_the_bot, test_a_restart_never_replays_an_update,
               test_a_long_retry_after_is_not_waited_out, test_windows_device_names_are_saved_as_files,
               test_the_poll_slows_down_when_telegram_asks, test_an_odd_getupdates_answer_never_stops_the_loop,
               test_the_statistics_do_not_grow_forever, test_the_moodle_probe, test_cambridge_one_assignments,
               test_cambridge_signs_in_by_itself, test_log_files_stay_small,
               test_portal_session_is_reused_for_15_minutes, test_portal_answers_from_the_saved_copy,
               test_replaced_files_move_to_old_versions, test_settings_are_checked_before_anything_runs,
               test_archive_root_defaults_to_documents_university, test_state_backup_lives_in_its_own_repository,
               test_sign_in_needs_no_form_when_microsoft_still_remembers_you,
               test_doctor_says_what_to_fix, test_no_personal_files_are_tracked,
               test_the_readme_translations_say_the_same, test_the_repository_is_ready_for_students):
        _chat = _package("uninotifier.bot")
        for _name, _original in _CHAT_ORIGINALS.items():      # a test that forgets to restore a stub cannot leak it into the next
            setattr(_chat, _name, _original)
        _chat._answered.clear()           # the tests reuse made-up button ids; Telegram's are unique
        _chat._last_press.clear()
        _chat.UI_FILE = _REMINDER_TMP / "ui_state.json"       # which message is the menu: never the real file
        _chat.reset_menu()
        _chat.MONITOR = _REMINDER_TMP / "never-run-the-real-monitor.py"     # /check in a test must not start a real check
        shutil.rmtree(_chat.CACHE_DIR, ignore_errors=True)    # a portal answer saved by one test is not the next one's
        importlib.import_module("uninotifier.portal").drop_session()   # ... nor its portal session
        _chat.ANSWER_ASYNC = False          # answers are sent in line, so a test can read them at once (test_buttons_are_answered_first covers the worker)
        _chat.DEBOUNCE_SECONDS = 0          # tests press the same button back to back; test_buttons_are_answered_first checks the debounce
        fn()
        if GUARD_HITS:                     # caught even when the code under test swallowed the refusal
            raise SystemExit(f"FAIL {fn.__name__} tried to reach the outside world: {GUARD_HITS}")
        print("ok", fn.__name__)
